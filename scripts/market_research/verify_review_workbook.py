"""Read-only checks of the artifact-authored market-content review workbook.

This verifier never saves a workbook or creates approval evidence. It compares
every immutable/editable initial value with the candidate row inventory, verifies
cached formula views, and checks basic readability/navigation structures.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl

HEADERS = (
    "country_code",
    "content_version",
    "locale",
    "json_pointer",
    "original_value",
    "edited_value",
)
SHEETS = (
    "审核导览",
    "国别总览",
    "中文阅读",
    "内容编辑",
    "来源依据",
    "待确认事项",
    "包信息",
)
LEVEL_LABELS = {"insufficient": "信息不足", "low": "低", "medium": "中", "high": "高"}
OVERVIEW_POINTER = re.compile(r"^/locales/(zh-CN|en)/(title|paragraphs/(?:0|[1-9]\d*)|disclaimer)$")
LEGACY_POINTER = re.compile(r"^/locales/(?:zh-CN|en)/(?:analysis|report)/.+$")
OVERVIEW_HEADERS = (
    "ISO3",
    "国家",
    "新能源市场概述标题",
    "开篇正文",
    "中文正文字符（去空白）",  # noqa: RUF001 - exact Chinese workbook header
    "自然段数",
    "待确认项",
    "修改行数",
    "内容版本",
    "数据截至",
    "阅读说明",
)


def review_profile(review: dict[str, Any]) -> str:
    """Recognize one immutable inventory; never reinterpret old leaves as new ones."""
    rows = review.get("content_rows")
    if (
        not isinstance(rows, list)
        or not rows
        or any(
            not isinstance(row, list)
            or len(row) != 6
            or any(not isinstance(value, str) for value in row)
            for row in rows
        )
    ):
        raise ValueError("review content must contain six literal string columns")
    if all(OVERVIEW_POINTER.fullmatch(row[3]) for row in rows):
        return "overview"
    if all(LEGACY_POINTER.fullmatch(row[3]) for row in rows):
        return "legacy_analysis_report"
    raise ValueError("mixed or unknown review pointers; legacy content cannot be converted")


def character_count(text: str) -> int:
    """The public contract counts Unicode characters, excluding whitespace."""
    return len(re.sub(r"\s", "", text))


def excel_length(text: str, *, exclude_whitespace: bool = False) -> int:
    """Excel LEN helpers use UTF-16 units; retain the separate contract count."""
    if exclude_whitespace:
        text = re.sub(r"\s", "", text)
    return len(text.encode("utf-16-le")) // 2


def overview_statistics(review: dict[str, Any]) -> dict[str, dict[str, Any]]:
    """Validate and derive the new overview view from its canonical rows only."""
    if review_profile(review) != "overview":
        raise ValueError("legacy reviews are read-only; no overview conversion is available")
    if review.get("content_headers") != list(HEADERS):
        raise ValueError("the six import-contract headers must remain unchanged")
    metadata = review.get("metadata", {})
    if metadata.get("schema_version") != "navigator.market-review.v1" or not re.fullmatch(
        r"[a-f0-9]{64}", str(metadata.get("candidate_sha256", ""))
    ):
        raise ValueError("overview review must be hash-bound")
    rows = review["content_rows"]
    country_rows = review.get("countries", [])
    countries = {item.get("code", item.get("country_code")): item for item in country_rows}
    if (
        not countries
        or len(countries) != len(country_rows)
        or len(countries) != metadata.get("country_count")
        or "CHN" in countries
        or set(countries) != {row[0] for row in rows}
    ):
        raise ValueError("outbound overview country inventory differs")
    labels = review.get("row_labels")
    if labels is not None and len(labels) != len(rows):
        raise ValueError("overview row labels must be complete")
    inventory: dict[tuple[str, str], list[str]] = {}
    for index, row in enumerate(rows):
        code, version, locale, pointer, original, edited = row
        match = OVERVIEW_POINTER.fullmatch(pointer)
        assert match is not None
        if (
            locale != match[1]
            or not version.startswith(f"OVERVIEW-{code}-")
            or original != edited
            or not original.strip()
            or (code, pointer) in inventory
        ):
            raise ValueError("overview locale, version, frozen value or pointer differs")
        inventory[code, pointer] = row
        if labels is not None:
            label = labels[index]
            if any(
                label.get(key) != value
                for key, value in (
                    ("country_code", code),
                    ("locale", locale),
                    ("json_pointer", pointer),
                    ("layer", "overview"),
                )
            ):
                raise ValueError("overview row label differs from its immutable row")
    for field in ("evidence_rows", "gap_rows"):
        if not isinstance(review.get(field), list) or any(
            item.get("country_code", item.get("iso3")) not in countries for item in review[field]
        ):
            raise ValueError("research rows must be separate and in the same country scope")
    result = {}
    for code, item in sorted(countries.items()):
        if not isinstance(code, str) or not re.fullmatch(r"[A-Z]{3}", code):
            raise ValueError("invalid overview country code")
        date.fromisoformat(item["as_of"])
        counts = []
        body = ""
        for locale in ("zh-CN", "en"):
            prefix = f"/locales/{locale}/"
            local = [row for row in rows if row[0] == code and row[2] == locale]
            paragraphs = [row for row in local if row[3].startswith(prefix + "paragraphs/")]
            if (
                not 6 <= len(paragraphs) <= 8
                or (code, prefix + "title") not in inventory
                or (code, prefix + "disclaimer") not in inventory
                or any(row[1] != item["content_version"] for row in local)
                or any(
                    (code, prefix + f"paragraphs/{index}") not in inventory
                    for index in range(len(paragraphs))
                )
            ):
                raise ValueError("overview title, disclaimer, paragraphs or version incomplete")
            counts.append(len(paragraphs))
            if locale == "zh-CN":
                body = "".join(row[5] for row in paragraphs)
        if counts[0] != counts[1] or counts[0] != item["paragraph_count"]:
            raise ValueError("bilingual overview paragraph counts differ")
        length = character_count(body)
        if not 1500 <= length <= 2000 or length != item["overview_zh_chars"]:
            raise ValueError("Chinese overview body count differs")
        result[code] = {
            "country_name": item["country_name"],
            "title": inventory[code, "/locales/zh-CN/title"][5],
            "opening": inventory[code, "/locales/zh-CN/paragraphs/0"][5],
            "overview_zh_chars": length,
            "excel_body_chars": excel_length(body, exclude_whitespace=True),
            "paragraph_count": counts[0],
            "content_version": item["content_version"],
            "as_of": item["as_of"],
        }
    return result


def literal(value: Any) -> str:
    """Match the author's literal-only private evidence cells."""
    if value is None:
        return ""
    if isinstance(value, list):
        return "；".join(literal(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def workbook_metadata(metadata: dict[str, Any]) -> dict[str, Any]:
    """The interchange permits an ISO creation date without inventing precision."""
    return {**metadata, "created_at": metadata["created_at"].split("T")[0]}


def reading_row(row: list[str]) -> bool:
    """Match the author's source-free Chinese reading view inventory."""
    pointer = row[3]
    return (
        row[2] == "zh-CN"
        and re.search(r"/(?:id|business|technology|verified_on)$", pointer) is None
        and re.search(r"/(?:businesses|technologies)/\d+$", pointer) is None
    )


def verify(review_path: Path, workbook_path: Path) -> dict[str, Any]:
    review = json.loads(review_path.read_text(encoding="utf-8"))
    profile = review_profile(review)
    statistics = overview_statistics(review) if profile == "overview" else {}
    expected = review["content_rows"]
    initial_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    formulas = openpyxl.load_workbook(workbook_path, read_only=False, data_only=False)
    values = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    failures = []
    warnings = []
    try:
        if tuple(formulas.sheetnames) != SHEETS:
            failures.append("sheet inventory/order differs")
        edit = formulas["内容编辑"]
        actual_rows = list(edit.iter_rows(min_row=2, max_col=6, values_only=True))
        if tuple(cell.value for cell in edit[1][:6]) != HEADERS:
            failures.append("canonical import headers differ")
        if actual_rows != [tuple(row) for row in expected]:
            failures.append("initial content row inventory is not an exact match")
        if any(cell.data_type in {"f", "e"} for row in edit.iter_rows(max_col=6) for cell in row):
            failures.append("canonical content contains formulas/errors")
        meta = dict(formulas["包信息"].iter_rows(min_row=2, max_col=2, values_only=True))
        if meta != workbook_metadata(review["metadata"]):
            failures.append("package metadata differs")
        counts = Counter(row[0] for row in expected)
        if "CHN" in counts or len(counts) != review["metadata"]["country_count"]:
            failures.append("outbound country scope differs")
        if values["审核导览"]["A6"].value != len(counts):
            failures.append("country KPI formula cache is missing or incorrect")
        if values["审核导览"]["C6"].value != len(expected):
            failures.append("content-row KPI formula cache is missing or incorrect")
        if values["审核导览"]["E6"].value != len(review["evidence_rows"]):
            failures.append("evidence KPI formula cache is missing or incorrect")
        if values["审核导览"]["G6"].value != len(review["gap_rows"]):
            failures.append("gap KPI formula cache is missing or incorrect")
        reading_expected = [row for row in expected if reading_row(row)]
        reading_actual = list(
            values["中文阅读"].iter_rows(min_row=2, min_col=5, max_col=6, values_only=True)
        )
        reading_indexes = [index + 2 for index, row in enumerate(expected) if reading_row(row)]
        if len(reading_expected) != len(reading_actual):
            failures.append("Chinese reading view inventory differs")
        for index, (row, actual) in enumerate(zip(reading_expected, reading_actual, strict=False)):
            text = LEVEL_LABELS.get(row[5], row[5]) if row[3].endswith("/level") else row[5]
            if actual != (text, reading_indexes[index]):
                failures.append(f"Chinese reading cache mismatch at row {index + 2}")
                break
        for index, row in enumerate(
            values["内容编辑"].iter_rows(min_row=2, min_col=11, max_col=12, values_only=True)
        ):
            # LEN counts Unicode UTF-16 code units in Excel, not only Hanzi.
            units = excel_length(expected[index][5])
            if row != (units, 0):
                failures.append(f"edit helper formula mismatch at row {index + 2}")
                break
        if profile == "overview":
            _verify_overview(review, formulas, values, statistics, failures)
        for sheet in formulas:
            if sheet.sheet_view.showGridLines is not False:
                failures.append(f"navigation/formatting missing in {sheet.title}")
            if not sheet.freeze_panes:
                warnings.append(
                    f"{sheet.title}: exporter did not retain requested freeze panes; "
                    "table filters and readable headers remain available"
                )
            if sheet.title != "审核导览" and not sheet.tables:
                failures.append(f"filter table missing in {sheet.title}")
        error_cells = []
        for sheet in values:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "e":
                        error_cells.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
        if error_cells:
            failures.append(f"spreadsheet formula errors: {error_cells[:20]}")
        report = {
            "status": "passed" if not failures else "failed",
            "content_profile": profile,
            "legacy_read_only": profile == "legacy_analysis_report",
            "workbook_sha256": initial_hash,
            "country_count": len(counts),
            "content_rows": len(expected),
            "chinese_reading_rows": len(reading_expected),
            "evidence_rows": len(review["evidence_rows"]),
            "gap_rows": len(review["gap_rows"]),
            "sheet_count": len(formulas.sheetnames),
            "formula_error_count": len(error_cells),
            "overview_zh_chars": {
                code: item["overview_zh_chars"] for code, item in statistics.items()
            },
            "failures": failures,
            "warnings": warnings,
            "published": False,
            "workbook_modified": False,
        }
    finally:
        formulas.close()
        values.close()
    if hashlib.sha256(workbook_path.read_bytes()).hexdigest() != initial_hash:
        raise RuntimeError("read-only workbook verification changed the file")
    return report


def _verify_overview(
    review: dict[str, Any],
    formulas: Any,
    values: Any,
    statistics: dict[str, dict[str, Any]],
    failures: list[str],
) -> None:
    """Verify live overview views and literal private sources, without changing them."""
    expected = review["content_rows"]
    row_indexes = {(row[0], row[3]): index + 2 for index, row in enumerate(expected)}
    edit = formulas["内容编辑"]
    if tuple(cell.value for cell in formulas["国别总览"][1]) != OVERVIEW_HEADERS:
        failures.append("overview headers retain a legacy or unknown view")
    for index, row in enumerate(expected, 2):
        if values["内容编辑"][f"O{index}"].value != excel_length(row[5], exclude_whitespace=True):
            failures.append(f"non-whitespace character formula mismatch at row {index}")
            break
        if edit[f"K{index}"].value != f"=LEN(F{index})" or edit[f"L{index}"].value != (
            f"=IF(EXACT(E{index},F{index}),0,1)"
        ):
            failures.append(f"editable helper formula was replaced at row {index}")
            break
        if edit[f"O{index}"].data_type != "f":
            failures.append(f"overview character helper is not a live formula at row {index}")
            break
    gap_counts = Counter(item["country_code"] for item in review["gap_rows"])
    for index, (code, item) in enumerate(statistics.items(), 2):
        actual = tuple(values["国别总览"].iter_rows(min_row=index, max_row=index, max_col=10))
        expected_summary = (
            code,
            item["country_name"],
            item["title"],
            item["opening"],
            item["excel_body_chars"],
            item["paragraph_count"],
            gap_counts[code],
            0,
            item["content_version"],
            item["as_of"],
        )
        if not actual or tuple(cell.value for cell in actual[0]) != expected_summary:
            failures.append(f"overview summary cache differs for {code}")
        for column, pointer in (("C", "title"), ("D", "paragraphs/0")):
            source_row = row_indexes[code, f"/locales/zh-CN/{pointer}"]
            if formulas["国别总览"][f"{column}{index}"].value != (f"='内容编辑'!F{source_row}"):
                failures.append(f"overview text is not linked to edited_value for {code}")
        if any(formulas["国别总览"][f"{column}{index}"].data_type != "f" for column in "EFGH"):
            failures.append(f"overview counts are not live formulas for {code}")
        paragraph_rows = [
            entry_index + 2
            for entry_index, entry in enumerate(expected)
            if entry[0] == code
            and entry[2] == "zh-CN"
            and entry[3].startswith("/locales/zh-CN/paragraphs/")
        ]
        for column, function, source_column in (("E", "SUM", "O"), ("F", "COUNTA", "F")):
            references = ",".join(
                f"'内容编辑'!{source_column}{number}" for number in paragraph_rows
            )
            if formulas["国别总览"][f"{column}{index}"].value != f"={function}({references})":
                failures.append(f"overview {column} count is not bound to its exact paragraphs")
    readable = [(index + 2, row) for index, row in enumerate(expected) if reading_row(row)]
    for index, (source_row, _row) in enumerate(readable, 2):
        if formulas["中文阅读"][f"E{index}"].value != f"='内容编辑'!F{source_row}":
            failures.append(f"overview reading row {index} is not linked to edited_value")
            break
    evidence_expected = [
        tuple(
            literal(value)
            for value in (
                item["country_code"],
                item.get("evidence_id", item.get("id")),
                item.get("title"),
                item.get("country_name", statistics[item["country_code"]]["country_name"]),
                item.get("url", item.get("official_url")),
                item.get("locator"),
                item.get("checked_on"),
                item.get("verification"),
                item.get("source_path", item.get("local_path")),
                item.get("sha256", item.get("actual_sha256")),
                item.get("notes", item.get("limitation")),
            )
        )
        for item in review["evidence_rows"]
    ]
    gaps_expected = [
        tuple(
            literal(value)
            for value in (
                item["country_code"],
                item.get("country_name", statistics[item["country_code"]]["country_name"]),
                item.get("topic", item.get("field", "项目核实")),
                item.get("text_zh", item.get("zh", item.get("text"))),
                item.get("text_en", item.get("en")),
                item.get(
                    "notes",
                    item.get(
                        "action",
                        item.get("next_action", "取得补充结果后，按需要同步修改中英文概述。"),
                    ),
                ),
            )
        )
        for item in review["gap_rows"]
    ]
    for sheet_name, expected_rows, columns in (
        ("来源依据", evidence_expected, 11),
        ("待确认事项", gaps_expected, 6),
    ):
        actual_rows = [
            tuple(literal(cell) for cell in row)
            for row in values[sheet_name].iter_rows(min_row=2, max_col=columns, values_only=True)
            if any(cell is not None for cell in row)
        ]
        if actual_rows != expected_rows:
            failures.append(f"private research inventory differs in {sheet_name}")
        if any(
            cell.data_type in {"f", "e"}
            for row in formulas[sheet_name].iter_rows(min_row=2, max_col=columns)
            for cell in row
        ):
            failures.append(f"private research cells are not literal text in {sheet_name}")
    private_tokens = [
        str(item[key])
        for item in review["evidence_rows"]
        for key in ("url", "source_path", "sha256")
        if item.get(key)
    ]
    if any(token in row[5] for row in expected for token in private_tokens):
        failures.append("private source metadata leaked into public overview content")
    legacy_labels = ("一分钟摘要", "六项分析", "完整报告", "进入判断概览", "通用核实清单占比")
    if any(
        label in str(cell.value)
        for name in ("审核导览", "国别总览", "中文阅读")
        for row in values[name].iter_rows()
        for cell in row
        for label in legacy_labels
    ):
        failures.append("overview reading views retain retired analysis/report labels")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("review_data", type=Path)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    report = verify(args.review_data, args.workbook)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
