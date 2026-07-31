from __future__ import annotations

import hashlib
import json
import os
import shutil
import tempfile
from copy import copy
from dataclasses import replace
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .baseline import D0_SHEETS, extract_contracts, sha256_file
from .d0_candidates import build_core_entity_evidence, build_core_field_evidence
from .d0_contract_resolution import validate_d0_contract_resolution
from .models import CheckResult
from .paths import RepositoryPaths
from .validation import validate_structure
from .workbook import load_read_only, normalized_value

_FIELD_SHEET = "核心字段冻结"
_FIELD_CONTRACT = "fields"
_FIELD_ID = "字段编号"
_UNIT_HEADER = "单位"


def _payload_sha256(payload: Any) -> str:
    serialized = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    return hashlib.sha256(serialized).hexdigest()


def _text(record: dict[str, Any], field: str) -> str:
    return str(record.get(field) or "").strip()


def _records(contracts: dict[str, Any], name: str) -> list[dict[str, Any]]:
    return cast(list[dict[str, Any]], contracts.get(name, []))


def _index(records: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    return {_text(record, _FIELD_ID): record for record in records if _text(record, _FIELD_ID)}


def _headers(workbook_path: Path) -> dict[str, dict[int, str]]:
    workbook = load_read_only(workbook_path, data_only=False)
    try:
        result: dict[str, dict[int, str]] = {}
        for sheet_name, contract_name in D0_SHEETS.items():
            row = next(workbook[sheet_name].iter_rows(min_row=4, max_row=4), ())
            result[contract_name] = {
                int(cell.column): str(cell.value).strip()
                for cell in row
                if cell.value is not None and str(cell.value).strip()
            }
        return result
    finally:
        workbook.close()


def _validate_headers(
    checks: list[CheckResult],
    source_workbook: Path,
    candidate_workbook: Path,
) -> dict[str, int]:
    source_headers = _headers(source_workbook)
    candidate_headers = _headers(candidate_workbook)
    for contract_name in D0_SHEETS.values():
        before = source_headers[contract_name]
        after = candidate_headers[contract_name]
        if contract_name != _FIELD_CONTRACT:
            if after != before:
                checks.append(
                    CheckResult(
                        code="D0_CHANGE_UNAUTHORIZED_HEADER_CHANGE",
                        message=f"{contract_name} headers changed outside the reviewed proposal",
                        location=contract_name,
                    )
                )
            continue

        expected = dict(before)
        expected[max(before, default=0) + 1] = _UNIT_HEADER
        if after != expected:
            checks.append(
                CheckResult(
                    code="D0_CHANGE_FIELD_HEADER_INVALID",
                    message=(
                        "fields must preserve every existing header and append exactly one 单位 "
                        "column"
                    ),
                    location=_FIELD_CONTRACT,
                )
            )
    return {header: column for column, header in candidate_headers.get(_FIELD_CONTRACT, {}).items()}


def _expected_key_marker(current: Any) -> str:
    marker = str(current or "").strip()
    if "主键" in marker:
        return marker
    return f"{marker}/主键" if marker else "主键"


def _unit_value(resolution: dict[str, Any]) -> str:
    if resolution.get("unit_applicability") == "not_applicable":
        return "不适用"
    return _text(resolution, "unit_code")


def _reviewed_changes(
    resolution: dict[str, Any],
) -> tuple[set[str], list[dict[str, Any]], dict[str, str]]:
    key_field_ids: set[str] = set()
    added_fields: list[dict[str, Any]] = []
    entity_items = resolution.get("entity_primary_key_resolutions")
    if isinstance(entity_items, list):
        for item in entity_items:
            if not isinstance(item, dict):
                continue
            action = item.get("action")
            if action in {"promote_existing_field", "model_composite_key"}:
                raw_ids = item.get("primary_key_field_ids")
                if isinstance(raw_ids, list):
                    key_field_ids.update(
                        str(value).strip() for value in raw_ids if str(value).strip()
                    )
            elif action == "add_primary_key_field":
                contract = item.get("proposed_field_contract")
                if isinstance(contract, dict):
                    added_fields.append(contract)

    unit_values: dict[str, str] = {}
    unit_items = resolution.get("field_unit_resolutions")
    if isinstance(unit_items, list):
        for item in unit_items:
            if isinstance(item, dict) and _text(item, "field_id"):
                unit_values[_text(item, "field_id")] = _unit_value(item)
    return key_field_ids, added_fields, unit_values


def _expected_fields(
    original_fields: list[dict[str, Any]],
    resolution: dict[str, Any],
) -> list[dict[str, Any]]:
    key_field_ids, added_fields, unit_values = _reviewed_changes(resolution)
    expected: list[dict[str, Any]] = []
    for original in original_fields:
        record = dict(original)
        field_id = _text(record, _FIELD_ID)
        record[_UNIT_HEADER] = unit_values.get(field_id, "")
        if field_id in key_field_ids:
            record["唯一/索引"] = _expected_key_marker(record.get("唯一/索引"))
        expected.append(record)
    expected.extend(dict(record) for record in added_fields)
    return expected


def _changed_fields(before: dict[str, Any], after: dict[str, Any]) -> list[str]:
    return sorted(
        field for field in set(before) | set(after) if before.get(field) != after.get(field)
    )


def _validate_contract_delta(
    checks: list[CheckResult],
    original: dict[str, Any],
    candidate: dict[str, Any],
    resolution: dict[str, Any],
) -> None:
    for contract_name in D0_SHEETS.values():
        if contract_name == _FIELD_CONTRACT:
            continue
        if _records(candidate, contract_name) != _records(original, contract_name):
            checks.append(
                CheckResult(
                    code="D0_CHANGE_UNAUTHORIZED_CONTRACT_CHANGE",
                    message=f"{contract_name} changed outside the reviewed AC-001/002 proposal",
                    location=contract_name,
                )
            )

    original_fields = _records(original, _FIELD_CONTRACT)
    actual_fields = _records(candidate, _FIELD_CONTRACT)
    expected_fields = _expected_fields(original_fields, resolution)
    original_ids = [_text(record, _FIELD_ID) for record in original_fields]
    expected_ids = [_text(record, _FIELD_ID) for record in expected_fields]
    actual_ids = [_text(record, _FIELD_ID) for record in actual_fields]
    if actual_ids != expected_ids:
        checks.append(
            CheckResult(
                code="D0_CHANGE_FIELD_ID_SEQUENCE_INVALID",
                message=(
                    "Candidate field IDs must preserve all frozen rows in order and append only "
                    "the reviewed new primary-key fields"
                ),
                location=_FIELD_CONTRACT,
            )
        )

    expected_index = _index(expected_fields)
    actual_index = _index(actual_fields)
    missing = sorted(set(expected_index) - set(actual_index))
    unexpected = sorted(set(actual_index) - set(expected_index))
    if missing or unexpected:
        checks.append(
            CheckResult(
                code="D0_CHANGE_FIELD_ID_SET_INVALID",
                message=f"Field ID delta differs; missing={missing}, unexpected={unexpected}",
                location=_FIELD_CONTRACT,
            )
        )
    for field_id in sorted(set(expected_index) & set(actual_index)):
        expected = expected_index[field_id]
        actual = actual_index[field_id]
        if actual != expected:
            checks.append(
                CheckResult(
                    code="D0_CHANGE_FIELD_ROW_MISMATCH",
                    message=(
                        f"{field_id} differs from the reviewed row definition: "
                        f"{', '.join(_changed_fields(expected, actual))}"
                    ),
                    location=f"{_FIELD_CONTRACT}.{field_id}",
                )
            )

    if actual_ids[: len(original_ids)] != original_ids:
        checks.append(
            CheckResult(
                code="D0_CHANGE_FROZEN_FIELD_ORDER_CHANGED",
                message="Existing frozen field rows were removed, inserted into, or reordered",
                location=_FIELD_CONTRACT,
            )
        )


def _row_lookup(workbook_path: Path) -> dict[str, int]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        sheet = workbook[_FIELD_SHEET]
        header_columns = {
            str(cell.value).strip(): cell.column
            for cell in sheet[4]
            if cell.value is not None and str(cell.value).strip()
        }
        identifier_column = header_columns.get(_FIELD_ID)
        if identifier_column is None:
            return {}
        return {
            str(sheet.cell(row, identifier_column).value).strip(): row
            for row in range(5, sheet.max_row + 1)
            if sheet.cell(row, identifier_column).value is not None
            and str(sheet.cell(row, identifier_column).value).strip()
        }
    finally:
        workbook.close()


def _cell_values(workbook_path: Path) -> tuple[list[str], dict[str, dict[tuple[int, int], Any]]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        values: dict[str, dict[tuple[int, int], Any]] = {}
        for sheet in workbook.worksheets:
            cells: dict[tuple[int, int], Any] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    value = normalized_value(cell.value)
                    if (
                        value is not None
                        and str(value).strip()
                        and cell.row is not None
                        and cell.column is not None
                    ):
                        cells[(int(cell.row), int(cell.column))] = value
            values[sheet.title] = cells
        return list(workbook.sheetnames), values
    finally:
        workbook.close()


def _validate_cell_delta(
    checks: list[CheckResult],
    source_workbook: Path,
    candidate_workbook: Path,
    resolution: dict[str, Any],
    candidate_headers: dict[str, int],
) -> None:
    source_sheets, source_values = _cell_values(source_workbook)
    candidate_sheets, candidate_values = _cell_values(candidate_workbook)
    if source_sheets != candidate_sheets:
        checks.append(
            CheckResult(
                code="D0_CHANGE_SHEET_SET_INVALID",
                message="Candidate workbook sheet names or order differ from the frozen workbook",
                location=str(candidate_workbook),
            )
        )
        return

    source_rows = _row_lookup(source_workbook)
    candidate_rows = _row_lookup(candidate_workbook)
    key_field_ids, added_fields, unit_values = _reviewed_changes(resolution)
    allowed: set[tuple[str, int, int]] = set()
    unit_column = candidate_headers.get(_UNIT_HEADER)
    key_column = candidate_headers.get("唯一/索引")
    if unit_column is not None:
        allowed.add((_FIELD_SHEET, 4, unit_column))
        for field_id in unit_values:
            row = candidate_rows.get(field_id)
            if row is not None:
                allowed.add((_FIELD_SHEET, row, unit_column))
    if key_column is not None:
        for field_id in key_field_ids:
            row = source_rows.get(field_id)
            if row is not None:
                allowed.add((_FIELD_SHEET, row, key_column))
    for contract in added_fields:
        row = candidate_rows.get(_text(contract, _FIELD_ID))
        if row is not None:
            allowed.update(
                (_FIELD_SHEET, row, column)
                for column in candidate_headers.values()
                if (row, column) not in source_values.get(_FIELD_SHEET, {})
            )

    for sheet_name in source_sheets:
        before = source_values.get(sheet_name, {})
        after = candidate_values.get(sheet_name, {})
        changed = {
            coordinate
            for coordinate in set(before) | set(after)
            if before.get(coordinate) != after.get(coordinate)
            and (sheet_name, coordinate[0], coordinate[1]) not in allowed
        }
        if changed:
            preview = ", ".join(f"R{row}C{column}" for row, column in sorted(changed)[:10])
            checks.append(
                CheckResult(
                    code="D0_CHANGE_UNAUTHORIZED_CELL_CHANGE",
                    message=(
                        f"{sheet_name} contains {len(changed)} unreviewed cell value/formula "
                        f"change(s): {preview}"
                    ),
                    location=sheet_name,
                )
            )


def assess_d0_baseline_change(
    paths: RepositoryPaths,
    resolution: dict[str, Any],
    candidate_workbook: Path,
) -> dict[str, Any]:
    checks = validate_d0_contract_resolution(paths, resolution)
    source_workbook = paths.d0_workbook.resolve()
    candidate_path = candidate_workbook.resolve()
    report: dict[str, Any] = {
        "schema_version": 1,
        "stage": "D0",
        "assessment_scope": "D0 AC-001/002 candidate workbook baseline-change verification",
        "automated_assessment_only": True,
        "does_not_activate_baseline": True,
        "source_workbook": {
            "path": source_workbook.relative_to(paths.root).as_posix(),
            "sha256": sha256_file(source_workbook),
        },
        "candidate_workbook": {"path": str(candidate_path), "sha256": None},
        "resolution_sha256": _payload_sha256(resolution),
        "core_entity_assessment": None,
        "core_field_assessment": None,
        "candidate_ready_for_formal_baseline_review": False,
        "checks": [],
        "warning": (
            "Passing this verifier does not modify or approve the frozen D0 baseline, does not "
            "complete D0/D4, and does not authorize user-facing development."
        ),
    }

    if candidate_path == source_workbook:
        checks.append(
            CheckResult(
                code="D0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN",
                message="The authoritative D0 workbook cannot be used as the change target",
                location=str(candidate_path),
            )
        )
    if not candidate_path.is_file():
        checks.append(
            CheckResult(
                code="D0_CHANGE_WORKBOOK_MISSING",
                message="Candidate D0 workbook does not exist",
                location=str(candidate_path),
            )
        )
    if checks:
        report["checks"] = [check.to_dict() for check in checks]
        return report

    candidate_hash = sha256_file(candidate_path)
    cast(dict[str, Any], report["candidate_workbook"])["sha256"] = candidate_hash
    if candidate_hash == report["source_workbook"]["sha256"]:
        checks.append(
            CheckResult(
                code="D0_CHANGE_WORKBOOK_UNCHANGED",
                message="Candidate workbook is byte-identical to the frozen D0 workbook",
                location=str(candidate_path),
            )
        )

    candidate_paths = replace(paths, d0_workbook=candidate_path)
    try:
        checks.extend(validate_structure(candidate_paths))
        candidate_headers = _validate_headers(checks, source_workbook, candidate_path)
        original_contracts = extract_contracts(paths, data_only=False)
        candidate_contracts = extract_contracts(candidate_paths, data_only=False)
        _validate_cell_delta(
            checks,
            source_workbook,
            candidate_path,
            resolution,
            candidate_headers,
        )
    except Exception as error:
        checks.append(
            CheckResult(
                code="D0_CHANGE_WORKBOOK_INVALID",
                message=f"Cannot read candidate D0 workbook: {error}",
                location=str(candidate_path),
            )
        )
        report["checks"] = [check.to_dict() for check in checks]
        return report

    _validate_contract_delta(checks, original_contracts, candidate_contracts, resolution)
    entity_assessment = build_core_entity_evidence(candidate_contracts)
    field_assessment = build_core_field_evidence(candidate_contracts)
    report["core_entity_assessment"] = entity_assessment
    report["core_field_assessment"] = field_assessment
    if entity_assessment["machine_status"] != "pass":
        checks.append(
            CheckResult(
                code="D0_CHANGE_ENTITY_EVIDENCE_NOT_READY",
                message="Candidate workbook still fails the AC-001 machine evidence",
                location="D0-AC-001",
            )
        )
    if field_assessment["machine_status"] != "pass":
        checks.append(
            CheckResult(
                code="D0_CHANGE_FIELD_EVIDENCE_NOT_READY",
                message="Candidate workbook still fails the AC-002 machine evidence",
                location="D0-AC-002",
            )
        )

    report["checks"] = [check.to_dict() for check in checks]
    report["candidate_ready_for_formal_baseline_review"] = not checks
    return report


def load_and_assess_d0_baseline_change(
    paths: RepositoryPaths,
    resolution_path: Path,
    candidate_workbook: Path,
) -> dict[str, Any]:
    try:
        payload = json.loads(resolution_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as error:
        return _invalid_resolution_report(
            f"Cannot read resolution packet: {error}", resolution_path
        )
    if not isinstance(payload, dict):
        return _invalid_resolution_report(
            "Resolution packet root must be an object", resolution_path
        )
    return assess_d0_baseline_change(paths, payload, candidate_workbook)


def _invalid_resolution_report(message: str, path: Path) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "stage": "D0",
        "assessment_scope": "D0 AC-001/002 candidate workbook baseline-change verification",
        "automated_assessment_only": True,
        "does_not_activate_baseline": True,
        "candidate_ready_for_formal_baseline_review": False,
        "checks": [
            CheckResult(
                code="D0_CHANGE_RESOLUTION_INVALID",
                message=message,
                location=str(path),
            ).to_dict()
        ],
    }


def _editable_headers(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): int(cell.column)
        for cell in sheet[4]
        if cell.value is not None and str(cell.value).strip() and cell.column is not None
    }


def _editable_field_rows(sheet: Worksheet, identifier_column: int) -> dict[str, int]:
    return {
        str(sheet.cell(row, identifier_column).value).strip(): row
        for row in range(5, sheet.max_row + 1)
        if sheet.cell(row, identifier_column).value is not None
        and str(sheet.cell(row, identifier_column).value).strip()
    }


def _copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _apply_resolution_to_workbook(
    candidate_workbook: Path,
    resolution: dict[str, Any],
) -> None:
    workbook = load_workbook(candidate_workbook, read_only=False, data_only=False)
    try:
        sheet = workbook[_FIELD_SHEET]
        headers = _editable_headers(sheet)
        identifier_column = headers[_FIELD_ID]
        field_rows = _editable_field_rows(sheet, identifier_column)
        last_original_row = max(field_rows.values())

        unit_column = max(headers.values()) + 1
        unit_header = sheet.cell(4, unit_column)
        _copy_cell_style(sheet.cell(4, unit_column - 1), unit_header)
        unit_header.value = _UNIT_HEADER
        headers[_UNIT_HEADER] = unit_column
        previous_letter = get_column_letter(unit_column - 1)
        unit_letter = get_column_letter(unit_column)
        sheet.column_dimensions[unit_letter].width = sheet.column_dimensions[previous_letter].width

        key_field_ids, added_fields, unit_values = _reviewed_changes(resolution)
        for field_id, value in unit_values.items():
            row = field_rows[field_id]
            unit_cell = sheet.cell(row, unit_column)
            _copy_cell_style(sheet.cell(row, unit_column - 1), unit_cell)
            unit_cell.value = value

        key_column = headers["唯一/索引"]
        for field_id in key_field_ids:
            row = field_rows[field_id]
            key_cell = sheet.cell(row, key_column)
            key_cell.value = _expected_key_marker(key_cell.value)

        for offset, contract in enumerate(added_fields, start=1):
            row = last_original_row + offset
            if sheet.row_dimensions[last_original_row].height is not None:
                sheet.row_dimensions[row].height = sheet.row_dimensions[last_original_row].height
            for field, column in headers.items():
                target = sheet.cell(row, column)
                _copy_cell_style(sheet.cell(last_original_row, column), target)
                target.value = contract.get(field)
        workbook.save(candidate_workbook)
    finally:
        workbook.close()


def _generation_report(
    paths: RepositoryPaths,
    resolution: dict[str, Any],
    output_path: Path,
    checks: list[CheckResult],
) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "stage": "D0",
        "operation": "generate validated AC-001/002 candidate workbook",
        "automated_assessment_only": True,
        "does_not_activate_baseline": True,
        "source_workbook": {
            "path": paths.d0_workbook.relative_to(paths.root).as_posix(),
            "sha256": sha256_file(paths.d0_workbook),
        },
        "candidate_workbook": {"path": str(output_path), "sha256": None},
        "resolution_sha256": _payload_sha256(resolution),
        "candidate_written": False,
        "candidate_ready_for_formal_baseline_review": False,
        "checks": [check.to_dict() for check in checks],
        "warning": (
            "Generation never overwrites or activates the authoritative D0 workbook. "
            "A written candidate still requires formal baseline review."
        ),
    }


def generate_d0_candidate_workbook(
    paths: RepositoryPaths,
    resolution: dict[str, Any],
    candidate_output: Path,
) -> dict[str, Any]:
    output_path = candidate_output.resolve()
    checks = validate_d0_contract_resolution(paths, resolution)
    source_path = paths.d0_workbook.resolve()
    if output_path == source_path:
        checks.append(
            CheckResult(
                code="D0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN",
                message="The authoritative D0 workbook cannot be overwritten",
                location=str(output_path),
            )
        )
    try:
        output_path.relative_to((paths.root / "doc").resolve())
    except ValueError:
        pass
    else:
        checks.append(
            CheckResult(
                code="D0_CHANGE_OUTPUT_LOCATION_FORBIDDEN",
                message="Candidate workbooks must be written outside the authoritative doc tree",
                location=str(output_path),
            )
        )
    if output_path.suffix.lower() != ".xlsx":
        checks.append(
            CheckResult(
                code="D0_CHANGE_OUTPUT_EXTENSION_INVALID",
                message="Candidate output must use the .xlsx extension",
                location=str(output_path),
            )
        )
    if output_path.exists():
        checks.append(
            CheckResult(
                code="D0_CHANGE_OUTPUT_EXISTS",
                message="Candidate output already exists and will not be overwritten",
                location=str(output_path),
            )
        )
    if checks:
        return _generation_report(paths, resolution, output_path, checks)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_path.parent,
            prefix=f".{output_path.stem}.",
            suffix=".xlsx",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        shutil.copy2(source_path, temporary_path)
        _apply_resolution_to_workbook(temporary_path, resolution)
        assessment = assess_d0_baseline_change(paths, resolution, temporary_path)
        assessment["operation"] = "generate validated AC-001/002 candidate workbook"
        assessment["candidate_written"] = False
        cast(dict[str, Any], assessment["candidate_workbook"])["path"] = str(output_path)
        if assessment.get("candidate_ready_for_formal_baseline_review") is not True:
            return assessment
        os.replace(temporary_path, output_path)
        temporary_path = None
        assessment["candidate_written"] = True
        return assessment
    except Exception as error:
        return _generation_report(
            paths,
            resolution,
            output_path,
            [
                CheckResult(
                    code="D0_CHANGE_GENERATION_FAILED",
                    message=f"Could not generate candidate workbook: {error}",
                    location=str(output_path),
                )
            ],
        )
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def load_and_generate_d0_candidate_workbook(
    paths: RepositoryPaths,
    resolution_path: Path,
    candidate_output: Path,
) -> dict[str, Any]:
    try:
        payload = json.loads(resolution_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as error:
        report = _invalid_resolution_report(
            f"Cannot read resolution packet: {error}", resolution_path
        )
        report["operation"] = "generate validated AC-001/002 candidate workbook"
        report["candidate_written"] = False
        return report
    if not isinstance(payload, dict):
        report = _invalid_resolution_report(
            "Resolution packet root must be an object", resolution_path
        )
        report["operation"] = "generate validated AC-001/002 candidate workbook"
        report["candidate_written"] = False
        return report
    return generate_d0_candidate_workbook(paths, payload, candidate_output)
