from __future__ import annotations

import csv
import hashlib
import hmac
import json
import math
import os
import re
import stat
import tempfile
from collections import Counter, defaultdict
from datetime import UTC, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook

from .baseline import sha256_file
from .basic60_evidence import (
    EvidenceCollectionError,
    build_d3_stage_payloads,
    collect_machine_evidence,
    validate_machine_evidence,
    write_d3_stage_payloads,
)
from .d2_collection import STOP_SIGNALS
from .d3_processing import PIPELINE_DEFINITIONS
from .models import CheckResult
from .paths import RepositoryPaths

SCHEMA_VERSION = 1
PROFILE_ID = "basic60_private"
RELEASE_ID = "BASIC60-PRIVATE-R1"
FORMAL_GATE_STATUS = "pending"
DECISION_ID = "PBD-BASIC60-PRIVATE-001"
PRIVATE_OUTCOMES = {"private_trial_ready", "not_ready", "revoked"}

EXPECTED_COUNTRIES = 60
EXPECTED_MACRO_ROWS = 300
EXPECTED_ENERGY_ROWS = 60
EXPECTED_AVAILABLE_VALUES = 2279
EXPECTED_PENDING_VALUES = 61
EXPECTED_SAMPLE_SIZE = 180
EXPECTED_DERIVED_VALUES = 120
FUTURE_DEEP_DIVE_PRIORITY_COUNTRY_CODES = ("IDN", "VNM", "SAU", "ZAF", "BRA")

RAW_MANIFEST_NAME = "basic60_raw_manifest.json"
PROFILE_NAME = "basic60_private_profile.json"
ACCEPTANCE_TEMPLATE_NAME = "basic60_private_acceptance.template.json"
ASSESSMENT_NAME = "basic60_private_assessment.json"
SEED_NAME = "basic60_seed.json"
SAMPLE_CANDIDATE_NAME = "basic60_sample_candidate.json"
SAMPLE_REVIEW_TEMPLATE_NAMES = (
    "basic60_sample_review_1.template.json",
    "basic60_sample_review_2.template.json",
)
ITEM_REVIEW_TEMPLATE_NAME = "basic60_item_review.template.json"
SOURCE_MATRIX_NAME = "basic60_country_domain_matrix.json"
READY_SEED_NAME = "basic60_seed.private_trial_ready.json"
RELEASE_AUTHORIZATION_NAME = "basic60_release_authorization.json"
RUNTIME_ATTESTATION_NAME = "basic60_runtime_attestation.json"
RUNTIME_ATTESTATION_KEY_ENV = "BASIC60_RUNTIME_ATTESTATION_KEY"
RUNTIME_ATTESTATION_KEY_MIN_BYTES = 32
RUNTIME_ATTESTATION_FORBIDDEN_KEYS = frozenset(
    {
        "replace-me",
        "replace-with-at-least-32-random-bytes",
        "test-runtime-attestation-key-with-32-bytes",
        "basic60-test-runtime-attestation-key-32-bytes",
    }
)
RUNTIME_ATTESTATION_FORBIDDEN_MARKERS = (
    "changeme",
    "example-value",
    "placeholder",
    "replace-with",
)
SOURCE_REGISTRY_TEMPLATE_NAME = "basic60_source_registry.template.json"
COUNTRY_SOURCE_URI_MANIFEST_NAME = "basic60_country_source_uri_manifest.json"
AI_USAGE_POLICY_NAME = "basic60_ai_usage_policy.json"
USAGE_AMENDMENT_TEMPLATE_NAME = "basic60_private_usage_amendment.template.json"
HUMAN_APPROVAL_TEMPLATE_NAME = "basic60_human_approval.template.json"
D1_SCOPE_AUTHORIZATION_TEMPLATE_NAME = "basic60_d1_scope_authorization.template.json"
L0_BATCH_MANIFEST_SCHEMA = "basic60.l0-batch-manifest.v2"
L0_DIFF_REPORT_SCHEMA = "basic60.l0-diff-report.v1"
RETRIEVAL_LOG_SCHEMA = "basic60.source-retrieval-log.v1"
SOURCE_DIFF_SCHEMA = "basic60.source-diff.v1"
SAMPLE_CANDIDATE_SCHEMA = "basic60.sample-candidate.v1"
SAMPLE_REVIEW_MANIFEST_SCHEMA = "basic60.sample-review-manifest.v2"
SOURCE_EVIDENCE_SCHEMA = "basic60.source-evidence.v5"
D1_AUTHORIZED_SCOPE_SCHEMA = "basic60.d1-authorized-scope.v2"
D1_SCOPE_AUTHORIZATION_RECORD_TYPE = "basic60_d1_scope_authorization"
D1_MACHINE_CONTENT_BINDING_SCHEMA = "basic60.d1-machine-content-binding.v2"
D1_MACHINE_BASELINE_ANCHOR_SCHEMA = "basic60.d1-machine-baseline-anchor.v1"
D1_MACHINE_BASELINE_ROOT = PurePosixPath("d1-machine-baselines")
D1_TRUSTED_VOLUME_REGISTRY_SCHEMA = "basic60.trusted-volume-registry.v1"
D1_TRUSTED_VOLUME_REGISTRY_PATH = Path("/etc/navigator/basic60-volume-registry.json")
D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS = frozenset({0})
D1_REAPPROVAL_TRIGGERS = (
    "provider_changed",
    "source_url_changed",
    "terms_or_license_changed",
    "dataset_structure_changed",
    "usage_boundary_changed",
)
REQUIRED_D1_SCOPE_AUTHORIZATION_ID = "B60-D1-SCOPE-AUTH-20260826-R2"
EVIDENCE_OBJECT_STORAGE_MODES = {"external_l0", "repository_public"}
REPLAY_MANIFEST_SCHEMA = "basic60.replay-manifest.v2"
SOURCE_MATRIX_SCHEMA = "basic60.country-domain-matrix.v1"
COUNTRY_SOURCE_URI_MANIFEST_SCHEMA = "basic60.country-source-uri-manifest.v1"
ITEM_REVIEW_MANIFEST_SCHEMA = "basic60.item-review-manifest.v2"
L0_BATCH_MANIFEST_NAMES = {
    "BASIC60-BATCH-PROFILE": "basic60-batch-profile.l0-manifest.json",
    "BASIC60-BATCH-MACRO": "basic60-batch-macro.l0-manifest.json",
    "BASIC60-BATCH-ENERGY": "basic60-batch-energy.l0-manifest.json",
}

PBD_CANDIDATE_PATH = PurePosixPath(
    "data/governance/candidates/basic60_private_baseline_change.2026-08-25.json"
)
PBD_CONFIRMATION_TEMPLATE_PATH = PurePosixPath(
    "data/governance/review/basic60_private_baseline_confirmation.template.2026-08-25.json"
)
PBD_DECISION_PATH = PurePosixPath(
    "data/governance/review/basic60_private_baseline_decision.2026-08-25.json"
)
PBD_CLARIFICATION_PATH = PurePosixPath(
    "data/governance/review/basic60_private_internal_review_clarification.2026-08-26.json"
)
USAGE_AMENDMENT_ID = "PBD-BASIC60-PRIVATE-001-A1"
USAGE_AMENDMENT_PATH = PurePosixPath(
    "data/governance/review/basic60_private_usage_amendment.2026-08-26.json"
)
USAGE_AMENDMENT_SCHEMA = "basic60.usage-amendment.v1"
AI_USAGE_POLICY_SCHEMA = "basic60.ai-usage-policy.v2"
EXCEL_REVIEW_GATE_SCHEMA = "basic60.excel-review-gate.v1"
EXCEL_REVIEW_PAYLOAD_SCHEMA = "basic60.excel-review-payload.v1"
EXCEL_REVIEW_WORKBOOK_SCHEMA = "basic60.excel-review-workbook.v1"
EXCEL_REVIEW_APPROVAL_SCHEMA = "basic60.excel-review-approval.v1"
EXCEL_REVIEW_MODE = "single_workbook_single_project_approver"
EXCEL_REVIEW_SCOPE = "all_basic60_data_and_sources"
EXCEL_REVIEW_OUTPUT_PREFIX = PurePosixPath("outputs/basic60")
EXCEL_REVIEW_MACHINE_SHEET = "机器核对"
REQUESTED_DECISION = "approve_basic60_private_trial"

MANUAL_ALLOWED_USES = (
    "internal_learning_and_exchange",
    "private_display",
    "internal_ai_processing",
    "local_model_processing",
    "controlled_external_model_processing",
)
MANUAL_PROHIBITED_USES = ("public_release", "model_training")
SOURCE_PERMISSION_FIELDS = frozenset(
    {
        "license_scope",
        "ai_processing",
        "local_model_processing",
        "controlled_external_model_processing",
        "cloud_processing",
        "model_training",
    }
)

PROFILE_CSV = PurePosixPath("global_sources/60_country_profiles.csv")
METRIC_CSV = PurePosixPath("global_sources/60_country_macro_energy.csv")
COLLECTION_MANIFEST = PurePosixPath("collection_manifest_60.json")

PROFILE_METRICS: tuple[tuple[str, str, str], ...] = (
    ("population_total", "population", "COUNT"),
    ("land_area_sq_km", "area_sq_km", "KM2"),
)
MACRO_METRICS: tuple[tuple[str, str, str], ...] = (
    ("gdp_current_usd", "gdp_current_usd", "USD"),
    ("gdp_growth_pct", "gdp_growth_pct", "PERCENT"),
    ("gdp_per_capita_current_usd", "gdp_per_capita_current_usd", "USD_PER_PERSON"),
    ("inflation_cpi_pct", "inflation_cpi_pct", "PERCENT"),
    (
        "official_exchange_rate_lcu_per_usd",
        "official_exchange_rate_lcu_per_usd",
        "LCU_PER_USD",
    ),
    ("fdi_net_inflows_usd", "fdi_net_inflows_usd", "USD"),
)
ENERGY_METRICS: tuple[tuple[str, str, str, str], ...] = (
    (
        "electricity_installed_capacity_mw",
        "electricity_installed_capacity_mw",
        "MW",
        "electricity_installed_capacity_year",
    ),
    (
        "electricity_generation_gwh",
        "electricity_generation_gwh",
        "GWH",
        "electricity_generation_year",
    ),
    ("renewable_capacity_mw", "renewable_capacity_mw", "MW", "renewable_capacity_year"),
    (
        "renewable_generation_gwh",
        "renewable_generation_gwh",
        "GWH",
        "renewable_generation_year",
    ),
    (
        "renewable_share_capacity_pct",
        "renewable_share_capacity_pct",
        "PERCENT",
        "renewable_share_capacity_year",
    ),
    (
        "renewable_share_generation_pct",
        "renewable_share_generation_pct",
        "PERCENT",
        "renewable_share_generation_year",
    ),
)

METRIC_LABELS: dict[str, tuple[str, str, str, str, int]] = {
    "population_total": ("人口", "Population", "identity", "COUNT", 0),
    "land_area_sq_km": ("陆地面积", "Land area", "identity", "KM2", 2),
    "gdp_current_usd": ("国内生产总值", "GDP", "macro", "USD", 2),
    "gdp_growth_pct": ("GDP增长率", "GDP growth", "macro", "PERCENT", 4),
    "gdp_per_capita_current_usd": (
        "人均GDP",
        "GDP per capita",
        "macro",
        "USD_PER_PERSON",
        2,
    ),
    "inflation_cpi_pct": ("CPI通胀率", "CPI inflation", "macro", "PERCENT", 4),
    "official_exchange_rate_lcu_per_usd": (
        "官方汇率",
        "Official exchange rate",
        "macro",
        "LCU_PER_USD",
        6,
    ),
    "fdi_net_inflows_usd": ("FDI净流入", "FDI net inflows", "macro", "USD", 2),
    "electricity_installed_capacity_mw": (
        "电力装机容量",
        "Installed capacity",
        "energy",
        "MW",
        2,
    ),
    "electricity_generation_gwh": (
        "发电量",
        "Electricity generation",
        "energy",
        "GWH",
        2,
    ),
    "renewable_capacity_mw": (
        "可再生能源装机",
        "Renewable capacity",
        "energy",
        "MW",
        2,
    ),
    "renewable_generation_gwh": (
        "可再生能源发电量",
        "Renewable generation",
        "energy",
        "GWH",
        2,
    ),
    "renewable_share_capacity_pct": (
        "可再生能源装机占比",
        "Renewable capacity share",
        "energy",
        "PERCENT",
        3,
    ),
    "renewable_share_generation_pct": (
        "可再生能源发电占比",
        "Renewable generation share",
        "energy",
        "PERCENT",
        3,
    ),
    "electricity_demand_gwh": (
        "电力需求",
        "Electricity demand",
        "energy",
        "GWH",
        2,
    ),
}

DOMAIN_IDS = ("country_profile", "macroeconomic", "energy")
BATCH_SPECS: tuple[tuple[str, str, PurePosixPath], ...] = (
    ("BASIC60-BATCH-PROFILE", "country_profile", PROFILE_CSV),
    ("BASIC60-BATCH-MACRO", "macroeconomic", METRIC_CSV),
    ("BASIC60-BATCH-ENERGY", "energy", METRIC_CSV),
)
SOURCE_CONTRACT_SPECS: dict[str, dict[str, Any]] = {
    "country_identity": {
        "dataset_path": PROFILE_CSV,
        "record_selector": None,
        "source_uri_mode": "country_source_manifest",
        "source_uri_field": "sources[].url",
        "source_uri_row_key_fields": ("iso3", "provider", "fields"),
    },
    "macro": {
        "dataset_path": METRIC_CSV,
        "record_selector": {"field": "record_type", "value": "macro_annual"},
        "source_uri_mode": "csv_column",
        "source_uri_field": "macro_source_url",
        "source_uri_row_key_fields": ("iso3", "record_type", "year"),
    },
    "energy": {
        "dataset_path": METRIC_CSV,
        "record_selector": {"field": "record_type", "value": "energy_latest"},
        "source_uri_mode": "csv_column",
        "source_uri_field": "energy_source_url",
        "source_uri_row_key_fields": ("iso3", "record_type"),
    },
}
SOURCE_TERMS_URI_REQUIREMENTS: dict[str, tuple[str, ...]] = {
    "country_identity": (
        "https://data.worldbank.org/summary-terms-of-use",
        "https://data.iana.org/time-zones/tzdb/LICENSE",
        "https://www.iso.org/fr/footer-links/obp-tems-of-use.html",
        (
            "https://raw.githubusercontent.com/restcountries/restcountries/"
            "ff97f12bb61f8cbc389228d53445f64102cc305f/LICENSE"
        ),
    ),
    "macro": ("https://data.worldbank.org/summary-terms-of-use",),
    "energy": ("https://www.irena.org/terms-and-conditions",),
}
SOURCE_TERMS_CAPTURE_METHODS = {
    "official_https_download",
    "official_page_rendered_dom_browser_capture",
}
PIPELINE_IDS = tuple(str(item["pipeline_id"]) for item in PIPELINE_DEFINITIONS)
REPLAY_MANIFEST_NAMES = {
    pipeline_id: f"basic60-{pipeline_id.lower()}.replay-manifest.json"
    for pipeline_id in PIPELINE_IDS
}
FINAL_APPROVAL_ROLES = (
    "数据负责人",
    "数据质量负责人",
    "合规负责人",
    "数据工程负责人",
    "项目批准人",
)
BATCH_APPROVAL_ROLE = "授权数据审核人"
AUDIT_ONLY_FIELD_NAMES = frozenset(
    {
        "source_ref",
        "source_snapshot_ref",
        "raw_record_ref",
        "source_text",
    }
)
PLACEHOLDER_PATTERN = re.compile(
    r"^(?:pending|todo|tbd|n/?a|none|null|unknown|待.*|.*姓名.*|reviewer|approver|codex)$",
    re.IGNORECASE,
)
SHA256_PATTERN = re.compile(r"^[0-9a-f]{64}$")
VOLUME_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._:-]{0,127}$")
D1_SCOPE_AUTHORIZATION_ID_PATTERN = re.compile(r"^B60-D1-SCOPE-[A-Z0-9][A-Z0-9._-]{0,95}$")
L0_MANIFEST_FIELDS = frozenset(
    {
        "schema_version",
        "profile_id",
        "release_id",
        "batch_id",
        "run_id",
        "copy_status",
        "batch_status",
        "domain_id",
        "raw_root_sha256",
        "source_subject_sha256",
        "source_contract_sha256",
        "approved_primary_source_id",
        "approved_alternative_source_id",
        "source_dataset_path",
        "source_sha256",
        "source_byte_size",
        "volume_id",
        "object_key",
        "object_sha256",
        "object_byte_size",
        "copied",
        "reused",
        "readback_verified",
        "immutable",
        "read_only",
        "posix_mode",
        "diff_status",
        "current_snapshot",
        "diff_report",
        "checked_at",
        "source_snapshot_comparisons",
        "watermark_advanced",
        "watermark",
        "access_signal",
        "materialized_at",
    }
)
SAMPLE_REVIEW_MANIFEST_FIELDS = frozenset(
    {
        "schema_version",
        "profile_id",
        "release_id",
        "sample_candidate",
        "sample_size",
        "country_count",
        "per_country_count",
        "coverage",
        "review_slot",
        "template_only",
        "reviewer_name",
        "reviewed_at",
        "status",
        "items",
        "items_sha256",
    }
)


def _owner_usage_authorization() -> dict[str, str]:
    return {
        "field_scope": "all_basic60_normalized_data_fields",
        "internal_learning_and_exchange": "allowed_by_manual_decision",
        "private_display": "allowed_by_manual_decision",
        "internal_ai_processing": "allowed_by_manual_decision",
        "local_model_processing": "allowed_by_manual_decision",
        "controlled_external_model_processing": "allowed_by_manual_decision",
        "ordinary_ui_api_source_labels": "omitted_by_manual_decision",
        "public_release": "prohibited",
        "model_training": "prohibited",
    }


def _v1_runtime_capabilities() -> dict[str, bool]:
    return {
        "ai_routes_enabled": False,
        "external_model_calls_enabled": False,
        "local_model_service_enabled": False,
        "vector_database_enabled": False,
        "embedding_enabled": False,
        "reranking_enabled": False,
        "full_text_search_enabled": False,
    }


def _usage_boundaries() -> dict[str, str]:
    return {
        "store_structured": "allowed",
        "private_display": "allowed_without_source_labels_by_manual_decision",
        "private_export": "prohibited_except_controlled_model_payload",
        "internal_ai_processing": "allowed_by_manual_decision",
        "local_model_processing": "allowed_by_manual_decision",
        "controlled_external_model_processing": "allowed_by_manual_decision",
        "ai_index": "deferred_to_v2",
        "model_training": "prohibited",
        "v1_ai_runtime": "disabled",
    }


def _normalized_data_field_paths(seed: dict[str, Any]) -> list[str]:
    countries = seed.get("countries")
    if not isinstance(countries, list) or not countries:
        raise ValueError("Basic60 seed must contain countries for the manual usage field scope")
    discovered: set[str] = set()

    def visit(value: Any, path: str) -> None:
        if isinstance(value, dict):
            for key in sorted(value):
                if key in AUDIT_ONLY_FIELD_NAMES:
                    continue
                visit(value[key], f"{path}.{key}" if path else key)
            return
        if isinstance(value, list):
            for item in value:
                visit(item, f"{path}[]")
            return
        discovered.add(path)

    for country in countries:
        visit(country, "countries[]")
    if not discovered:
        raise ValueError("Basic60 manual usage field scope cannot be empty")
    return sorted(discovered)


def _manual_usage_authorization(field_paths: list[str]) -> dict[str, Any]:
    return {
        "permission_basis": USAGE_AMENDMENT_ID,
        "permission_authority": "project_owner_manual_decision",
        "field_scope": {
            "mode": "all_normalized_basic60_fields",
            "field_count": len(field_paths),
            "fields_sha256": _canonical_sha256(field_paths),
        },
        "allowed_uses": list(MANUAL_ALLOWED_USES),
        "prohibited_uses": list(MANUAL_PROHIBITED_USES),
        "v1_runtime_capabilities": _v1_runtime_capabilities(),
    }


def build_ai_usage_policy(
    seed: dict[str, Any],
    *,
    seed_reference: dict[str, str],
    usage_amendment_reference: dict[str, str] | None,
) -> dict[str, Any]:
    """Project the A1 decision over every normalized field in one seed artifact.

    The returned artifact is deliberately machine-generated scope evidence. It is
    not a second permission decision and contains no per-field permission states.
    """

    if set(seed_reference) != {"path", "sha256"} or not SHA256_PATTERN.fullmatch(
        str(seed_reference.get("sha256") or "")
    ):
        raise ValueError("AI usage projection requires an exact seed artifact reference")
    if usage_amendment_reference is not None and (
        set(usage_amendment_reference) != {"path", "sha256"}
        or usage_amendment_reference.get("path") != USAGE_AMENDMENT_PATH.as_posix()
        or not SHA256_PATTERN.fullmatch(str(usage_amendment_reference.get("sha256") or ""))
    ):
        raise ValueError("AI usage projection requires the canonical A1 artifact reference")

    field_paths = _normalized_data_field_paths(seed)
    return {
        "schema_version": AI_USAGE_POLICY_SCHEMA,
        "record_type": "basic60_ai_usage_scope_projection",
        "projection_type": "machine_generated_all_field_scope_projection",
        "generator": "navigator-data",
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "permission_basis": USAGE_AMENDMENT_ID,
        "usage_amendment": usage_amendment_reference,
        "seed": seed_reference,
        "field_scope": {
            "mode": "all_normalized_basic60_fields",
            "field_count": len(field_paths),
            "fields_sha256": _canonical_sha256(field_paths),
            "fields": field_paths,
        },
        "owner_authorization_sha256": _canonical_sha256(_owner_usage_authorization()),
        "v1_runtime_capabilities": _v1_runtime_capabilities(),
    }


def _canonical_sha256(payload: Any) -> str:
    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _resolve(paths: RepositoryPaths, value: Path) -> Path:
    return value.resolve() if value.is_absolute() else (paths.root / value).resolve()


def _display_path(paths: RepositoryPaths, value: Path) -> str:
    try:
        return value.resolve().relative_to(paths.root.resolve()).as_posix()
    except ValueError:
        return str(value.resolve())


def _json_object(path: Path) -> dict[str, Any]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"JSON root must be an object: {path}")
    return payload


def _csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.DictReader(source)
        if reader.fieldnames is None:
            raise ValueError(f"CSV header is missing: {path}")
        return [dict(item) for item in reader]


def _csv_fieldnames(path: Path) -> list[str]:
    with path.open("r", encoding="utf-8-sig", newline="") as source:
        reader = csv.reader(source)
        try:
            fieldnames = next(reader)
        except StopIteration as error:
            raise ValueError(f"CSV header is missing: {path}") from error
    if not fieldnames or any(not str(value).strip() for value in fieldnames):
        raise ValueError(f"CSV header contains a blank field: {path}")
    return [str(value) for value in fieldnames]


def _canonical_source_uri(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError("Source URI is blank")
    try:
        parsed = urlsplit(text)
    except ValueError as error:
        raise ValueError("Source URI cannot be parsed") from error
    scheme = parsed.scheme.lower()
    if not scheme:
        raise ValueError("Source URI has no scheme")
    if scheme in {"http", "https"}:
        if parsed.username is not None or parsed.password is not None or parsed.hostname is None:
            raise ValueError("Source URI authority is invalid")
        try:
            hostname = parsed.hostname.encode("idna").decode("ascii").lower().rstrip(".")
        except UnicodeError as error:
            raise ValueError("Source URI hostname is invalid") from error
        try:
            port = parsed.port
        except ValueError as error:
            raise ValueError("Source URI port is invalid") from error
        default_port = (scheme == "http" and port == 80) or (scheme == "https" and port == 443)
        host = f"[{hostname}]" if ":" in hostname else hostname
        netloc = host if port is None or default_port else f"{host}:{port}"
        query = urlencode(sorted(parse_qsl(parsed.query, keep_blank_values=True)), doseq=True)
        return urlunsplit((scheme, netloc, parsed.path or "/", query, parsed.fragment))
    if scheme == "urn" and not parsed.netloc and parsed.path:
        return urlunsplit((scheme, "", parsed.path, parsed.query, parsed.fragment))
    raise ValueError("Source URI scheme is not approved for Basic60")


def build_basic60_country_source_uri_manifest(raw_dir: Path) -> dict[str, Any]:
    source_pattern = "countries/*/00_country_profile/country_profile.json"
    paths = sorted(raw_dir.glob(source_pattern), key=lambda item: item.as_posix())
    profile_iso3_codes = {
        str(row.get("iso3") or "").strip().upper() for row in _csv_rows(raw_dir / PROFILE_CSV)
    }
    if len(profile_iso3_codes) != EXPECTED_COUNTRIES or any(
        not re.fullmatch(r"[A-Z]{3}", value) for value in profile_iso3_codes
    ):
        raise ValueError("Country source URI manifest requires the exact 60-country CSV set")
    entries: list[dict[str, Any]] = []
    iso3_codes: set[str] = set()
    for path in paths:
        payload = _json_object(path)
        iso3 = str(payload.get("iso3") or "").strip().upper()
        directory_iso3 = path.parent.parent.name
        sources = payload.get("sources")
        if directory_iso3 != iso3 or not re.fullmatch(r"[A-Z]{3}", iso3) or iso3 in iso3_codes:
            raise ValueError(f"Country source manifest has an invalid or duplicate ISO3: {path}")
        if not isinstance(sources, list) or not sources:
            raise ValueError(f"Country source manifest has no source rows: {path}")
        iso3_codes.add(iso3)
        for source in sources:
            if not isinstance(source, dict):
                raise ValueError(f"Country source row is not an object: {path}")
            provider = str(source.get("source_name") or "").strip()
            fields = source.get("fields")
            normalized_fields = (
                sorted({str(field).strip() for field in fields if str(field).strip()})
                if isinstance(fields, list)
                else []
            )
            if not provider or not normalized_fields:
                raise ValueError(f"Country source row lacks provider or fields: {path}")
            entries.append(
                {
                    "iso3": iso3,
                    "provider": provider,
                    "fields": normalized_fields,
                    "url": _canonical_source_uri(source.get("url")),
                }
            )
    if (
        len(paths) != EXPECTED_COUNTRIES
        or len(iso3_codes) != EXPECTED_COUNTRIES
        or iso3_codes != profile_iso3_codes
    ):
        raise ValueError("Country source URI manifest requires exactly 60 country profiles")
    entries.sort(
        key=lambda item: (
            str(item["iso3"]),
            str(item["provider"]),
            tuple(item["fields"]),
            str(item["url"]),
        )
    )
    return {
        "schema_version": COUNTRY_SOURCE_URI_MANIFEST_SCHEMA,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "source_pattern": source_pattern,
        "country_file_count": len(paths),
        "entry_count": len(entries),
        "entries": entries,
        "entries_sha256": _canonical_sha256(entries),
    }


def _csv_source_uri_fingerprint(
    raw_dir: Path,
    *,
    dataset_path: PurePosixPath,
    selector: dict[str, str],
    uri_field: str,
    row_key_fields: tuple[str, ...],
) -> tuple[int, str]:
    rows = _csv_rows(raw_dir.joinpath(*dataset_path.parts))
    selected = [row for row in rows if row.get(selector["field"]) == selector["value"]]
    entries: list[dict[str, Any]] = []
    for row in selected:
        iso3 = str(row.get("iso3") or "").strip().upper()
        if not re.fullmatch(r"[A-Z]{3}", iso3):
            raise ValueError(f"Source URI row has an invalid ISO3 in {dataset_path}")
        row_key = {field: str(row.get(field) or "").strip() for field in row_key_fields}
        if any(not value for value in row_key.values()):
            raise ValueError(f"Source URI row has an incomplete identity in {dataset_path}")
        entries.append({"row_key": row_key, "url": _canonical_source_uri(row.get(uri_field))})
    entries.sort(
        key=lambda item: (
            tuple(str(item["row_key"][field]) for field in row_key_fields),
            str(item["url"]),
        )
    )
    return len(selected), _canonical_sha256(entries)


def _terms_snapshot_objects(evidence: Any) -> list[dict[str, Any]]:
    if not isinstance(evidence, dict):
        return []
    snapshots = evidence.get("terms_snapshots")
    if isinstance(snapshots, list):
        return [item for item in snapshots if isinstance(item, dict)]
    snapshot = evidence.get("terms_snapshot")
    return [snapshot] if isinstance(snapshot, dict) else []


def _terms_fingerprint_sha256(terms_uri: Any, snapshots: Any) -> str | None:
    snapshot_rows = snapshots if isinstance(snapshots, list) else []
    entries: list[dict[str, str]] = []
    try:
        canonical_terms_uri = _canonical_source_uri(terms_uri)
        for snapshot in snapshot_rows:
            if not isinstance(snapshot, dict):
                return None
            digest = str(snapshot.get("sha256") or "")
            if not SHA256_PATTERN.fullmatch(digest):
                return None
            entries.append(
                {
                    "retrieval_uri": _canonical_source_uri(snapshot.get("retrieval_uri")),
                    "sha256": digest,
                }
            )
    except ValueError:
        return None
    if not entries:
        return None
    entries.sort(key=lambda item: (item["retrieval_uri"], item["sha256"]))
    return _canonical_sha256({"terms_uri": canonical_terms_uri, "snapshots": entries})


def _source_evidence_for_fingerprint(
    paths: RepositoryPaths,
    source: dict[str, Any],
) -> dict[str, Any] | None:
    reference = source.get("evidence")
    if not isinstance(reference, dict):
        return None
    value = reference.get("path")
    digest = reference.get("sha256")
    if not _valid_relative_path(value, PurePosixPath("data/basic60/evidence")):
        return None
    target = paths.root.joinpath(*PurePosixPath(str(value)).parts)
    try:
        if (
            target.is_symlink()
            or not target.resolve(strict=True).is_file()
            or not isinstance(digest, str)
            or not SHA256_PATTERN.fullmatch(digest)
            or sha256_file(target) != digest
        ):
            return None
        return _json_object(target)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        return None


def _source_admission_fingerprint(
    raw_dir: Path,
    source: dict[str, Any],
    *,
    country_source_manifest: dict[str, Any],
    terms_evidence: dict[str, Any] | None,
) -> dict[str, Any]:
    domain = str(source.get("data_domain") or "")
    spec = SOURCE_CONTRACT_SPECS.get(domain)
    if spec is None:
        raise ValueError(f"Unsupported Basic60 source domain: {domain}")
    dataset_path = spec["dataset_path"]
    assert isinstance(dataset_path, PurePosixPath)
    fieldnames = _csv_fieldnames(raw_dir.joinpath(*dataset_path.parts))
    uri_mode = str(spec["source_uri_mode"])
    if uri_mode == "country_source_manifest":
        uri_count = int(country_source_manifest["entry_count"])
        uri_sha256 = str(country_source_manifest["entries_sha256"])
    else:
        selector = spec["record_selector"]
        uri_field = str(spec["source_uri_field"])
        row_key_fields = spec["source_uri_row_key_fields"]
        assert isinstance(selector, dict)
        assert isinstance(row_key_fields, tuple)
        uri_count, uri_sha256 = _csv_source_uri_fingerprint(
            raw_dir,
            dataset_path=dataset_path,
            selector=selector,
            uri_field=uri_field,
            row_key_fields=row_key_fields,
        )
    return {
        "dataset_path": dataset_path.as_posix(),
        "record_selector": spec["record_selector"],
        "schema_field_count": len(fieldnames),
        "schema_fingerprint_sha256": _canonical_sha256(
            {"dataset_path": dataset_path.as_posix(), "fieldnames": fieldnames}
        ),
        "source_uri_mode": uri_mode,
        "source_uri_field": spec["source_uri_field"],
        "source_uri_row_key_fields": list(spec["source_uri_row_key_fields"]),
        "source_uri_entry_count": uri_count,
        "source_uri_fingerprint_sha256": uri_sha256,
        "terms_fingerprint_sha256": _terms_fingerprint_sha256(
            source.get("terms_uri"),
            _terms_snapshot_objects(terms_evidence),
        ),
    }


def _number_is_valid(value: str) -> bool:
    try:
        return math.isfinite(float(value))
    except ValueError:
        return False


def _json_list(value: str) -> list[str]:
    try:
        payload = json.loads(value)
    except json.JSONDecodeError:
        return []
    if not isinstance(payload, list):
        return []
    return [str(item).strip() for item in payload if str(item).strip()]


def _check(
    checks: list[CheckResult],
    condition: bool,
    code: str,
    message: str,
    location: str | None = None,
) -> None:
    if not condition:
        checks.append(CheckResult(code=code, message=message, location=location))


def inspect_basic60_raw(raw_dir: Path) -> tuple[dict[str, Any], list[CheckResult]]:
    checks: list[CheckResult] = []
    summary: dict[str, Any] = {
        "country_count": 0,
        "macro_annual_record_count": 0,
        "energy_latest_record_count": 0,
        "available_observation_count": 0,
        "pending_observation_count": 0,
    }
    if not raw_dir.is_dir():
        checks.append(
            CheckResult(
                code="B60_RAW_DIRECTORY_MISSING",
                message="Basic60 raw-material directory does not exist",
                location=str(raw_dir),
            )
        )
        return summary, checks

    required = (COLLECTION_MANIFEST, PROFILE_CSV, METRIC_CSV)
    missing = [path.as_posix() for path in required if not (raw_dir / path).is_file()]
    if missing:
        checks.append(
            CheckResult(
                code="B60_RAW_REQUIRED_FILE_MISSING",
                message=f"Missing required Basic60 files: {', '.join(missing)}",
                location=str(raw_dir),
            )
        )
        return summary, checks

    try:
        source_manifest = _json_object(raw_dir / COLLECTION_MANIFEST)
        profiles = _csv_rows(raw_dir / PROFILE_CSV)
        metric_rows = _csv_rows(raw_dir / METRIC_CSV)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        checks.append(
            CheckResult(
                code="B60_RAW_PARSE_FAILED",
                message=str(error),
                location=str(raw_dir),
            )
        )
        return summary, checks

    macro_rows = [item for item in metric_rows if item.get("record_type") == "macro_annual"]
    energy_rows = [item for item in metric_rows if item.get("record_type") == "energy_latest"]
    unknown_rows = [
        item
        for item in metric_rows
        if item.get("record_type") not in {"macro_annual", "energy_latest"}
    ]
    summary.update(
        {
            "country_count": len(profiles),
            "macro_annual_record_count": len(macro_rows),
            "energy_latest_record_count": len(energy_rows),
        }
    )

    profile_isos = [item.get("iso3", "").strip() for item in profiles]
    macro_isos = [item.get("iso3", "").strip() for item in macro_rows]
    energy_isos = [item.get("iso3", "").strip() for item in energy_rows]
    profile_set = set(profile_isos)

    _check(
        checks,
        len(profiles) == EXPECTED_COUNTRIES,
        "B60_PROFILE_COUNT_MISMATCH",
        f"Expected {EXPECTED_COUNTRIES} country profiles, found {len(profiles)}",
        PROFILE_CSV.as_posix(),
    )
    _check(
        checks,
        len(profile_set) == EXPECTED_COUNTRIES and "" not in profile_set,
        "B60_PROFILE_ISO_INVALID",
        "Country profiles require 60 unique non-empty ISO3 codes",
        PROFILE_CSV.as_posix(),
    )
    _check(
        checks,
        len(macro_rows) == EXPECTED_MACRO_ROWS,
        "B60_MACRO_COUNT_MISMATCH",
        f"Expected {EXPECTED_MACRO_ROWS} macro rows, found {len(macro_rows)}",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        len(energy_rows) == EXPECTED_ENERGY_ROWS,
        "B60_ENERGY_COUNT_MISMATCH",
        f"Expected {EXPECTED_ENERGY_ROWS} energy rows, found {len(energy_rows)}",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        not unknown_rows,
        "B60_RECORD_TYPE_INVALID",
        f"Found {len(unknown_rows)} unsupported metric record types",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        set(macro_isos) == profile_set and set(energy_isos) == profile_set,
        "B60_COUNTRY_SET_MISMATCH",
        "Profile, macro, and energy country sets must match exactly",
        METRIC_CSV.as_posix(),
    )

    required_profile_fields = (
        "iso2",
        "iso3",
        "country_name_zh",
        "country_name_en",
        "official_name_en",
        "capital",
        "admin_level_1_count",
        "admin_level_1_type",
        "population",
        "population_year",
        "official_languages",
        "currency_code",
        "currency_name",
        "time_zones",
        "area_sq_km",
        "area_year",
        "region",
        "collected_at",
        "review_status",
    )
    available = 0
    pending = 0
    for index, item in enumerate(profiles, start=2):
        missing_fields = [
            field for field in required_profile_fields if not item.get(field, "").strip()
        ]
        if missing_fields:
            checks.append(
                CheckResult(
                    code="B60_PROFILE_FIELD_MISSING",
                    message=f"Missing required profile fields: {', '.join(missing_fields)}",
                    location=f"{PROFILE_CSV.as_posix()}:{index}",
                )
            )
        if item.get("review_status") != "machine_collected_pending_human_review":
            checks.append(
                CheckResult(
                    code="B60_RAW_REVIEW_STATUS_INVALID",
                    message="Raw records must remain machine_collected_pending_human_review",
                    location=f"{PROFILE_CSV.as_posix()}:{index}",
                )
            )
        for _metric_code, field, _unit in PROFILE_METRICS:
            value = item.get(field, "").strip()
            if value and _number_is_valid(value):
                available += 1
            else:
                pending += 1
                if value:
                    checks.append(
                        CheckResult(
                            code="B60_NUMERIC_VALUE_INVALID",
                            message=f"{field} is not a finite number",
                            location=f"{PROFILE_CSV.as_posix()}:{index}",
                        )
                    )

    years_by_country: dict[str, set[int]] = defaultdict(set)
    lka_2024_rate_blank = False
    for index, item in enumerate(macro_rows, start=2):
        iso3 = item.get("iso3", "").strip()
        try:
            year = int(item.get("year", ""))
        except ValueError:
            year = -1
        years_by_country[iso3].add(year)
        if not item.get("macro_source_url", "").strip():
            checks.append(
                CheckResult(
                    code="B60_SOURCE_REF_MISSING",
                    message="Macro record requires macro_source_url",
                    location=f"{METRIC_CSV.as_posix()}:{index}",
                )
            )
        if item.get("review_status") != "machine_collected_pending_human_review":
            checks.append(
                CheckResult(
                    code="B60_RAW_REVIEW_STATUS_INVALID",
                    message="Raw records must remain machine_collected_pending_human_review",
                    location=f"{METRIC_CSV.as_posix()}:{index}",
                )
            )
        for _metric_code, field, _unit in MACRO_METRICS:
            value = item.get(field, "").strip()
            if value:
                if _number_is_valid(value):
                    available += 1
                else:
                    checks.append(
                        CheckResult(
                            code="B60_NUMERIC_VALUE_INVALID",
                            message=f"{field} is not a finite number",
                            location=f"{METRIC_CSV.as_posix()}:{index}",
                        )
                    )
            else:
                pending += 1
                if iso3 == "LKA" and year == 2024 and field == "official_exchange_rate_lcu_per_usd":
                    lka_2024_rate_blank = True

    for iso3 in sorted(profile_set):
        if years_by_country.get(iso3) != {2020, 2021, 2022, 2023, 2024}:
            checks.append(
                CheckResult(
                    code="B60_MACRO_YEARS_MISMATCH",
                    message=f"{iso3} must contain exactly 2020-2024 macro years",
                    location=METRIC_CSV.as_posix(),
                )
            )

    demand_pending_countries: set[str] = set()
    metric_header_offset = len(macro_rows) + 2
    for index, item in enumerate(energy_rows, start=metric_header_offset):
        if not item.get("energy_source_url", "").strip():
            checks.append(
                CheckResult(
                    code="B60_SOURCE_REF_MISSING",
                    message="Energy record requires energy_source_url",
                    location=f"{METRIC_CSV.as_posix()}:{index}",
                )
            )
        for _metric_code, field, _unit, year_field in ENERGY_METRICS:
            value = item.get(field, "").strip()
            source_year = item.get(year_field, "").strip()
            if value and source_year and _number_is_valid(value) and source_year.isdigit():
                available += 1
            else:
                pending += 1
                checks.append(
                    CheckResult(
                        code="B60_ENERGY_VALUE_INVALID",
                        message=f"{field} requires a finite value and explicit source year",
                        location=f"{METRIC_CSV.as_posix()}:{index}",
                    )
                )
        demand = item.get("electricity_demand_gwh", "").strip()
        demand_year = item.get("electricity_demand_year", "").strip()
        if not demand and not demand_year:
            pending += 1
            demand_pending_countries.add(item.get("iso3", "").strip())
        else:
            if demand and _number_is_valid(demand) and demand_year.isdigit():
                available += 1
            checks.append(
                CheckResult(
                    code="B60_DEMAND_MUST_REMAIN_PENDING",
                    message="V1 requires electricity demand to remain explicitly pending",
                    location=f"{METRIC_CSV.as_posix()}:{index}",
                )
            )

    summary["available_observation_count"] = available
    summary["pending_observation_count"] = pending
    _check(
        checks,
        demand_pending_countries == profile_set,
        "B60_DEMAND_PENDING_COVERAGE_MISMATCH",
        "All 60 countries must retain a pending electricity-demand observation",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        lka_2024_rate_blank,
        "B60_LKA_RATE_PENDING_MISSING",
        "LKA 2024 official exchange rate must remain explicitly pending",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        available == EXPECTED_AVAILABLE_VALUES,
        "B60_AVAILABLE_COUNT_MISMATCH",
        f"Expected {EXPECTED_AVAILABLE_VALUES} available values, found {available}",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        pending == EXPECTED_PENDING_VALUES,
        "B60_PENDING_COUNT_MISMATCH",
        f"Expected {EXPECTED_PENDING_VALUES} pending values, found {pending}",
        METRIC_CSV.as_posix(),
    )
    _check(
        checks,
        source_manifest.get("country_count") == EXPECTED_COUNTRIES
        and source_manifest.get("macro_annual_record_count") == EXPECTED_MACRO_ROWS
        and source_manifest.get("energy_latest_record_count") == EXPECTED_ENERGY_ROWS,
        "B60_SOURCE_MANIFEST_COUNT_MISMATCH",
        "Collection manifest counts do not match the Basic60 profile",
        COLLECTION_MANIFEST.as_posix(),
    )
    return summary, checks


def build_raw_manifest(paths: RepositoryPaths, raw_dir: Path) -> dict[str, Any]:
    if not raw_dir.is_dir():
        raise ValueError(f"Basic60 raw-material directory does not exist: {raw_dir}")
    digest = hashlib.sha256()
    included_count = 0
    included_bytes = 0
    excluded_ads: list[str] = []
    unsafe_entries: list[str] = []
    scoped_files: list[dict[str, Any]] = []
    scoped_paths = {PROFILE_CSV.as_posix(), METRIC_CSV.as_posix(), COLLECTION_MANIFEST.as_posix()}

    for path in sorted(raw_dir.rglob("*"), key=lambda value: value.relative_to(raw_dir).as_posix()):
        relative = path.relative_to(raw_dir).as_posix()
        if path.is_symlink():
            unsafe_entries.append(relative)
            continue
        if not path.is_file():
            continue
        if relative.endswith(":Zone.Identifier"):
            excluded_ads.append(relative)
            continue
        before = path.stat()
        file_hash = sha256_file(path)
        after = path.stat()
        if before.st_size != after.st_size or before.st_mtime_ns != after.st_mtime_ns:
            raise ValueError(f"Raw file changed while hashing: {relative}")
        digest.update(f"{relative}\0{after.st_size}\0{file_hash}\n".encode())
        included_count += 1
        included_bytes += after.st_size
        if relative in scoped_paths:
            scoped_files.append({"path": relative, "byte_size": after.st_size, "sha256": file_hash})

    return {
        "schema_version": SCHEMA_VERSION,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "source_directory": _display_path(paths, raw_dir),
        "algorithm": "sha256(sorted(relative_path\\0byte_size\\0file_sha256\\n))",
        "root_sha256": digest.hexdigest(),
        "included_file_count": included_count,
        "included_byte_count": included_bytes,
        "excluded_zone_identifier_count": len(excluded_ads),
        "excluded_zone_identifier_paths": excluded_ads,
        "unsafe_entry_count": len(unsafe_entries),
        "unsafe_entries": unsafe_entries,
        "scoped_files": sorted(scoped_files, key=lambda item: str(item["path"])),
    }


def _metric_value(
    *,
    metric_code: str,
    raw_value: str,
    unit: str,
    period: str | None,
    source_ref: str,
    null_reason: str | None = None,
) -> dict[str, Any]:
    value = raw_value.strip()
    return {
        "metric_code": metric_code,
        "period": period,
        "original_value": value or None,
        "normalized_value": value or None,
        "unit": unit,
        "value_status": "available" if value else "pending",
        "null_reason": None if value else null_reason or "source_value_missing",
        "source_ref": source_ref,
        "quality_status": "machine_collected_pending_human_review",
    }


def _capitals(value: str) -> list[dict[str, Any]]:
    role_map = {
        "administrative": "administrative",
        "legislative": "legislative",
        "judicial": "judicial",
        "official": "official",
        "de facto": "de_facto",
    }
    capitals: list[dict[str, Any]] = []
    for order, part in enumerate((item.strip() for item in value.split(";")), start=1):
        match = re.fullmatch(r"(.+?)\s*\(([^)]+)\)", part)
        name = match.group(1).strip() if match else part
        label = match.group(2).strip().lower() if match else "official"
        capitals.append(
            {
                "name": name,
                "role": role_map.get(label, "official"),
                "display_order": order,
            }
        )
    return capitals


def build_basic60_seed(raw_dir: Path, raw_manifest: dict[str, Any]) -> dict[str, Any]:
    profiles = _csv_rows(raw_dir / PROFILE_CSV)
    metric_rows = _csv_rows(raw_dir / METRIC_CSV)
    rows_by_country: dict[str, list[dict[str, str]]] = defaultdict(list)
    for row in metric_rows:
        rows_by_country[row.get("iso3", "").strip()].append(row)

    countries: list[dict[str, Any]] = []
    for profile in sorted(profiles, key=lambda item: int(item["order"])):
        iso3 = profile["iso3"].strip()
        localized = [
            {
                "field_code": "short_name",
                "locale": "zh-CN",
                "text": profile["country_name_zh"].strip(),
                "preferred": True,
                "translation_status": "pending_human_review",
            },
            {
                "field_code": "short_name",
                "locale": "en",
                "text": profile["country_name_en"].strip(),
                "preferred": True,
                "translation_status": "source",
            },
            {
                "field_code": "official_name",
                "locale": "en",
                "text": profile["official_name_en"].strip(),
                "preferred": True,
                "translation_status": "source",
            },
        ]
        localized.extend(
            {
                "field_code": "local_name",
                "locale": "und",
                "text": name,
                "preferred": index == 0,
                "translation_status": "source",
            }
            for index, name in enumerate(_json_list(profile.get("local_names", "")))
        )
        metrics = [
            _metric_value(
                metric_code=metric_code,
                raw_value=profile[field],
                unit=unit,
                period=profile["population_year"]
                if metric_code == "population_total"
                else profile["area_year"],
                source_ref="SRC-WORLD-BANK-WDI",
            )
            for metric_code, field, unit in PROFILE_METRICS
        ]
        for row in rows_by_country[iso3]:
            if row["record_type"] == "macro_annual":
                for metric_code, field, unit in MACRO_METRICS:
                    metrics.append(
                        _metric_value(
                            metric_code=metric_code,
                            raw_value=row[field],
                            unit=unit,
                            period=row["year"],
                            source_ref="SRC-WORLD-BANK-WDI",
                            null_reason="not_available_from_approved_snapshot",
                        )
                    )
            elif row["record_type"] == "energy_latest":
                for metric_code, field, unit, year_field in ENERGY_METRICS:
                    metrics.append(
                        _metric_value(
                            metric_code=metric_code,
                            raw_value=row[field],
                            unit=unit,
                            period=row[year_field] or None,
                            source_ref="SRC-IRENASTAT",
                        )
                    )
                metrics.append(
                    _metric_value(
                        metric_code="electricity_demand_gwh",
                        raw_value=row["electricity_demand_gwh"],
                        unit="GWH",
                        period=row["electricity_demand_year"] or None,
                        source_ref="SRC-ENERGY-DEMAND-PENDING",
                        null_reason="no_reliable_uniform_public_value",
                    )
                )

        countries.append(
            {
                "id": f"COUNTRY-{iso3}",
                "iso2": profile["iso2"].strip(),
                "iso3": iso3,
                "region_code": profile["region"].strip(),
                "coverage_status": "covered",
                "coverage_level": "Basic",
                "opportunity_level": "pending",
                "policy_friendliness_level": "pending",
                "risk_assessment_status": "unknown",
                "risk_level": None,
                "localized_texts": localized,
                "capitals": _capitals(profile["capital"].strip()),
                "admin_structure": {
                    "level": 1,
                    "unit_type": profile["admin_level_1_type"].strip(),
                    "count": int(profile["admin_level_1_count"]),
                },
                "languages": _json_list(profile["official_languages"]),
                "currency": {
                    "code": profile["currency_code"].strip(),
                    "name": profile["currency_name"].strip(),
                },
                "time_zones": _json_list(profile["time_zones"]),
                "metrics": metrics,
                "collected_at": profile["collected_at"].strip(),
            }
        )

    return {
        "schema_version": SCHEMA_VERSION,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "release_profile": PROFILE_ID,
        "formal_gate_status": FORMAL_GATE_STATUS,
        "publication_status": "candidate_not_approved",
        "source_root_sha256": raw_manifest["root_sha256"],
        "countries": countries,
    }


def _period_bounds(period: Any, fallback_date: str) -> tuple[str, str, str]:
    label = str(period or "").strip()
    if len(label) == 4 and label.isdigit():
        return f"{label}-01-01", f"{label}-12-31", label
    return fallback_date, fallback_date, label or "pending"


def _provenance_fields(
    source_ref: str,
    snapshot_ref: str,
    raw_record_ref: str,
) -> dict[str, str]:
    return {
        "source_ref": source_ref,
        "source_snapshot_ref": snapshot_ref,
        "raw_record_ref": raw_record_ref,
    }


def _sample_record_identity(iso3: str, metric: dict[str, Any]) -> tuple[str, str, str, str]:
    return (
        iso3,
        str(metric.get("metric_code") or ""),
        str(metric.get("period") or "not_applicable"),
        str(metric.get("source_ref") or ""),
    )


def build_basic60_sample_candidate(
    seed: dict[str, Any],
    *,
    seed_artifact_sha256: str,
) -> dict[str, Any]:
    if not SHA256_PATTERN.fullmatch(seed_artifact_sha256):
        raise ValueError("Basic60 sample candidate requires the exact seed artifact SHA-256")
    countries = seed.get("countries")
    if not isinstance(countries, list) or len(countries) != EXPECTED_COUNTRIES:
        raise ValueError("Basic60 sample candidate requires exactly 60 seed countries")

    available_by_country: dict[str, list[dict[str, Any]]] = {}
    all_available: list[tuple[str, dict[str, Any]]] = []
    for country in countries:
        if not isinstance(country, dict):
            raise ValueError("Basic60 sample country must be an object")
        iso3 = str(country.get("iso3") or "")
        metrics = country.get("metrics")
        metric_rows = metrics if isinstance(metrics, list) else []
        available = [
            metric
            for metric in metric_rows
            if isinstance(metric, dict) and metric.get("value_status") == "available"
        ]
        if not iso3 or len(available) < 3:
            raise ValueError(f"Basic60 sample requires at least three available values for {iso3}")
        available.sort(key=lambda metric: _sample_record_identity(iso3, metric))
        available_by_country[iso3] = available
        all_available.extend((iso3, metric) for metric in available)
    if len(available_by_country) != EXPECTED_COUNTRIES:
        raise ValueError("Basic60 sample countries must have unique ISO3 codes")

    required_coverage = {
        "metric_codes": sorted({str(metric["metric_code"]) for _iso3, metric in all_available}),
        "periods": sorted(
            {str(metric.get("period") or "not_applicable") for _iso3, metric in all_available}
        ),
        "source_refs": sorted({str(metric["source_ref"]) for _iso3, metric in all_available}),
    }
    selected: dict[str, list[dict[str, Any]]] = {iso3: [] for iso3 in sorted(available_by_country)}
    selected_identities: set[tuple[str, str, str, str]] = set()

    coverage_tokens = [
        (field, value)
        for field in ("metric_codes", "periods", "source_refs")
        for value in required_coverage[field]
    ]
    metric_field = {
        "metric_codes": "metric_code",
        "periods": "period",
        "source_refs": "source_ref",
    }
    for coverage_field, value in coverage_tokens:
        candidates: list[tuple[str, tuple[str, str, str, str], dict[str, Any]]] = []
        for iso3, metrics in available_by_country.items():
            if len(selected[iso3]) >= 3:
                continue
            for metric in metrics:
                identity = _sample_record_identity(iso3, metric)
                actual_value = (
                    identity[2]
                    if coverage_field == "periods"
                    else str(metric.get(metric_field[coverage_field]) or "")
                )
                if actual_value != value or identity in selected_identities:
                    continue
                score = hashlib.sha256(
                    f"{RELEASE_ID}|coverage|{coverage_field}|{value}|{'|'.join(identity)}".encode()
                ).hexdigest()
                candidates.append((score, identity, metric))
        if not candidates:
            raise ValueError(f"Basic60 sample cannot cover {coverage_field}={value}")
        _score, identity, metric = min(candidates, key=lambda item: item[0])
        selected[identity[0]].append(metric)
        selected_identities.add(identity)

    for iso3, metrics in available_by_country.items():
        remaining = [
            metric
            for metric in metrics
            if _sample_record_identity(iso3, metric) not in selected_identities
        ]
        remaining.sort(
            key=lambda metric: hashlib.sha256(
                f"{RELEASE_ID}|fill|{'|'.join(_sample_record_identity(iso3, metric))}".encode()
            ).hexdigest()
        )
        for metric in remaining[: 3 - len(selected[iso3])]:
            selected[iso3].append(metric)
            selected_identities.add(_sample_record_identity(iso3, metric))
        if len(selected[iso3]) != 3:
            raise ValueError(f"Basic60 sample could not select three records for {iso3}")

    records: list[dict[str, Any]] = []
    for iso3, metrics in sorted(selected.items()):
        for metric in sorted(metrics, key=lambda item: _sample_record_identity(iso3, item)):
            identity = _sample_record_identity(iso3, metric)
            sample_digest = hashlib.sha256("|".join(identity).encode()).hexdigest()[:16]
            records.append(
                {
                    "sample_id": f"B60-SAMPLE-{sample_digest}",
                    "iso3": iso3,
                    "metric_code": identity[1],
                    "period": identity[2],
                    "source_ref": identity[3],
                    "unit": metric.get("unit"),
                    "original_value": metric.get("original_value"),
                    "normalized_value": metric.get("normalized_value"),
                    "value_status": "available",
                }
            )
    actual_coverage = {
        "metric_codes": sorted({str(item["metric_code"]) for item in records}),
        "periods": sorted({str(item["period"]) for item in records}),
        "source_refs": sorted({str(item["source_ref"]) for item in records}),
    }
    if (
        len(records) != EXPECTED_SAMPLE_SIZE
        or len({str(item["sample_id"]) for item in records}) != EXPECTED_SAMPLE_SIZE
        or actual_coverage != required_coverage
    ):
        raise ValueError("Basic60 frozen sample does not satisfy size or coverage invariants")
    return {
        "schema_version": SAMPLE_CANDIDATE_SCHEMA,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "seed_artifact_sha256": seed_artifact_sha256,
        "selection_algorithm": "coverage_first_sha256_rank_then_per_country_sha256_fill.v1",
        "sample_size": EXPECTED_SAMPLE_SIZE,
        "country_count": EXPECTED_COUNTRIES,
        "per_country_count": 3,
        "coverage": required_coverage,
        "records": records,
    }


def build_basic60_sample_review_manifest(
    sample_candidate_reference: dict[str, str],
    sample_candidate: dict[str, Any],
    *,
    review_slot: int,
) -> dict[str, Any]:
    """Build one blank reviewer copy; this function never manufactures review outcomes."""

    if review_slot not in {1, 2}:
        raise ValueError("Sample review slot must be 1 or 2")
    records = sample_candidate.get("records")
    if not isinstance(records, list):
        raise ValueError("Sample candidate records are required")
    sample_ids = {str(item.get("sample_id") or "") for item in records if isinstance(item, dict)}
    if "" in sample_ids or len(sample_ids) != EXPECTED_SAMPLE_SIZE:
        raise ValueError("Sample candidate must contain 180 unique sample IDs")
    items = [{"sample_id": sample_id, "outcome": None} for sample_id in sorted(sample_ids)]
    return {
        "schema_version": SAMPLE_REVIEW_MANIFEST_SCHEMA,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "sample_candidate": sample_candidate_reference,
        "sample_size": sample_candidate.get("sample_size"),
        "country_count": sample_candidate.get("country_count"),
        "per_country_count": sample_candidate.get("per_country_count"),
        "coverage": sample_candidate.get("coverage"),
        "review_slot": review_slot,
        "template_only": True,
        "reviewer_name": None,
        "reviewed_at": None,
        "status": "pending",
        "items": items,
        "items_sha256": _canonical_sha256(items),
    }


def _metric_review_id(iso3: str, metric: dict[str, Any]) -> str:
    identity = _sample_record_identity(iso3, metric)
    return f"B60-OBS-{hashlib.sha256('|'.join(identity).encode()).hexdigest()[:20]}"


def _expected_item_review_ids(seed: dict[str, Any]) -> dict[str, list[str]]:
    chinese_names: list[str] = []
    pending_values: list[str] = []
    derived_values: list[str] = []
    derived_codes = {"renewable_share_capacity_pct", "renewable_share_generation_pct"}
    countries = seed.get("countries")
    if not isinstance(countries, list):
        return {
            "chinese_names": [],
            "pending_values": [],
            "derived_values": [],
        }
    for country in countries:
        if not isinstance(country, dict):
            continue
        iso3 = str(country.get("iso3") or "")
        chinese_names.append(f"B60-NAME-{iso3}-ZH")
        metrics = country.get("metrics")
        for metric in metrics if isinstance(metrics, list) else []:
            if not isinstance(metric, dict):
                continue
            item_id = _metric_review_id(iso3, metric)
            if metric.get("value_status") == "pending":
                pending_values.append(item_id)
            if metric.get("metric_code") in derived_codes:
                derived_values.append(item_id)
    return {
        "chinese_names": sorted(chinese_names),
        "pending_values": sorted(pending_values),
        "derived_values": sorted(derived_values),
    }


def _decimal_text(value: Decimal) -> str:
    normalized = format(value, "f").rstrip("0").rstrip(".")
    return normalized or "0"


def _item_review_rows(seed: dict[str, Any], *, blank: bool) -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = {
        "chinese_names": [],
        "pending_values": [],
        "derived_values": [],
    }
    countries = seed.get("countries")
    for country in countries if isinstance(countries, list) else []:
        if not isinstance(country, dict):
            continue
        iso3 = str(country.get("iso3") or "")
        localized = country.get("localized_texts")
        chinese_name = next(
            (
                str(item.get("text") or "")
                for item in (localized if isinstance(localized, list) else [])
                if isinstance(item, dict)
                if item.get("field_code") == "short_name" and item.get("locale") == "zh-CN"
            ),
            "",
        )
        result["chinese_names"].append(
            {
                "item_id": f"B60-NAME-{iso3}-ZH",
                "iso3": iso3,
                "observed_value": None if blank else chinese_name,
                "outcome": None if blank else "approved",
            }
        )
        metrics = country.get("metrics")
        metric_rows = [
            item
            for item in (metrics if isinstance(metrics, list) else [])
            if isinstance(item, dict)
        ]
        metric_index = {str(item.get("metric_code") or ""): item for item in metric_rows}
        for metric in metric_rows:
            if metric.get("value_status") == "pending":
                result["pending_values"].append(
                    {
                        "item_id": _metric_review_id(iso3, metric),
                        "iso3": iso3,
                        "metric_code": metric.get("metric_code"),
                        "period": metric.get("period"),
                        "observed_value_status": None if blank else metric.get("value_status"),
                        "observed_null_reason": None if blank else metric.get("null_reason"),
                        "outcome": None if blank else "approved",
                    }
                )
            code = str(metric.get("metric_code") or "")
            dependencies = {
                "renewable_share_capacity_pct": (
                    "renewable_capacity_mw",
                    "electricity_installed_capacity_mw",
                ),
                "renewable_share_generation_pct": (
                    "renewable_generation_gwh",
                    "electricity_generation_gwh",
                ),
            }.get(code)
            if dependencies is None:
                continue
            numerator = metric_index.get(dependencies[0], {})
            denominator = metric_index.get(dependencies[1], {})
            try:
                numerator_value = Decimal(str(numerator.get("normalized_value")))
                denominator_value = Decimal(str(denominator.get("normalized_value")))
                recomputed = _decimal_text(numerator_value / denominator_value * Decimal(100))
            except (InvalidOperation, ZeroDivisionError):
                recomputed = ""
            result["derived_values"].append(
                {
                    "item_id": _metric_review_id(iso3, metric),
                    "iso3": iso3,
                    "metric_code": code,
                    "period": metric.get("period"),
                    "numerator_metric_code": dependencies[0],
                    "numerator_value": None if blank else numerator.get("normalized_value"),
                    "denominator_metric_code": dependencies[1],
                    "denominator_value": None if blank else denominator.get("normalized_value"),
                    "reported_value": None if blank else metric.get("normalized_value"),
                    "recomputed_value": None if blank else recomputed,
                    "outcome": None if blank else "approved",
                }
            )
    for rows in result.values():
        rows.sort(key=lambda item: str(item["item_id"]))
    return result


def build_basic60_item_review_manifest(
    seed_reference: dict[str, str],
    seed: dict[str, Any],
) -> dict[str, Any]:
    """Build the single blank D4 item-review copy for the one acceptance meeting."""

    rows = _item_review_rows(seed, blank=True)
    groups = [
        {
            "review_type": review_type,
            "item_count": len(rows[review_type]),
            "items": rows[review_type],
        }
        for review_type in ("chinese_names", "pending_values", "derived_values")
    ]
    return {
        "schema_version": ITEM_REVIEW_MANIFEST_SCHEMA,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "seed_artifact": seed_reference,
        "template_only": True,
        "reviewer_name": None,
        "reviewed_at": None,
        "status": "pending",
        "groups": groups,
        "groups_sha256": _canonical_sha256(groups),
    }


def build_basic60_source_matrix_candidate(raw_dir: Path) -> dict[str, Any]:
    rows = _csv_rows(raw_dir / PROFILE_CSV)
    iso3_codes = sorted(str(row.get("iso3") or "") for row in rows)
    if len(iso3_codes) != EXPECTED_COUNTRIES or len(set(iso3_codes)) != EXPECTED_COUNTRIES:
        raise ValueError("Source matrix requires exactly 60 unique ISO3 codes")
    if not set(FUTURE_DEEP_DIVE_PRIORITY_COUNTRY_CODES).issubset(iso3_codes):
        raise ValueError("Source matrix is missing a future deep-dive priority country")
    matrix_rows = [
        {"iso3": iso3, "domain_id": domain_id} for iso3 in iso3_codes for domain_id in DOMAIN_IDS
    ]
    return {
        "schema_version": SOURCE_MATRIX_SCHEMA,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "coverage_level": "Basic",
        "country_count": EXPECTED_COUNTRIES,
        "domain_count": len(DOMAIN_IDS),
        "row_count": EXPECTED_COUNTRIES * len(DOMAIN_IDS),
        "rows": matrix_rows,
        "rows_sha256": _canonical_sha256(matrix_rows),
    }


def _derived_source_matrix_sha256(matrix: dict[str, Any], domains: list[dict[str, Any]]) -> str:
    domain_index = {str(item.get("domain_id")): item for item in domains}
    rows = [
        {
            "iso3": item["iso3"],
            "domain_id": item["domain_id"],
            "primary_source_id": domain_index[item["domain_id"]].get("primary_source_id"),
            "alternative_source_id": domain_index[item["domain_id"]].get("alternative_source_id"),
        }
        for item in matrix["rows"]
    ]
    return _canonical_sha256(rows)


def build_basic60_replay_manifests(
    raw_manifest: dict[str, Any],
    *,
    runtime_seed_sha256: str,
    d3_artifacts: dict[str, Any],
    stage_payloads: dict[str, dict[str, Any]],
) -> dict[str, dict[str, Any]]:
    if not SHA256_PATTERN.fullmatch(runtime_seed_sha256):
        raise ValueError("Replay manifests require the exact runtime seed file hash")
    stage_references = d3_artifacts.get("stages")
    replay = d3_artifacts.get("replay")
    if not isinstance(stage_references, dict) or not isinstance(replay, dict):
        raise ValueError("D3 artifacts require three stage references and replay hashes")
    first_hashes = replay.get("first_output_sha256")
    second_hashes = replay.get("second_output_sha256")
    if not isinstance(first_hashes, dict) or not isinstance(second_hashes, dict):
        raise ValueError("D3 artifacts require two independently computed stage hash maps")
    result: dict[str, dict[str, Any]] = {}
    for pipeline_id, stage_id in zip(PIPELINE_IDS, ("parse", "standardize", "entity"), strict=True):
        stage = stage_payloads.get(stage_id)
        reference = stage_references.get(stage_id)
        if not isinstance(stage, dict) or not isinstance(reference, dict):
            raise ValueError(f"Missing D3 runtime stage: {stage_id}")
        quality = stage.get("quality_observations")
        if not isinstance(quality, dict):
            raise ValueError(f"Missing D3 quality observations: {stage_id}")
        open_anomalies = sum(len(value) for value in quality.values() if isinstance(value, list))
        first_hash = str(first_hashes.get(stage_id) or "")
        second_hash = str(second_hashes.get(stage_id) or "")
        if (
            not SHA256_PATTERN.fullmatch(first_hash)
            or not SHA256_PATTERN.fullmatch(second_hash)
            or first_hash != second_hash
            or first_hash != _canonical_sha256(stage)
        ):
            raise ValueError(f"D3 independent replay mismatch: {stage_id}")
        records = stage.get("records")
        record_count = len(records) if isinstance(records, list) else -1
        result[pipeline_id] = {
            "schema_version": REPLAY_MANIFEST_SCHEMA,
            "profile_id": PROFILE_ID,
            "release_id": RELEASE_ID,
            "pipeline_id": pipeline_id,
            "stage_id": stage_id,
            "status": "succeeded" if open_anomalies == 0 else "needs_review",
            "raw_root_sha256": raw_manifest["root_sha256"],
            "runtime_seed_sha256": runtime_seed_sha256,
            "runtime_stage": reference,
            "stage_payload_sha256": first_hash,
            "implementation_sha256": stage.get("implementation_sha256"),
            "input_payload_sha256": stage.get("input_payload_sha256"),
            "record_count": record_count,
            "original_values_sha256": stage.get("original_values_sha256"),
            "source_fields_sha256": stage.get("source_fields_sha256"),
            "quality_observations": quality,
            "open_anomaly_count": open_anomalies,
            "independent_replay_sha256": {
                "first": first_hash,
                "second": second_hash,
            },
        }
    return result


def build_basic60_ready_seed(
    candidate_seed: dict[str, Any],
    raw_manifest: dict[str, Any],
    source_registry: dict[str, Any],
    acceptance_bundle: dict[str, Any],
    *,
    release_bundle_sha256: str,
    validation_report_sha256: str,
    reviewed_at: str,
) -> dict[str, Any]:
    """Finalize the fail-closed candidate into the strict API seed only after validation."""

    if not SHA256_PATTERN.fullmatch(release_bundle_sha256) or not SHA256_PATTERN.fullmatch(
        validation_report_sha256
    ):
        raise ValueError("Ready seed requires release-bundle and validation-report SHA-256 values")
    if not _aware_timestamp(reviewed_at):
        raise ValueError("Ready seed reviewed_at must be timezone-aware")
    countries = candidate_seed.get("countries")
    sources = source_registry.get("sources")
    domains = acceptance_bundle.get("source_admission", {}).get("domains")
    if (
        not isinstance(countries, list)
        or not isinstance(sources, list)
        or not isinstance(domains, list)
    ):
        raise ValueError("Candidate countries and approved source admission are required")

    source_index = {str(item.get("source_ref")): item for item in sources if isinstance(item, dict)}
    domain_key = {
        "country_profile": "identity",
        "macroeconomic": "macro",
        "energy": "energy",
    }
    primary_refs: dict[str, str] = {}
    for item in domains:
        if not isinstance(item, dict):
            continue
        normalized = domain_key.get(str(item.get("domain_id")))
        source_ref = str(item.get("primary_source_id") or "")
        if normalized and source_ref in source_index:
            primary_refs[normalized] = source_ref
    if set(primary_refs) != {"identity", "macro", "energy"}:
        raise ValueError("Approved primary sources are incomplete")

    scoped_files = {
        str(item.get("path")): item
        for item in raw_manifest.get("scoped_files", [])
        if isinstance(item, dict)
    }
    profile_file = scoped_files.get(PROFILE_CSV.as_posix())
    metric_file = scoped_files.get(METRIC_CSV.as_posix())
    if not isinstance(profile_file, dict) or not isinstance(metric_file, dict):
        raise ValueError("Raw manifest is missing Basic60 scoped file hashes")

    raw_records: list[dict[str, Any]] = []
    provenance_by_domain: dict[str, dict[str, str]] = {}
    file_by_domain = {
        "identity": (PROFILE_CSV, profile_file),
        "macro": (METRIC_CSV, metric_file),
        "energy": (METRIC_CSV, metric_file),
    }
    for domain, source_ref in primary_refs.items():
        source = source_index[source_ref]
        snapshots = source.get("snapshots")
        if not isinstance(snapshots, list) or not snapshots or not isinstance(snapshots[0], dict):
            raise ValueError(f"Primary source has no approved snapshot: {source_ref}")
        snapshot_ref = str(snapshots[0].get("snapshot_ref") or "")
        relative_path, file_item = file_by_domain[domain]
        raw_record_ref = f"RAW-BASIC60-{domain.upper()}"
        raw_records.append(
            {
                "record_ref": raw_record_ref,
                "source_ref": source_ref,
                "source_snapshot_ref": snapshot_ref,
                "country_code": None,
                "object_key": relative_path.as_posix(),
                "payload_sha256": file_item["sha256"],
                "media_type": "text/csv",
            }
        )
        provenance_by_domain[domain] = _provenance_fields(
            source_ref,
            snapshot_ref,
            raw_record_ref,
        )

    policy_reference = acceptance_bundle.get("ai_usage_policy")
    seed_reference = acceptance_bundle.get("runtime_seed")
    amendment_reference = acceptance_bundle.get("usage_amendment")
    if not (
        isinstance(policy_reference, dict)
        and set(policy_reference) == {"path", "sha256"}
        and SHA256_PATTERN.fullmatch(str(policy_reference.get("sha256") or ""))
        and isinstance(seed_reference, dict)
        and set(seed_reference) == {"path", "sha256"}
        and SHA256_PATTERN.fullmatch(str(seed_reference.get("sha256") or ""))
        and isinstance(amendment_reference, dict)
        and amendment_reference.get("path") == USAGE_AMENDMENT_PATH.as_posix()
        and SHA256_PATTERN.fullmatch(str(amendment_reference.get("sha256") or ""))
    ):
        raise ValueError("Ready seed requires exact A1, AI-policy, and candidate-seed bindings")
    field_paths = _normalized_data_field_paths(candidate_seed)
    manual_usage_authorization = _manual_usage_authorization(field_paths)

    collected_dates = [str(item.get("collected_at") or "") for item in countries]
    source_cutoff = max(value for value in collected_dates if value)
    ready_countries: list[dict[str, Any]] = []
    for country in countries:
        if not isinstance(country, dict):
            raise ValueError("Candidate country must be an object")
        identity_provenance = provenance_by_domain["identity"]
        localized_texts = []
        for item in country.get("localized_texts", []):
            if not isinstance(item, dict):
                continue
            text = str(item.get("text") or "")
            field_code = str(item.get("field_code") or "")
            translation_status = "reviewed" if field_code == "short_name" else "source"
            localized_texts.append(
                {
                    "field_code": field_code,
                    "locale": item.get("locale"),
                    "text": text,
                    "preferred": bool(item.get("preferred")),
                    "translation_status": translation_status,
                    "source_text_sha256": hashlib.sha256(text.encode()).hexdigest(),
                    "valid_from": None,
                    "valid_to": None,
                    **identity_provenance,
                }
            )
        capitals = [
            {
                "name": item["name"],
                "role": item["role"],
                "display_order": item["display_order"],
                "valid_from": None,
                "valid_to": None,
                **identity_provenance,
            }
            for item in country.get("capitals", [])
            if isinstance(item, dict)
        ]
        languages = [
            {
                "code": f"und-{index:02d}",
                "name_en": name,
                "name_local": None,
                "status": "official",
                **identity_provenance,
            }
            for index, name in enumerate(country.get("languages", []), start=1)
        ]
        currency = country.get("currency")
        currencies = (
            [
                {
                    "code": currency["code"],
                    "name_en": currency["name"],
                    "legal_tender": True,
                    "valid_from": None,
                    "valid_to": None,
                    **identity_provenance,
                }
            ]
            if isinstance(currency, dict)
            else []
        )
        timezones = [
            {
                "iana_code": value,
                "primary": index == 0,
                **identity_provenance,
            }
            for index, value in enumerate(country.get("time_zones", []))
        ]
        admin = country.get("admin_structure")
        admin_structures = (
            [
                {
                    "admin_level": admin["level"],
                    "unit_type": admin["unit_type"],
                    "unit_count": admin["count"],
                    "as_of_year": int(str(country["collected_at"])[:4]),
                    "status": "reviewed",
                    **identity_provenance,
                }
            ]
            if isinstance(admin, dict)
            else []
        )
        metrics = []
        for metric in country.get("metrics", []):
            if not isinstance(metric, dict):
                continue
            metric_code = str(metric.get("metric_code") or "")
            definition = METRIC_LABELS.get(metric_code)
            if definition is None:
                raise ValueError(f"Unknown Basic60 metric code: {metric_code}")
            domain = definition[2]
            start, end, label = _period_bounds(metric.get("period"), source_cutoff)
            available = metric.get("value_status") == "available"
            metrics.append(
                {
                    "metric_code": metric_code,
                    "period_start": start,
                    "period_end": end,
                    "period_label": label,
                    "original_value": metric.get("original_value"),
                    "original_unit": metric.get("unit"),
                    "normalized_value": metric.get("normalized_value") if available else None,
                    "normalized_unit": metric.get("unit"),
                    "value_status": "available" if available else "pending",
                    "null_reason": None if available else metric.get("null_reason"),
                    "quality_status": "reviewed",
                    "freshness_status": "current" if available else "pending",
                    "reviewed_at": reviewed_at,
                    **provenance_by_domain[domain],
                }
            )
        ready_countries.append(
            {
                "iso3": country["iso3"],
                "iso2": country["iso2"],
                "region_code": country["region_code"],
                "coverage_status": "covered",
                "coverage_level": "Basic",
                "opportunity_level": "pending",
                "policy_friendliness_level": "pending",
                "risk_assessment_status": "unknown",
                "risk_level": None,
                "last_reviewed_at": reviewed_at,
                "localized_texts": localized_texts,
                "capitals": capitals,
                "languages": languages,
                "currencies": currencies,
                "timezones": timezones,
                "admin_structures": admin_structures,
                "metrics": metrics,
            }
        )

    metric_definitions = [
        {
            "metric_code": metric_code,
            "label_zh": values[0],
            "label_en": values[1],
            "data_domain": values[2],
            "canonical_unit": values[3],
            "period_type": "latest" if values[2] == "energy" else "year",
            "display_precision": values[4],
            "normalization_rule": "Preserve source value; normalize only the declared unit",
        }
        for metric_code, values in METRIC_LABELS.items()
    ]
    return {
        "schema_version": "basic60.seed.v1",
        "release": {
            "release_id": RELEASE_ID,
            "release_profile": PROFILE_ID,
            "formal_gate_status": FORMAL_GATE_STATUS,
            "status": "private_trial_ready",
            "as_of": source_cutoff,
            "source_cutoff": source_cutoff,
            "release_bundle_sha256": release_bundle_sha256,
            "validation_report_sha256": validation_report_sha256,
            "counts": {
                "countries": EXPECTED_COUNTRIES,
                "macro_rows": EXPECTED_MACRO_ROWS,
                "energy_rows": EXPECTED_ENERGY_ROWS,
                "available_metric_values": EXPECTED_AVAILABLE_VALUES,
                "pending_metric_values": EXPECTED_PENDING_VALUES,
            },
        },
        "sources": [
            {
                key: value
                for key, value in source.items()
                if key
                not in {
                    "admission_fingerprint",
                    "evidence",
                    "license_basis_sha256s",
                    *SOURCE_PERMISSION_FIELDS,
                }
            }
            for source in sources
            if isinstance(source, dict)
        ],
        "manual_usage_authorization": manual_usage_authorization,
        "ai_usage_policy": {
            "projection_type": "machine_generated_all_field_scope_projection",
            "permission_basis": USAGE_AMENDMENT_ID,
            "policy_reference": policy_reference,
            "seed_sha256": seed_reference["sha256"],
            "field_count": manual_usage_authorization["field_scope"]["field_count"],
            "fields_sha256": manual_usage_authorization["field_scope"]["fields_sha256"],
        },
        "raw_records": raw_records,
        "metric_definitions": metric_definitions,
        "countries": ready_countries,
    }


def build_basic60_excel_review_payload(
    raw_dir: Path,
    seed: dict[str, Any],
    *,
    seed_sha256: str,
    source_uri_manifest: dict[str, Any],
    source_uri_manifest_sha256: str,
) -> dict[str, Any]:
    """Build the canonical data-and-source subject shown in the one review workbook.

    Formatting and reviewer input are intentionally excluded. The payload covers the
    exact raw rows that feed Basic60 plus the normalized source bindings displayed in
    the workbook, so a later workbook approval cannot be reused after data or source
    drift.
    """

    if not SHA256_PATTERN.fullmatch(seed_sha256) or not SHA256_PATTERN.fullmatch(
        source_uri_manifest_sha256
    ):
        raise ValueError("Excel review payload requires seed and source-manifest SHA-256 values")
    if (
        seed.get("release_id") != RELEASE_ID
        or seed.get("profile_id") != PROFILE_ID
        or seed.get("formal_gate_status") != FORMAL_GATE_STATUS
    ):
        raise ValueError("Excel review payload requires the Basic60 private candidate seed")
    expected_source_manifest = build_basic60_country_source_uri_manifest(raw_dir)
    if source_uri_manifest != expected_source_manifest:
        raise ValueError("Country source URI manifest drifted from the current raw package")

    summary, raw_checks = inspect_basic60_raw(raw_dir)
    if raw_checks:
        raise ValueError("Excel review payload requires a clean Basic60 raw package")
    countries = sorted(
        _csv_rows(raw_dir / PROFILE_CSV),
        key=lambda item: str(item.get("iso3") or ""),
    )
    metric_rows = _csv_rows(raw_dir / METRIC_CSV)
    macro_rows = sorted(
        (item for item in metric_rows if item.get("record_type") == "macro_annual"),
        key=lambda item: (str(item.get("iso3") or ""), str(item.get("year") or "")),
    )
    energy_rows = sorted(
        (item for item in metric_rows if item.get("record_type") == "energy_latest"),
        key=lambda item: str(item.get("iso3") or ""),
    )

    country_entries = source_uri_manifest.get("entries")
    if not isinstance(country_entries, list):
        raise ValueError("Country source URI manifest entries are required")
    normalized_country_entries = sorted(
        (
            {
                "scope": "country_profile",
                "iso3": str(item.get("iso3") or ""),
                "provider": str(item.get("provider") or ""),
                "dataset": str(item.get("provider") or ""),
                "fields": sorted(str(value) for value in item.get("fields", [])),
                "url": str(item.get("url") or ""),
            }
            for item in country_entries
            if isinstance(item, dict)
        ),
        key=lambda item: (
            item["iso3"],
            item["provider"],
            item["fields"],
            item["url"],
        ),
    )
    country_sequence: defaultdict[str, int] = defaultdict(int)
    source_bindings: list[dict[str, Any]] = []
    for country_binding in normalized_country_entries:
        iso3 = str(country_binding["iso3"])
        country_sequence[iso3] += 1
        source_bindings.append(
            {
                "source_detail_id": f"B60-COUNTRY-{iso3}-{country_sequence[iso3]:02d}",
                **country_binding,
            }
        )

    macro_fields = sorted({source_field for _metric_code, source_field, _unit in MACRO_METRICS})
    for macro_row in macro_rows:
        iso3 = str(macro_row.get("iso3") or "")
        if any(binding["source_detail_id"] == f"B60-MACRO-{iso3}" for binding in source_bindings):
            continue
        source_bindings.append(
            {
                "source_detail_id": f"B60-MACRO-{iso3}",
                "scope": "macroeconomic",
                "iso3": iso3,
                "provider": "World Bank",
                "dataset": "World Development Indicators",
                "fields": macro_fields,
                "url": str(macro_row.get("macro_source_url") or ""),
            }
        )

    energy_urls = sorted(
        {
            str(item.get("energy_source_url") or "")
            for item in energy_rows
            if item.get("energy_source_url")
        }
    )
    if len(energy_urls) != 1:
        raise ValueError("Basic60 Excel review requires one shared IRENA energy source URI")
    source_bindings.append(
        {
            "source_detail_id": "B60-ENERGY-IRENA",
            "scope": "energy",
            "iso3": "ALL",
            "provider": "IRENA",
            "dataset": "IRENASTAT",
            "fields": sorted(
                {source_field for _metric_code, source_field, _unit, _year_field in ENERGY_METRICS}
            ),
            "url": energy_urls[0],
        }
    )
    source_bindings.sort(
        key=lambda item: (
            str(item["scope"]),
            str(item["iso3"]),
            str(item["provider"]),
            str(item["source_detail_id"]),
        )
    )
    return {
        "schema_version": EXCEL_REVIEW_PAYLOAD_SCHEMA,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "raw_root_sha256": seed.get("source_root_sha256"),
        "seed_sha256": seed_sha256,
        "source_uri_manifest_sha256": source_uri_manifest_sha256,
        "source_uri_entries_sha256": source_uri_manifest.get("entries_sha256"),
        "counts": {
            **summary,
            "source_binding_count": len(source_bindings),
        },
        "countries": countries,
        "macro_rows": macro_rows,
        "energy_rows": energy_rows,
        "source_bindings": source_bindings,
    }


def build_basic60_excel_review_gate(
    paths: RepositoryPaths,
    raw_dir: Path,
    seed_path: Path,
    source_uri_manifest_path: Path,
    workbook_path: Path,
) -> dict[str, Any]:
    """Bind one generated workbook to the complete Basic60 data/source payload."""

    resolved_workbook = workbook_path.resolve(strict=True)
    output_root = (paths.root / EXCEL_REVIEW_OUTPUT_PREFIX).resolve()
    if (
        workbook_path.is_symlink()
        or not resolved_workbook.is_file()
        or not resolved_workbook.is_relative_to(output_root)
        or resolved_workbook.suffix.casefold() != ".xlsx"
    ):
        raise ValueError("Basic60 review workbook must be a safe .xlsx under outputs/basic60")
    seed = _json_object(seed_path)
    source_uri_manifest = _json_object(source_uri_manifest_path)
    payload = build_basic60_excel_review_payload(
        raw_dir,
        seed,
        seed_sha256=sha256_file(seed_path),
        source_uri_manifest=source_uri_manifest,
        source_uri_manifest_sha256=sha256_file(source_uri_manifest_path),
    )
    summary = payload["counts"]
    return {
        "schema_version": EXCEL_REVIEW_GATE_SCHEMA,
        "review_mode": EXCEL_REVIEW_MODE,
        "review_scope": EXCEL_REVIEW_SCOPE,
        "workbook": _reference(resolved_workbook, paths),
        "canonical_payload_sha256": _canonical_sha256(payload),
        "raw_root_sha256": payload["raw_root_sha256"],
        "runtime_seed": _reference(seed_path, paths),
        "source_uri_manifest": _reference(source_uri_manifest_path, paths),
        "machine_counts": {
            "country_count": summary["country_count"],
            "macro_annual_record_count": summary["macro_annual_record_count"],
            "energy_latest_record_count": summary["energy_latest_record_count"],
            "available_observation_count": summary["available_observation_count"],
            "pending_observation_count": summary["pending_observation_count"],
            "source_binding_count": summary["source_binding_count"],
        },
        "approval": None,
        "formal_gate_status": FORMAL_GATE_STATUS,
        "does_not_complete_formal_d1_d4": True,
        "does_not_authorize_p0_or_production": True,
    }


def prepare_basic60_excel_review_bundle(
    paths: RepositoryPaths,
    raw_dir: Path,
    seed_path: Path,
    source_uri_manifest_path: Path,
    input_bundle_path: Path,
    workbook_path: Path,
    output_bundle_path: Path,
) -> tuple[Path, dict[str, Any]]:
    """Create a new, non-overwriting single-workbook private-review packet.

    PBD/A1 and machine provenance are retained from the input packet. Legacy D2 batch
    close and D4 sample/signature sections are omitted because the frozen workbook and
    one project-approver decision replace those private-trial review gates.
    """

    review_root = (paths.root / "data/basic60/review").resolve()
    resolved_output = output_bundle_path.resolve(strict=False)
    if not resolved_output.is_relative_to(review_root):
        raise ValueError("Single-workbook review bundle must remain under data/basic60/review")
    if output_bundle_path.exists():
        raise ValueError(f"Refusing to overwrite existing review bundle: {output_bundle_path}")
    bundle = _json_object(input_bundle_path)
    source = bundle.get("source_admission")
    if not isinstance(source, dict):
        raise ValueError("Input packet requires Basic60 source admission metadata")
    retained_source_keys = {
        "source_registry",
        "ai_usage_policy",
        "country_domain_matrix",
        "derived_matrix_sha256",
        "domains",
        "usage_boundaries",
    }
    bundle["source_admission"] = {
        key: value for key, value in source.items() if key in retained_source_keys
    }
    bundle.pop("collection", None)
    bundle.pop("acceptance", None)
    bundle["single_excel_review"] = build_basic60_excel_review_gate(
        paths,
        raw_dir,
        seed_path,
        source_uri_manifest_path,
        workbook_path,
    )
    _write_json(output_bundle_path, bundle)
    return output_bundle_path, bundle


def build_profile() -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "delivery_mode": "access_controlled_private_trial",
        "formal_gate_status": FORMAL_GATE_STATUS,
        "allowed_outcomes": sorted(PRIVATE_OUTCOMES),
        "does_not_complete_formal_d1_d4": True,
        "does_not_authorize_p0_or_production": True,
        "scope": {
            "countries": EXPECTED_COUNTRIES,
            "coverage_level": "Basic",
            "future_deep_dive_priority_country_codes": list(
                FUTURE_DEEP_DIVE_PRIORITY_COUNTRY_CODES
            ),
            "included_domains": list(DOMAIN_IDS),
            "excluded_domains": [
                "policy",
                "risk",
                "project",
                "tender",
                "partner",
                "search",
                "ai",
            ],
            "read_only_routes": [
                "GET /api/v1/countries",
                "GET /api/v1/countries/{code}",
            ],
        },
        "expected_counts": {
            "countries": EXPECTED_COUNTRIES,
            "macro_annual_records": EXPECTED_MACRO_ROWS,
            "energy_latest_records": EXPECTED_ENERGY_ROWS,
            "available_observations": EXPECTED_AVAILABLE_VALUES,
            "pending_observations": EXPECTED_PENDING_VALUES,
        },
        "owner_data_usage_authorization": _owner_usage_authorization(),
        "disabled_v1_runtime_capabilities": _v1_runtime_capabilities(),
        "unimplemented_v1_features": {
            "policy_publication": True,
            "policy_indexing": True,
            "ai_processing": True,
            "model_training": True,
            "cloud_processing": True,
            "vector_database": True,
            "full_text_search": True,
        },
    }


def build_source_registry_template(
    raw_dir: Path,
    country_source_uri_manifest: dict[str, str],
) -> dict[str, Any]:
    current_country_manifest = build_basic60_country_source_uri_manifest(raw_dir)
    sources: list[dict[str, Any]] = [
        {
            "source_ref": None,
            "provider": None,
            "dataset": None,
            "data_domain": domain,
            "source_role": "primary",
            "status": "under_review",
            "terms_uri": None,
            "evidence_sha256": None,
            "evidence": None,
            "access_method": None,
            "rate_limit": None,
            "snapshots": [],
        }
        for domain in ("country_identity", "macro", "energy")
    ]
    for source in sources:
        source["admission_fingerprint"] = _source_admission_fingerprint(
            raw_dir,
            source,
            country_source_manifest=current_country_manifest,
            terms_evidence=None,
        )
    return {
        "schema_version": 1,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "template_only": True,
        "source_declaration": "official_public_data",
        "update_policy": {
            "mode": "manual_incremental",
            "reapproval_triggers": list(D1_REAPPROVAL_TRIGGERS),
        },
        "country_source_uri_manifest": country_source_uri_manifest,
        "sources": sources,
    }


def build_human_approval_template() -> dict[str, Any]:
    return {
        "schema_version": 1,
        "record_type": "basic60_human_approval",
        "evidence_id": None,
        "decision": "pending",
        "person_name": None,
        "role": None,
        "signed_at": None,
        "subject_sha256": None,
        "template_only": True,
        "warning": "Only the actual named reviewer may complete this hash-bound evidence copy.",
    }


def build_d1_scope_authorization_template() -> dict[str, Any]:
    """Return a truthful D1 human authorization template.

    The authorization binds only the stable source scope.  Exact registry,
    evidence, terms-snapshot, URI, and schema hashes are bound later by the
    machine-content binding and are deliberately not represented as having
    been signed by the human approver.
    """

    return {
        "schema_version": SCHEMA_VERSION,
        "record_type": D1_SCOPE_AUTHORIZATION_RECORD_TYPE,
        "evidence_id": None,
        "authorization_type": "human_scope_authorization",
        "decision": "pending",
        "person_name": None,
        "role": "项目批准人",
        "signed_at": None,
        "transcribed_at": None,
        "authorization_basis": None,
        "scope_authorization_id": None,
        "scope_subject_sha256": None,
        "exact_content_hash_signed": False,
        "template_only": True,
        "warning": (
            "This records a human authorization of stable scope only. Exact content hashes "
            "must be added later by a machine binding and must never be described as signed "
            "at the human authorization time."
        ),
    }


def build_usage_amendment_template() -> dict[str, Any]:
    """Return the unsigned A1 owner-authorization template.

    A1 changes data-use permission only. It deliberately leaves every V1 AI,
    search, vector, and model runtime capability disabled.
    """

    return {
        "schema_version": USAGE_AMENDMENT_SCHEMA,
        "record_type": "basic60_private_usage_amendment",
        "amendment_id": USAGE_AMENDMENT_ID,
        "amends_decision_id": DECISION_ID,
        "status": "pending_project_approver",
        "effective_from": None,
        "owner_authorization": _owner_usage_authorization(),
        "source_label_policy": {
            "ordinary_ui_api": "omit_source_labels_and_provenance",
            "backend_terms_and_provenance": "retain_for_audit_only",
        },
        "processing_controls": {
            "private_internal_purpose_only": True,
            "approved_model_configuration_required": True,
            "provider_training_allowed": False,
            "provider_retention_allowed": False,
            "raw_terms_or_source_evidence_allowed_in_model_payload": False,
        },
        "v1_runtime_capabilities": _v1_runtime_capabilities(),
        "does_not_complete": [
            "D2-P",
            "D4-P",
            "formal_D1_D4",
            "P0",
            "private_trial_ready",
            "production_release",
        ],
        "approval": {
            "person_name": None,
            "role": "项目批准人",
            "decision": "pending",
            "signed_at": None,
            "authorization_basis": None,
            "evidence": None,
        },
        "template_only": True,
    }


def build_d1_authorized_scope(
    bundle: dict[str, Any],
    source_registry: dict[str, Any],
) -> dict[str, Any]:
    """Project the stable, human-authorized portion of D1 source admission."""

    source = bundle.get("source_admission")
    if not isinstance(source, dict):
        raise ValueError("Basic60 source admission is required")
    authorization_id = str(source.get("scope_authorization_id") or "")
    if authorization_id != REQUIRED_D1_SCOPE_AUTHORIZATION_ID:
        raise ValueError(
            "Basic60 usage-boundary amendment requires the new D1 R2 scope authorization"
        )
    domains = source.get("domains")
    registry_sources = source_registry.get("sources")
    if not isinstance(domains, list) or not isinstance(registry_sources, list):
        raise ValueError("Basic60 domains and source registry entries are required")
    indexed = {
        str(item.get("source_ref")): item
        for item in registry_sources
        if isinstance(item, dict) and str(item.get("source_ref") or "")
    }
    domain_map = {
        "country_profile": "country_identity",
        "macroeconomic": "macro",
        "energy": "energy",
    }
    primary_sources: list[dict[str, Any]] = []
    for domain_id in DOMAIN_IDS:
        domain = next(
            (
                item
                for item in domains
                if isinstance(item, dict) and item.get("domain_id") == domain_id
            ),
            None,
        )
        if domain is None:
            raise ValueError(f"Missing Basic60 domain: {domain_id}")
        source_ref = str(domain.get("primary_source_id") or "")
        registry_source = indexed.get(source_ref)
        if registry_source is None:
            raise ValueError(f"Missing Basic60 primary source: {source_ref or domain_id}")
        primary_sources.append(
            {
                "domain_id": domain_id,
                "data_domain": domain_map[domain_id],
                "source_ref": source_ref,
                "provider": registry_source.get("provider"),
                "dataset": registry_source.get("dataset"),
                "terms_uri": registry_source.get("terms_uri"),
                "source_uri_control": "exact_machine_content_binding",
                "schema_control": "exact_machine_content_binding",
                "terms_snapshot_control": "exact_machine_content_binding",
            }
        )
    return {
        "schema_version": D1_AUTHORIZED_SCOPE_SCHEMA,
        "release_id": bundle.get("release_id"),
        "profile_id": bundle.get("profile_id"),
        "baseline_clarification": bundle.get("baseline_clarification"),
        "usage_amendment": bundle.get("usage_amendment"),
        "source_declaration": "official_public_data",
        "scope_authorization_id": authorization_id,
        "review_unit": "provider_dataset",
        "primary_source_count": 3,
        "alternative_sources_required": False,
        "primary_sources": primary_sources,
        "usage_boundaries": source.get("usage_boundaries"),
        "manual_update_policy": {
            "mode": "manual_incremental",
            "routine_value_updates_reuse_d1_admission": True,
            "d1_reapproval_triggers": list(D1_REAPPROVAL_TRIGGERS),
        },
        "machine_baseline_policy": {
            "mode": "single_use_append_only",
            "authorization_anchor_key": (
                D1_MACHINE_BASELINE_ROOT / "authorizations" / f"{authorization_id}.json"
            ).as_posix(),
            "content_object_mode": "content_addressed_read_only",
            "new_authorization_required_for": list(D1_REAPPROVAL_TRIGGERS),
        },
        "authorization_statement": (
            "Human authorization covers the stable source scope and A1 all-field manual data-"
            "usage decision while V1 AI runtime capabilities remain disabled; exact content "
            "hashes are machine-bound after capture."
        ),
    }


def _reference(path: Path, paths: RepositoryPaths) -> dict[str, str]:
    return {"path": _display_path(paths, path), "sha256": sha256_file(path)}


def _approval_slot(role: str) -> dict[str, Any]:
    return {
        "role": role,
        "person_name": None,
        "decision": "pending",
        "signed_at": None,
        "subject_sha256": None,
        "evidence": None,
    }


def _d1_scope_approval_slot() -> dict[str, Any]:
    return {
        "approval_type": "human_scope_authorization",
        "role": "项目批准人",
        "person_name": None,
        "decision": "pending",
        "signed_at": None,
        "scope_authorization_id": None,
        "scope_subject_sha256": None,
        "exact_content_hash_signed": False,
        "evidence": None,
    }


def build_acceptance_template(
    paths: RepositoryPaths,
    raw_dir: Path,
    profile_path: Path,
    ai_usage_policy_path: Path,
    source_matrix_path: Path,
    manifest_path: Path,
    raw_manifest: dict[str, Any],
    seed_path: Path,
    sample_candidate_path: Path,
    sample_review_template_paths: tuple[Path, Path],
    item_review_template_path: Path,
    replay_manifest_paths: dict[str, Path],
    summary: dict[str, Any],
) -> dict[str, Any]:
    scoped = {str(item["path"]): item for item in raw_manifest["scoped_files"]}
    batches = []
    for batch_id, domain_id, relative_path in BATCH_SPECS:
        item = scoped[relative_path.as_posix()]
        batches.append(
            {
                "batch_id": batch_id,
                "domain_id": domain_id,
                "dataset_path": relative_path.as_posix(),
                "dataset_sha256": item["sha256"],
                "run_id": "legacy_import_run",
                "run_status": "pending",
                "access_signal": "none",
                "input_immutable": True,
                "diff_status": "pending",
                "output_manifest": None,
                "status": "pending",
            }
        )
    replay_payloads = {
        pipeline_id: _json_object(replay_manifest_paths[pipeline_id])
        for pipeline_id in PIPELINE_IDS
    }
    processing_open_anomalies = sum(
        int(replay_payloads[pipeline_id].get("open_anomaly_count") or 0)
        for pipeline_id in PIPELINE_IDS
    )
    return {
        "schema_version": SCHEMA_VERSION,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "formal_gate_status": FORMAL_GATE_STATUS,
        "template_only": True,
        "profile": _reference(profile_path, paths),
        "raw_manifest": _reference(manifest_path, paths),
        "raw_root_sha256": raw_manifest["root_sha256"],
        "runtime_seed": _reference(seed_path, paths),
        "baseline_candidate": _reference(paths.root / PBD_CANDIDATE_PATH, paths),
        "baseline_confirmation_template": _reference(
            paths.root / PBD_CONFIRMATION_TEMPLATE_PATH,
            paths,
        ),
        "baseline_confirmation": None,
        "baseline_decision": None,
        "baseline_clarification": _reference(
            paths.root / PBD_CLARIFICATION_PATH,
            paths,
        ),
        "usage_amendment": None,
        "ai_usage_policy": _reference(ai_usage_policy_path, paths),
        "source_admission": {
            "source_registry": None,
            "ai_usage_policy": _reference(ai_usage_policy_path, paths),
            "scope_authorization_id": None,
            "authorized_scope": None,
            "machine_content_binding": None,
            "country_domain_matrix": _reference(source_matrix_path, paths),
            "derived_matrix_sha256": None,
            "domains": [
                {
                    "domain_id": domain_id,
                    "primary_source_id": None,
                    "alternative_source_id": None,
                    "status": "pending",
                }
                for domain_id in DOMAIN_IDS
            ],
            "usage_boundaries": _usage_boundaries(),
            "approvals": [_d1_scope_approval_slot()],
        },
        "collection": {
            "input_directory": _display_path(paths, raw_dir),
            "input_immutable": True,
            "batches": batches,
            "incidents": [],
            "approval": _approval_slot(BATCH_APPROVAL_ROLE),
        },
        "processing": {
            "pipelines": [
                {
                    "pipeline_id": pipeline_id,
                    "status": replay_payloads[pipeline_id].get("status"),
                    "replay_manifest": _reference(replay_manifest_paths[pipeline_id], paths),
                }
                for pipeline_id in PIPELINE_IDS
            ],
            "original_values_preserved": True,
            "source_text_preserved": True,
            "open_anomaly_count": processing_open_anomalies,
            "open_high_risk_anomaly_count": processing_open_anomalies,
        },
        "acceptance": {
            "machine_counts": summary,
            "sample_candidate": _reference(sample_candidate_path, paths),
            "review_templates": {
                "sample_reviews": [
                    _reference(path, paths) for path in sample_review_template_paths
                ],
                "item_review": _reference(item_review_template_path, paths),
            },
            "chinese_names_reviewed": 0,
            "sample_size": EXPECTED_SAMPLE_SIZE,
            "sample_accurate_count": None,
            "duplicate_count": None,
            "pending_values_reviewed": 0,
            "derived_values_reviewed": 0,
            "unresolved_p0_data_issue_count": None,
            "artifacts": {
                "sample_reviews": [None, None],
                "item_review": None,
                "machine_evidence": None,
            },
            "sample_approvals": [
                _approval_slot("数据质量复核人"),
                _approval_slot("数据质量复核人"),
            ],
            "final_approvals": [_approval_slot(role) for role in FINAL_APPROVAL_ROLES],
        },
        "revocation": None,
    }


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.tmp")
    temporary.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _absolute_without_resolving(path: Path) -> Path:
    return path if path.is_absolute() else Path.cwd() / path


def _has_symlink_component(path: Path) -> bool:
    absolute = _absolute_without_resolving(path)
    current = Path(absolute.anchor)
    for part in absolute.parts[1:]:
        current /= part
        if current.is_symlink():
            return True
    return False


def _paths_overlap(first: Path, second: Path) -> bool:
    return first == second or first.is_relative_to(second) or second.is_relative_to(first)


def _validated_l0_directory(paths: RepositoryPaths, raw_dir: Path, l0_dir: Path) -> Path:
    if not l0_dir.is_absolute():
        raise ValueError("B60_L0_TARGET_INVALID: --l0-dir must be an absolute path")
    if _has_symlink_component(l0_dir):
        raise ValueError("B60_L0_TARGET_SYMLINK: --l0-dir cannot contain symbolic links")
    resolved = l0_dir.resolve(strict=False)
    if len(resolved.parts) <= 2 or resolved == Path.home().resolve():
        raise ValueError("B60_L0_TARGET_BROAD: --l0-dir cannot be a broad filesystem root")
    repository_root = paths.root.resolve()
    raw_root = raw_dir.resolve(strict=True)
    if _paths_overlap(resolved, repository_root) or _paths_overlap(resolved, raw_root):
        raise ValueError(
            "B60_L0_TARGET_OVERLAP: --l0-dir must be outside both the repository and raw input"
        )
    if resolved.exists() and not resolved.is_dir():
        raise ValueError("B60_L0_TARGET_INVALID: --l0-dir must be a directory")
    object_root = resolved / "sha256"
    if object_root.is_symlink() or (object_root.exists() and not object_root.is_dir()):
        raise ValueError("B60_L0_TARGET_INVALID: sha256 object root is unsafe")
    return resolved


def _validated_evidence_directory(paths: RepositoryPaths, evidence_dir: Path) -> Path:
    if not evidence_dir.is_absolute():
        raise ValueError("B60_L0_EVIDENCE_DIR_INVALID: --evidence-dir must be an absolute path")
    if _has_symlink_component(evidence_dir):
        raise ValueError(
            "B60_L0_EVIDENCE_DIR_SYMLINK: --evidence-dir cannot contain symbolic links"
        )
    resolved = evidence_dir.resolve(strict=False)
    approved_root = (paths.root / "data/basic60/evidence").resolve(strict=False)
    if not resolved.is_relative_to(approved_root):
        raise ValueError(
            "B60_L0_EVIDENCE_DIR_INVALID: evidence must remain under data/basic60/evidence"
        )
    if resolved.exists() and not resolved.is_dir():
        raise ValueError("B60_L0_EVIDENCE_DIR_INVALID: --evidence-dir must be a directory")
    return resolved


def _verify_l0_object(path: Path, expected_hash: str, expected_size: int) -> None:
    if path.is_symlink() or not path.is_file():
        raise ValueError(f"B60_L0_OBJECT_UNSAFE: content object is not a regular file: {path}")
    actual = path.stat()
    if actual.st_size != expected_size or sha256_file(path) != expected_hash:
        raise ValueError(f"B60_L0_OBJECT_MISMATCH: existing object does not match source: {path}")


def _materialize_content_object(
    source: Path,
    destination: Path,
    *,
    expected_hash: str,
    expected_size: int,
) -> bool:
    """Return True when copied, False when an existing or raced object was reused."""

    if destination.exists() or destination.is_symlink():
        _verify_l0_object(destination, expected_hash, expected_size)
        destination.chmod(0o444)
        if stat.S_IMODE(destination.stat().st_mode) != 0o444:
            raise ValueError(f"B60_L0_OBJECT_NOT_READ_ONLY: cannot set 0444: {destination}")
        return False

    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{expected_hash}.",
        suffix=".tmp",
        dir=destination.parent,
    )
    temporary = Path(temporary_name)
    copied = False
    try:
        source_before = source.stat()
        digest = hashlib.sha256()
        byte_count = 0
        with source.open("rb") as source_stream, os.fdopen(file_descriptor, "wb") as target_stream:
            while chunk := source_stream.read(1024 * 1024):
                target_stream.write(chunk)
                digest.update(chunk)
                byte_count += len(chunk)
            target_stream.flush()
            os.fsync(target_stream.fileno())
        source_after = source.stat()
        if (
            source_before.st_size != source_after.st_size
            or source_before.st_mtime_ns != source_after.st_mtime_ns
            or byte_count != expected_size
            or digest.hexdigest() != expected_hash
        ):
            raise ValueError(f"B60_L0_SOURCE_DRIFT: source changed while copying: {source}")
        temporary.chmod(0o444)
        try:
            os.link(temporary, destination)
            copied = True
        except FileExistsError:
            _verify_l0_object(destination, expected_hash, expected_size)
        destination.chmod(0o444)
        _verify_l0_object(destination, expected_hash, expected_size)
        if stat.S_IMODE(destination.stat().st_mode) != 0o444:
            raise ValueError(f"B60_L0_OBJECT_NOT_READ_ONLY: cannot set 0444: {destination}")
        return copied
    finally:
        temporary.unlink(missing_ok=True)


def _atomic_write_json_no_overwrite(path: Path, payload: dict[str, Any]) -> None:
    encoded = (json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode(
        "utf-8"
    )
    file_descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.name}.",
        suffix=".tmp",
        dir=path.parent,
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(file_descriptor, "wb") as destination:
            destination.write(encoded)
            destination.flush()
            os.fsync(destination.fileno())
        try:
            os.link(temporary, path)
        except FileExistsError as error:
            raise ValueError(f"B60_L0_EVIDENCE_EXISTS: refusing to overwrite {path}") from error
    finally:
        temporary.unlink(missing_ok=True)


def prepare_basic60_private(
    paths: RepositoryPaths,
    raw_dir: Path,
    runtime_dir: Path,
    candidate_dir: Path,
) -> tuple[list[Path], dict[str, Any]]:
    summary, checks = inspect_basic60_raw(raw_dir)
    if checks:
        assessment = _assessment("not_ready", summary, checks)
        assessment_path = candidate_dir / ASSESSMENT_NAME
        _write_json(assessment_path, assessment)
        return [assessment_path], assessment

    raw_manifest = build_raw_manifest(paths, raw_dir)
    if raw_manifest["unsafe_entry_count"]:
        checks = [
            CheckResult(
                code="B60_RAW_UNSAFE_ENTRY",
                message="Raw material contains symlinks or other unsafe entries",
                location=str(raw_dir),
            )
        ]
        assessment = _assessment("not_ready", summary, checks)
        assessment_path = candidate_dir / ASSESSMENT_NAME
        _write_json(assessment_path, assessment)
        return [assessment_path], assessment

    seed = build_basic60_seed(raw_dir, raw_manifest)
    seed_path = runtime_dir / SEED_NAME
    sample_candidate_path = runtime_dir / SAMPLE_CANDIDATE_NAME
    sample_review_template_paths: tuple[Path, Path] = (
        runtime_dir / SAMPLE_REVIEW_TEMPLATE_NAMES[0],
        runtime_dir / SAMPLE_REVIEW_TEMPLATE_NAMES[1],
    )
    item_review_template_path = runtime_dir / ITEM_REVIEW_TEMPLATE_NAME
    profile_path = candidate_dir / PROFILE_NAME
    ai_usage_policy_path = candidate_dir / AI_USAGE_POLICY_NAME
    usage_amendment_template_path = candidate_dir / USAGE_AMENDMENT_TEMPLATE_NAME
    source_matrix_path = candidate_dir / SOURCE_MATRIX_NAME
    country_source_uri_manifest_path = candidate_dir / COUNTRY_SOURCE_URI_MANIFEST_NAME
    source_registry_path = candidate_dir / SOURCE_REGISTRY_TEMPLATE_NAME
    human_approval_path = candidate_dir / HUMAN_APPROVAL_TEMPLATE_NAME
    d1_scope_authorization_path = candidate_dir / D1_SCOPE_AUTHORIZATION_TEMPLATE_NAME
    manifest_path = candidate_dir / RAW_MANIFEST_NAME
    acceptance_path = candidate_dir / ACCEPTANCE_TEMPLATE_NAME
    assessment_path = candidate_dir / ASSESSMENT_NAME

    _write_json(seed_path, seed)
    replayed_seed = build_basic60_seed(raw_dir, raw_manifest)
    if _canonical_sha256(replayed_seed) != _canonical_sha256(seed):
        raise ValueError("Basic60 normalization is not deterministic across two replays")
    sample_candidate = build_basic60_sample_candidate(
        seed,
        seed_artifact_sha256=sha256_file(seed_path),
    )
    _write_json(sample_candidate_path, sample_candidate)
    for review_slot, review_path in enumerate(sample_review_template_paths, start=1):
        _write_json(
            review_path,
            build_basic60_sample_review_manifest(
                _reference(sample_candidate_path, paths),
                sample_candidate,
                review_slot=review_slot,
            ),
        )
    _write_json(
        item_review_template_path,
        build_basic60_item_review_manifest(_reference(seed_path, paths), seed),
    )
    _write_json(profile_path, build_profile())
    usage_amendment_path = paths.root / USAGE_AMENDMENT_PATH
    usage_amendment_reference = (
        _reference(usage_amendment_path, paths) if usage_amendment_path.is_file() else None
    )
    _write_json(
        ai_usage_policy_path,
        build_ai_usage_policy(
            seed,
            seed_reference=_reference(seed_path, paths),
            usage_amendment_reference=usage_amendment_reference,
        ),
    )
    _write_json(usage_amendment_template_path, build_usage_amendment_template())
    _write_json(source_matrix_path, build_basic60_source_matrix_candidate(raw_dir))
    _write_json(
        country_source_uri_manifest_path,
        build_basic60_country_source_uri_manifest(raw_dir),
    )
    _write_json(
        source_registry_path,
        build_source_registry_template(
            raw_dir,
            _reference(country_source_uri_manifest_path, paths),
        ),
    )
    _write_json(human_approval_path, build_human_approval_template())
    _write_json(d1_scope_authorization_path, build_d1_scope_authorization_template())
    raw_manifest["seed"] = _reference(seed_path, paths)
    _write_json(manifest_path, raw_manifest)
    d3_artifacts = write_d3_stage_payloads(
        seed,
        raw_manifest,
        repository_root=paths.root,
        output_directory=runtime_dir / "d3",
    )
    stage_payloads = build_d3_stage_payloads(seed, raw_manifest)
    replay_payloads = build_basic60_replay_manifests(
        raw_manifest,
        runtime_seed_sha256=sha256_file(seed_path),
        d3_artifacts=d3_artifacts,
        stage_payloads=stage_payloads,
    )
    replay_manifest_paths = {
        pipeline_id: candidate_dir / REPLAY_MANIFEST_NAMES[pipeline_id]
        for pipeline_id in PIPELINE_IDS
    }
    for pipeline_id, replay_path in replay_manifest_paths.items():
        _write_json(replay_path, replay_payloads[pipeline_id])
    acceptance = build_acceptance_template(
        paths,
        raw_dir,
        profile_path,
        ai_usage_policy_path,
        source_matrix_path,
        manifest_path,
        raw_manifest,
        seed_path,
        sample_candidate_path,
        sample_review_template_paths,
        item_review_template_path,
        replay_manifest_paths,
        summary,
    )
    _write_json(acceptance_path, acceptance)
    initial_checks = [
        CheckResult(
            code="B60_BASELINE_DECISION_MISSING",
            message="Approved PBD-BASIC60-PRIVATE-001 decision and confirmation are required",
            location="baseline_decision",
        ),
        CheckResult(
            code="B60_USAGE_AMENDMENT_MISSING",
            message=("Approved PBD-BASIC60-PRIVATE-001-A1 manual data-usage amendment is required"),
            location="usage_amendment",
        ),
        CheckResult(
            code="B60_EXCEL_REVIEW_NOT_PREPARED",
            message=(
                "Generate the frozen Basic60 data/source workbook and prepare its one-approval "
                "private-review packet"
            ),
            location="prepare-basic60-review",
        ),
    ]
    assessment = _assessment("not_ready", summary, initial_checks, raw_manifest=raw_manifest)
    _write_json(assessment_path, assessment)
    return (
        [
            seed_path,
            sample_candidate_path,
            *sample_review_template_paths,
            item_review_template_path,
            *[
                paths.root / d3_artifacts["stages"][stage_id]["path"]
                for stage_id in ("parse", "standardize", "entity")
            ],
            profile_path,
            ai_usage_policy_path,
            usage_amendment_template_path,
            source_matrix_path,
            country_source_uri_manifest_path,
            source_registry_path,
            human_approval_path,
            d1_scope_authorization_path,
            manifest_path,
            *[replay_manifest_paths[pipeline_id] for pipeline_id in PIPELINE_IDS],
            acceptance_path,
            assessment_path,
        ],
        assessment,
    )


def _valid_relative_path(value: Any, prefix: PurePosixPath) -> bool:
    if not isinstance(value, str) or not value or "\\" in value:
        return False
    pure = PurePosixPath(value)
    return not pure.is_absolute() and ".." not in pure.parts and pure.is_relative_to(prefix)


def _validate_file_reference(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
    *,
    code: str,
    location: str,
    prefix: PurePosixPath | None = None,
) -> Path | None:
    if not isinstance(reference, dict):
        checks.append(
            CheckResult(code=code, message="File reference is required", location=location)
        )
        return None
    relative = reference.get("path")
    expected_hash = reference.get("sha256")
    if not isinstance(relative, str) or not relative or "\\" in relative:
        checks.append(
            CheckResult(code=code, message="Path must be repository-relative", location=location)
        )
        return None
    pure = PurePosixPath(relative)
    if pure.is_absolute() or ".." in pure.parts or (prefix and not pure.is_relative_to(prefix)):
        checks.append(
            CheckResult(
                code=code,
                message="Path escapes its approved directory",
                location=location,
            )
        )
        return None
    target = paths.root.joinpath(*pure.parts)
    try:
        resolved = target.resolve(strict=True)
    except OSError:
        checks.append(
            CheckResult(code=code, message="Referenced file does not exist", location=location)
        )
        return None
    if (
        target.is_symlink()
        or not resolved.is_relative_to(paths.root.resolve())
        or not resolved.is_file()
    ):
        checks.append(
            CheckResult(code=code, message="Referenced file is unsafe", location=location)
        )
        return None
    if not isinstance(expected_hash, str) or not SHA256_PATTERN.fullmatch(expected_hash):
        checks.append(
            CheckResult(code=code, message="Reference requires a SHA-256", location=location)
        )
        return None
    if sha256_file(resolved) != expected_hash:
        checks.append(
            CheckResult(
                code=code,
                message="Referenced file hash does not match",
                location=location,
            )
        )
        return None
    return resolved


def _subject_payloads(bundle: dict[str, Any]) -> dict[str, Any]:
    source = bundle.get("source_admission")
    collection = bundle.get("collection")
    processing = bundle.get("processing")
    acceptance = bundle.get("acceptance")
    baseline = bundle.get("baseline_decision")
    source_scope_payload = {
        "release_id": bundle.get("release_id"),
        "profile_id": bundle.get("profile_id"),
        "authorized_scope": source.get("authorized_scope") if isinstance(source, dict) else None,
    }
    source_payload = {
        "release_id": bundle.get("release_id"),
        "baseline_clarification": bundle.get("baseline_clarification"),
        "usage_amendment": bundle.get("usage_amendment"),
        "source_registry": source.get("source_registry") if isinstance(source, dict) else None,
        "country_domain_matrix": source.get("country_domain_matrix")
        if isinstance(source, dict)
        else None,
        "derived_matrix_sha256": source.get("derived_matrix_sha256")
        if isinstance(source, dict)
        else None,
        "domains": source.get("domains") if isinstance(source, dict) else None,
        "usage_boundaries": source.get("usage_boundaries") if isinstance(source, dict) else None,
    }
    collection_payload = {
        "release_id": bundle.get("release_id"),
        "raw_root_sha256": bundle.get("raw_root_sha256"),
        "source_subject_sha256": _canonical_sha256(source_payload),
        "input_immutable": collection.get("input_immutable")
        if isinstance(collection, dict)
        else None,
        "batches": collection.get("batches") if isinstance(collection, dict) else None,
        "incidents": collection.get("incidents") if isinstance(collection, dict) else None,
    }
    acceptance_payload = {
        "release_id": bundle.get("release_id"),
        "raw_root_sha256": bundle.get("raw_root_sha256"),
        "runtime_seed": bundle.get("runtime_seed"),
        "profile": bundle.get("profile"),
        "baseline_decision": baseline,
        "baseline_confirmation": bundle.get("baseline_confirmation"),
        "usage_amendment": bundle.get("usage_amendment"),
        "ai_usage_policy": bundle.get("ai_usage_policy"),
        "source_subject_sha256": _canonical_sha256(source_payload),
        "collection_subject_sha256": _canonical_sha256(collection_payload),
        "processing": processing,
        "acceptance": {
            key: value
            for key, value in acceptance.items()
            if key not in {"sample_approvals", "final_approvals"}
        }
        if isinstance(acceptance, dict)
        else None,
    }
    revocation_payload = {
        "release_id": bundle.get("release_id"),
        "raw_root_sha256": bundle.get("raw_root_sha256"),
        "runtime_seed": bundle.get("runtime_seed"),
    }
    payloads = {
        "source_scope": source_scope_payload,
        "source": source_payload,
        "collection": collection_payload,
        "acceptance": acceptance_payload,
        "revocation": revocation_payload,
    }
    artifacts = acceptance.get("artifacts") if isinstance(acceptance, dict) else None
    sample_reviews = artifacts.get("sample_reviews") if isinstance(artifacts, dict) else None
    for index in range(2):
        review_reference = (
            sample_reviews[index]
            if isinstance(sample_reviews, list) and index < len(sample_reviews)
            else None
        )
        payloads[f"sample_review_{index + 1}"] = {
            "release_id": bundle.get("release_id"),
            "sample_candidate": acceptance.get("sample_candidate")
            if isinstance(acceptance, dict)
            else None,
            "sample_review": review_reference,
        }
    return payloads


def approval_subjects(bundle: dict[str, Any]) -> dict[str, str]:
    return {key: _canonical_sha256(value) for key, value in _subject_payloads(bundle).items()}


def _d1_machine_content_inventory(source_registry: dict[str, Any]) -> list[dict[str, Any]]:
    sources = source_registry.get("sources")
    if not isinstance(sources, list):
        raise ValueError("Basic60 source registry entries are required")
    inventory: list[dict[str, Any]] = []
    for item in sources:
        if not isinstance(item, dict):
            raise ValueError("Basic60 source registry entries must be objects")
        source_ref = str(item.get("source_ref") or "")
        evidence = item.get("evidence")
        if not source_ref or not isinstance(evidence, dict):
            raise ValueError("Basic60 source evidence reference is required")
        inventory.append(
            {
                "source_ref": source_ref,
                "provider": item.get("provider"),
                "dataset": item.get("dataset"),
                "data_domain": item.get("data_domain"),
                "source_role": item.get("source_role"),
                "terms_uri": item.get("terms_uri"),
                "evidence": evidence,
                "evidence_sha256": item.get("evidence_sha256"),
                "access_method": item.get("access_method"),
                "rate_limit": item.get("rate_limit"),
                "snapshots": item.get("snapshots"),
                "admission_fingerprint": item.get("admission_fingerprint"),
            }
        )
    inventory.sort(key=lambda item: item["source_ref"])
    return inventory


def build_d1_machine_content_binding(
    bundle: dict[str, Any],
    source_registry: dict[str, Any],
    *,
    generated_at: str,
) -> dict[str, Any]:
    """Bind exact post-capture D1 content without representing a human signature."""

    if not _aware_timestamp(generated_at):
        raise ValueError("D1 machine-content binding requires a timezone-aware generated_at")
    source = bundle.get("source_admission")
    if not isinstance(source, dict) or not isinstance(source.get("source_registry"), dict):
        raise ValueError("Basic60 source-registry reference is required")
    authorization_id = str(source.get("scope_authorization_id") or "")
    if not D1_SCOPE_AUTHORIZATION_ID_PATTERN.fullmatch(authorization_id):
        raise ValueError("D1 machine binding requires a valid scope_authorization_id")
    approvals = source.get("approvals")
    approval = approvals[0] if isinstance(approvals, list) and len(approvals) == 1 else None
    if not isinstance(approval, dict) or not isinstance(approval.get("evidence"), dict):
        raise ValueError("D1 machine binding requires scope-authorization evidence")
    subjects = approval_subjects(bundle)
    claims: dict[str, Any] = {
        "schema_version": D1_MACHINE_CONTENT_BINDING_SCHEMA,
        "binding_type": "machine_generated_content_binding",
        "generator": "navigator-data",
        "generated_at": generated_at,
        "human_signature": False,
        "release_id": bundle.get("release_id"),
        "profile_id": bundle.get("profile_id"),
        "scope_authorization_id": authorization_id,
        "scope_authorization_evidence": approval["evidence"],
        "scope_subject_sha256": subjects["source_scope"],
        "content_subject_sha256": subjects["source"],
        "source_registry": source["source_registry"],
        "sources": _d1_machine_content_inventory(source_registry),
        "statement": (
            "Machine binding of exact registry, source evidence, terms snapshots, URI, and "
            "schema content; A1 is the sole human usage-permission decision and this binding "
            "is not a second decision or human signature."
        ),
    }
    claims["binding_sha256"] = _canonical_sha256(claims)
    return claims


def _d1_machine_baseline_anchor_path(authorization_id: str) -> PurePosixPath:
    if not D1_SCOPE_AUTHORIZATION_ID_PATTERN.fullmatch(authorization_id):
        raise ValueError("Basic60 D1 scope_authorization_id is missing or invalid")
    return D1_MACHINE_BASELINE_ROOT / "authorizations" / f"{authorization_id}.json"


def _serialized_json_artifact(payload: dict[str, Any]) -> bytes:
    return (json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n").encode(
        "utf-8"
    )


def _write_json_exclusive_read_only(path: Path, payload: dict[str, Any]) -> None:
    """Atomically publish a JSON artifact once, without an overwrite path."""

    path.parent.mkdir(parents=True, exist_ok=True)
    encoded = _serialized_json_artifact(payload)
    temporary_fd, temporary_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
    )
    temporary_path = Path(temporary_name)
    try:
        with os.fdopen(temporary_fd, "wb") as handle:
            handle.write(encoded)
            handle.flush()
            os.fsync(handle.fileno())
        temporary_path.chmod(0o444)
        try:
            os.link(temporary_path, path)
        except FileExistsError as error:
            raise ValueError(
                "B60_D1_MACHINE_BASELINE_ALREADY_BOUND: scope authorization already has "
                "an immutable machine baseline"
            ) from error
    finally:
        temporary_path.unlink(missing_ok=True)


def _trusted_path_chain(path: Path) -> tuple[Path, ...]:
    """Return an absolute path and every ancestor, rooted at the filesystem root."""

    if not path.is_absolute():
        raise ValueError("trusted paths must be absolute")
    current = Path(path.anchor)
    chain = [current]
    for part in path.parts[1:]:
        current /= part
        chain.append(current)
    return tuple(chain)


def _runtime_account_can_write(path: Path) -> bool:
    """Report effective write access for the ordinary validator/runtime account."""

    return os.access(path, os.W_OK)


def _validate_trusted_path_chain(
    path: Path,
    *,
    target_type: str,
    require_runtime_read_only: bool,
    require_target_mode_read_only: bool,
    error_code: str,
) -> None:
    """Validate trusted ownership and replacement resistance through every ancestor."""

    if target_type not in {"directory", "file"}:
        raise ValueError("trusted-path target_type is invalid")
    if not path.is_absolute() or path.resolve(strict=False) != path:
        raise ValueError(f"{error_code}: trusted path is not canonical")
    chain = _trusted_path_chain(path)
    for index, component in enumerate(chain):
        is_target = index == len(chain) - 1
        try:
            metadata = component.lstat()
        except OSError as error:
            raise ValueError(
                f"{error_code}: trusted path component is unavailable: {component}"
            ) from error
        component_is_directory = stat.S_ISDIR(metadata.st_mode)
        component_is_file = stat.S_ISREG(metadata.st_mode)
        valid_type = component_is_directory and (not is_target or target_type == "directory")
        if is_target and target_type == "file":
            valid_type = component_is_file
        mode = stat.S_IMODE(metadata.st_mode)
        structurally_trusted = bool(
            valid_type
            and metadata.st_uid in D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS
            and not stat.S_ISLNK(metadata.st_mode)
            and (not component_is_directory or mode & 0o022 == 0)
            and (not (is_target and require_target_mode_read_only) or mode & 0o222 == 0)
        )
        runtime_read_only = not (
            require_runtime_read_only and _runtime_account_can_write(component)
        )
        if not structurally_trusted or not runtime_read_only:
            raise ValueError(
                f"{error_code}: path and every ancestor must be trusted, non-symlink, and "
                f"non-writable to the ordinary runtime/update account: {component}"
            )


def _trusted_d1_volume_root(
    volume_id: str,
    *,
    require_runtime_read_only: bool,
) -> Path:
    """Resolve a D1 volume ID through the fixed, administrator-owned trust registry."""

    registry_path = D1_TRUSTED_VOLUME_REGISTRY_PATH
    try:
        _validate_trusted_path_chain(
            registry_path,
            target_type="file",
            require_runtime_read_only=require_runtime_read_only,
            require_target_mode_read_only=True,
            error_code="B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID",
        )
        resolved_registry = registry_path.resolve(strict=True)
    except (OSError, ValueError) as error:
        if str(error).startswith("B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID"):
            raise
        raise ValueError(
            "B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID: registry is unavailable"
        ) from error
    try:
        registry = _json_object(resolved_registry)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        raise ValueError(
            "B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID: registry must be a JSON object"
        ) from error
    volumes = registry.get("volumes")
    if (
        set(registry) != {"schema_version", "volumes"}
        or registry.get("schema_version") != D1_TRUSTED_VOLUME_REGISTRY_SCHEMA
        or not isinstance(volumes, list)
    ):
        raise ValueError("B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID: registry schema is invalid")
    indexed: dict[str, Path] = {}
    for item in volumes:
        if not isinstance(item, dict) or set(item) != {"volume_id", "canonical_root", "status"}:
            raise ValueError("B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID: volume entry is invalid")
        item_id = str(item.get("volume_id") or "")
        root_value = str(item.get("canonical_root") or "")
        root = Path(root_value)
        if (
            item.get("status") != "active"
            or not VOLUME_ID_PATTERN.fullmatch(item_id)
            or item_id in indexed
            or not root.is_absolute()
            or _has_symlink_component(root)
            or root.resolve(strict=False) != root
        ):
            raise ValueError("B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID: volume entry is unsafe")
        indexed[item_id] = root
    try:
        trusted_root = indexed[volume_id]
    except KeyError as error:
        raise ValueError(
            "B60_D1_TRUSTED_VOLUME_UNKNOWN: volume_id is not in the trusted registry"
        ) from error
    _validate_trusted_path_chain(
        trusted_root,
        target_type="directory",
        require_runtime_read_only=require_runtime_read_only,
        require_target_mode_read_only=False,
        error_code="B60_D1_TRUSTED_VOLUME_ROOT_UNSAFE",
    )
    return trusted_root


def write_d1_machine_baseline_evidence(
    paths: RepositoryPaths,
    bundle: dict[str, Any],
    source_registry: dict[str, Any],
    *,
    l0_dir: Path,
    volume_id: str,
    generated_at: str,
) -> dict[str, Any]:
    """Create and seal one external-L0 baseline for a D1 scope authorization."""

    source = bundle.get("source_admission")
    if not isinstance(source, dict):
        raise ValueError("Basic60 source admission is required")
    if not l0_dir.is_absolute() or _has_symlink_component(l0_dir):
        raise ValueError("B60_D1_LEDGER_INVALID: external L0 directory must be absolute and safe")
    resolved_l0 = l0_dir.resolve(strict=False)
    if (
        len(resolved_l0.parts) <= 2
        or resolved_l0 == Path.home().resolve()
        or _paths_overlap(resolved_l0, paths.root.resolve())
    ):
        raise ValueError("B60_D1_LEDGER_INVALID: external L0 directory is broad or overlaps repo")
    if not VOLUME_ID_PATTERN.fullmatch(volume_id):
        raise ValueError("B60_D1_LEDGER_INVALID: volume_id is invalid")
    trusted_root = _trusted_d1_volume_root(
        volume_id,
        require_runtime_read_only=False,
    )
    if resolved_l0 != trusted_root:
        raise ValueError(
            "B60_D1_TRUSTED_VOLUME_MISMATCH: l0_dir does not match the registered canonical root"
        )
    authorization_id = str(source.get("scope_authorization_id") or "")
    anchor_key = _d1_machine_baseline_anchor_path(authorization_id)
    ledger_root = resolved_l0.joinpath(*D1_MACHINE_BASELINE_ROOT.parts)
    authorization_root = ledger_root / "authorizations"
    object_root = ledger_root / "sha256"
    anchor_path = resolved_l0.joinpath(*anchor_key.parts)
    if anchor_path.exists() or anchor_path.is_symlink():
        raise ValueError(
            "B60_D1_MACHINE_BASELINE_ALREADY_BOUND: scope authorization already has an "
            "immutable machine baseline"
        )
    for directory in (ledger_root, authorization_root, object_root):
        if directory.is_symlink() or (directory.exists() and not directory.is_dir()):
            raise ValueError("B60_D1_LEDGER_INVALID: ledger directory is unsafe")
    ledger_root.mkdir(parents=True, exist_ok=True)
    authorization_root.mkdir(exist_ok=True)
    object_root.mkdir(exist_ok=True)
    writer_uid = os.geteuid()
    if any(
        directory.stat().st_uid not in D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS
        or writer_uid not in D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS
        or stat.S_IMODE(directory.stat().st_mode) & 0o200 == 0
        or not os.access(directory, os.W_OK)
        for directory in (ledger_root, authorization_root, object_root)
    ):
        raise ValueError(
            "B60_D1_LEDGER_SEALED: D1 admin must explicitly unseal the ledger for a new "
            "human-authorized revision"
        )
    binding = build_d1_machine_content_binding(
        bundle,
        source_registry,
        generated_at=generated_at,
    )
    object_sha256 = hashlib.sha256(_serialized_json_artifact(binding)).hexdigest()
    object_key = D1_MACHINE_BASELINE_ROOT / "sha256" / f"{object_sha256}.json"
    object_path = resolved_l0.joinpath(*object_key.parts)
    if object_path.exists():
        if sha256_file(object_path) != object_sha256:
            raise ValueError("Existing D1 machine-baseline object is invalid")
    else:
        _write_json_exclusive_read_only(object_path, binding)
    object_reference = {
        "storage_mode": "external_l0",
        "volume_id": volume_id,
        "object_key": object_key.as_posix(),
        "sha256": sha256_file(object_path),
        "byte_size": object_path.stat().st_size,
    }
    anchor = {
        "schema_version": D1_MACHINE_BASELINE_ANCHOR_SCHEMA,
        "record_type": "basic60_d1_machine_baseline_anchor",
        "scope_authorization_id": authorization_id,
        "scope_subject_sha256": approval_subjects(bundle)["source_scope"],
        "machine_baseline": object_reference,
        "anchored_at": generated_at,
        "append_only": True,
        "human_signature": False,
        "statement": (
            "Single-use machine baseline anchor; exact post-capture hashes are not a human "
            "signature and cannot be rebound under this authorization ID."
        ),
    }
    _write_json_exclusive_read_only(anchor_path, anchor)
    for directory in (object_root, authorization_root, ledger_root):
        directory.chmod(0o555)
    return {
        "storage_mode": "external_l0",
        "volume_id": volume_id,
        "scope_authorization_id": authorization_id,
        "object_key": anchor_key.as_posix(),
        "sha256": sha256_file(anchor_path),
        "byte_size": anchor_path.stat().st_size,
    }


def _valid_person(value: Any) -> bool:
    text = str(value or "").strip()
    return bool(text) and not PLACEHOLDER_PATTERN.fullmatch(text)


def _aware_timestamp(value: Any) -> bool:
    text = str(value or "").strip()
    if not text:
        return False
    try:
        parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return False
    return parsed.tzinfo is not None and parsed.utcoffset() is not None


def _validate_approval(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    approval: Any,
    *,
    expected_role: str,
    subject_sha256: str,
    location: str,
) -> str | None:
    if not isinstance(approval, dict):
        checks.append(
            CheckResult(
                code="B60_HUMAN_APPROVAL_MISSING",
                message=f"Approval from {expected_role} is required",
                location=location,
            )
        )
        return None
    person = approval.get("person_name")
    valid = (
        approval.get("role") == expected_role
        and _valid_person(person)
        and approval.get("decision") == "approved"
        and _aware_timestamp(approval.get("signed_at"))
        and approval.get("subject_sha256") == subject_sha256
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_HUMAN_APPROVAL_INVALID",
                message=(
                    f"{expected_role} approval is incomplete or not bound to the current subject"
                ),
                location=location,
            )
        )
        return None
    evidence_path = _validate_file_reference(
        checks,
        paths,
        approval.get("evidence"),
        code="B60_APPROVAL_EVIDENCE_INVALID",
        location=f"{location}.evidence",
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if evidence_path is None:
        return None
    try:
        evidence = _json_object(evidence_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_APPROVAL_EVIDENCE_INVALID",
                message="Approval evidence must be a JSON object",
                location=f"{location}.evidence",
            )
        )
        return None
    evidence_valid = (
        evidence.get("schema_version") == SCHEMA_VERSION
        and evidence.get("record_type") == "basic60_human_approval"
        and evidence.get("template_only") is False
        and evidence.get("decision") == "approved"
        and evidence.get("person_name") == person
        and evidence.get("role") == expected_role
        and evidence.get("signed_at") == approval.get("signed_at")
        and evidence.get("subject_sha256") == subject_sha256
        and evidence.get("evidence_id") == approval.get("evidence", {}).get("evidence_id")
    )
    if not evidence_valid:
        checks.append(
            CheckResult(
                code="B60_APPROVAL_EVIDENCE_MISMATCH",
                message="Approval fields do not match their hash-bound evidence record",
                location=f"{location}.evidence",
            )
        )
        return None
    return str(person)


def _validate_d1_scope_authorization(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    source: dict[str, Any],
    *,
    scope_subject_sha256: str,
    content_subject_sha256: str,
) -> dict[str, Any] | None:
    approvals = source.get("approvals")
    if not isinstance(approvals, list) or len(approvals) != 1:
        checks.append(
            CheckResult(
                code="B60_SOURCE_APPROVAL_COUNT_INVALID",
                message="D1 source admission requires exactly one human scope authorization",
                location="source_admission.approvals",
            )
        )
    approval = approvals[0] if isinstance(approvals, list) and approvals else None
    if not isinstance(approval, dict):
        checks.append(
            CheckResult(
                code="B60_D1_SCOPE_AUTHORIZATION_MISSING",
                message="D1 requires a human authorization of the stable source scope",
                location="source_admission.approvals.0",
            )
        )
        return None
    if (
        approval.get("subject_sha256") == content_subject_sha256
        or approval.get("scope_subject_sha256") == content_subject_sha256
    ):
        checks.append(
            CheckResult(
                code="B60_D1_LEGACY_PRECISE_SIGNATURE_REJECTED",
                message=(
                    "A human authorization cannot be represented as signing exact registry or "
                    "post-capture evidence hashes"
                ),
                location="source_admission.approvals.0",
            )
        )
    person = approval.get("person_name")
    authorization_id = str(source.get("scope_authorization_id") or "")
    valid = (
        approval.get("approval_type") == "human_scope_authorization"
        and approval.get("role") == "项目批准人"
        and _valid_person(person)
        and approval.get("decision") == "approved"
        and _aware_timestamp(approval.get("signed_at"))
        and authorization_id == REQUIRED_D1_SCOPE_AUTHORIZATION_ID
        and approval.get("scope_authorization_id") == authorization_id
        and approval.get("scope_subject_sha256") == scope_subject_sha256
        and approval.get("exact_content_hash_signed") is False
        and "subject_sha256" not in approval
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_D1_SCOPE_AUTHORIZATION_INVALID",
                message=(
                    "The A1 usage-boundary change requires the new D1 R2 human authorization "
                    "of stable scope and must explicitly exclude exact content hashes"
                ),
                location="source_admission.approvals.0",
            )
        )
        return None
    if person != "kevin":
        checks.append(
            CheckResult(
                code="B60_SOURCE_PROJECT_APPROVER_INVALID",
                message="Basic60 D1 project approver must be kevin",
                location="source_admission.approvals.0",
            )
        )
    evidence_path = _validate_file_reference(
        checks,
        paths,
        approval.get("evidence"),
        code="B60_APPROVAL_EVIDENCE_INVALID",
        location="source_admission.approvals.0.evidence",
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if evidence_path is None:
        return None
    try:
        evidence = _json_object(evidence_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_APPROVAL_EVIDENCE_INVALID",
                message="D1 scope-authorization evidence must be a JSON object",
                location="source_admission.approvals.0.evidence",
            )
        )
        return None
    transcribed_at = evidence.get("transcribed_at")
    chronology_valid = False
    if _aware_timestamp(approval.get("signed_at")) and _aware_timestamp(transcribed_at):
        signed = datetime.fromisoformat(str(approval["signed_at"]).replace("Z", "+00:00"))
        transcribed = datetime.fromisoformat(str(transcribed_at).replace("Z", "+00:00"))
        chronology_valid = transcribed >= signed
    evidence_valid = (
        evidence.get("schema_version") == SCHEMA_VERSION
        and evidence.get("record_type") == D1_SCOPE_AUTHORIZATION_RECORD_TYPE
        and evidence.get("authorization_type") == "human_scope_authorization"
        and evidence.get("template_only") is False
        and evidence.get("decision") == "approved"
        and evidence.get("person_name") == person
        and evidence.get("role") == "项目批准人"
        and evidence.get("signed_at") == approval.get("signed_at")
        and evidence.get("scope_authorization_id") == authorization_id
        and evidence.get("scope_subject_sha256") == scope_subject_sha256
        and evidence.get("exact_content_hash_signed") is False
        and evidence.get("evidence_id") == approval.get("evidence", {}).get("evidence_id")
        and evidence.get("authorization_basis") == "explicit_user_confirmation_in_project_task"
        and chronology_valid
        and "subject_sha256" not in evidence
    )
    if not evidence_valid:
        checks.append(
            CheckResult(
                code="B60_D1_SCOPE_AUTHORIZATION_EVIDENCE_MISMATCH",
                message=(
                    "D1 evidence must truthfully record the earlier scope authorization and "
                    "a later transcription without claiming an exact-content signature"
                ),
                location="source_admission.approvals.0.evidence",
            )
        )
        return None
    return approval


def _latest_d1_content_capture_at(
    paths: RepositoryPaths,
    source_registry: dict[str, Any],
) -> datetime | None:
    captured: list[datetime] = []
    sources = source_registry.get("sources")
    if not isinstance(sources, list):
        return None
    for item in sources:
        if not isinstance(item, dict):
            continue
        for snapshot in item.get("snapshots", []):
            if isinstance(snapshot, dict) and _aware_timestamp(snapshot.get("captured_at")):
                captured.append(
                    datetime.fromisoformat(str(snapshot["captured_at"]).replace("Z", "+00:00"))
                )
        evidence = _source_evidence_for_fingerprint(paths, item)
        if not isinstance(evidence, dict):
            continue
        values: list[Any] = [evidence.get("captured_at")]
        terms_snapshots = evidence.get("terms_snapshots")
        if isinstance(terms_snapshots, list):
            values.extend(
                snapshot.get("captured_at")
                for snapshot in terms_snapshots
                if isinstance(snapshot, dict)
            )
        for value in values:
            if _aware_timestamp(value):
                captured.append(datetime.fromisoformat(str(value).replace("Z", "+00:00")))
    return max(captured) if captured else None


def _validate_d1_two_stage_admission(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    bundle: dict[str, Any],
    source: dict[str, Any],
    source_registry: dict[str, Any] | None,
    l0_dir: Path | None,
) -> None:
    subjects = approval_subjects(bundle)
    authorized_scope = source.get("authorized_scope")
    if source_registry is None:
        expected_scope = None
    else:
        try:
            expected_scope = build_d1_authorized_scope(bundle, source_registry)
        except ValueError:
            expected_scope = None
    if not isinstance(authorized_scope, dict) or authorized_scope != expected_scope:
        checks.extend(
            [
                CheckResult(
                    code="B60_D1_AUTHORIZED_SCOPE_INVALID",
                    message=(
                        "Authorized D1 scope must match the three current provider/dataset "
                        "identities, terms locations, update rules, clarification, and boundaries"
                    ),
                    location="source_admission.authorized_scope",
                ),
                CheckResult(
                    code="B60_D1_REAPPROVAL_REQUIRED",
                    message=(
                        "Provider, source URL, terms, schema, or usage-boundary changes require "
                        "a new human scope authorization"
                    ),
                    location="source_admission",
                ),
            ]
        )
    approval = _validate_d1_scope_authorization(
        checks,
        paths,
        source,
        scope_subject_sha256=subjects["source_scope"],
        content_subject_sha256=subjects["source"],
    )
    binding_reference = source.get("machine_content_binding")
    if not isinstance(binding_reference, dict):
        checks.append(
            CheckResult(
                code="B60_D1_MACHINE_CONTENT_BINDING_MISSING",
                message=(
                    "D1 requires a post-capture machine binding for exact registry, evidence, "
                    "terms, URI, schema, and boundary content"
                ),
                location="source_admission.machine_content_binding",
            )
        )
        return
    authorization_id = str(source.get("scope_authorization_id") or "")
    try:
        expected_anchor_key = _d1_machine_baseline_anchor_path(authorization_id)
    except ValueError:
        expected_anchor_key = None
    reference_fields_valid = set(binding_reference) == {
        "storage_mode",
        "volume_id",
        "scope_authorization_id",
        "object_key",
        "sha256",
        "byte_size",
    }
    anchor_digest = str(binding_reference.get("sha256") or "")
    anchor_size = binding_reference.get("byte_size")
    reference_valid = bool(
        reference_fields_valid
        and binding_reference.get("storage_mode") == "external_l0"
        and VOLUME_ID_PATTERN.fullmatch(str(binding_reference.get("volume_id") or ""))
        and binding_reference.get("scope_authorization_id") == authorization_id
        and expected_anchor_key is not None
        and binding_reference.get("object_key") == expected_anchor_key.as_posix()
        and SHA256_PATTERN.fullmatch(anchor_digest)
        and isinstance(anchor_size, int)
        and not isinstance(anchor_size, bool)
        and anchor_size > 0
    )
    if l0_dir is None or not reference_valid:
        checks.extend(
            [
                CheckResult(
                    code="B60_D1_MACHINE_CONTENT_BINDING_INVALID",
                    message="D1 machine baseline requires a canonical external-L0 reference",
                    location="source_admission.machine_content_binding",
                ),
                CheckResult(
                    code="B60_D1_REAPPROVAL_REQUIRED",
                    message=(
                        "Each scope_authorization_id has exactly one canonical external-L0 "
                        "anchor; another reference requires a new human authorization"
                    ),
                    location="source_admission",
                ),
            ]
        )
        return
    volume_id = str(binding_reference.get("volume_id") or "")
    try:
        trusted_root = _trusted_d1_volume_root(
            volume_id,
            require_runtime_read_only=True,
        )
    except ValueError as error:
        message = str(error)
        if message.startswith("B60_D1_TRUSTED_VOLUME_UNKNOWN"):
            code = "B60_D1_TRUSTED_VOLUME_UNKNOWN"
        elif message.startswith("B60_D1_TRUSTED_VOLUME_ROOT_UNSAFE"):
            code = "B60_D1_TRUSTED_VOLUME_ROOT_UNSAFE"
        else:
            code = "B60_D1_TRUSTED_VOLUME_REGISTRY_INVALID"
        checks.append(
            CheckResult(
                code=code,
                message=message,
                location="source_admission.machine_content_binding.volume_id",
            )
        )
        return
    if l0_dir.resolve(strict=False) != trusted_root:
        checks.extend(
            [
                CheckResult(
                    code="B60_D1_TRUSTED_VOLUME_MISMATCH",
                    message="--l0-dir does not match the registered canonical root for volume_id",
                    location="source_admission.machine_content_binding.volume_id",
                ),
                CheckResult(
                    code="B60_D1_REAPPROVAL_REQUIRED",
                    message="An alternate L0 root cannot reuse an existing D1 authorization",
                    location="source_admission",
                ),
            ]
        )
        return
    assert expected_anchor_key is not None
    ledger_root = l0_dir.joinpath(*D1_MACHINE_BASELINE_ROOT.parts)
    authorization_root = ledger_root / "authorizations"
    object_root = ledger_root / "sha256"
    protected_directories = (ledger_root, authorization_root, object_root)
    try:
        directories_protected = all(
            not directory.is_symlink()
            and directory.resolve(strict=True).is_dir()
            and directory.resolve(strict=True).is_relative_to(l0_dir.resolve(strict=True))
            and directory.stat().st_uid in D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS
            and stat.S_IMODE(directory.stat().st_mode) & 0o222 == 0
            and not _runtime_account_can_write(directory)
            for directory in protected_directories
        )
    except OSError:
        directories_protected = False
    if not directories_protected:
        checks.append(
            CheckResult(
                code="B60_D1_MACHINE_BASELINE_LEDGER_UNSEALED",
                message=(
                    "D1 ledger root and anchor/object parent directories must be non-writable "
                    "to the validation/runtime account"
                ),
                location="source_admission.machine_content_binding",
            )
        )
        return
    anchor_path = l0_dir.joinpath(*expected_anchor_key.parts)
    try:
        anchor_resolved = anchor_path.resolve(strict=True)
        anchor_safe = bool(
            not anchor_path.is_symlink()
            and anchor_resolved.is_file()
            and anchor_resolved.is_relative_to(ledger_root.resolve(strict=True))
            and anchor_resolved.stat().st_uid in D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS
            and stat.S_IMODE(anchor_resolved.stat().st_mode) == 0o444
            and not _runtime_account_can_write(anchor_resolved)
            and anchor_resolved.stat().st_size == anchor_size
            and sha256_file(anchor_resolved) == anchor_digest
        )
    except OSError:
        anchor_safe = False
    if not anchor_safe:
        checks.append(
            CheckResult(
                code="B60_D1_MACHINE_CONTENT_BINDING_INVALID",
                message="External D1 machine-baseline anchor is missing, mutable, or changed",
                location="source_admission.machine_content_binding",
            )
        )
        return
    try:
        anchor = _json_object(anchor_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        anchor = {}
    object_reference = anchor.get("machine_baseline")
    object_reference_dict = object_reference if isinstance(object_reference, dict) else {}
    object_digest = str(object_reference_dict.get("sha256") or "")
    expected_object_key = (D1_MACHINE_BASELINE_ROOT / "sha256" / f"{object_digest}.json").as_posix()
    object_size = object_reference_dict.get("byte_size")
    object_key = str(object_reference_dict.get("object_key") or "")
    object_path = l0_dir.joinpath(*PurePosixPath(object_key).parts)
    object_reference_valid = bool(
        set(object_reference_dict)
        == {"storage_mode", "volume_id", "object_key", "sha256", "byte_size"}
        and object_reference_dict.get("storage_mode") == "external_l0"
        and object_reference_dict.get("volume_id") == binding_reference.get("volume_id")
        and SHA256_PATTERN.fullmatch(object_digest)
        and object_key == expected_object_key
        and isinstance(object_size, int)
        and not isinstance(object_size, bool)
        and object_size > 0
    )
    try:
        object_resolved = object_path.resolve(strict=True)
        object_safe = bool(
            object_reference_valid
            and not object_path.is_symlink()
            and object_resolved.is_file()
            and object_resolved.is_relative_to(object_root.resolve(strict=True))
            and object_resolved.stat().st_uid in D1_TRUSTED_VOLUME_REGISTRY_ALLOWED_OWNER_UIDS
            and stat.S_IMODE(object_resolved.stat().st_mode) == 0o444
            and not _runtime_account_can_write(object_resolved)
            and object_resolved.stat().st_size == object_size
            and sha256_file(object_resolved) == object_digest
        )
    except OSError:
        object_safe = False
    if not object_safe:
        checks.extend(
            [
                CheckResult(
                    code="B60_D1_MACHINE_CONTENT_BINDING_INVALID",
                    message="D1 machine-baseline object must be content-addressed and read-only",
                    location="source_admission.machine_content_binding.machine_baseline",
                ),
                CheckResult(
                    code="B60_D1_REAPPROVAL_REQUIRED",
                    message="Invalid machine baseline cannot be rebound under an existing scope",
                    location="source_admission",
                ),
            ]
        )
        return
    try:
        binding = _json_object(object_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        binding = {}
    generated_at = binding.get("generated_at")
    expected_binding: dict[str, Any] | None = None
    if source_registry is not None and _aware_timestamp(generated_at):
        try:
            expected_binding = build_d1_machine_content_binding(
                bundle,
                source_registry,
                generated_at=str(generated_at),
            )
        except ValueError:
            expected_binding = None
    anchor_valid = (
        anchor.get("schema_version") == D1_MACHINE_BASELINE_ANCHOR_SCHEMA
        and anchor.get("record_type") == "basic60_d1_machine_baseline_anchor"
        and anchor.get("scope_authorization_id") == authorization_id
        and anchor.get("scope_subject_sha256") == subjects["source_scope"]
        and anchor.get("machine_baseline") == object_reference
        and anchor.get("anchored_at") == generated_at
        and anchor.get("append_only") is True
        and anchor.get("human_signature") is False
    )
    if expected_binding is None or binding != expected_binding or not anchor_valid:
        checks.extend(
            [
                CheckResult(
                    code="B60_D1_MACHINE_CONTENT_BINDING_INVALID",
                    message=(
                        "D1 machine binding does not match the exact current registry, evidence, "
                        "terms, URI, schema, scope, and boundaries"
                    ),
                    location="source_admission.machine_content_binding",
                ),
                CheckResult(
                    code="B60_D1_REAPPROVAL_REQUIRED",
                    message=(
                        "Exact D1 content changed after machine binding; re-establish the binding "
                        "under a current human scope authorization"
                    ),
                    location="source_admission",
                ),
            ]
        )
        return
    assert source_registry is not None
    generated = datetime.fromisoformat(str(generated_at).replace("Z", "+00:00"))
    latest_capture = _latest_d1_content_capture_at(paths, source_registry)
    signed: datetime | None = None
    transcribed: datetime | None = None
    if approval is not None:
        signed = datetime.fromisoformat(str(approval["signed_at"]).replace("Z", "+00:00"))
        evidence_reference = approval.get("evidence")
        if isinstance(evidence_reference, dict):
            value = evidence_reference.get("path")
            if _valid_relative_path(value, PurePosixPath("data/basic60/evidence")):
                evidence_path = paths.root.joinpath(*PurePosixPath(str(value)).parts)
                try:
                    evidence = _json_object(evidence_path)
                except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
                    evidence = {}
                if _aware_timestamp(evidence.get("transcribed_at")):
                    transcribed = datetime.fromisoformat(
                        str(evidence["transcribed_at"]).replace("Z", "+00:00")
                    )
    if (
        (latest_capture is not None and generated < latest_capture)
        or (signed is not None and generated < signed)
        or (transcribed is not None and generated < transcribed)
    ):
        checks.append(
            CheckResult(
                code="B60_D1_MACHINE_CONTENT_BINDING_PREMATURE",
                message=(
                    "Machine content binding must be generated after the human scope record and "
                    "all exact source/terms captures"
                ),
                location="source_admission.machine_content_binding.generated_at",
            )
        )


def _validate_baseline_decision(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    decision_reference: Any,
    confirmation_reference: Any,
    candidate_reference: Any,
    confirmation_template_reference: Any,
) -> None:
    canonical_candidate = _reference(paths.root / PBD_CANDIDATE_PATH, paths)
    canonical_template = _reference(paths.root / PBD_CONFIRMATION_TEMPLATE_PATH, paths)
    if (
        candidate_reference != canonical_candidate
        or confirmation_template_reference != canonical_template
    ):
        checks.append(
            CheckResult(
                code="B60_BASELINE_AUTHORITY_CHAIN_CHANGED",
                message="Bundle must retain the canonical PBD candidate and confirmation template",
                location="baseline_candidate",
            )
        )
    try:
        candidate = _json_object(paths.root / PBD_CANDIDATE_PATH)
        confirmation_template = _json_object(paths.root / PBD_CONFIRMATION_TEMPLATE_PATH)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_BASELINE_CANDIDATE_INVALID",
                message="Canonical PBD candidate or confirmation template is unreadable",
                location=PBD_CANDIDATE_PATH.as_posix(),
            )
        )
        return
    candidate_valid = (
        candidate.get("candidate_type") == "project_baseline_change_candidate"
        and candidate.get("decision_id") == DECISION_ID
        and candidate.get("requested_decision") == REQUESTED_DECISION
        and candidate.get("status") == "pending_named_project_approver_review"
        and candidate.get("does_not_complete_formal_d1_d4") is True
        and candidate.get("does_not_complete_p0") is True
        and candidate.get("does_not_authorize_production") is True
        and candidate.get("does_not_authorize_public_release") is True
        and candidate.get("template_only") is True
    )
    if not candidate_valid:
        checks.append(
            CheckResult(
                code="B60_BASELINE_CANDIDATE_INVALID",
                message="Canonical PBD candidate does not match the Basic60 private-trial contract",
                location=PBD_CANDIDATE_PATH.as_posix(),
            )
        )
    if confirmation_template.get("candidate") != canonical_candidate:
        checks.append(
            CheckResult(
                code="B60_BASELINE_CONFIRMATION_TEMPLATE_INVALID",
                message="PBD confirmation template is not bound to the canonical candidate",
                location=PBD_CONFIRMATION_TEMPLATE_PATH.as_posix(),
            )
        )

    if decision_reference is None:
        checks.append(
            CheckResult(
                code="B60_BASELINE_DECISION_MISSING",
                message="Approved PBD-BASIC60-PRIVATE-001 decision is required",
                location="baseline_decision",
            )
        )
        return
    decision_path = _validate_file_reference(
        checks,
        paths,
        decision_reference,
        code="B60_BASELINE_DECISION_INVALID",
        location="baseline_decision",
        prefix=PurePosixPath("data/governance/review"),
    )
    if decision_path is None:
        return
    try:
        decision = _json_object(decision_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_BASELINE_DECISION_INVALID",
                message="Baseline decision must be a JSON object",
                location="baseline_decision",
            )
        )
        return
    scope = decision.get("scope")
    approval = decision.get("approval")
    approval_dict = approval if isinstance(approval, dict) else {}
    decision_valid = (
        decision.get("decision_id") == DECISION_ID
        and decision.get("status") == "approved_and_effective"
        and _aware_timestamp(decision.get("effective_from"))
        and scope
        == {
            "release_profile": PROFILE_ID,
            "private_trial_only": True,
            "real_data_enabled": True,
            "policies_enabled": False,
            "ai_enabled": False,
            "external_model_calls_enabled": False,
        }
        and bool(approval_dict)
        and approval_dict.get("person_id") == "kevin"
        and approval_dict.get("role") == "project_approver"
        and _aware_timestamp(approval_dict.get("signed_at"))
        and isinstance(approval_dict.get("signature_sha256"), str)
        and SHA256_PATTERN.fullmatch(str(approval_dict["signature_sha256"]))
    )
    if not decision_valid:
        checks.append(
            CheckResult(
                code="B60_BASELINE_DECISION_INVALID",
                message="Baseline decision does not authorize the bounded Basic60 private trial",
                location=_display_path(paths, decision_path),
            )
        )
        return
    effective_from = datetime.fromisoformat(str(decision["effective_from"]).replace("Z", "+00:00"))
    signed_at = datetime.fromisoformat(str(approval_dict["signed_at"]).replace("Z", "+00:00"))
    if signed_at > effective_from:
        checks.append(
            CheckResult(
                code="B60_BASELINE_DECISION_INVALID",
                message="PBD approval cannot be signed after its effective time",
                location=_display_path(paths, decision_path),
            )
        )
        return
    confirmation_path = _validate_file_reference(
        checks,
        paths,
        confirmation_reference,
        code="B60_BASELINE_CONFIRMATION_INVALID",
        location="baseline_confirmation",
        prefix=PurePosixPath("data/governance/evidence"),
    )
    if confirmation_path is None:
        return
    try:
        confirmation = _json_object(confirmation_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_BASELINE_CONFIRMATION_INVALID",
                message="Candidate and confirmation must be JSON objects",
                location="baseline_decision",
            )
        )
        return
    confirmation_valid = (
        confirmation.get("confirmation_type") == "project_baseline_change_review"
        and confirmation.get("decision_id") == DECISION_ID
        and confirmation.get("decision") == REQUESTED_DECISION
        and confirmation.get("template_only") is False
        and confirmation.get("candidate") == canonical_candidate
        and confirmation.get("authorized_transcription") is True
        and confirmation.get("authorized_actions")
        == {
            "write_decision_record": True,
            "activate_basic60_private_profile": True,
            "start_real_data_private_trial": True,
        }
        and _aware_timestamp(confirmation.get("reviewed_at"))
        and isinstance(confirmation.get("reviewer_signature"), dict)
        and confirmation["reviewer_signature"].get("person_name") == "kevin"
        and confirmation["reviewer_signature"].get("role") == "项目批准人"
        and confirmation["reviewer_signature"].get("signed_at") == confirmation.get("reviewed_at")
        and approval_dict.get("signed_at") == confirmation.get("reviewed_at")
        and approval_dict.get("signature_sha256") == sha256_file(confirmation_path)
    )
    if not confirmation_valid:
        checks.append(
            CheckResult(
                code="B60_BASELINE_CONFIRMATION_INVALID",
                message="Decision is not bound to the exact named approval confirmation",
                location=_display_path(paths, confirmation_path),
            )
        )


def _validate_baseline_clarification(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    clarification_reference: Any,
    candidate_reference: Any,
    confirmation_reference: Any,
    decision_reference: Any,
) -> None:
    location = "baseline_clarification"
    if not (
        isinstance(clarification_reference, dict)
        and clarification_reference.get("path") == PBD_CLARIFICATION_PATH.as_posix()
    ):
        checks.append(
            CheckResult(
                code="B60_BASELINE_CLARIFICATION_INVALID",
                message="Bundle must bind the canonical approved D1 review clarification",
                location=location,
            )
        )
        return
    clarification_path = _validate_file_reference(
        checks,
        paths,
        clarification_reference,
        code="B60_BASELINE_CLARIFICATION_INVALID",
        location=location,
        prefix=PurePosixPath("data/governance/review"),
    )
    if clarification_path is None:
        return
    try:
        clarification = _json_object(clarification_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        clarification = {}
    effective_from = clarification.get("effective_from")
    approval = clarification.get("approval")
    valid = (
        clarification.get("schema_version") == 1
        and clarification.get("record_type") == "basic60_private_internal_review_clarification"
        and clarification.get("decision_id") == DECISION_ID
        and clarification.get("status") == "approved_and_effective"
        and _aware_timestamp(effective_from)
        and clarification.get("basis")
        == {
            "candidate": candidate_reference,
            "confirmation": confirmation_reference,
            "decision": decision_reference,
        }
        and clarification.get("owner_confirmations")
        == {
            "external_information_accuracy_confirmed": True,
            "collected_data_origin": "official_public_sources",
            "manual_data_updates_allowed": True,
            "simplified_internal_review_authorized": True,
        }
        and clarification.get("d1_internal_review")
        == {
            "review_unit": "provider_dataset",
            "required_primary_source_count": 3,
            "alternative_sources_required": False,
            "per_country_manual_review_required": False,
            "approval_mode": "single_consolidated_project_approver",
            "required_approver": {
                "person_name": "kevin",
                "role": "项目批准人",
            },
            "machine_checks_retained": [
                "official_source_reference",
                "terms_and_license_evidence",
                "snapshot_hash",
                "attribution",
                "usage_boundary",
            ],
        }
        and clarification.get("manual_update_policy")
        == {
            "mode": "manual_incremental",
            "routine_value_updates_reuse_d1_admission": True,
            "required_each_update": [
                "new_snapshot_hash",
                "difference_report",
                "source_reference",
                "operator_record",
            ],
            "d1_reapproval_triggers": [
                "source_url_changed",
                "terms_or_license_changed",
                "dataset_structure_changed",
            ],
        }
        and clarification.get("v1_usage_boundaries")
        == {
            "private_structured_storage": "allowed",
            "private_display": "allowed",
            "public_release": "prohibited",
            "private_export": "prohibited",
            "ai_indexing": "prohibited",
            "model_training": "prohibited",
            "cloud_model_processing": "prohibited",
        }
        and clarification.get("does_not_complete")
        == [
            "D2-P",
            "D3-P",
            "D4-P",
            "formal_D1_D4",
            "P0",
            "private_trial_ready",
            "production_release",
        ]
        and isinstance(approval, dict)
        and approval
        == {
            "person_name": "kevin",
            "role": "项目批准人",
            "approved_at": effective_from,
            "authorization_basis": "explicit_user_confirmation_in_project_task",
        }
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_BASELINE_CLARIFICATION_INVALID",
                message=(
                    "D1 simplification must retain the approved three-primary, single-project-"
                    "approver, manual-update, and V1 usage-boundary clarification"
                ),
                location=location,
            )
        )


def _validate_usage_amendment(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
) -> None:
    location = "usage_amendment"
    if not (
        isinstance(reference, dict) and reference.get("path") == USAGE_AMENDMENT_PATH.as_posix()
    ):
        checks.append(
            CheckResult(
                code="B60_USAGE_AMENDMENT_MISSING",
                message=(
                    "Bundle must bind the canonical PBD-BASIC60-PRIVATE-001-A1 manual "
                    "usage amendment"
                ),
                location=location,
            )
        )
        return
    amendment_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_USAGE_AMENDMENT_INVALID",
        location=location,
        prefix=PurePosixPath("data/governance/review"),
    )
    if amendment_path is None:
        return
    try:
        amendment = _json_object(amendment_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        amendment = {}
    effective_from = amendment.get("effective_from")
    approval = amendment.get("approval")
    evidence = approval.get("evidence") if isinstance(approval, dict) else None
    evidence_path = _validate_file_reference(
        checks,
        paths,
        evidence,
        code="B60_USAGE_AMENDMENT_EVIDENCE_INVALID",
        location=f"{location}.approval.evidence",
        prefix=PurePosixPath("data/governance/evidence"),
    )
    signed_at = approval.get("signed_at") if isinstance(approval, dict) else None
    chronology_valid = False
    if _aware_timestamp(effective_from) and _aware_timestamp(signed_at):
        chronology_valid = datetime.fromisoformat(str(signed_at).replace("Z", "+00:00")) <= (
            datetime.fromisoformat(str(effective_from).replace("Z", "+00:00"))
        )
    valid = (
        set(amendment) == set(build_usage_amendment_template())
        and amendment.get("schema_version") == USAGE_AMENDMENT_SCHEMA
        and amendment.get("record_type") == "basic60_private_usage_amendment"
        and amendment.get("amendment_id") == USAGE_AMENDMENT_ID
        and amendment.get("amends_decision_id") == DECISION_ID
        and amendment.get("status") == "approved_and_effective"
        and _aware_timestamp(effective_from)
        and amendment.get("owner_authorization") == _owner_usage_authorization()
        and amendment.get("source_label_policy")
        == {
            "ordinary_ui_api": "omit_source_labels_and_provenance",
            "backend_terms_and_provenance": "retain_for_audit_only",
        }
        and amendment.get("processing_controls")
        == {
            "private_internal_purpose_only": True,
            "approved_model_configuration_required": True,
            "provider_training_allowed": False,
            "provider_retention_allowed": False,
            "raw_terms_or_source_evidence_allowed_in_model_payload": False,
        }
        and amendment.get("v1_runtime_capabilities") == _v1_runtime_capabilities()
        and amendment.get("does_not_complete")
        == [
            "D2-P",
            "D4-P",
            "formal_D1_D4",
            "P0",
            "private_trial_ready",
            "production_release",
        ]
        and amendment.get("template_only") is False
        and isinstance(approval, dict)
        and set(approval)
        == {
            "person_name",
            "role",
            "decision",
            "signed_at",
            "authorization_basis",
            "evidence",
        }
        and approval.get("person_name") == "kevin"
        and approval.get("role") == "项目批准人"
        and approval.get("decision") == "approved"
        and approval.get("authorization_basis") == "explicit_user_confirmation_in_project_task"
        and evidence_path is not None
        and chronology_valid
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_USAGE_AMENDMENT_INVALID",
                message=(
                    "A1 must be the exact kevin-approved all-field manual permission, retain "
                    "backend traceability, and keep every V1 AI runtime capability disabled"
                ),
                location=location,
            )
        )


def _validate_ai_usage_policy(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    bundle_reference: Any,
    source_reference: Any,
    seed: dict[str, Any],
    seed_reference: Any,
    usage_amendment_reference: Any,
) -> None:
    location = "ai_usage_policy"
    if bundle_reference != source_reference:
        checks.append(
            CheckResult(
                code="B60_AI_USAGE_POLICY_BINDING_INVALID",
                message="Bundle and D1 source admission must bind the same field-level policy",
                location=location,
            )
        )
        return
    policy_path = _validate_file_reference(
        checks,
        paths,
        bundle_reference,
        code="B60_AI_USAGE_POLICY_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/candidates"),
    )
    if policy_path is None:
        return
    try:
        policy = _json_object(policy_path)
        expected = build_ai_usage_policy(
            seed,
            seed_reference=seed_reference,
            usage_amendment_reference=usage_amendment_reference,
        )
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        policy = {}
        expected = {"invalid": True}
    if policy != expected:
        checks.append(
            CheckResult(
                code="B60_AI_USAGE_POLICY_INVALID",
                message=(
                    "Machine projection must bind the exact A1 and candidate seed, cover every "
                    "normalized Basic60 field, and contain no second permission decision"
                ),
                location=location,
            )
        )


def _validate_profile_reference(
    checks: list[CheckResult], paths: RepositoryPaths, reference: Any
) -> None:
    profile_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_PROFILE_INVALID",
        location="profile",
        prefix=PurePosixPath("data/basic60/candidates"),
    )
    if profile_path is None:
        return
    try:
        profile = _json_object(profile_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        profile = {}
    if profile != build_profile():
        checks.append(
            CheckResult(
                code="B60_PROFILE_CONTENT_INVALID",
                message=(
                    "Basic60 profile must preserve Basic coverage and the exact five-country "
                    "future deep-dive priority subset"
                ),
                location="profile",
            )
        )


def _validate_source_matrix(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    source: Any,
) -> None:
    if not isinstance(source, dict):
        return
    location = "source_admission.country_domain_matrix"
    matrix_path = _validate_file_reference(
        checks,
        paths,
        source.get("country_domain_matrix"),
        code="B60_SOURCE_MATRIX_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/candidates"),
    )
    if matrix_path is None:
        return
    try:
        matrix = _json_object(matrix_path)
        expected = build_basic60_source_matrix_candidate(raw_dir)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        checks.append(
            CheckResult(code="B60_SOURCE_MATRIX_INVALID", message=str(error), location=location)
        )
        return
    if matrix != expected:
        checks.append(
            CheckResult(
                code="B60_SOURCE_MATRIX_CONTENT_INVALID",
                message="Country-domain matrix must be the deterministic 60 by 3 Basic matrix",
                location=location,
            )
        )
        return
    domains = source.get("domains")
    if not isinstance(domains, list):
        return
    try:
        expected_derived_hash = _derived_source_matrix_sha256(matrix, domains)
    except (KeyError, TypeError):
        expected_derived_hash = None
    if source.get("derived_matrix_sha256") != expected_derived_hash:
        checks.append(
            CheckResult(
                code="B60_SOURCE_MATRIX_BINDING_INVALID",
                message="Approved source selection is not bound to all 180 matrix rows",
                location="source_admission.derived_matrix_sha256",
            )
        )


def _validate_domains(checks: list[CheckResult], source: Any) -> None:
    if not isinstance(source, dict):
        checks.append(
            CheckResult(
                code="B60_SOURCE_ADMISSION_INVALID",
                message="Source admission section is required",
                location="source_admission",
            )
        )
        return
    domains = source.get("domains")
    if not isinstance(domains, list):
        checks.append(
            CheckResult(
                code="B60_SOURCE_DOMAINS_INVALID",
                message="Source domains must be a list",
                location="source_admission.domains",
            )
        )
        return
    indexed = {str(item.get("domain_id")): item for item in domains if isinstance(item, dict)}
    if set(indexed) != set(DOMAIN_IDS) or len(domains) != len(DOMAIN_IDS):
        checks.append(
            CheckResult(
                code="B60_SOURCE_DOMAINS_INVALID",
                message="Source admission must contain exactly the three Basic60 domains",
                location="source_admission.domains",
            )
        )
    for domain_id in DOMAIN_IDS:
        item = indexed.get(domain_id, {})
        primary = str(item.get("primary_source_id") or "").strip()
        if (
            not primary
            or "alternative_source_id" not in item
            or item.get("alternative_source_id") is not None
            or item.get("status") != "active"
        ):
            checks.append(
                CheckResult(
                    code="B60_SOURCE_DOMAIN_NOT_ACTIVE",
                    message=(
                        f"{domain_id} requires one active primary source and a null "
                        "alternative_source_id"
                    ),
                    location=f"source_admission.domains.{domain_id}",
                )
            )
    boundaries = source.get("usage_boundaries")
    expected = _usage_boundaries()
    if boundaries != expected:
        checks.append(
            CheckResult(
                code="B60_SOURCE_BOUNDARY_INVALID",
                message=(
                    "A1 must authorize all-field private display and controlled model processing "
                    "while the V1 AI runtime remains disabled"
                ),
                location="source_admission.usage_boundaries",
            )
        )


def _validate_evidence_snapshot_object(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
    *,
    l0_dir: Path | None,
    location: str,
    official_terms: bool = False,
) -> dict[str, Any] | None:
    if not isinstance(reference, dict):
        checks.append(
            CheckResult(
                code="B60_EVIDENCE_OBJECT_INVALID",
                message="Evidence object reference is required",
                location=location,
            )
        )
        return None
    storage_mode = reference.get("storage_mode")
    common = {
        "storage_mode",
        "sha256",
        "byte_size",
        "captured_at",
        "retrieval_uri",
        "media_type",
    }
    if official_terms:
        common |= {"canonical_url", "capture_method"}
    expected_fields = (
        common | {"volume_id", "object_key"}
        if storage_mode == "external_l0"
        else common | {"path", "public_and_redacted"}
    )
    valid_metadata = (
        storage_mode in EVIDENCE_OBJECT_STORAGE_MODES
        and set(reference) == expected_fields
        and SHA256_PATTERN.fullmatch(str(reference.get("sha256") or "")) is not None
        and isinstance(reference.get("byte_size"), int)
        and not isinstance(reference.get("byte_size"), bool)
        and reference["byte_size"] >= 0
        and _aware_timestamp(reference.get("captured_at"))
        and str(reference.get("retrieval_uri") or "").strip()
        and str(reference.get("media_type") or "").strip()
    )
    if official_terms:
        try:
            retrieval_uri = _canonical_source_uri(reference.get("retrieval_uri"))
            canonical_url = _canonical_source_uri(reference.get("canonical_url"))
        except (UnicodeError, ValueError):
            retrieval_uri = None
            canonical_url = None
        valid_metadata = bool(
            valid_metadata
            and storage_mode == "external_l0"
            and reference.get("capture_method") in SOURCE_TERMS_CAPTURE_METHODS
            and retrieval_uri == canonical_url
            and str(reference.get("media_type") or "").split(";", 1)[0].strip().lower()
            in {"text/html", "text/plain"}
        )
    object_path: Path | None = None
    if storage_mode == "external_l0":
        digest = str(reference.get("sha256") or "")
        object_key = str(reference.get("object_key") or "")
        valid_metadata = bool(
            valid_metadata
            and l0_dir is not None
            and VOLUME_ID_PATTERN.fullmatch(str(reference.get("volume_id") or ""))
            and object_key == f"sha256/{digest}"
        )
        if l0_dir is not None and object_key:
            object_path = l0_dir / "sha256" / digest
    elif storage_mode == "repository_public":
        valid_metadata = bool(valid_metadata and reference.get("public_and_redacted") is True)
        path_value = reference.get("path")
        if _valid_relative_path(path_value, PurePosixPath("data/basic60/evidence")):
            pure = PurePosixPath(str(path_value))
            object_path = paths.root.joinpath(*pure.parts)
    if not valid_metadata or object_path is None:
        checks.append(
            CheckResult(
                code="B60_EVIDENCE_OBJECT_INVALID",
                message=(
                    "Official terms evidence must be a protected external L0 original"
                    if official_terms
                    else (
                        "Evidence object must be an external L0 object or an explicitly public, "
                        "redacted repository object"
                    )
                ),
                location=location,
            )
        )
        return None
    try:
        resolved = object_path.resolve(strict=True)
        actual_size = resolved.stat().st_size
        actual_mode = stat.S_IMODE(resolved.stat().st_mode)
    except OSError:
        checks.append(
            CheckResult(
                code="B60_EVIDENCE_OBJECT_MISSING",
                message="Referenced evidence object cannot be opened",
                location=location,
            )
        )
        return None
    repository_public = storage_mode == "repository_public"
    safe = (
        not object_path.is_symlink()
        and resolved.is_file()
        and (
            resolved.is_relative_to((paths.root / "data/basic60/evidence").resolve())
            if repository_public
            else l0_dir is not None and resolved.is_relative_to((l0_dir / "sha256").resolve())
        )
        and actual_size == reference.get("byte_size")
        and sha256_file(resolved) == reference.get("sha256")
        and (repository_public or actual_mode == 0o444)
    )
    if not safe:
        checks.append(
            CheckResult(
                code="B60_EVIDENCE_OBJECT_CONTENT_INVALID",
                message="Evidence object hash, size, location, or read-only state changed",
                location=location,
            )
        )
        return None
    return reference


def _validate_source_evidence(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    source_ref: str,
    item: dict[str, Any],
    *,
    l0_dir: Path | None,
) -> None:
    location = f"source_admission.source_registry.{source_ref}.evidence"
    evidence_path = _validate_file_reference(
        checks,
        paths,
        item.get("evidence"),
        code="B60_SOURCE_EVIDENCE_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if evidence_path is None:
        return
    evidence_reference = item.get("evidence")
    if not isinstance(evidence_reference, dict) or item.get(
        "evidence_sha256"
    ) != evidence_reference.get("sha256"):
        checks.append(
            CheckResult(
                code="B60_SOURCE_EVIDENCE_HASH_INVALID",
                message="Source evidence_sha256 must equal the referenced evidence file hash",
                location=location,
            )
        )
        return
    try:
        evidence = _json_object(evidence_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_SOURCE_EVIDENCE_INVALID",
                message="Source evidence must be a JSON object",
                location=location,
            )
        )
        return
    expected_fields = {
        "schema_version",
        "source_ref",
        "terms_snapshots",
        "captured_at",
    }
    snapshots = item.get("snapshots")
    captured_at = (
        snapshots[0].get("captured_at")
        if isinstance(snapshots, list) and snapshots and isinstance(snapshots[0], dict)
        else None
    )
    raw_terms_snapshots = evidence.get("terms_snapshots")
    terms_snapshots: list[dict[str, Any]] = []
    if isinstance(raw_terms_snapshots, list):
        for index, reference in enumerate(raw_terms_snapshots):
            validated = _validate_evidence_snapshot_object(
                checks,
                paths,
                reference,
                l0_dir=l0_dir,
                location=f"{location}.terms_snapshots.{index}",
                official_terms=True,
            )
            if isinstance(validated, dict):
                terms_snapshots.append(validated)
    else:
        checks.append(
            CheckResult(
                code="B60_SOURCE_TERMS_ORIGINALS_INVALID",
                message="Source evidence must bind a list of official terms originals",
                location=f"{location}.terms_snapshots",
            )
        )
    required_uris = SOURCE_TERMS_URI_REQUIREMENTS.get(str(item.get("data_domain")), ())
    try:
        actual_uris = Counter(
            _canonical_source_uri(snapshot.get("canonical_url")) for snapshot in terms_snapshots
        )
        expected_uris = Counter(_canonical_source_uri(uri) for uri in required_uris)
    except (UnicodeError, ValueError):
        actual_uris = Counter()
        expected_uris = Counter(required_uris)
    if actual_uris != expected_uris or len(terms_snapshots) != len(required_uris):
        checks.append(
            CheckResult(
                code="B60_SOURCE_TERMS_ORIGINALS_INVALID",
                message=(
                    "Terms originals do not cover the exact official providers required for "
                    f"{item.get('data_domain')}"
                ),
                location=f"{location}.terms_snapshots",
            )
        )
    valid = (
        set(evidence) == expected_fields
        and evidence.get("schema_version") == SOURCE_EVIDENCE_SCHEMA
        and evidence.get("source_ref") == source_ref
        and evidence.get("captured_at") == captured_at
        and _aware_timestamp(evidence.get("captured_at"))
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_SOURCE_EVIDENCE_CONTENT_INVALID",
                message=(
                    "Source audit evidence does not match its registry entry and exact terms "
                    "snapshot originals"
                ),
                location=location,
            )
        )


def _validate_country_source_uri_manifest(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    registry: dict[str, Any],
) -> dict[str, Any] | None:
    location = "source_admission.source_registry.country_source_uri_manifest"
    manifest_path = _validate_file_reference(
        checks,
        paths,
        registry.get("country_source_uri_manifest"),
        code="B60_SOURCE_URI_MANIFEST_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/candidates"),
    )
    try:
        current = build_basic60_country_source_uri_manifest(raw_dir)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        checks.append(
            CheckResult(
                code="B60_SOURCE_URI_MANIFEST_INVALID",
                message=str(error),
                location=location,
            )
        )
        return None
    if manifest_path is None:
        return current
    try:
        approved = _json_object(manifest_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        approved = {}
    if approved != current:
        checks.append(
            CheckResult(
                code="B60_SOURCE_URI_MANIFEST_CHANGED",
                message=(
                    "Country-profile provider/field/source-URI manifest changed after D1 approval"
                ),
                location=location,
            )
        )
    return current


def _validate_source_registry(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    source: Any,
    *,
    l0_dir: Path | None = None,
) -> dict[str, Any] | None:
    if not isinstance(source, dict):
        return None
    registry_path = _validate_file_reference(
        checks,
        paths,
        source.get("source_registry"),
        code="B60_SOURCE_REGISTRY_INVALID",
        location="source_admission.source_registry",
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if registry_path is None:
        return None
    try:
        registry = _json_object(registry_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_SOURCE_REGISTRY_INVALID",
                message="Source registry evidence must be a JSON object",
                location="source_admission.source_registry",
            )
        )
        return None
    sources = registry.get("sources")
    declaration_valid = (
        registry.get("source_declaration") == "official_public_data"
        and registry.get("update_policy")
        == {
            "mode": "manual_incremental",
            "reapproval_triggers": list(D1_REAPPROVAL_TRIGGERS),
        }
        and set(registry)
        == {
            "schema_version",
            "profile_id",
            "release_id",
            "template_only",
            "source_declaration",
            "update_policy",
            "country_source_uri_manifest",
            "sources",
        }
    )
    if (
        registry.get("template_only") is not False
        or not isinstance(sources, list)
        or not declaration_valid
    ):
        checks.append(
            CheckResult(
                code="B60_SOURCE_REGISTRY_INVALID",
                message=(
                    "Copy the source registry template and retain the official-public-source "
                    "manual-update declaration"
                ),
                location="source_admission.source_registry",
            )
        )
        return None
    country_source_manifest = _validate_country_source_uri_manifest(
        checks,
        paths,
        raw_dir,
        registry,
    )
    indexed = {str(item.get("source_ref")): item for item in sources if isinstance(item, dict)}
    required_roles = {
        (domain, role)
        for domain in ("country_identity", "macro", "energy")
        for role in ("primary",)
    }
    actual_roles = {
        (str(item.get("data_domain")), str(item.get("source_role"))) for item in indexed.values()
    }
    if (
        actual_roles != required_roles
        or len(sources) != len(required_roles)
        or len(indexed) != len(required_roles)
    ):
        checks.append(
            CheckResult(
                code="B60_SOURCE_ALTERNATIVES_INCOMPLETE",
                message="Each Basic60 domain requires exactly one active primary dataset",
                location="source_admission.source_registry",
            )
        )
    for source_ref, item in indexed.items():
        snapshots = item.get("snapshots")
        source_fields_valid = set(item) == {
            "source_ref",
            "provider",
            "dataset",
            "data_domain",
            "source_role",
            "status",
            "terms_uri",
            "evidence_sha256",
            "evidence",
            "access_method",
            "rate_limit",
            "snapshots",
            "admission_fingerprint",
        }
        stored_fingerprint = item.get("admission_fingerprint")
        try:
            expected_fingerprint = _source_admission_fingerprint(
                raw_dir,
                item,
                country_source_manifest=country_source_manifest or {},
                terms_evidence=_source_evidence_for_fingerprint(paths, item),
            )
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError, KeyError):
            expected_fingerprint = {}
        fingerprint_location = (
            f"source_admission.source_registry.sources.{source_ref}.admission_fingerprint"
        )
        fingerprint_keys_valid = isinstance(stored_fingerprint, dict) and set(
            stored_fingerprint
        ) == set(expected_fingerprint)
        if not fingerprint_keys_valid:
            checks.append(
                CheckResult(
                    code="B60_SOURCE_ADMISSION_FINGERPRINT_INVALID",
                    message=f"Source admission fingerprint is missing or malformed: {source_ref}",
                    location=fingerprint_location,
                )
            )
        stored_dict = stored_fingerprint if isinstance(stored_fingerprint, dict) else {}
        metadata_fields = {
            "dataset_path",
            "record_selector",
            "source_uri_mode",
            "source_uri_field",
        }
        if any(stored_dict.get(key) != expected_fingerprint.get(key) for key in metadata_fields):
            checks.append(
                CheckResult(
                    code="B60_SOURCE_ADMISSION_FINGERPRINT_INVALID",
                    message=f"Source admission fingerprint metadata changed: {source_ref}",
                    location=fingerprint_location,
                )
            )
        if any(
            stored_dict.get(key) != expected_fingerprint.get(key)
            for key in ("schema_field_count", "schema_fingerprint_sha256")
        ):
            checks.append(
                CheckResult(
                    code="B60_SOURCE_SCHEMA_FINGERPRINT_MISMATCH",
                    message=f"Raw CSV header/schema changed after D1 approval: {source_ref}",
                    location=fingerprint_location,
                )
            )
        if any(
            stored_dict.get(key) != expected_fingerprint.get(key)
            for key in ("source_uri_entry_count", "source_uri_fingerprint_sha256")
        ):
            checks.append(
                CheckResult(
                    code="B60_SOURCE_URI_FINGERPRINT_MISMATCH",
                    message=f"Source URI set changed after D1 approval: {source_ref}",
                    location=fingerprint_location,
                )
            )
        terms_fingerprint = expected_fingerprint.get("terms_fingerprint_sha256")
        if (
            not isinstance(terms_fingerprint, str)
            or not SHA256_PATTERN.fullmatch(terms_fingerprint)
            or stored_dict.get("terms_fingerprint_sha256") != terms_fingerprint
        ):
            checks.append(
                CheckResult(
                    code="B60_SOURCE_TERMS_FINGERPRINT_MISMATCH",
                    message=f"Terms URI or approved terms snapshots changed: {source_ref}",
                    location=fingerprint_location,
                )
            )
        fingerprint_valid = bool(
            fingerprint_keys_valid
            and stored_dict == expected_fingerprint
            and isinstance(terms_fingerprint, str)
            and SHA256_PATTERN.fullmatch(terms_fingerprint)
        )
        valid = (
            bool(source_ref)
            and item.get("status") == "active"
            and source_fields_valid
            and str(item.get("provider") or "").strip()
            and str(item.get("dataset") or "").strip()
            and str(item.get("terms_uri") or "").strip()
            and SHA256_PATTERN.fullmatch(str(item.get("evidence_sha256") or ""))
            and str(item.get("access_method") or "").strip()
            and str(item.get("rate_limit") or "").strip()
            and isinstance(snapshots, list)
            and bool(snapshots)
            and fingerprint_valid
            and all(
                isinstance(snapshot, dict)
                and str(snapshot.get("snapshot_ref") or "").strip()
                and _aware_timestamp(snapshot.get("captured_at"))
                and SHA256_PATTERN.fullmatch(str(snapshot.get("content_sha256") or ""))
                and str(snapshot.get("retrieval_uri") or "").strip()
                and str(snapshot.get("media_type") or "").strip()
                for snapshot in snapshots
            )
        )
        if not valid:
            checks.append(
                CheckResult(
                    code="B60_SOURCE_REGISTRY_ENTRY_INVALID",
                    message=(
                        "Source registry audit entry is incomplete or contains non-audit "
                        f"permission fields: {source_ref}"
                    ),
                    location="source_admission.source_registry",
                )
            )
        _validate_source_evidence(checks, paths, source_ref, item, l0_dir=l0_dir)
    domains = source.get("domains")
    domain_map = {
        "country_profile": "country_identity",
        "macroeconomic": "macro",
        "energy": "energy",
    }
    if isinstance(domains, list):
        for item in domains:
            if not isinstance(item, dict):
                continue
            expected_domain = domain_map.get(str(item.get("domain_id")))
            for field, role in (("primary_source_id", "primary"),):
                registered = indexed.get(str(item.get(field)))
                if (
                    registered is None
                    or registered.get("data_domain") != expected_domain
                    or registered.get("source_role") != role
                ):
                    checks.append(
                        CheckResult(
                            code="B60_SOURCE_DOMAIN_BINDING_INVALID",
                            message=f"{item.get('domain_id')} {role} source is not in the registry",
                            location="source_admission.domains",
                        )
                    )
    d1_change_codes = {
        "B60_SOURCE_URI_MANIFEST_CHANGED",
        "B60_SOURCE_SCHEMA_FINGERPRINT_MISMATCH",
        "B60_SOURCE_URI_FINGERPRINT_MISMATCH",
        "B60_SOURCE_TERMS_FINGERPRINT_MISMATCH",
    }
    if any(item.code in d1_change_codes for item in checks) and not any(
        item.code == "B60_D1_REAPPROVAL_REQUIRED" for item in checks
    ):
        checks.append(
            CheckResult(
                code="B60_D1_REAPPROVAL_REQUIRED",
                message=(
                    "Provider/source URL, terms, schema, or usage-boundary drift requires a "
                    "new human D1 scope authorization"
                ),
                location="source_admission",
            )
        )
    return registry


def _basic60_l0_authorization_preflight(
    paths: RepositoryPaths,
    raw_dir: Path,
    manifest_path: Path,
    bundle_path: Path,
    l0_dir: Path,
) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    checks: list[CheckResult] = []
    repository_root = paths.root.resolve()
    review_root = (paths.root / "data/basic60/review").resolve(strict=False)
    for name, candidate in (("raw", raw_dir), ("manifest", manifest_path), ("bundle", bundle_path)):
        if _has_symlink_component(candidate):
            checks.append(
                CheckResult(
                    code="B60_L0_INPUT_SYMLINK",
                    message=f"{name} path cannot contain symbolic links",
                    location=str(candidate),
                )
            )
    if not raw_dir.is_dir():
        checks.append(
            CheckResult(
                code="B60_L0_RAW_INVALID",
                message="Raw input directory does not exist",
                location=str(raw_dir),
            )
        )
    for name, candidate in (("manifest", manifest_path), ("bundle", bundle_path)):
        try:
            resolved = candidate.resolve(strict=True)
        except OSError:
            checks.append(
                CheckResult(
                    code="B60_L0_INPUT_MISSING",
                    message=f"{name} file does not exist",
                    location=str(candidate),
                )
            )
            continue
        if not resolved.is_relative_to(repository_root) or not resolved.is_file():
            checks.append(
                CheckResult(
                    code="B60_L0_INPUT_PATH_INVALID",
                    message=f"{name} must be a regular repository file",
                    location=str(candidate),
                )
            )
        if name == "bundle" and not resolved.is_relative_to(review_root):
            checks.append(
                CheckResult(
                    code="B60_L0_BUNDLE_NOT_REVIEW_COPY",
                    message="D1-approved bundle must be a copy under data/basic60/review",
                    location=str(candidate),
                )
            )
    if checks:
        details = "; ".join(f"{item.code}: {item.message}" for item in checks)
        raise ValueError(f"B60_L0_PREFLIGHT_FAILED: {details}")

    try:
        bundle = _json_object(bundle_path)
        stored_manifest = _json_object(manifest_path)
        _summary, raw_checks = inspect_basic60_raw(raw_dir)
        checks.extend(raw_checks)
        current_manifest = build_raw_manifest(paths, raw_dir)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        raise ValueError(f"B60_L0_PREFLIGHT_FAILED: {error}") from error

    if not (
        bundle.get("schema_version") == SCHEMA_VERSION
        and bundle.get("profile_id") == PROFILE_ID
        and bundle.get("release_id") == RELEASE_ID
        and bundle.get("formal_gate_status") == FORMAL_GATE_STATUS
        and bundle.get("template_only") is False
    ):
        checks.append(
            CheckResult(
                code="B60_L0_BUNDLE_HEADER_INVALID",
                message="Materialization requires a completed Basic60 D1 review copy",
                location="bundle",
            )
        )
    if bundle.get("revocation") is not None:
        checks.append(
            CheckResult(
                code="B60_L0_REVOKED",
                message="A revoked Basic60 packet cannot authorize L0 materialization",
                location="revocation",
            )
        )
    for field in (
        "root_sha256",
        "included_file_count",
        "included_byte_count",
        "excluded_zone_identifier_count",
        "unsafe_entry_count",
        "scoped_files",
    ):
        if stored_manifest.get(field) != current_manifest.get(field):
            checks.append(
                CheckResult(
                    code="B60_RAW_MANIFEST_DRIFT",
                    message=f"Raw manifest field changed: {field}",
                    location=_display_path(paths, manifest_path),
                )
            )
    if current_manifest.get("unsafe_entry_count") != 0:
        checks.append(
            CheckResult(
                code="B60_RAW_UNSAFE_ENTRY",
                message="Raw material contains unsafe symlink entries",
                location=str(raw_dir),
            )
        )
    if bundle.get("raw_root_sha256") != current_manifest.get("root_sha256"):
        checks.append(
            CheckResult(
                code="B60_BUNDLE_RAW_BINDING_INVALID",
                message="D1 bundle is not bound to the current raw root",
                location="raw_root_sha256",
            )
        )
    try:
        expected_manifest_reference = _reference(manifest_path, paths)
    except (OSError, ValueError) as error:
        checks.append(
            CheckResult(
                code="B60_BUNDLE_MANIFEST_BINDING_INVALID",
                message=str(error),
                location="raw_manifest",
            )
        )
    else:
        if bundle.get("raw_manifest") != expected_manifest_reference:
            checks.append(
                CheckResult(
                    code="B60_BUNDLE_MANIFEST_BINDING_INVALID",
                    message="D1 bundle is not bound to the exact raw manifest",
                    location="raw_manifest",
                )
            )

    _validate_baseline_decision(
        checks,
        paths,
        bundle.get("baseline_decision"),
        bundle.get("baseline_confirmation"),
        bundle.get("baseline_candidate"),
        bundle.get("baseline_confirmation_template"),
    )
    _validate_baseline_clarification(
        checks,
        paths,
        bundle.get("baseline_clarification"),
        bundle.get("baseline_candidate"),
        bundle.get("baseline_confirmation"),
        bundle.get("baseline_decision"),
    )
    source = bundle.get("source_admission")
    _validate_profile_reference(checks, paths, bundle.get("profile"))
    _validate_domains(checks, source)
    _validate_source_matrix(checks, paths, raw_dir, source)
    source_registry = _validate_source_registry(
        checks,
        paths,
        raw_dir,
        source,
        l0_dir=l0_dir,
    )
    if isinstance(source, dict):
        _validate_d1_two_stage_admission(
            checks,
            paths,
            bundle,
            source,
            source_registry,
            l0_dir,
        )
    else:
        checks.append(
            CheckResult(
                code="B60_SOURCE_ADMISSION_INVALID",
                message="Source admission section is required",
                location="source_admission",
            )
        )
    if checks:
        details = "; ".join(f"{item.code}: {item.message}" for item in checks)
        raise ValueError(f"B60_L0_AUTHORIZATION_FAILED: {details}")
    if source_registry is None:
        raise ValueError("B60_L0_AUTHORIZATION_FAILED: source registry is unavailable")
    return bundle, current_manifest, source_registry


def materialize_basic60_l0(
    paths: RepositoryPaths,
    raw_dir: Path,
    manifest_path: Path,
    bundle_path: Path,
    l0_dir: Path,
    evidence_dir: Path,
    volume_id: str,
) -> list[Path]:
    """Materialize only D1-approved Basic60 CSV inputs into an external immutable L0."""

    resolved_l0 = _validated_l0_directory(paths, raw_dir, l0_dir)
    bundle, current_manifest, source_registry = _basic60_l0_authorization_preflight(
        paths,
        raw_dir,
        manifest_path,
        bundle_path,
        resolved_l0,
    )
    if not VOLUME_ID_PATTERN.fullmatch(volume_id):
        raise ValueError(
            "B60_L0_VOLUME_ID_INVALID: volume-id must be a non-secret stable identifier"
        )
    resolved_evidence = _validated_evidence_directory(paths, evidence_dir)

    evidence_paths = [
        resolved_evidence / L0_BATCH_MANIFEST_NAMES[batch_id]
        for batch_id, _domain_id, _dataset_path in BATCH_SPECS
    ]
    for path in evidence_paths:
        if path.exists() or path.is_symlink():
            raise ValueError(f"B60_L0_EVIDENCE_EXISTS: refusing to overwrite {path}")

    scoped = {str(item["path"]): item for item in current_manifest["scoped_files"]}
    object_plans: dict[str, tuple[Path, int, Path]] = {}
    for _batch_id, _domain_id, dataset_path in BATCH_SPECS:
        source = raw_dir.joinpath(*dataset_path.parts)
        if source.is_symlink() or not source.is_file():
            raise ValueError(f"B60_L0_SOURCE_UNSAFE: approved source is not a file: {source}")
        expected = scoped.get(dataset_path.as_posix())
        if not isinstance(expected, dict):
            raise ValueError(
                f"B60_L0_SOURCE_NOT_SCOPED: dataset is absent from raw manifest: {dataset_path}"
            )
        digest = str(expected.get("sha256") or "")
        byte_size = int(expected.get("byte_size") or -1)
        if not SHA256_PATTERN.fullmatch(digest) or byte_size < 0:
            raise ValueError(f"B60_L0_SOURCE_INVALID: invalid scoped hash for {dataset_path}")
        destination = resolved_l0 / "sha256" / digest
        if destination.exists() or destination.is_symlink():
            _verify_l0_object(destination, digest, byte_size)
        object_plans.setdefault(digest, (source, byte_size, destination))

    object_root = resolved_l0 / "sha256"
    object_root.mkdir(parents=True, exist_ok=True, mode=0o755)
    resolved_evidence.mkdir(parents=True, exist_ok=True, mode=0o755)
    copied_by_digest: dict[str, bool] = {}
    for digest, (source, byte_size, destination) in object_plans.items():
        copied_by_digest[digest] = _materialize_content_object(
            source,
            destination,
            expected_hash=digest,
            expected_size=byte_size,
        )

    domains = {
        str(item.get("domain_id")): item
        for item in bundle["source_admission"]["domains"]
        if isinstance(item, dict)
    }
    registry_sources = {
        str(item.get("source_ref")): item
        for item in source_registry.get("sources", [])
        if isinstance(item, dict)
    }
    source_subject = approval_subjects(bundle)["source"]
    materialized_at = datetime.now(UTC).isoformat()
    newly_claimed: set[str] = set()
    manifests: list[dict[str, Any]] = []
    for batch_id, domain_id, dataset_path in BATCH_SPECS:
        source_item = scoped[dataset_path.as_posix()]
        digest = str(source_item["sha256"])
        byte_size = int(source_item["byte_size"])
        copied = copied_by_digest[digest] and digest not in newly_claimed
        newly_claimed.add(digest)
        domain = domains[domain_id]
        comparisons = []
        for field, role in (("primary_source_id", "primary"),):
            source_ref = str(domain[field])
            registry_source = registry_sources[source_ref]
            snapshots = registry_source["snapshots"]
            source_contract = registry_source["admission_fingerprint"]
            comparisons.append(
                {
                    "source_ref": source_ref,
                    "source_role": role,
                    "approved_snapshot_ref": snapshots[0]["snapshot_ref"],
                    "approved_terms_fingerprint_sha256": source_contract[
                        "terms_fingerprint_sha256"
                    ],
                    "current_terms_snapshots": [],
                    "current_terms_retrieval_logs": [],
                    "current_terms_fingerprint_sha256": None,
                    "comparison_status": "pending",
                }
            )
        manifests.append(
            {
                "schema_version": L0_BATCH_MANIFEST_SCHEMA,
                "profile_id": PROFILE_ID,
                "release_id": RELEASE_ID,
                "batch_id": batch_id,
                "run_id": "legacy_import_run",
                "copy_status": "succeeded",
                "batch_status": "materialized_not_closed",
                "domain_id": domain_id,
                "raw_root_sha256": current_manifest["root_sha256"],
                "source_subject_sha256": source_subject,
                "source_contract_sha256": _canonical_sha256(
                    registry_sources[str(domain["primary_source_id"])]["admission_fingerprint"]
                ),
                "approved_primary_source_id": domains[domain_id]["primary_source_id"],
                "approved_alternative_source_id": domains[domain_id]["alternative_source_id"],
                "source_dataset_path": dataset_path.as_posix(),
                "source_sha256": digest,
                "source_byte_size": byte_size,
                "volume_id": volume_id,
                "object_key": f"sha256/{digest}",
                "object_sha256": digest,
                "object_byte_size": byte_size,
                "copied": copied,
                "reused": not copied,
                "readback_verified": True,
                "immutable": True,
                "read_only": True,
                "posix_mode": "0444",
                "diff_status": "pending_external_snapshot_comparison",
                "current_snapshot": {"sha256": digest, "byte_size": byte_size},
                "diff_report": None,
                "checked_at": None,
                "source_snapshot_comparisons": comparisons,
                "watermark_advanced": False,
                "watermark": None,
                "access_signal": "none",
                "materialized_at": materialized_at,
            }
        )
    for destination, payload in zip(evidence_paths, manifests, strict=True):
        _atomic_write_json_no_overwrite(destination, payload)
    return evidence_paths


def collect_basic60_machine_evidence(
    paths: RepositoryPaths,
    seed_path: Path,
    manifest_path: Path,
    bundle_path: Path,
    operation_logs_path: Path,
    api_observations_path: Path,
    route_probes_path: Path,
    output_path: Path,
) -> Path:
    """Aggregate real runner/HTTP observations; never synthesize a passing observation."""

    review_root = (paths.root / "data/basic60/review").resolve(strict=False)
    bundle_resolved = bundle_path.resolve(strict=True)
    if (
        bundle_path.is_symlink()
        or not bundle_resolved.is_file()
        or not bundle_resolved.is_relative_to(review_root)
    ):
        raise ValueError("B60_MACHINE_BUNDLE_INVALID: bundle must be a D1-D3 review copy")
    runtime_root = (paths.root / "runtime/basic60").resolve(strict=False)
    output_resolved = output_path.resolve(strict=False)
    if (
        output_path.is_symlink()
        or not output_resolved.is_relative_to(runtime_root)
        or output_resolved.parent == runtime_root
    ):
        raise ValueError(
            "B60_MACHINE_OUTPUT_INVALID: output must be a specific path below runtime/basic60"
        )
    for input_path in (
        seed_path,
        manifest_path,
        operation_logs_path,
        api_observations_path,
        route_probes_path,
    ):
        if input_path.is_symlink() or not input_path.resolve(strict=True).is_file():
            raise ValueError(f"B60_MACHINE_INPUT_INVALID: unsafe input file: {input_path}")

    bundle = _json_object(bundle_path)
    operation_logs = _json_object(operation_logs_path)
    api_observations = _json_object(api_observations_path)
    route_probes_value = json.loads(route_probes_path.read_text(encoding="utf-8"))
    if not isinstance(route_probes_value, list):
        raise ValueError("B60_MACHINE_INPUT_INVALID: route probes must be a JSON list")
    stage_references: dict[str, Any] = {}
    first_hashes: dict[str, str] = {}
    second_hashes: dict[str, str] = {}
    processing = bundle.get("processing")
    pipelines = processing.get("pipelines") if isinstance(processing, dict) else None
    pipeline_index = {
        str(item.get("pipeline_id")): item
        for item in (pipelines if isinstance(pipelines, list) else [])
        if isinstance(item, dict)
    }
    for pipeline_id, stage_id in zip(PIPELINE_IDS, ("parse", "standardize", "entity"), strict=True):
        item = pipeline_index.get(pipeline_id)
        replay_reference = item.get("replay_manifest") if isinstance(item, dict) else None
        if not isinstance(replay_reference, dict) or not isinstance(
            replay_reference.get("path"), str
        ):
            raise ValueError(f"B60_MACHINE_D3_INVALID: missing {pipeline_id} replay reference")
        if not _valid_relative_path(
            replay_reference["path"], PurePosixPath("data/basic60/candidates")
        ):
            raise ValueError(f"B60_MACHINE_D3_INVALID: unsafe {pipeline_id} replay reference")
        replay_path = paths.root.joinpath(*PurePosixPath(replay_reference["path"]).parts)
        if (
            replay_path.is_symlink()
            or not replay_path.resolve(strict=True).is_file()
            or sha256_file(replay_path) != replay_reference.get("sha256")
        ):
            raise ValueError(f"B60_MACHINE_D3_INVALID: changed {pipeline_id} replay manifest")
        replay = _json_object(replay_path)
        hashes = replay.get("independent_replay_sha256")
        if not isinstance(hashes, dict):
            raise ValueError(f"B60_MACHINE_D3_INVALID: missing {pipeline_id} replay hashes")
        stage_references[stage_id] = replay.get("runtime_stage")
        first_hashes[stage_id] = str(hashes.get("first") or "")
        second_hashes[stage_id] = str(hashes.get("second") or "")
    try:
        payload = collect_machine_evidence(
            repository_root=paths.root,
            seed_path=seed_path,
            raw_manifest_path=manifest_path,
            bundle_path=bundle_path,
            d3_artifacts={
                "stages": stage_references,
                "replay": {
                    "first_output_sha256": first_hashes,
                    "second_output_sha256": second_hashes,
                },
            },
            operation_logs=operation_logs,
            api_observations=api_observations,
            route_probes=route_probes_value,
        )
    except EvidenceCollectionError as error:
        raise ValueError(f"B60_MACHINE_COLLECTION_FAILED: {error}") from error
    if output_resolved.exists():
        if output_resolved.is_symlink() or _json_object(output_resolved) != payload:
            raise ValueError("B60_MACHINE_OUTPUT_EXISTS: refusing to replace changed evidence")
        return output_resolved
    output_resolved.parent.mkdir(parents=True, exist_ok=True)
    _atomic_write_json_no_overwrite(output_resolved, payload)
    return output_resolved


def _validate_retrieval_log(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
    *,
    comparison: dict[str, Any],
    snapshot: dict[str, Any],
    registry_source: dict[str, Any],
    location: str,
) -> bool:
    log_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_RETRIEVAL_LOG_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if log_path is None:
        return False
    try:
        log = _json_object(log_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        log = {}
    request = log.get("request")
    response = log.get("response")
    valid = (
        set(log) == {"schema_version", "source_ref", "source_role", "request", "response"}
        and log.get("schema_version") == RETRIEVAL_LOG_SCHEMA
        and log.get("source_ref") == comparison.get("source_ref")
        and log.get("source_role") == comparison.get("source_role")
        and isinstance(request, dict)
        and set(request) == {"method", "uri", "requested_at", "access_method", "rate_limit"}
        and request.get("method") == "GET"
        and request.get("uri") == snapshot.get("retrieval_uri")
        and _aware_timestamp(request.get("requested_at"))
        and request.get("access_method") == registry_source.get("access_method")
        and request.get("rate_limit") == registry_source.get("rate_limit")
        and isinstance(response, dict)
        and set(response)
        == {
            "status_code",
            "retrieved_at",
            "media_type",
            "content_sha256",
            "byte_size",
            "object_key",
            "access_signal",
        }
        and response.get("status_code") == 200
        and _aware_timestamp(response.get("retrieved_at"))
        and response.get("retrieved_at") == snapshot.get("captured_at")
        and str(response.get("media_type") or "").strip() == snapshot.get("media_type")
        and response.get("content_sha256") == snapshot.get("sha256")
        and response.get("byte_size") == snapshot.get("byte_size")
        and response.get("object_key") == snapshot.get("object_key")
        and response.get("access_signal") == "none"
    )
    if valid:
        assert isinstance(request, dict)
        assert isinstance(response, dict)
        requested = datetime.fromisoformat(str(request["requested_at"]).replace("Z", "+00:00"))
        retrieved = datetime.fromisoformat(str(response["retrieved_at"]).replace("Z", "+00:00"))
        valid = requested <= retrieved
    if not valid:
        checks.append(
            CheckResult(
                code="B60_RETRIEVAL_LOG_CONTENT_INVALID",
                message="Retrieval log does not prove a successful source request and L0 write",
                location=location,
            )
        )
    return bool(valid)


def _validate_source_diff(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
    *,
    batch_id: str,
    comparison: dict[str, Any],
    location: str,
) -> tuple[int, int, int] | None:
    diff_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_SOURCE_DIFF_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if diff_path is None:
        return None
    try:
        diff = _json_object(diff_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        diff = {}
    items = diff.get("items")
    item_rows = items if isinstance(items, list) else []
    resolved = sum(
        isinstance(item, dict)
        and item.get("outcome") == "resolved"
        and str(item.get("resolution") or "").strip() != ""
        for item in item_rows
    )
    unresolved = len(item_rows) - resolved
    valid = (
        set(diff)
        == {
            "schema_version",
            "batch_id",
            "source_ref",
            "approved_snapshot_sha256",
            "current_snapshot_sha256",
            "items",
            "items_sha256",
        }
        and diff.get("schema_version") == SOURCE_DIFF_SCHEMA
        and diff.get("batch_id") == batch_id
        and diff.get("source_ref") == comparison.get("source_ref")
        and diff.get("approved_snapshot_sha256") == comparison.get("approved_snapshot_sha256")
        and diff.get("current_snapshot_sha256")
        == comparison.get("current_snapshot", {}).get("sha256")
        and len(item_rows) > 0
        and len({str(item.get("difference_id")) for item in item_rows if isinstance(item, dict)})
        == len(item_rows)
        and all(
            isinstance(item, dict)
            and set(item)
            == {
                "difference_id",
                "before_sha256",
                "after_sha256",
                "resolution",
                "outcome",
            }
            and str(item.get("difference_id") or "").strip()
            and SHA256_PATTERN.fullmatch(str(item.get("before_sha256") or ""))
            and SHA256_PATTERN.fullmatch(str(item.get("after_sha256") or ""))
            and item.get("outcome") == "resolved"
            and str(item.get("resolution") or "").strip()
            for item in item_rows
        )
        and diff.get("items_sha256") == _canonical_sha256(item_rows)
        and unresolved == 0
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_SOURCE_DIFF_CONTENT_INVALID",
                message="Reconciled source diff must enumerate and resolve every difference",
                location=location,
            )
        )
        return None
    return len(item_rows), resolved, unresolved


def _validate_l0_batch_manifest(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    batch: dict[str, Any],
    *,
    batch_id: str,
    domain_id: str,
    dataset_path: PurePosixPath,
    raw_root_sha256: str,
    source_subject_sha256: str,
    source_domains: dict[str, dict[str, Any]],
    source_registry: dict[str, Any] | None,
    l0_dir: Path | None,
) -> dict[str, Any] | None:
    location = f"collection.batches.{batch_id}.output_manifest"
    manifest_path = _validate_file_reference(
        checks,
        paths,
        batch.get("output_manifest"),
        code="B60_BATCH_OUTPUT_EVIDENCE_INVALID",
        location=location,
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if manifest_path is None:
        return None
    try:
        manifest = _json_object(manifest_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_BATCH_OUTPUT_SCHEMA_INVALID",
                message="L0 output manifest must be a JSON object",
                location=location,
            )
        )
        return None
    if set(manifest) != L0_MANIFEST_FIELDS:
        checks.append(
            CheckResult(
                code="B60_BATCH_OUTPUT_SCHEMA_INVALID",
                message="L0 output manifest does not match basic60.l0-batch-manifest.v1",
                location=location,
            )
        )
        return None
    source_path = raw_dir.joinpath(*dataset_path.parts)
    expected_hash = sha256_file(source_path) if source_path.is_file() else None
    expected_size = source_path.stat().st_size if source_path.is_file() else None
    basic_object_valid = False
    if l0_dir is not None and expected_hash is not None and expected_size is not None:
        basic_object = l0_dir / "sha256" / expected_hash
        try:
            basic_object_valid = (
                not basic_object.is_symlink()
                and basic_object.is_file()
                and basic_object.stat().st_size == expected_size
                and stat.S_IMODE(basic_object.stat().st_mode) == 0o444
                and sha256_file(basic_object) == expected_hash
            )
        except OSError:
            basic_object_valid = False
    domain = source_domains.get(domain_id, {})
    checked_at = manifest.get("checked_at")
    diff_status = manifest.get("diff_status")
    comparisons = manifest.get("source_snapshot_comparisons")
    registry_rows = source_registry.get("sources", []) if isinstance(source_registry, dict) else []
    registry_sources = {
        str(item.get("source_ref")): item for item in registry_rows if isinstance(item, dict)
    }
    expected_refs = {
        str(domain.get("primary_source_id")): "primary",
    }
    comparison_valid = isinstance(comparisons, list) and len(comparisons) == 1
    comparison_rows = comparisons if isinstance(comparisons, list) else []
    comparison_index = {
        str(item.get("source_ref")): item for item in comparison_rows if isinstance(item, dict)
    }
    if set(comparison_index) != set(expected_refs):
        comparison_valid = False
    comparison_statuses: list[str] = []
    difference_count = 0
    resolved_difference_count = 0
    unresolved_difference_count = 0
    current_hashes: list[str] = []
    source_contract_hashes: list[str] = []
    approved_source_contract_hash = ""
    for source_ref, role in expected_refs.items():
        comparison = comparison_index.get(source_ref, {})
        registry_source = registry_sources.get(source_ref, {})
        approved_source_contract = registry_source.get("admission_fingerprint")
        try:
            current_source_contract = _source_admission_fingerprint(
                raw_dir,
                registry_source,
                country_source_manifest=build_basic60_country_source_uri_manifest(raw_dir),
                terms_evidence=_source_evidence_for_fingerprint(paths, registry_source),
            )
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError, KeyError):
            current_source_contract = {}
        approved_source_contract_hash = (
            _canonical_sha256(approved_source_contract)
            if isinstance(approved_source_contract, dict)
            else ""
        )
        current_source_contract_hash = (
            _canonical_sha256(current_source_contract) if current_source_contract else ""
        )
        source_contract_hashes.append(current_source_contract_hash)
        snapshots = registry_source.get("snapshots")
        approved_snapshot = (
            snapshots[0]
            if isinstance(snapshots, list) and snapshots and isinstance(snapshots[0], dict)
            else {}
        )
        snapshot = comparison.get("current_snapshot")
        snapshot_dict = snapshot if isinstance(snapshot, dict) else {}
        snapshot_valid = _validate_evidence_snapshot_object(
            checks,
            paths,
            snapshot,
            l0_dir=l0_dir,
            location=f"{location}.source_snapshot_comparisons.{source_ref}.current_snapshot",
        )
        current_hash = str(snapshot_dict.get("sha256") or "")
        current_hashes.append(current_hash)
        derived_status = (
            "no_change"
            if current_hash and current_hash == approved_snapshot.get("content_sha256")
            else "reconciled"
        )
        status_value = str(comparison.get("comparison_status") or "")
        comparison_statuses.append(derived_status)
        current_terms_snapshots = comparison.get("current_terms_snapshots")
        current_terms_rows = (
            current_terms_snapshots if isinstance(current_terms_snapshots, list) else []
        )
        current_terms_retrieval_logs = comparison.get("current_terms_retrieval_logs")
        current_terms_log_rows = (
            current_terms_retrieval_logs if isinstance(current_terms_retrieval_logs, list) else []
        )
        current_terms_valid = bool(current_terms_rows) and len(current_terms_log_rows) == len(
            current_terms_rows
        )
        for index, terms_snapshot in enumerate(current_terms_rows):
            captured_at = (
                terms_snapshot.get("captured_at") if isinstance(terms_snapshot, dict) else None
            )
            capture_is_current = captured_at == checked_at
            if not capture_is_current:
                checks.append(
                    CheckResult(
                        code="B60_SOURCE_TERMS_SNAPSHOT_STALE",
                        message="Current terms snapshot was not captured during this D2 comparison",
                        location=(
                            f"{location}.source_snapshot_comparisons.{source_ref}."
                            f"current_terms_snapshots.{index}.captured_at"
                        ),
                    )
                )
            terms_log_valid = index < len(current_terms_log_rows) and _validate_retrieval_log(
                checks,
                paths,
                current_terms_log_rows[index] if index < len(current_terms_log_rows) else None,
                comparison=comparison,
                snapshot=terms_snapshot if isinstance(terms_snapshot, dict) else {},
                registry_source=registry_source,
                location=(
                    f"{location}.source_snapshot_comparisons.{source_ref}."
                    f"current_terms_retrieval_logs.{index}"
                ),
            )
            current_terms_valid = bool(
                _validate_evidence_snapshot_object(
                    checks,
                    paths,
                    terms_snapshot,
                    l0_dir=l0_dir,
                    location=(
                        f"{location}.source_snapshot_comparisons.{source_ref}."
                        f"current_terms_snapshots.{index}"
                    ),
                    official_terms=True,
                )
                and capture_is_current
                and terms_log_valid
                and current_terms_valid
            )
        current_terms_fingerprint = _terms_fingerprint_sha256(
            registry_source.get("terms_uri"),
            current_terms_rows,
        )
        approved_terms_fingerprint = (
            approved_source_contract.get("terms_fingerprint_sha256")
            if isinstance(approved_source_contract, dict)
            else None
        )
        if current_terms_fingerprint != approved_terms_fingerprint:
            checks.append(
                CheckResult(
                    code="B60_SOURCE_TERMS_CHANGED",
                    message=f"Official terms changed after D1 approval: {source_ref}",
                    location=(
                        f"{location}.source_snapshot_comparisons.{source_ref}."
                        "current_terms_fingerprint_sha256"
                    ),
                )
            )
        fields_valid = set(comparison) == {
            "source_ref",
            "source_role",
            "approved_snapshot_ref",
            "approved_snapshot_sha256",
            "approved_terms_fingerprint_sha256",
            "current_snapshot",
            "current_terms_snapshots",
            "current_terms_retrieval_logs",
            "current_terms_fingerprint_sha256",
            "retrieval_log",
            "diff_artifact",
            "comparison_status",
            "checked_at",
        }
        retrieval_valid = _validate_retrieval_log(
            checks,
            paths,
            comparison.get("retrieval_log"),
            comparison=comparison,
            snapshot=snapshot_dict,
            registry_source=registry_source,
            location=f"{location}.source_snapshot_comparisons.{source_ref}.retrieval_log",
        )
        diff_valid = comparison.get("diff_artifact") is None
        if derived_status == "reconciled":
            diff_counts = _validate_source_diff(
                checks,
                paths,
                comparison.get("diff_artifact"),
                batch_id=batch_id,
                comparison=comparison,
                location=f"{location}.source_snapshot_comparisons.{source_ref}.diff_artifact",
            )
            diff_valid = diff_counts is not None
            if diff_counts is not None:
                differences, resolved_differences, unresolved_differences = diff_counts
                difference_count += differences
                resolved_difference_count += resolved_differences
                unresolved_difference_count += unresolved_differences
        comparison_valid = comparison_valid and bool(
            fields_valid
            and comparison.get("source_role") == role
            and comparison.get("approved_snapshot_ref") == approved_snapshot.get("snapshot_ref")
            and comparison.get("approved_snapshot_sha256")
            == approved_snapshot.get("content_sha256")
            and comparison.get("approved_terms_fingerprint_sha256") == approved_terms_fingerprint
            and snapshot_valid is not None
            and snapshot_dict.get("storage_mode") == "external_l0"
            and snapshot_dict.get("volume_id") == manifest.get("volume_id")
            and retrieval_valid
            and current_terms_valid
            and comparison.get("current_terms_fingerprint_sha256")
            == current_terms_fingerprint
            == approved_terms_fingerprint
            and diff_valid
            and status_value == derived_status
            and comparison.get("checked_at") == checked_at
        )
    comparison_valid = comparison_valid and (
        (diff_status == "no_change" and set(comparison_statuses) == {"no_change"})
        or (diff_status == "reconciled" and "reconciled" in comparison_statuses)
    )

    diff_report_path = _validate_file_reference(
        checks,
        paths,
        manifest.get("diff_report"),
        code="B60_BATCH_DIFF_REPORT_INVALID",
        location=f"{location}.diff_report",
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    diff_report_valid = False
    if diff_report_path is not None:
        try:
            diff_report = _json_object(diff_report_path)
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
            diff_report = {}
        diff_report_valid = (
            set(diff_report)
            == {
                "schema_version",
                "batch_id",
                "run_id",
                "raw_root_sha256",
                "source_sha256",
                "checked_at",
                "diff_status",
                "difference_count",
                "resolved_difference_count",
                "unresolved_difference_count",
                "comparison_count",
                "source_snapshot_comparisons_sha256",
            }
            and diff_report.get("schema_version") == L0_DIFF_REPORT_SCHEMA
            and diff_report.get("batch_id") == batch_id
            and diff_report.get("run_id") == "legacy_import_run"
            and diff_report.get("raw_root_sha256") == raw_root_sha256
            and diff_report.get("source_sha256") == expected_hash
            and diff_report.get("checked_at") == checked_at
            and diff_report.get("diff_status") == diff_status
            and diff_report.get("difference_count") == difference_count
            and diff_report.get("resolved_difference_count") == resolved_difference_count
            and diff_report.get("unresolved_difference_count") == unresolved_difference_count == 0
            and diff_report.get("comparison_count") == 1
            and diff_report.get("source_snapshot_comparisons_sha256")
            == _canonical_sha256(comparisons)
        )
    derived_watermark = _canonical_sha256(
        {
            "source_sha256": expected_hash,
            "current_snapshot_sha256s": sorted(current_hashes),
            "diff_status": diff_status,
            "difference_count": difference_count,
        }
    )
    valid = (
        manifest.get("schema_version") == L0_BATCH_MANIFEST_SCHEMA
        and manifest.get("profile_id") == PROFILE_ID
        and manifest.get("release_id") == RELEASE_ID
        and manifest.get("batch_id") == batch_id
        and manifest.get("run_id") == batch.get("run_id") == "legacy_import_run"
        and manifest.get("copy_status") == "succeeded"
        and manifest.get("batch_status") == batch.get("status") == "closed"
        and manifest.get("domain_id") == domain_id
        and manifest.get("raw_root_sha256") == raw_root_sha256
        and manifest.get("source_subject_sha256") == source_subject_sha256
        and len(source_contract_hashes) == 1
        and manifest.get("source_contract_sha256")
        == source_contract_hashes[0]
        == approved_source_contract_hash
        and manifest.get("approved_primary_source_id") == domain.get("primary_source_id")
        and manifest.get("approved_alternative_source_id") == domain.get("alternative_source_id")
        and manifest.get("source_dataset_path") == dataset_path.as_posix()
        and manifest.get("source_sha256") == expected_hash == batch.get("dataset_sha256")
        and manifest.get("source_byte_size") == expected_size
        and VOLUME_ID_PATTERN.fullmatch(str(manifest.get("volume_id") or ""))
        and manifest.get("object_key") == f"sha256/{expected_hash}"
        and manifest.get("object_sha256") == expected_hash
        and manifest.get("object_byte_size") == expected_size
        and basic_object_valid
        and isinstance(manifest.get("copied"), bool)
        and isinstance(manifest.get("reused"), bool)
        and manifest.get("copied") is not manifest.get("reused")
        and manifest.get("readback_verified") is True
        and manifest.get("immutable") is True
        and manifest.get("read_only") is True
        and manifest.get("posix_mode") == "0444"
        and manifest.get("diff_status") == batch.get("diff_status")
        and manifest.get("diff_status") in {"no_change", "reconciled"}
        and manifest.get("current_snapshot")
        == {"sha256": expected_hash, "byte_size": expected_size}
        and _aware_timestamp(checked_at)
        and comparison_valid
        and diff_report_valid
        and manifest.get("watermark_advanced") is True
        and manifest.get("watermark")
        == {"committed_sha256": derived_watermark, "committed_at": checked_at}
        and manifest.get("access_signal") == batch.get("access_signal") == "none"
        and _aware_timestamp(manifest.get("materialized_at"))
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_BATCH_OUTPUT_CONTENT_INVALID",
                message=f"{batch_id} L0 manifest does not match the current approved source",
                location=location,
            )
        )
        return None
    return manifest


def _validate_collection(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    collection: Any,
    *,
    raw_root_sha256: str,
    source_subject_sha256: str,
    source_admission: Any,
    source_registry: dict[str, Any] | None,
    l0_dir: Path | None,
) -> None:
    if not isinstance(collection, dict):
        checks.append(
            CheckResult(
                code="B60_COLLECTION_INVALID",
                message="Collection section is required",
                location="collection",
            )
        )
        return
    if collection.get("input_immutable") is not True:
        checks.append(
            CheckResult(
                code="B60_COLLECTION_NOT_IMMUTABLE",
                message="Raw Basic60 input must remain immutable",
                location="collection.input_immutable",
            )
        )
    batches = collection.get("batches")
    indexed = (
        {
            str(item.get("batch_id")): item
            for item in batches
            if isinstance(batches, list) and isinstance(item, dict)
        }
        if isinstance(batches, list)
        else {}
    )
    if len(indexed) != len(BATCH_SPECS):
        checks.append(
            CheckResult(
                code="B60_BATCH_SET_INVALID",
                message="Exactly three Basic60 collection batches are required",
                location="collection.batches",
            )
        )
    completed_claim = any(
        item.get("status") == "closed" or item.get("run_status") == "succeeded"
        for item in indexed.values()
    )
    resolved_l0: Path | None = None
    if completed_claim:
        if l0_dir is None:
            checks.append(
                CheckResult(
                    code="B60_L0_DIRECTORY_REQUIRED",
                    message="Closed Basic60 batches require --l0-dir for object revalidation",
                    location="collection.batches",
                )
            )
        else:
            try:
                resolved_l0 = _validated_l0_directory(paths, raw_dir, l0_dir)
            except (OSError, ValueError) as error:
                checks.append(
                    CheckResult(
                        code="B60_L0_DIRECTORY_INVALID",
                        message=str(error),
                        location=str(l0_dir),
                    )
                )
    source_domain_rows = (
        source_admission.get("domains", []) if isinstance(source_admission, dict) else []
    )
    source_domains = {
        str(item.get("domain_id")): item for item in source_domain_rows if isinstance(item, dict)
    }
    output_manifests: dict[str, dict[str, Any]] = {}
    for batch_id, domain_id, dataset_path in BATCH_SPECS:
        item = indexed.get(batch_id, {})
        raw_path = raw_dir.joinpath(*dataset_path.parts)
        expected_hash = sha256_file(raw_path) if raw_path.is_file() else None
        valid = (
            item.get("domain_id") == domain_id
            and item.get("dataset_path") == dataset_path.as_posix()
            and item.get("dataset_sha256") == expected_hash
            and item.get("run_status") == "succeeded"
            and item.get("access_signal") == "none"
            and item.get("input_immutable") is True
            and item.get("diff_status") in {"no_change", "reconciled"}
            and item.get("status") == "closed"
        )
        if not valid:
            checks.append(
                CheckResult(
                    code="B60_BATCH_NOT_CLOSED",
                    message=f"{batch_id} is incomplete or changed",
                    location=f"collection.batches.{batch_id}",
                )
            )
        manifest = _validate_l0_batch_manifest(
            checks,
            paths,
            raw_dir,
            item,
            batch_id=batch_id,
            domain_id=domain_id,
            dataset_path=dataset_path,
            raw_root_sha256=raw_root_sha256,
            source_subject_sha256=source_subject_sha256,
            source_domains=source_domains,
            source_registry=source_registry,
            l0_dir=resolved_l0,
        )
        if manifest is not None:
            output_manifests[batch_id] = manifest
    if len(output_manifests) == len(BATCH_SPECS):
        volume_ids = {str(item["volume_id"]) for item in output_manifests.values()}
        if len(volume_ids) != 1:
            checks.append(
                CheckResult(
                    code="B60_BATCH_OUTPUT_VOLUME_MISMATCH",
                    message="All Basic60 L0 batch manifests must use one external volume",
                    location="collection.batches",
                )
            )
        macro = output_manifests["BASIC60-BATCH-MACRO"]
        energy = output_manifests["BASIC60-BATCH-ENERGY"]
        object_fields = ("volume_id", "object_key", "object_sha256", "object_byte_size")
        if any(macro[field] != energy[field] for field in object_fields):
            checks.append(
                CheckResult(
                    code="B60_BATCH_SHARED_OBJECT_INVALID",
                    message="Macro and energy manifests must reference the same content object",
                    location="collection.batches",
                )
            )
    incidents = collection.get("incidents")
    if not isinstance(incidents, list):
        checks.append(
            CheckResult(
                code="B60_INCIDENTS_INVALID",
                message="Collection incidents must be a list",
                location="collection.incidents",
            )
        )
    else:
        stop_signals = {name for name, rule in STOP_SIGNALS.items() if rule["must_stop"]}
        for index, item in enumerate(incidents):
            if not isinstance(item, dict):
                checks.append(
                    CheckResult(
                        code="B60_INCIDENT_INVALID",
                        message="Incident must be an object",
                        location=f"collection.incidents.{index}",
                    )
                )
                continue
            signal = item.get("access_signal")
            if signal in stop_signals and (
                item.get("automation_stopped") is not True
                or item.get("watermark_advanced") is not False
                or item.get("status") != "resolved"
            ):
                checks.append(
                    CheckResult(
                        code="B60_STOP_SIGNAL_NOT_ENFORCED",
                        message="Access-control stop signal was not safely closed",
                        location=f"collection.incidents.{index}",
                    )
                )


def _validate_processing(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    processing: Any,
    raw_manifest: dict[str, Any] | None,
    seed_path: Path,
    seed: dict[str, Any] | None,
) -> None:
    if not isinstance(processing, dict):
        checks.append(
            CheckResult(
                code="B60_PROCESSING_INVALID",
                message="Processing section is required",
                location="processing",
            )
        )
        return
    pipelines = processing.get("pipelines")
    indexed = (
        {
            str(item.get("pipeline_id")): item
            for item in pipelines
            if isinstance(pipelines, list) and isinstance(item, dict)
        }
        if isinstance(pipelines, list)
        else {}
    )
    if set(indexed) != set(PIPELINE_IDS):
        checks.append(
            CheckResult(
                code="B60_PIPELINE_SET_INVALID",
                message="Parse, standardize, and entity pipelines are all required",
                location="processing.pipelines",
            )
        )
    replay_payloads: dict[str, dict[str, Any]] = {}
    for pipeline_id in PIPELINE_IDS:
        item = indexed.get(pipeline_id, {})
        if item.get("status") != "succeeded":
            checks.append(
                CheckResult(
                    code="B60_PIPELINE_NOT_REPLAYED",
                    message=f"{pipeline_id} has not succeeded",
                    location=f"processing.pipelines.{pipeline_id}",
                )
            )
        replay_path = _validate_file_reference(
            checks,
            paths,
            item.get("replay_manifest"),
            code="B60_PIPELINE_EVIDENCE_INVALID",
            location=f"processing.pipelines.{pipeline_id}.replay_manifest",
            prefix=PurePosixPath("data/basic60/candidates"),
        )
        if replay_path is None:
            continue
        try:
            replay = _json_object(replay_path)
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
            checks.append(
                CheckResult(
                    code="B60_PIPELINE_REPLAY_MANIFEST_INVALID",
                    message=str(error),
                    location=f"processing.pipelines.{pipeline_id}.replay_manifest",
                )
            )
            continue
        replay_payloads[pipeline_id] = replay
    if (
        raw_manifest is not None
        and isinstance(seed, dict)
        and seed_path.is_file()
        and len(replay_payloads) == len(PIPELINE_IDS)
    ):
        try:
            first = build_d3_stage_payloads(seed, raw_manifest)
            second = build_d3_stage_payloads(seed, raw_manifest)
            stage_references: dict[str, Any] = {}
            stage_payloads: dict[str, dict[str, Any]] = {}
            for pipeline_id, stage_id in zip(
                PIPELINE_IDS, ("parse", "standardize", "entity"), strict=True
            ):
                replay = replay_payloads[pipeline_id]
                stage_reference = replay.get("runtime_stage")
                stage_path = _validate_file_reference(
                    checks,
                    paths,
                    stage_reference,
                    code="B60_PIPELINE_STAGE_INVALID",
                    location=f"processing.pipelines.{pipeline_id}.runtime_stage",
                    prefix=PurePosixPath("runtime/basic60/d3"),
                )
                stage_references[stage_id] = stage_reference
                if stage_path is None:
                    continue
                if (
                    not isinstance(stage_reference, dict)
                    or set(stage_reference) != {"path", "sha256", "byte_size"}
                    or stage_reference.get("byte_size") != stage_path.stat().st_size
                ):
                    raise ValueError(f"{stage_id} runtime stage reference is not exact")
                actual_stage = _json_object(stage_path)
                expected_stage = first[stage_id]
                if actual_stage != expected_stage or expected_stage != second[stage_id]:
                    raise ValueError(
                        f"{stage_id} runtime stage does not match an independent replay"
                    )
                stage_payloads[stage_id] = actual_stage
            if len(stage_payloads) != 3:
                raise ValueError("All three D3 runtime stages are required")
            replay_hashes = {
                stage_id: _canonical_sha256(payload) for stage_id, payload in first.items()
            }
            expected_replays = build_basic60_replay_manifests(
                raw_manifest,
                runtime_seed_sha256=sha256_file(seed_path),
                d3_artifacts={
                    "stages": stage_references,
                    "replay": {
                        "first_output_sha256": replay_hashes,
                        "second_output_sha256": {
                            stage_id: _canonical_sha256(payload)
                            for stage_id, payload in second.items()
                        },
                    },
                },
                stage_payloads=stage_payloads,
            )
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
            checks.append(
                CheckResult(
                    code="B60_PIPELINE_REPLAY_CONTENT_INVALID",
                    message=str(error),
                    location="processing.pipelines",
                )
            )
        else:
            derived_open_anomalies = sum(
                int(expected_replays[pipeline_id]["open_anomaly_count"])
                for pipeline_id in PIPELINE_IDS
            )
            for pipeline_id in PIPELINE_IDS:
                if (
                    replay_payloads[pipeline_id] != expected_replays[pipeline_id]
                    or indexed.get(pipeline_id, {}).get("status")
                    != expected_replays[pipeline_id]["status"]
                ):
                    checks.append(
                        CheckResult(
                            code="B60_PIPELINE_REPLAY_CONTENT_INVALID",
                            message=(
                                f"{pipeline_id} is not bound to the real runtime stage and "
                                "independent replay"
                            ),
                            location=f"processing.pipelines.{pipeline_id}.replay_manifest",
                        )
                    )
            if (
                processing.get("open_anomaly_count") != derived_open_anomalies
                or processing.get("open_high_risk_anomaly_count") != derived_open_anomalies
            ):
                checks.append(
                    CheckResult(
                        code="B60_PROCESSING_ANOMALY_COUNT_INVALID",
                        message="D3 anomaly counts must be recalculated from runtime stages",
                        location="processing",
                    )
                )
    if (
        processing.get("original_values_preserved") is not True
        or processing.get("source_text_preserved") is not True
    ):
        checks.append(
            CheckResult(
                code="B60_PROCESSING_SOURCE_NOT_PRESERVED",
                message="D3 invariants require original values and source text to be preserved",
                location="processing",
            )
        )
    if (
        processing.get("open_anomaly_count") != 0
        or processing.get("open_high_risk_anomaly_count") != 0
    ):
        checks.append(
            CheckResult(
                code="B60_OPEN_ANOMALIES",
                message="All Basic60 anomalies must be closed before private trial",
                location="processing",
            )
        )


def _validate_sample_candidate(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    seed_path: Path,
    seed: dict[str, Any],
    reference: Any,
) -> dict[str, Any] | None:
    candidate_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_SAMPLE_CANDIDATE_INVALID",
        location="acceptance.sample_candidate",
        prefix=PurePosixPath("runtime/basic60"),
    )
    if candidate_path is None:
        return None
    try:
        candidate = _json_object(candidate_path)
        expected = build_basic60_sample_candidate(
            seed,
            seed_artifact_sha256=sha256_file(seed_path),
        )
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        checks.append(
            CheckResult(
                code="B60_SAMPLE_CANDIDATE_INVALID",
                message=str(error),
                location="acceptance.sample_candidate",
            )
        )
        return None
    if _canonical_sha256(candidate) != _canonical_sha256(expected):
        checks.append(
            CheckResult(
                code="B60_SAMPLE_CANDIDATE_DRIFT",
                message="Frozen sample is not the deterministic 60-country candidate",
                location="acceptance.sample_candidate",
            )
        )
        return None
    return candidate


def _validate_sample_review_manifest(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    acceptance: dict[str, Any],
    reference: Any,
    sample_candidate: dict[str, Any] | None,
    *,
    review_slot: int,
) -> tuple[str, dict[str, str]] | None:
    location = f"acceptance.artifacts.sample_reviews.{review_slot - 1}"
    manifest_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_SAMPLE_REVIEW_MANIFEST_INVALID",
        location=location,
        prefix=PurePosixPath("runtime/basic60/review"),
    )
    if manifest_path is None or sample_candidate is None:
        return None
    try:
        manifest = _json_object(manifest_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_SAMPLE_REVIEW_MANIFEST_INVALID",
                message="Sample review manifest must be a JSON object",
                location=location,
            )
        )
        return None
    if set(manifest) != SAMPLE_REVIEW_MANIFEST_FIELDS:
        checks.append(
            CheckResult(
                code="B60_SAMPLE_REVIEW_MANIFEST_INVALID",
                message="Sample review manifest does not match its frozen schema",
                location=location,
            )
        )
        return None
    items = manifest.get("items")
    candidate_records = sample_candidate.get("records")
    candidate_rows = candidate_records if isinstance(candidate_records, list) else []
    expected_sample_ids = {
        str(item.get("sample_id") or "") for item in candidate_rows if isinstance(item, dict)
    }
    item_rows = items if isinstance(items, list) else []
    indexed = {str(item.get("sample_id")): item for item in item_rows if isinstance(item, dict)}
    item_results_valid = (
        len(item_rows) == EXPECTED_SAMPLE_SIZE
        and len(indexed) == EXPECTED_SAMPLE_SIZE
        and set(indexed) == expected_sample_ids
    )
    outcomes: dict[str, str] = {}
    if item_results_valid:
        for sample_id in sorted(expected_sample_ids):
            item = indexed[sample_id]
            valid_row = set(item) == {"sample_id", "outcome"} and item.get("outcome") in {
                "accurate",
                "inaccurate",
            }
            if not valid_row:
                item_results_valid = False
                break
            outcomes[sample_id] = str(item["outcome"])
    valid = (
        manifest.get("schema_version") == SAMPLE_REVIEW_MANIFEST_SCHEMA
        and manifest.get("profile_id") == PROFILE_ID
        and manifest.get("release_id") == RELEASE_ID
        and manifest.get("sample_candidate") == acceptance.get("sample_candidate")
        and manifest.get("sample_size") == EXPECTED_SAMPLE_SIZE
        and manifest.get("country_count") == EXPECTED_COUNTRIES
        and manifest.get("per_country_count") == 3
        and manifest.get("coverage") == sample_candidate.get("coverage")
        and manifest.get("review_slot") == review_slot
        and manifest.get("template_only") is False
        and _valid_person(manifest.get("reviewer_name"))
        and _aware_timestamp(manifest.get("reviewed_at"))
        and manifest.get("status") == "completed"
        and item_results_valid
        and manifest.get("items_sha256") == _canonical_sha256(item_rows)
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_SAMPLE_REVIEW_CONTENT_INVALID",
                message="Sample review summary is not bound to the frozen candidate and results",
                location=location,
            )
        )
        return None
    return str(manifest["reviewer_name"]), outcomes


def _seed_duplicate_count(seed: dict[str, Any] | None) -> int:
    if not isinstance(seed, dict):
        return -1
    identities: list[tuple[str, str, str, str]] = []
    countries = seed.get("countries")
    for country in countries if isinstance(countries, list) else []:
        if not isinstance(country, dict):
            continue
        iso3 = str(country.get("iso3") or "")
        metrics = country.get("metrics")
        metric_rows = metrics if isinstance(metrics, list) else []
        identities.extend(
            _sample_record_identity(iso3, metric)
            for metric in metric_rows
            if isinstance(metric, dict)
        )
    return len(identities) - len(set(identities))


def _validate_review_templates(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    acceptance: dict[str, Any],
    sample_candidate: dict[str, Any] | None,
    seed: dict[str, Any] | None,
    seed_reference: Any,
) -> None:
    templates = acceptance.get("review_templates")
    if not isinstance(templates, dict) or set(templates) != {"sample_reviews", "item_review"}:
        checks.append(
            CheckResult(
                code="B60_REVIEW_TEMPLATES_INVALID",
                message="Acceptance must bind exactly two sample templates and one item template",
                location="acceptance.review_templates",
            )
        )
        return
    sample_references = templates.get("sample_reviews")
    if not isinstance(sample_references, list) or len(sample_references) != 2:
        checks.append(
            CheckResult(
                code="B60_REVIEW_TEMPLATES_INVALID",
                message="Two independent frozen sample-review templates are required",
                location="acceptance.review_templates.sample_reviews",
            )
        )
        sample_references = [None, None]
    for index in range(2):
        path = _validate_file_reference(
            checks,
            paths,
            sample_references[index] if index < len(sample_references) else None,
            code="B60_REVIEW_TEMPLATE_INVALID",
            location=f"acceptance.review_templates.sample_reviews.{index}",
            prefix=PurePosixPath("runtime/basic60"),
        )
        if path is None or sample_candidate is None:
            continue
        try:
            actual = _json_object(path)
            expected = build_basic60_sample_review_manifest(
                acceptance["sample_candidate"],
                sample_candidate,
                review_slot=index + 1,
            )
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
            checks.append(
                CheckResult(
                    code="B60_REVIEW_TEMPLATE_INVALID",
                    message=str(error),
                    location=f"acceptance.review_templates.sample_reviews.{index}",
                )
            )
            continue
        if actual != expected:
            checks.append(
                CheckResult(
                    code="B60_REVIEW_TEMPLATE_DRIFT",
                    message="Sample-review template is not the deterministic blank copy",
                    location=f"acceptance.review_templates.sample_reviews.{index}",
                )
            )
    item_path = _validate_file_reference(
        checks,
        paths,
        templates.get("item_review"),
        code="B60_REVIEW_TEMPLATE_INVALID",
        location="acceptance.review_templates.item_review",
        prefix=PurePosixPath("runtime/basic60"),
    )
    if item_path is not None and isinstance(seed, dict) and isinstance(seed_reference, dict):
        try:
            actual = _json_object(item_path)
            expected = build_basic60_item_review_manifest(seed_reference, seed)
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
            checks.append(
                CheckResult(
                    code="B60_REVIEW_TEMPLATE_INVALID",
                    message=str(error),
                    location="acceptance.review_templates.item_review",
                )
            )
        else:
            if actual != expected:
                checks.append(
                    CheckResult(
                        code="B60_REVIEW_TEMPLATE_DRIFT",
                        message="Item-review template is not the deterministic blank copy",
                        location="acceptance.review_templates.item_review",
                    )
                )


def _validate_item_review_manifest(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
    *,
    seed: dict[str, Any] | None,
    seed_reference: Any,
) -> dict[str, int] | None:
    location = "acceptance.artifacts.item_review"
    manifest_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_ITEM_REVIEW_MANIFEST_INVALID",
        location=location,
        prefix=PurePosixPath("runtime/basic60/review"),
    )
    if manifest_path is None or not isinstance(seed, dict):
        return None
    try:
        manifest = _json_object(manifest_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        manifest = {}
    expected_fields = {
        "schema_version",
        "profile_id",
        "release_id",
        "seed_artifact",
        "template_only",
        "reviewer_name",
        "reviewed_at",
        "status",
        "groups",
        "groups_sha256",
    }
    groups = manifest.get("groups")
    group_rows = groups if isinstance(groups, list) else []
    group_index = {
        str(item.get("review_type")): item for item in group_rows if isinstance(item, dict)
    }
    expected_rows = _item_review_rows(seed, blank=False)
    expected_types = ("chinese_names", "pending_values", "derived_values")
    groups_valid = len(group_rows) == 3 and set(group_index) == set(expected_types)
    for review_type in expected_types:
        group = group_index.get(review_type, {})
        actual_items = group.get("items")
        actual_rows = actual_items if isinstance(actual_items, list) else []
        expected_items = expected_rows[review_type]
        groups_valid = groups_valid and bool(
            set(group) == {"review_type", "item_count", "items"}
            and group.get("item_count") == len(expected_items)
            and actual_rows == expected_items
            and all(item.get("outcome") == "approved" for item in actual_rows)
        )
    derived_items = group_index.get("derived_values", {}).get("items", [])
    derived_math_valid = isinstance(derived_items, list) and all(
        isinstance(item, dict) and _derived_review_math_matches(item) for item in derived_items
    )
    valid = (
        set(manifest) == expected_fields
        and manifest.get("schema_version") == ITEM_REVIEW_MANIFEST_SCHEMA
        and manifest.get("profile_id") == PROFILE_ID
        and manifest.get("release_id") == RELEASE_ID
        and manifest.get("seed_artifact") == seed_reference
        and manifest.get("template_only") is False
        and _valid_person(manifest.get("reviewer_name"))
        and _aware_timestamp(manifest.get("reviewed_at"))
        and manifest.get("status") == "completed"
        and groups_valid
        and derived_math_valid
        and manifest.get("groups_sha256") == _canonical_sha256(group_rows)
    )
    if not valid:
        checks.append(
            CheckResult(
                code="B60_ITEM_REVIEW_CONTENT_INVALID",
                message="Item review must contain the exact seed-derived observations and math",
                location=location,
            )
        )
        return None
    return {review_type: len(expected_rows[review_type]) for review_type in expected_types}


def _derived_review_math_matches(item: dict[str, Any]) -> bool:
    try:
        numerator = Decimal(str(item.get("numerator_value")))
        denominator = Decimal(str(item.get("denominator_value")))
        recomputed = Decimal(str(item.get("recomputed_value")))
        reported = Decimal(str(item.get("reported_value")))
    except (InvalidOperation, ValueError):
        return False
    if denominator == 0 or recomputed != numerator / denominator * Decimal(100):
        return False
    return abs(reported - recomputed) <= Decimal("0.01")


def _validate_machine_evidence_artifact(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    reference: Any,
    *,
    summary: dict[str, Any],
    seed_reference: Any,
    source_registry: dict[str, Any] | None,
    bundle: dict[str, Any],
) -> None:
    """Replay the single aggregate report and bind it to the unchanged D1-D3 packet."""

    del summary
    location = "acceptance.artifacts.machine_evidence"
    report_path = _validate_file_reference(
        checks,
        paths,
        reference,
        code="B60_MACHINE_EVIDENCE_INVALID",
        location=location,
        prefix=PurePosixPath("runtime/basic60"),
    )
    if report_path is None:
        return
    try:
        report = _json_object(report_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        report = {}
    if report.get("schema_version") != "basic60.machine-evidence.v2":
        checks.append(
            CheckResult(
                code="B60_MACHINE_EVIDENCE_CONTENT_INVALID",
                message="Legacy or self-reported machine evidence is not accepted",
                location=location,
            )
        )
        return
    for issue in validate_machine_evidence(report, repository_root=paths.root):
        checks.append(
            CheckResult(
                code=issue.code,
                message=issue.message,
                location=f"{location}.{issue.location}",
            )
        )
    report_seed = report.get("seed_artifact")
    report_raw = report.get("raw_manifest")
    current_raw = bundle.get("raw_manifest")
    references_match = all(
        isinstance(report_reference, dict)
        and isinstance(current_reference, dict)
        and report_reference.get("path") == current_reference.get("path")
        and report_reference.get("sha256") == current_reference.get("sha256")
        for report_reference, current_reference in (
            (report_seed, seed_reference),
            (report_raw, current_raw),
        )
    )
    if not references_match:
        checks.append(
            CheckResult(
                code="B60_MACHINE_EVIDENCE_BINDING_INVALID",
                message="Machine report seed/raw references do not match the acceptance packet",
                location=location,
            )
        )
    preflight_reference = report.get("bundle_artifact")
    preflight_path: Path | None = None
    if isinstance(preflight_reference, dict):
        preflight_path = _validate_file_reference(
            checks,
            paths,
            preflight_reference,
            code="B60_MACHINE_PREFLIGHT_BUNDLE_INVALID",
            location=f"{location}.bundle_artifact",
            prefix=PurePosixPath("data/basic60/review"),
        )
    if preflight_path is not None:
        try:
            preflight = _json_object(preflight_path)
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
            preflight = {}
        immutable_sections = (
            "schema_version",
            "profile_id",
            "release_id",
            "formal_gate_status",
            "profile",
            "raw_manifest",
            "raw_root_sha256",
            "runtime_seed",
            "baseline_candidate",
            "baseline_confirmation",
            "baseline_decision",
            "baseline_clarification",
            "source_admission",
            "collection",
            "processing",
        )
        if any(preflight.get(key) != bundle.get(key) for key in immutable_sections):
            checks.append(
                CheckResult(
                    code="B60_MACHINE_PREFLIGHT_BUNDLE_DRIFT",
                    message="D1-D3 changed after the live machine run was collected",
                    location=f"{location}.bundle_artifact",
                )
            )
    operations = report.get("operations")
    source_revocation = (
        operations.get("source_revocation") if isinstance(operations, dict) else None
    )
    inputs = source_revocation.get("inputs") if isinstance(source_revocation, dict) else None
    revoked_source = inputs.get("source_ref") if isinstance(inputs, dict) else None
    registry_rows = source_registry.get("sources") if isinstance(source_registry, dict) else None
    active_refs = {
        str(item.get("source_ref"))
        for item in (registry_rows if isinstance(registry_rows, list) else [])
        if isinstance(item, dict)
        if item.get("status") == "active"
    }
    if revoked_source not in active_refs:
        checks.append(
            CheckResult(
                code="B60_MACHINE_REVOCATION_SOURCE_INVALID",
                message="Source-revocation drill must target an approved active Basic60 source",
                location=f"{location}.operations.source_revocation.inputs.source_ref",
            )
        )


def _validate_acceptance(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    acceptance: Any,
    summary: dict[str, Any],
    sample_candidate: dict[str, Any] | None,
    seed: dict[str, Any] | None,
    seed_reference: Any,
    source_registry: dict[str, Any] | None,
    bundle: dict[str, Any],
) -> tuple[str, str] | None:
    if not isinstance(acceptance, dict):
        checks.append(
            CheckResult(
                code="B60_ACCEPTANCE_INVALID",
                message="Acceptance section is required",
                location="acceptance",
            )
        )
        return None
    if acceptance.get("machine_counts") != summary:
        checks.append(
            CheckResult(
                code="B60_ACCEPTANCE_COUNTS_CHANGED",
                message="Acceptance counts must equal the current machine-recalculated counts",
                location="acceptance.machine_counts",
            )
        )
    artifacts = acceptance.get("artifacts")
    if not isinstance(artifacts, dict):
        artifacts = {}
    expected_artifacts = {"sample_reviews", "item_review", "machine_evidence"}
    if set(artifacts) != expected_artifacts:
        checks.append(
            CheckResult(
                code="B60_ACCEPTANCE_ARTIFACT_SET_INVALID",
                message="Only two human review files and one collected machine report are accepted",
                location="acceptance.artifacts",
            )
        )
    _validate_review_templates(
        checks,
        paths,
        acceptance,
        sample_candidate,
        seed,
        seed_reference,
    )
    sample_references = artifacts.get("sample_reviews")
    review_results: list[tuple[str, dict[str, str]]] = []
    if not isinstance(sample_references, list) or len(sample_references) != 2:
        checks.append(
            CheckResult(
                code="B60_SAMPLE_REVIEW_SET_INVALID",
                message="Two independent completed sample-review files are required",
                location="acceptance.artifacts.sample_reviews",
            )
        )
        sample_references = [None, None]
    for index in range(2):
        result = _validate_sample_review_manifest(
            checks,
            paths,
            acceptance,
            sample_references[index] if index < len(sample_references) else None,
            sample_candidate,
            review_slot=index + 1,
        )
        if result is not None:
            review_results.append(result)
    sample_reviewers: tuple[str, str] | None = None
    accurate_count = -1
    if len(review_results) == 2:
        first_name, first_outcomes = review_results[0]
        second_name, second_outcomes = review_results[1]
        if first_name == second_name:
            checks.append(
                CheckResult(
                    code="B60_SAMPLE_REVIEWERS_NOT_INDEPENDENT",
                    message="The frozen sample requires two distinct named reviewers",
                    location="acceptance.artifacts.sample_reviews",
                )
            )
        elif first_outcomes != second_outcomes:
            checks.append(
                CheckResult(
                    code="B60_SAMPLE_REVIEW_DISAGREEMENT",
                    message="All sample disagreements must be resolved before acceptance",
                    location="acceptance.artifacts.sample_reviews",
                )
            )
        else:
            sample_reviewers = (first_name, second_name)
            accurate_count = sum(value == "accurate" for value in first_outcomes.values())
    item_counts = _validate_item_review_manifest(
        checks,
        paths,
        artifacts.get("item_review"),
        seed=seed,
        seed_reference=seed_reference,
    )
    duplicate_count = _seed_duplicate_count(seed)
    derived_item_counts = item_counts or {}
    numeric_valid = (
        acceptance.get("chinese_names_reviewed")
        == derived_item_counts.get("chinese_names")
        == EXPECTED_COUNTRIES
        and acceptance.get("sample_size") == EXPECTED_SAMPLE_SIZE
        and acceptance.get("sample_accurate_count") == accurate_count
        and accurate_count >= math.ceil(EXPECTED_SAMPLE_SIZE * 0.95)
        and acceptance.get("duplicate_count") == duplicate_count
        and duplicate_count <= math.floor(EXPECTED_SAMPLE_SIZE * 0.01)
        and acceptance.get("pending_values_reviewed")
        == derived_item_counts.get("pending_values")
        == EXPECTED_PENDING_VALUES
        and acceptance.get("derived_values_reviewed")
        == derived_item_counts.get("derived_values")
        == EXPECTED_DERIVED_VALUES
        and acceptance.get("unresolved_p0_data_issue_count") == 0
    )
    if not numeric_valid:
        checks.append(
            CheckResult(
                code="B60_ACCEPTANCE_THRESHOLDS_NOT_MET",
                message=(
                    "Name, sample, pending, derived, duplicate, or P0-data thresholds are unmet"
                ),
                location="acceptance",
            )
        )
    if acceptance.get("duplicate_count") != duplicate_count:
        checks.append(
            CheckResult(
                code="B60_DUPLICATE_COUNT_INVALID",
                message="Duplicate count must equal the value recalculated from the runtime seed",
                location="acceptance.duplicate_count",
            )
        )
    _validate_machine_evidence_artifact(
        checks,
        paths,
        artifacts.get("machine_evidence"),
        summary=summary,
        seed_reference=seed_reference,
        source_registry=source_registry,
        bundle=bundle,
    )
    return sample_reviewers


def _validate_manifest_and_seed(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    manifest_path: Path,
    seed_path: Path,
    bundle: dict[str, Any],
) -> tuple[dict[str, Any] | None, dict[str, Any] | None, dict[str, Any] | None]:
    try:
        stored_manifest = _json_object(manifest_path)
        current_manifest = build_raw_manifest(paths, raw_dir)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        checks.append(
            CheckResult(
                code="B60_RAW_MANIFEST_INVALID",
                message=str(error),
                location=str(manifest_path),
            )
        )
        return None, None, None
    for field in (
        "root_sha256",
        "included_file_count",
        "included_byte_count",
        "excluded_zone_identifier_count",
        "unsafe_entry_count",
        "scoped_files",
    ):
        if stored_manifest.get(field) != current_manifest.get(field):
            checks.append(
                CheckResult(
                    code="B60_RAW_MANIFEST_DRIFT",
                    message=f"Raw manifest field changed: {field}",
                    location=_display_path(paths, manifest_path),
                )
            )
    if current_manifest.get("unsafe_entry_count") != 0:
        checks.append(
            CheckResult(
                code="B60_RAW_UNSAFE_ENTRY",
                message="Raw material contains unsafe symlink entries",
                location=str(raw_dir),
            )
        )
    if bundle.get("raw_root_sha256") != current_manifest.get("root_sha256"):
        checks.append(
            CheckResult(
                code="B60_BUNDLE_RAW_BINDING_INVALID",
                message="Acceptance bundle is not bound to the current raw root",
                location="raw_root_sha256",
            )
        )
    expected_manifest_ref = _reference(manifest_path, paths)
    if bundle.get("raw_manifest") != expected_manifest_ref:
        checks.append(
            CheckResult(
                code="B60_BUNDLE_MANIFEST_BINDING_INVALID",
                message="Acceptance bundle is not bound to the exact raw manifest",
                location="raw_manifest",
            )
        )
    if not seed_path.is_file():
        checks.append(
            CheckResult(
                code="B60_RUNTIME_SEED_MISSING",
                message="Local runtime Basic60 seed is missing",
                location=str(seed_path),
            )
        )
        return stored_manifest, None, None
    try:
        actual_seed = _json_object(seed_path)
        expected_seed = build_basic60_seed(raw_dir, current_manifest)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        checks.append(
            CheckResult(
                code="B60_RUNTIME_SEED_INVALID",
                message=str(error),
                location=str(seed_path),
            )
        )
        return stored_manifest, None, None
    if _canonical_sha256(actual_seed) != _canonical_sha256(expected_seed):
        checks.append(
            CheckResult(
                code="B60_RUNTIME_SEED_DRIFT",
                message="Runtime seed is not the deterministic normalization of current raw input",
                location=str(seed_path),
            )
        )
    expected_seed_ref = _reference(seed_path, paths)
    if stored_manifest.get("seed") != expected_seed_ref:
        checks.append(
            CheckResult(
                code="B60_MANIFEST_SEED_BINDING_INVALID",
                message="Raw manifest is not bound to the exact deterministic runtime seed",
                location="raw_manifest.seed",
            )
        )
    if bundle.get("runtime_seed") != expected_seed_ref:
        checks.append(
            CheckResult(
                code="B60_BUNDLE_SEED_BINDING_INVALID",
                message="Acceptance bundle is not bound to the exact local runtime seed",
                location="runtime_seed",
            )
        )
    sample_candidate: dict[str, Any] | None = None
    if "single_excel_review" not in bundle:
        acceptance = bundle.get("acceptance")
        sample_reference = (
            acceptance.get("sample_candidate") if isinstance(acceptance, dict) else None
        )
        sample_candidate = _validate_sample_candidate(
            checks,
            paths,
            seed_path,
            actual_seed,
            sample_reference,
        )
    return stored_manifest, actual_seed, sample_candidate


def _validate_source_registry_for_excel_review(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    source: Any,
) -> dict[str, Any] | None:
    """Validate source identity/provenance without making a license decision.

    Under the single-workbook path the project approver reviews the data and sources
    together. The system therefore verifies stable source identity and snapshot hashes,
    but it does not infer a second per-source permission state from terms metadata.
    """

    if not isinstance(source, dict):
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_SOURCE_REGISTRY_INVALID",
                message="Single-workbook review requires the Basic60 source registry",
                location="source_admission",
            )
        )
        return None
    registry_path = _validate_file_reference(
        checks,
        paths,
        source.get("source_registry"),
        code="B60_EXCEL_REVIEW_SOURCE_REGISTRY_INVALID",
        location="source_admission.source_registry",
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if registry_path is None:
        return None
    try:
        registry = _json_object(registry_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_SOURCE_REGISTRY_INVALID",
                message="Single-workbook source registry must be a JSON object",
                location="source_admission.source_registry",
            )
        )
        return None
    sources = registry.get("sources")
    if (
        registry.get("profile_id") != PROFILE_ID
        or registry.get("release_id") != RELEASE_ID
        or registry.get("template_only") is not False
        or registry.get("source_declaration") != "official_public_data"
        or not isinstance(sources, list)
    ):
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_SOURCE_REGISTRY_INVALID",
                message="Source registry identity or official-public declaration is invalid",
                location="source_admission.source_registry",
            )
        )
        return None
    current_manifest = _validate_country_source_uri_manifest(checks, paths, raw_dir, registry)
    indexed = {
        str(item.get("source_ref")): item
        for item in sources
        if isinstance(item, dict) and str(item.get("source_ref") or "")
    }
    expected_domains = {"country_identity", "macro", "energy"}
    actual_domains = {
        str(item.get("data_domain"))
        for item in indexed.values()
        if item.get("source_role") == "primary" and item.get("status") == "active"
    }
    entries_valid = (
        len(indexed) == 3
        and actual_domains == expected_domains
        and all(
            not SOURCE_PERMISSION_FIELDS.intersection(item)
            and str(item.get("provider") or "").strip()
            and str(item.get("dataset") or "").strip()
            and isinstance(item.get("snapshots"), list)
            and bool(item["snapshots"])
            and all(
                isinstance(snapshot, dict)
                and str(snapshot.get("snapshot_ref") or "").strip()
                and _aware_timestamp(snapshot.get("captured_at"))
                and SHA256_PATTERN.fullmatch(str(snapshot.get("content_sha256") or ""))
                for snapshot in item["snapshots"]
            )
            for item in indexed.values()
        )
    )
    if not entries_valid or current_manifest is None:
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_SOURCE_REGISTRY_INVALID",
                message=(
                    "Single-workbook review requires three active source identities with "
                    "stable snapshot hashes and no system-derived permission fields"
                ),
                location="source_admission.source_registry",
            )
        )
    domain_map = {
        "country_profile": "country_identity",
        "macroeconomic": "macro",
        "energy": "energy",
    }
    domains = source.get("domains")
    domain_bindings_valid = False
    if isinstance(domains, list) and len(domains) == len(domain_map):
        domain_bindings_valid = True
        for domain_id, data_domain in domain_map.items():
            domain = next(
                (
                    item
                    for item in domains
                    if isinstance(item, dict) and item.get("domain_id") == domain_id
                ),
                None,
            )
            registered = indexed.get(str(domain.get("primary_source_id") or "")) if domain else None
            if registered is None or registered.get("data_domain") != data_domain:
                domain_bindings_valid = False
                break
    if not domain_bindings_valid:
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_SOURCE_BINDING_INVALID",
                message=(
                    "Workbook source domains are not bound to the three active registry entries"
                ),
                location="source_admission.domains",
            )
        )
    return registry


def _excel_review_workbook_metadata(
    workbook_path: Path,
) -> dict[str, str | int]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    try:
        if EXCEL_REVIEW_MACHINE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Workbook is missing {EXCEL_REVIEW_MACHINE_SHEET}")
        worksheet = workbook[EXCEL_REVIEW_MACHINE_SHEET]
        metadata: dict[str, str | int] = {}
        for key_value, value in worksheet.iter_rows(
            min_col=1,
            max_col=2,
            values_only=True,
        ):
            if key_value is None:
                continue
            key = str(key_value).strip()
            if not key or key in metadata:
                raise ValueError("Workbook machine metadata contains a blank or duplicate key")
            if not isinstance(value, (str, int)):
                raise ValueError(f"Workbook machine metadata value is invalid: {key}")
            metadata[key] = value
        return metadata
    finally:
        workbook.close()


def _validate_single_excel_review(
    checks: list[CheckResult],
    paths: RepositoryPaths,
    raw_dir: Path,
    gate: Any,
    seed_path: Path,
    seed: dict[str, Any] | None,
    summary: dict[str, Any],
) -> dict[str, Any] | None:
    location = "single_excel_review"
    if not isinstance(gate, dict):
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_GATE_INVALID",
                message="Single-workbook review gate is required",
                location=location,
            )
        )
        return None
    header_valid = (
        gate.get("schema_version") == EXCEL_REVIEW_GATE_SCHEMA
        and gate.get("review_mode") == EXCEL_REVIEW_MODE
        and gate.get("review_scope") == EXCEL_REVIEW_SCOPE
        and gate.get("formal_gate_status") == FORMAL_GATE_STATUS
        and gate.get("does_not_complete_formal_d1_d4") is True
        and gate.get("does_not_authorize_p0_or_production") is True
    )
    if not header_valid or not isinstance(seed, dict):
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_GATE_INVALID",
                message="Single-workbook gate identity or private-only boundary is invalid",
                location=location,
            )
        )
        return None
    workbook_path = _validate_file_reference(
        checks,
        paths,
        gate.get("workbook"),
        code="B60_EXCEL_REVIEW_WORKBOOK_INVALID",
        location=f"{location}.workbook",
        prefix=EXCEL_REVIEW_OUTPUT_PREFIX,
    )
    source_manifest_path = _validate_file_reference(
        checks,
        paths,
        gate.get("source_uri_manifest"),
        code="B60_EXCEL_REVIEW_PAYLOAD_INVALID",
        location=f"{location}.source_uri_manifest",
        prefix=PurePosixPath("data/basic60/candidates"),
    )
    expected_seed_reference = _reference(seed_path, paths)
    if gate.get("runtime_seed") != expected_seed_reference:
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_PAYLOAD_INVALID",
                message="Review workbook is not bound to the exact candidate seed",
                location=f"{location}.runtime_seed",
            )
        )
    payload: dict[str, Any] | None = None
    if source_manifest_path is not None:
        try:
            source_manifest = _json_object(source_manifest_path)
            payload = build_basic60_excel_review_payload(
                raw_dir,
                seed,
                seed_sha256=sha256_file(seed_path),
                source_uri_manifest=source_manifest,
                source_uri_manifest_sha256=sha256_file(source_manifest_path),
            )
        except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
            checks.append(
                CheckResult(
                    code="B60_EXCEL_REVIEW_PAYLOAD_INVALID",
                    message=str(error),
                    location=location,
                )
            )
    expected_payload_hash = _canonical_sha256(payload) if payload is not None else None
    expected_counts = {
        "country_count": summary["country_count"],
        "macro_annual_record_count": summary["macro_annual_record_count"],
        "energy_latest_record_count": summary["energy_latest_record_count"],
        "available_observation_count": summary["available_observation_count"],
        "pending_observation_count": summary["pending_observation_count"],
        "source_binding_count": len(payload["source_bindings"]) if payload is not None else 0,
    }
    if (
        expected_payload_hash is None
        or gate.get("canonical_payload_sha256") != expected_payload_hash
        or gate.get("raw_root_sha256") != seed.get("source_root_sha256")
        or gate.get("machine_counts") != expected_counts
    ):
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_PAYLOAD_INVALID",
                message="Workbook gate does not match the current canonical data/source payload",
                location=location,
            )
        )
    if workbook_path is not None and source_manifest_path is not None and expected_payload_hash:
        try:
            metadata = _excel_review_workbook_metadata(workbook_path)
        except (OSError, ValueError, KeyError) as error:
            checks.append(
                CheckResult(
                    code="B60_EXCEL_REVIEW_WORKBOOK_INVALID",
                    message=str(error),
                    location=f"{location}.workbook",
                )
            )
        else:
            expected_metadata: dict[str, str | int] = {
                "schema_version": EXCEL_REVIEW_WORKBOOK_SCHEMA,
                "release_id": RELEASE_ID,
                "profile_id": PROFILE_ID,
                "canonical_payload_sha256": expected_payload_hash,
                "country_count": EXPECTED_COUNTRIES,
                "macro_row_count": EXPECTED_MACRO_ROWS,
                "energy_row_count": EXPECTED_ENERGY_ROWS,
                "available_observation_count": EXPECTED_AVAILABLE_VALUES,
                "pending_observation_count": EXPECTED_PENDING_VALUES,
                "seed_sha256": sha256_file(seed_path),
                "source_uri_manifest_sha256": sha256_file(source_manifest_path),
            }
            if metadata != expected_metadata:
                checks.append(
                    CheckResult(
                        code="B60_EXCEL_REVIEW_WORKBOOK_INVALID",
                        message="Workbook machine metadata does not match its canonical payload",
                        location=f"{location}.workbook.{EXCEL_REVIEW_MACHINE_SHEET}",
                    )
                )

    approval_reference = gate.get("approval")
    if approval_reference is None:
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_APPROVAL_PENDING",
                message=(
                    "kevin must approve the frozen Basic60 data-and-source workbook before "
                    "private publication and use"
                ),
                location=f"{location}.approval",
            )
        )
        return None
    approval_path = _validate_file_reference(
        checks,
        paths,
        approval_reference,
        code="B60_EXCEL_REVIEW_APPROVAL_INVALID",
        location=f"{location}.approval",
        prefix=PurePosixPath("data/basic60/evidence"),
    )
    if approval_path is None or workbook_path is None or expected_payload_hash is None:
        return None
    try:
        approval = _json_object(approval_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError):
        approval = {}
    valid_approval = (
        approval.get("schema_version") == EXCEL_REVIEW_APPROVAL_SCHEMA
        and approval.get("record_type") == "basic60_single_excel_review_approval"
        and str(approval.get("evidence_id") or "").strip()
        and approval.get("release_id") == RELEASE_ID
        and approval.get("profile_id") == PROFILE_ID
        and approval.get("decision") == "approved"
        and approval.get("person_name") == "kevin"
        and approval.get("role") == "项目批准人"
        and _aware_timestamp(approval.get("approved_at"))
        and approval.get("review_scope") == EXCEL_REVIEW_SCOPE
        and approval.get("authorized_outcome") == "private_trial_ready"
        and approval.get("workbook") == gate.get("workbook")
        and approval.get("canonical_payload_sha256") == expected_payload_hash
        and approval.get("formal_gate_status") == FORMAL_GATE_STATUS
        and approval.get("does_not_complete_formal_d1_d4") is True
        and approval.get("does_not_authorize_p0_or_production") is True
        and approval.get("authorization_basis") == "explicit_user_approval_after_workbook_review"
        and approval.get("template_only") is False
    )
    if not valid_approval:
        checks.append(
            CheckResult(
                code="B60_EXCEL_REVIEW_APPROVAL_INVALID",
                message=(
                    "Approval must be kevin's timezone-aware project-approver decision for all "
                    "Basic60 data and sources, bound to the exact workbook and payload hashes"
                ),
                location=f"{location}.approval",
            )
        )
        return None
    return approval


def validate_basic60_private(
    paths: RepositoryPaths,
    raw_dir: Path,
    manifest_path: Path,
    seed_path: Path,
    bundle: dict[str, Any],
    l0_dir: Path | None = None,
) -> tuple[str, dict[str, Any], list[CheckResult]]:
    summary, checks = inspect_basic60_raw(raw_dir)
    header_valid = (
        bundle.get("schema_version") == SCHEMA_VERSION
        and bundle.get("profile_id") == PROFILE_ID
        and bundle.get("release_id") == RELEASE_ID
        and bundle.get("formal_gate_status") == FORMAL_GATE_STATUS
    )
    if not header_valid:
        checks.append(
            CheckResult(
                code="B60_BUNDLE_HEADER_INVALID",
                message=(
                    "Bundle must target BASIC60-PRIVATE-R1/basic60_private "
                    "with formal gates pending"
                ),
                location="bundle",
            )
        )

    subjects = approval_subjects(bundle)
    revocation = bundle.get("revocation")
    if revocation is not None:
        revocation_checks: list[CheckResult] = []
        if not isinstance(revocation, dict) or not str(revocation.get("reason") or "").strip():
            revocation_checks.append(
                CheckResult(
                    code="B60_REVOCATION_INVALID",
                    message="Revocation requires a non-empty reason",
                    location="revocation",
                )
            )
        else:
            _validate_approval(
                revocation_checks,
                paths,
                revocation.get("approval"),
                expected_role="项目批准人",
                subject_sha256=subjects["revocation"],
                location="revocation.approval",
            )
        if not revocation_checks and header_valid:
            return "revoked", summary, []
        return "not_ready", summary, [*checks, *revocation_checks]

    if bundle.get("template_only") is not False:
        checks.append(
            CheckResult(
                code="B60_TEMPLATE_UNCOPIED",
                message="Copy the generated template before completing private-trial evidence",
                location="template_only",
            )
        )

    stored_manifest, actual_seed, sample_candidate = _validate_manifest_and_seed(
        checks,
        paths,
        raw_dir,
        manifest_path,
        seed_path,
        bundle,
    )
    _validate_profile_reference(checks, paths, bundle.get("profile"))
    _validate_baseline_decision(
        checks,
        paths,
        bundle.get("baseline_decision"),
        bundle.get("baseline_confirmation"),
        bundle.get("baseline_candidate"),
        bundle.get("baseline_confirmation_template"),
    )
    _validate_baseline_clarification(
        checks,
        paths,
        bundle.get("baseline_clarification"),
        bundle.get("baseline_candidate"),
        bundle.get("baseline_confirmation"),
        bundle.get("baseline_decision"),
    )
    _validate_usage_amendment(checks, paths, bundle.get("usage_amendment"))

    source = bundle.get("source_admission")
    if isinstance(source, dict) and isinstance(actual_seed, dict):
        _validate_ai_usage_policy(
            checks,
            paths,
            bundle.get("ai_usage_policy"),
            source.get("ai_usage_policy"),
            actual_seed,
            bundle.get("runtime_seed"),
            bundle.get("usage_amendment"),
        )
    else:
        checks.append(
            CheckResult(
                code="B60_AI_USAGE_POLICY_INVALID",
                message="AI usage policy requires a valid candidate seed and source admission",
                location="ai_usage_policy",
            )
        )
    if "single_excel_review" in bundle:
        _validate_domains(checks, source)
        _validate_source_matrix(checks, paths, raw_dir, source)
        _validate_source_registry_for_excel_review(checks, paths, raw_dir, source)
        _validate_processing(
            checks,
            paths,
            bundle.get("processing"),
            stored_manifest,
            seed_path,
            actual_seed,
        )
        _validate_single_excel_review(
            checks,
            paths,
            raw_dir,
            bundle.get("single_excel_review"),
            seed_path,
            actual_seed,
            summary,
        )
        status = "not_ready" if checks else "private_trial_ready"
        return status, summary, checks

    _validate_domains(checks, source)
    _validate_source_matrix(checks, paths, raw_dir, source)
    source_l0: Path | None = None
    if l0_dir is not None:
        try:
            source_l0 = _validated_l0_directory(paths, raw_dir, l0_dir)
        except (OSError, ValueError):
            source_l0 = None
    source_registry = _validate_source_registry(
        checks,
        paths,
        raw_dir,
        source,
        l0_dir=source_l0,
    )
    if isinstance(source, dict):
        _validate_d1_two_stage_admission(
            checks,
            paths,
            bundle,
            source,
            source_registry,
            source_l0,
        )

    collection = bundle.get("collection")
    _validate_collection(
        checks,
        paths,
        raw_dir,
        collection,
        raw_root_sha256=str(bundle.get("raw_root_sha256") or ""),
        source_subject_sha256=subjects["source"],
        source_admission=source,
        source_registry=source_registry,
        l0_dir=l0_dir,
    )
    if isinstance(collection, dict):
        _validate_approval(
            checks,
            paths,
            collection.get("approval"),
            expected_role=BATCH_APPROVAL_ROLE,
            subject_sha256=subjects["collection"],
            location="collection.approval",
        )

    _validate_processing(
        checks,
        paths,
        bundle.get("processing"),
        stored_manifest,
        seed_path,
        actual_seed,
    )
    acceptance = bundle.get("acceptance")
    reviewed_sample_people = _validate_acceptance(
        checks,
        paths,
        acceptance,
        summary,
        sample_candidate,
        actual_seed,
        bundle.get("runtime_seed"),
        source_registry,
        bundle,
    )
    if isinstance(acceptance, dict):
        sample_approvals = acceptance.get("sample_approvals")
        sample_people: list[str] = []
        for index in range(2):
            approval = (
                sample_approvals[index]
                if isinstance(sample_approvals, list) and index < len(sample_approvals)
                else None
            )
            person = _validate_approval(
                checks,
                paths,
                approval,
                expected_role="数据质量复核人",
                subject_sha256=subjects[f"sample_review_{index + 1}"],
                location=f"acceptance.sample_approvals.{index}",
            )
            if person:
                sample_people.append(person)
        if len(sample_people) == 2 and len(set(sample_people)) != 2:
            checks.append(
                CheckResult(
                    code="B60_SAMPLE_REVIEWERS_NOT_INDEPENDENT",
                    message="The two frozen-sample reviewers must be different people",
                    location="acceptance.sample_approvals",
                )
            )
        if reviewed_sample_people is not None and set(sample_people) != set(reviewed_sample_people):
            checks.append(
                CheckResult(
                    code="B60_SAMPLE_REVIEWER_SUBJECT_MISMATCH",
                    message=(
                        "Frozen-sample result reviewers must match the two hash-bound sample "
                        "approvers"
                    ),
                    location="acceptance.sample_approvals",
                )
            )
        final_approvals = acceptance.get("final_approvals")
        for index, role in enumerate(FINAL_APPROVAL_ROLES):
            approval = (
                final_approvals[index]
                if isinstance(final_approvals, list) and index < len(final_approvals)
                else None
            )
            final_person = _validate_approval(
                checks,
                paths,
                approval,
                expected_role=role,
                subject_sha256=subjects["acceptance"],
                location=f"acceptance.final_approvals.{index}",
            )
            if role == "项目批准人" and final_person is not None and final_person != "kevin":
                checks.append(
                    CheckResult(
                        code="B60_PROJECT_APPROVER_INVALID",
                        message="Basic60 final project approver must be kevin",
                        location=f"acceptance.final_approvals.{index}",
                    )
                )

    status = "not_ready" if checks else "private_trial_ready"
    return status, summary, checks


def _assessment(
    status: str,
    summary: dict[str, Any],
    checks: list[CheckResult],
    *,
    raw_manifest: dict[str, Any] | None = None,
    release_bundle_sha256: str | None = None,
) -> dict[str, Any]:
    if status not in PRIVATE_OUTCOMES:
        raise ValueError(f"Unsupported Basic60 private outcome: {status}")
    payload: dict[str, Any] = {
        "schema_version": SCHEMA_VERSION,
        "profile_id": PROFILE_ID,
        "release_id": RELEASE_ID,
        "status": status,
        "ready": status == "private_trial_ready",
        "formal_gate_status": FORMAL_GATE_STATUS,
        "does_not_complete_formal_d1_d4": True,
        "does_not_authorize_p0_or_production": True,
        "machine_counts": summary,
        "checks": [item.to_dict() for item in checks],
    }
    if raw_manifest is not None:
        payload["raw_root_sha256"] = raw_manifest.get("root_sha256")
        payload["raw_file_count"] = raw_manifest.get("included_file_count")
        payload["excluded_zone_identifier_count"] = raw_manifest.get(
            "excluded_zone_identifier_count"
        )
    if release_bundle_sha256 is not None:
        if not SHA256_PATTERN.fullmatch(release_bundle_sha256):
            raise ValueError("Release bundle hash must be a SHA-256")
        payload["release_bundle_sha256"] = release_bundle_sha256
    return payload


def load_and_assess_basic60_private(
    paths: RepositoryPaths,
    raw_dir: Path,
    manifest_path: Path,
    seed_path: Path,
    bundle_path: Path,
    l0_dir: Path | None = None,
) -> dict[str, Any]:
    try:
        bundle = _json_object(bundle_path)
    except (OSError, UnicodeDecodeError, json.JSONDecodeError, ValueError) as error:
        summary, raw_checks = inspect_basic60_raw(raw_dir)
        checks = [
            *raw_checks,
            CheckResult(
                code="B60_ACCEPTANCE_BUNDLE_MISSING",
                message=str(error),
                location=str(bundle_path),
            ),
        ]
        return _assessment("not_ready", summary, checks)
    status, summary, checks = validate_basic60_private(
        paths,
        raw_dir,
        manifest_path,
        seed_path,
        bundle,
        l0_dir,
    )
    return _assessment(
        status,
        summary,
        checks,
        release_bundle_sha256=sha256_file(bundle_path),
    )


def write_basic60_ready_seed(
    paths: RepositoryPaths,
    manifest_path: Path,
    candidate_seed_path: Path,
    bundle_path: Path,
    output_path: Path,
    *,
    validation_report_sha256: str,
) -> Path:
    """Write an API-importable seed only from an already validated ready bundle."""

    if output_path.resolve() == candidate_seed_path.resolve():
        raise ValueError("Ready seed must not overwrite the fail-closed candidate seed")
    runtime_root = (paths.root / "runtime/basic60").resolve()
    if not output_path.resolve().is_relative_to(runtime_root):
        raise ValueError("Ready seed output must remain under runtime/basic60")
    raw_manifest = _json_object(manifest_path)
    candidate_seed = _json_object(candidate_seed_path)
    bundle = _json_object(bundle_path)
    source_reference = bundle.get("source_admission", {}).get("source_registry")
    if not isinstance(source_reference, dict) or not isinstance(source_reference.get("path"), str):
        raise ValueError("Approved source registry reference is missing")
    pure = PurePosixPath(source_reference["path"])
    if (
        pure.is_absolute()
        or ".." in pure.parts
        or not pure.is_relative_to(PurePosixPath("data/basic60/evidence"))
    ):
        raise ValueError("Approved source registry reference is outside data/basic60/evidence")
    source_registry_path = paths.root.joinpath(*pure.parts)
    if sha256_file(source_registry_path) != source_reference.get("sha256"):
        raise ValueError("Approved source registry hash changed after validation")
    source_registry = _json_object(source_registry_path)
    single_excel_review = bundle.get("single_excel_review")
    if isinstance(single_excel_review, dict):
        approval_reference = single_excel_review.get("approval")
        if not isinstance(approval_reference, dict) or not isinstance(
            approval_reference.get("path"), str
        ):
            raise ValueError("Ready seed requires the validated workbook approval")
        approval_pure = PurePosixPath(str(approval_reference["path"]))
        if (
            approval_pure.is_absolute()
            or ".." in approval_pure.parts
            or not approval_pure.is_relative_to(PurePosixPath("data/basic60/evidence"))
        ):
            raise ValueError("Workbook approval evidence is outside data/basic60/evidence")
        approval_path = paths.root.joinpath(*approval_pure.parts)
        if sha256_file(approval_path) != approval_reference.get("sha256"):
            raise ValueError("Workbook approval evidence hash changed after validation")
        approval = _json_object(approval_path)
        signed_times = (
            [str(approval.get("approved_at"))]
            if _aware_timestamp(approval.get("approved_at"))
            else []
        )
    else:
        acceptance = bundle.get("acceptance")
        final_approvals = (
            acceptance.get("final_approvals") if isinstance(acceptance, dict) else None
        )
        approval_items = final_approvals if isinstance(final_approvals, list) else []
        signed_times = [
            str(item.get("signed_at"))
            for item in approval_items
            if isinstance(item, dict) and _aware_timestamp(item.get("signed_at"))
        ]
    if not signed_times:
        raise ValueError("Ready seed requires a completed private-trial approval")
    ready_seed = build_basic60_ready_seed(
        candidate_seed,
        raw_manifest,
        source_registry,
        bundle,
        release_bundle_sha256=sha256_file(bundle_path),
        validation_report_sha256=validation_report_sha256,
        reviewed_at=max(signed_times),
    )
    _write_json(output_path, ready_seed)
    return output_path


def _release_signature(approval: Any, role: str) -> dict[str, str]:
    if not isinstance(approval, dict) or not isinstance(approval.get("evidence"), dict):
        raise ValueError(f"Missing validated release signature for {role}")
    person = str(approval.get("person_name") or "")
    signed_at = str(approval.get("signed_at") or "")
    evidence_hash = str(approval["evidence"].get("sha256") or "")
    if not _valid_person(person) or not _aware_timestamp(signed_at):
        raise ValueError(f"Invalid validated release signature for {role}")
    if not SHA256_PATTERN.fullmatch(evidence_hash):
        raise ValueError(f"Release signature evidence hash is invalid for {role}")
    return {
        "person_id": person,
        "role": role,
        "signed_at": signed_at,
        "signature_sha256": evidence_hash,
    }


def write_basic60_release_authorization(
    paths: RepositoryPaths,
    bundle_path: Path,
    ready_seed_path: Path,
    output_path: Path,
    *,
    validation_report_sha256: str,
) -> Path:
    """Transcribe an API authorization from exact approvals already validated by the CLI."""

    runtime_root = (paths.root / "runtime/basic60").resolve()
    if not output_path.resolve().is_relative_to(runtime_root):
        raise ValueError("Release authorization output must remain under runtime/basic60")
    bundle = _json_object(bundle_path)
    single_excel_review = bundle.get("single_excel_review")
    if isinstance(single_excel_review, dict):
        approval_reference = single_excel_review.get("approval")
        if not isinstance(approval_reference, dict) or not isinstance(
            approval_reference.get("path"), str
        ):
            raise ValueError("Validated workbook approval reference is required")
        approval_pure = PurePosixPath(str(approval_reference["path"]))
        if (
            approval_pure.is_absolute()
            or ".." in approval_pure.parts
            or not approval_pure.is_relative_to(PurePosixPath("data/basic60/evidence"))
        ):
            raise ValueError("Workbook approval evidence is outside data/basic60/evidence")
        approval_path = paths.root.joinpath(*approval_pure.parts)
        if sha256_file(approval_path) != approval_reference.get("sha256"):
            raise ValueError("Workbook approval evidence hash changed after validation")
        approval = _json_object(approval_path)
        if (
            approval.get("person_name") != "kevin"
            or approval.get("role") != "项目批准人"
            or approval.get("decision") != "approved"
            or not _aware_timestamp(approval.get("approved_at"))
            or approval.get("workbook") != single_excel_review.get("workbook")
            or approval.get("canonical_payload_sha256")
            != single_excel_review.get("canonical_payload_sha256")
        ):
            raise ValueError("Validated workbook approval content changed")
        authorization = {
            "schema_version": "basic60.release-authorization.v1",
            "release_id": RELEASE_ID,
            "release_profile": PROFILE_ID,
            "status": "private_trial_ready",
            "formal_gate_status": FORMAL_GATE_STATUS,
            "release_bundle_sha256": sha256_file(bundle_path),
            "validation_report_sha256": validation_report_sha256,
            "seed_artifact_sha256": sha256_file(ready_seed_path),
            "approved_at": approval["approved_at"],
            "pbd_decision_id": DECISION_ID,
            "signoffs": {
                "single_excel_review": {
                    "status": "approved",
                    "review_scope": EXCEL_REVIEW_SCOPE,
                    "workbook": single_excel_review["workbook"],
                    "canonical_payload_sha256": single_excel_review["canonical_payload_sha256"],
                    "signature": {
                        "person_id": "kevin",
                        "role": "project_approver",
                        "signed_at": approval["approved_at"],
                        "signature_sha256": approval_reference["sha256"],
                    },
                }
            },
            "does_not_complete_formal_d1_d4": True,
            "does_not_authorize_p0_or_production": True,
        }
        _write_json(output_path, authorization)
        return output_path

    source = bundle.get("source_admission")
    collection = bundle.get("collection")
    acceptance = bundle.get("acceptance")
    if (
        not isinstance(source, dict)
        or not isinstance(collection, dict)
        or not isinstance(acceptance, dict)
    ):
        raise ValueError("Validated source, collection, and acceptance sections are required")

    source_approvals = source.get("approvals")
    sample_approvals = acceptance.get("sample_approvals")
    final_approvals = acceptance.get("final_approvals")
    if (
        not isinstance(source_approvals, list)
        or not isinstance(sample_approvals, list)
        or not isinstance(final_approvals, list)
    ):
        raise ValueError("Validated D1 and D4 approvals are required")
    if len(source_approvals) != 1:
        raise ValueError("Exactly one validated D1 project approver is required")
    source_by_role = {
        str(item.get("role")): item for item in source_approvals if isinstance(item, dict)
    }
    final_by_role = {
        str(item.get("role")): item for item in final_approvals if isinstance(item, dict)
    }
    if len(sample_approvals) != 2:
        raise ValueError("Exactly two validated quality reviewers are required")

    d1_signatures = [_release_signature(source_by_role.get("项目批准人"), "project_approver")]
    if d1_signatures[0]["person_id"] != "kevin":
        raise ValueError("Basic60 D1 project approver must be kevin")
    d2_signatures = [_release_signature(collection.get("approval"), "authorized_data_reviewer")]
    d4_signatures = [
        _release_signature(final_by_role.get("数据负责人"), "data_owner"),
        _release_signature(sample_approvals[0], "quality_reviewer_1"),
        _release_signature(sample_approvals[1], "quality_reviewer_2"),
        _release_signature(final_by_role.get("合规负责人"), "compliance_reviewer"),
        _release_signature(final_by_role.get("数据工程负责人"), "technical_reviewer"),
        _release_signature(final_by_role.get("项目批准人"), "project_approver"),
    ]
    if d4_signatures[-1]["person_id"] != "kevin":
        raise ValueError("Basic60 project approver must be kevin")
    signed_datetimes = [
        datetime.fromisoformat(item["signed_at"].replace("Z", "+00:00"))
        for item in [*d1_signatures, *d2_signatures, *d4_signatures]
    ]
    authorization = {
        "schema_version": "basic60.release-authorization.v1",
        "release_id": RELEASE_ID,
        "release_profile": PROFILE_ID,
        "status": "private_trial_ready",
        "release_bundle_sha256": sha256_file(bundle_path),
        "validation_report_sha256": validation_report_sha256,
        "seed_artifact_sha256": sha256_file(ready_seed_path),
        "approved_at": max(signed_datetimes).isoformat(),
        "pbd_decision_id": DECISION_ID,
        "signoffs": {
            "d1_source_admission": {"status": "approved", "signatures": d1_signatures},
            "d2_batch_close": {"status": "approved", "signatures": d2_signatures},
            "d4_private_trial_acceptance": {
                "status": "approved",
                "signatures": d4_signatures,
            },
        },
    }
    _write_json(output_path, authorization)
    return output_path


def _runtime_attestation_message(claims: dict[str, str]) -> bytes:
    """Return the stable message authenticated by the runtime trust key."""

    return json.dumps(
        claims,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")


def validate_runtime_attestation_key(trust_key: str | None) -> str:
    """Reject missing, weak, and repository-known example trust keys."""

    value = trust_key or ""
    folded = value.casefold()
    if value in RUNTIME_ATTESTATION_FORBIDDEN_KEYS or any(
        marker in folded for marker in RUNTIME_ATTESTATION_FORBIDDEN_MARKERS
    ):
        raise ValueError(
            f"{RUNTIME_ATTESTATION_KEY_ENV} must not use a known placeholder or example value"
        )
    if len(value.encode("utf-8")) < RUNTIME_ATTESTATION_KEY_MIN_BYTES:
        raise ValueError(
            f"{RUNTIME_ATTESTATION_KEY_ENV} must contain at least "
            f"{RUNTIME_ATTESTATION_KEY_MIN_BYTES} UTF-8 bytes"
        )
    return value


def write_basic60_runtime_attestation(
    paths: RepositoryPaths,
    bundle_path: Path,
    ready_seed_path: Path,
    release_authorization_path: Path,
    output_path: Path,
    *,
    validation_report_sha256: str,
    trust_key: str,
) -> Path:
    """Emit a machine HMAC proving the final artifacts came through the validator.

    This is an independent runtime trust-root check.  It authenticates machine
    validation output and must never be represented as a human signature.
    """

    runtime_root = (paths.root / "runtime/basic60").resolve()
    if not output_path.resolve().is_relative_to(runtime_root):
        raise ValueError("Runtime attestation output must remain under runtime/basic60")
    trust_key = validate_runtime_attestation_key(trust_key)
    if not SHA256_PATTERN.fullmatch(validation_report_sha256):
        raise ValueError("Validation report hash must be a SHA-256")

    bundle_sha256 = sha256_file(bundle_path)
    seed_sha256 = sha256_file(ready_seed_path)
    authorization_sha256 = sha256_file(release_authorization_path)
    seed = _json_object(ready_seed_path)
    authorization = _json_object(release_authorization_path)
    release = seed.get("release")
    if not isinstance(release, dict):
        raise ValueError("Ready seed release metadata is missing")

    expected_seed = {
        "release_id": RELEASE_ID,
        "release_profile": PROFILE_ID,
        "status": "private_trial_ready",
        "release_bundle_sha256": bundle_sha256,
        "validation_report_sha256": validation_report_sha256,
    }
    seed_mismatches = [
        field for field, expected in expected_seed.items() if release.get(field) != expected
    ]
    if seed_mismatches:
        raise ValueError(
            "Ready seed does not match runtime attestation inputs: " + ", ".join(seed_mismatches)
        )

    expected_authorization = {
        "release_id": RELEASE_ID,
        "release_profile": PROFILE_ID,
        "status": "private_trial_ready",
        "release_bundle_sha256": bundle_sha256,
        "validation_report_sha256": validation_report_sha256,
        "seed_artifact_sha256": seed_sha256,
    }
    authorization_mismatches = [
        field
        for field, expected in expected_authorization.items()
        if authorization.get(field) != expected
    ]
    if authorization_mismatches:
        raise ValueError(
            "Release authorization does not match runtime attestation inputs: "
            + ", ".join(authorization_mismatches)
        )

    claims = {
        "release_id": RELEASE_ID,
        "release_profile": PROFILE_ID,
        "validation_report_sha256": validation_report_sha256,
        "release_bundle_sha256": bundle_sha256,
        "seed_artifact_sha256": seed_sha256,
        "release_authorization_sha256": authorization_sha256,
    }
    mac_sha256 = hmac.new(
        trust_key.encode("utf-8"),
        _runtime_attestation_message(claims),
        hashlib.sha256,
    ).hexdigest()
    attestation = {
        "schema_version": "basic60.runtime-attestation.v1",
        "artifact_kind": "machine_validation_attestation",
        "algorithm": "HMAC-SHA256",
        "claims": claims,
        "mac_sha256": mac_sha256,
    }
    _write_json(output_path, attestation)
    return output_path
