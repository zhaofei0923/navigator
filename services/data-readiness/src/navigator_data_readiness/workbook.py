from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

CellValue = str | int | float | bool | None
RowRecord = dict[str, CellValue]


def load_read_only(path: Path) -> Any:
    return load_workbook(path, read_only=True, data_only=True)


def normalized_headers(values: Sequence[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        header = str(value).strip() if value is not None else f"column_{index}"
        seen[header] = seen.get(header, 0) + 1
        headers.append(header if seen[header] == 1 else f"{header}_{seen[header]}")
    return headers


def normalized_value(value: Any) -> CellValue:
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def table_records(
    sheet: Worksheet,
    *,
    header_row: int = 4,
    stop_after_blank_rows: int = 2,
) -> list[RowRecord]:
    rows = sheet.iter_rows(min_row=header_row, values_only=True)
    try:
        headers = normalized_headers(next(rows))
    except StopIteration:
        return []

    records: list[RowRecord] = []
    blank_rows = 0
    for raw_row in rows:
        values = [normalized_value(value) for value in raw_row[: len(headers)]]
        if all(value is None or str(value).strip() == "" for value in values):
            blank_rows += 1
            if blank_rows >= stop_after_blank_rows:
                break
            continue
        blank_rows = 0
        record: RowRecord = {
            headers[index]: value
            for index, value in enumerate(values)
            if value is not None and str(value).strip() != ""
        }
        records.append(record)
    return records


def rows_from_sections(
    sheet: Worksheet,
    sections: Iterable[tuple[str, int, int | None]],
) -> dict[str, list[RowRecord]]:
    result: dict[str, list[RowRecord]] = {}
    for name, header_row, end_row in sections:
        records = table_records(sheet, header_row=header_row, stop_after_blank_rows=1)
        if end_row is not None:
            records = records[: max(0, end_row - header_row)]
        result[name] = records
    return result
