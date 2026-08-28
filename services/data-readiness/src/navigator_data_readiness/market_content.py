"""One-review, independent country market-overview lifecycle.

This module never changes Basic60 seeds, formal gates, source material, or database rows.
Excel is read only. Publication needs an actual supplied human confirmation, not a
machine-generated approval. API readers use only the separately mounted published store.
"""

from __future__ import annotations

import argparse
import copy
import csv
import ctypes
import errno
import hashlib
import io
import json
import os
import re
import shutil
import sys
import tempfile
import zipfile
from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError

import openpyxl
from navigator_api.market_schemas import CountryMarketOverview
from navigator_api.market_storage import (
    MAX_BATCH_BYTES,
    VERSION_PATTERN,
    MarketContentError,
    MarketContentUnavailable,
    canonical_bytes,
    checked_path,
    checked_sha,
    content_sha256,
    ensure_no_symlinks,
    load_market_state,
    load_published_version,
    read_json,
)
from openpyxl.utils.exceptions import InvalidFileException

CANDIDATE_SCHEMA = "navigator.market-candidate.v1"
REVIEW_SCHEMA = "navigator.market-review.v1"
RECEIPT_SCHEMA = "navigator.market-review-import.v1"
CONFIRMATION_SCHEMA = "navigator.market-confirmation.v1"
EDIT_HEADERS = (
    "country_code",
    "content_version",
    "locale",
    "json_pointer",
    "original_value",
    "edited_value",
)
MARKET_COMMANDS = frozenset(
    {
        "validate-market-content",
        "import-market-review",
        "publish-market-content",
        "revoke-market-content",
        "rollback-market-content",
    }
)
MAX_REVIEW_ROWS = 250_000
PACKAGE_PATTERN = re.compile(r"^[A-Z][A-Z0-9-]{0,99}$")
_INDEX = r"(?:0|[1-9]\d*)"
_EDITABLE = re.compile(rf"^/locales/(?:zh-CN|en)/(?:title|paragraphs/{_INDEX}|disclaimer)$")


@dataclass(frozen=True)
class MarketCandidate:
    package_id: str
    countries: tuple[CountryMarketOverview, ...]

    def payload(self) -> dict[str, Any]:
        return {
            "schema_version": CANDIDATE_SCHEMA,
            "package_id": self.package_id,
            "countries": [
                country.model_dump(mode="json")
                for country in sorted(self.countries, key=lambda item: item.country_code)
            ],
        }

    @property
    def sha256(self) -> str:
        return content_sha256(self.payload())


def _now() -> str:
    return datetime.now(UTC).isoformat()


def _iso_time(value: object, field: str) -> str:
    if not isinstance(value, str):
        raise MarketContentError(f"{field} must be an ISO date or timezone-aware timestamp")
    try:
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            date.fromisoformat(value)
        else:
            timestamp = datetime.fromisoformat(value)
            if timestamp.tzinfo is None or timestamp.utcoffset() is None:
                raise ValueError("missing timezone")
    except ValueError as exc:
        raise MarketContentError(f"invalid {field}") from exc
    return value


def _time_before(first: str, second: str) -> bool:
    """Compare actual precision: date-only confirmations do not invent a signing time."""
    first_day = bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", first))
    second_day = bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", second))
    if first_day or second_day:
        return date.fromisoformat(first[:10]) < date.fromisoformat(second[:10])
    return datetime.fromisoformat(first) < datetime.fromisoformat(second)


def _country_code(value: str) -> str:
    code = value.upper()
    if not re.fullmatch(r"[A-Z]{3}", code) or code == "CHN":
        raise MarketContentError("invalid outbound country code")
    return code


def load_country_scope(profiles_path: Path) -> frozenset[str]:
    ensure_no_symlinks(profiles_path)
    with profiles_path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames or "iso3" not in reader.fieldnames:
            raise MarketContentError("country profiles require an iso3 column")
        codes = [row["iso3"] for row in reader]
    if (
        not codes
        or len(set(codes)) != len(codes)
        or any(not re.fullmatch(r"[A-Z]{3}", code or "") for code in codes)
    ):
        raise MarketContentError("country profiles contain invalid or duplicate ISO3 codes")
    return frozenset(code for code in codes if code != "CHN")


def _validated_candidate(
    package_id: object, payloads: list[dict[str, Any]], scope: frozenset[str]
) -> MarketCandidate:
    if not isinstance(package_id, str) or not PACKAGE_PATTERN.fullmatch(package_id):
        raise MarketContentError("invalid market candidate package_id")
    if not payloads or len(payloads) > len(scope):
        raise MarketContentError("candidate must contain at least one supported country")
    try:
        countries = tuple(CountryMarketOverview.model_validate(item) for item in payloads)
    except ValueError as exc:
        raise MarketContentError("candidate violates the public market-overview schema") from exc
    codes = [item.country_code for item in countries]
    if len(set(codes)) != len(codes) or not set(codes).issubset(scope):
        raise MarketContentError("candidate contains duplicate or unsupported countries")
    candidate = MarketCandidate(package_id, countries)
    if len(canonical_bytes(candidate.payload())) > MAX_BATCH_BYTES:
        raise MarketContentError("market candidate exceeds maximum batch size")
    return candidate


def load_market_candidate(candidate_dir: Path, profiles_path: Path) -> MarketCandidate:
    """Load a subset or complete outbound package, without mutating its files."""
    package = read_json(checked_path(candidate_dir, "package.json"))
    if (
        set(package) != {"schema_version", "package_id"}
        or package.get("schema_version") != CANDIDATE_SCHEMA
    ):
        raise MarketContentError("invalid candidate package.json")
    country_dir = checked_path(candidate_dir, "countries")
    if not country_dir.is_dir():
        raise MarketContentError("missing candidate countries directory")
    payloads = []
    for path in sorted(country_dir.iterdir()):
        if not re.fullmatch(r"[A-Z]{3}\.json", path.name):
            raise MarketContentError("countries directory permits only ISO3.json content objects")
        payload = read_json(path)
        if payload.get("country_code") != path.stem:
            raise MarketContentError("country filename and content disagree")
        payloads.append(payload)
    return _validated_candidate(
        package.get("package_id"), payloads, load_country_scope(profiles_path)
    )


def validate_market_content(candidate_dir: Path, profiles_path: Path) -> dict[str, Any]:
    candidate = load_market_candidate(candidate_dir, profiles_path)
    return {
        "status": "candidate_valid",
        "package_id": candidate.package_id,
        "candidate_sha256": candidate.sha256,
        "country_count": len(candidate.countries),
        "countries": sorted(item.country_code for item in candidate.countries),
        "published": False,
    }


def _pointer_part(value: str) -> str:
    return value.replace("~", "~0").replace("/", "~1")


def _string_leaves(value: object, pointer: str) -> Iterator[tuple[str, str]]:
    if isinstance(value, dict):
        for key, child in value.items():
            yield from _string_leaves(child, f"{pointer}/{_pointer_part(key)}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _string_leaves(child, f"{pointer}/{index}")
    elif isinstance(value, str):
        yield pointer, value
    else:
        raise MarketContentError("reviewable localized content must contain only string leaves")


def market_review_rows(candidate: MarketCandidate) -> list[tuple[str, str, str, str, str, str]]:
    """Authoring-neutral rows; the workbook producer controls its readable formatting."""
    rows = []
    for country in sorted(candidate.countries, key=lambda item: item.country_code):
        payload = country.model_dump(mode="json")
        for locale in ("zh-CN", "en"):
            for pointer, value in _string_leaves(payload["locales"][locale], f"/locales/{locale}"):
                rows.append(
                    (country.country_code, country.content_version, locale, pointer, value, value)
                )
    return rows


def _replace_leaf(payload: dict[str, Any], pointer: str, value: str) -> None:
    # Only a pointer already found in the exact original leaf inventory reaches this function.
    tokens = [part.replace("~1", "/").replace("~0", "~") for part in pointer.split("/")[1:]]
    current: Any = payload
    for token in tokens[:-1]:
        current = current[int(token)] if isinstance(current, list) else current[token]
    if isinstance(current, list):
        current[int(tokens[-1])] = value
    else:
        current[tokens[-1]] = value


def apply_market_review_rows(
    candidate: MarketCandidate,
    metadata: dict[str, object],
    rows: Sequence[Sequence[object]],
    scope: frozenset[str],
) -> MarketCandidate:
    """Apply a frozen row inventory. A changed original or structural edit fails closed."""
    if (
        metadata.get("schema_version") != REVIEW_SCHEMA
        or metadata.get("package_id") != candidate.package_id
        or metadata.get("candidate_sha256") != candidate.sha256
        or str(metadata.get("country_count")) != str(len(candidate.countries))
    ):
        raise MarketContentError("review workbook metadata does not match the original candidate")
    _iso_time(metadata.get("created_at"), "workbook created_at")
    expected = {(row[0], row[3]): row for row in market_review_rows(candidate)}
    payloads = {
        item.country_code: copy.deepcopy(item.model_dump(mode="json"))
        for item in candidate.countries
    }
    seen: set[tuple[str, str]] = set()
    for row in rows:
        if len(row) < 6:
            raise MarketContentError("review row is missing one of its six required columns")
        code, version, locale, pointer, original, edited = row[:6]
        if not all(isinstance(value, str) for value in (code, version, locale, pointer, original)):
            raise MarketContentError(
                "review identity, pointer and original value must be literal text"
            )
        key = (str(code), str(pointer))
        frozen = expected.get(key)
        if frozen is None or tuple(row[:5]) != frozen[:5] or key in seen:
            raise MarketContentError("review row is unknown, duplicated, or has a changed original")
        seen.add(key)
        if not isinstance(edited, str):
            raise MarketContentError("edited values must be literal text, including ISO dates")
        if edited != original and not _EDITABLE.fullmatch(str(pointer)):
            raise MarketContentError(
                "only overview titles, existing paragraphs and disclaimers are editable"
            )
        _replace_leaf(payloads[str(code)], str(pointer), edited)
    if seen != set(expected):
        raise MarketContentError("review workbook must retain every original localized content row")
    return _validated_candidate(candidate.package_id, list(payloads.values()), scope)


def _workbook_bytes(workbook_path: Path) -> bytes:
    ensure_no_symlinks(workbook_path)
    if not workbook_path.is_file() or workbook_path.stat().st_size > MAX_BATCH_BYTES:
        raise MarketContentError("missing or oversized review workbook")
    payload = workbook_path.read_bytes()
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 10_000
                or sum(item.file_size for item in entries) > 4 * MAX_BATCH_BYTES
            ):
                raise MarketContentError("review workbook exceeds safe decompression limits")
            if any(item.flag_bits & 1 for item in entries):
                raise MarketContentError("encrypted review workbooks are not supported")
    except zipfile.BadZipFile as exc:
        raise MarketContentError("review workbook must be an unencrypted .xlsx file") from exc
    return payload


def _read_workbook(payload: bytes) -> tuple[dict[str, object], list[tuple[object, ...]]]:
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except (
        KeyError,
        ValueError,
        OSError,
        zipfile.BadZipFile,
        InvalidFileException,
        ParseError,
    ) as exc:
        raise MarketContentError("invalid review workbook structure") from exc
    try:
        if not {"包信息", "内容编辑"}.issubset(workbook.sheetnames):
            raise MarketContentError("review workbook requires 包信息 and 内容编辑 worksheets")
        metadata_sheet = workbook["包信息"]
        edit_sheet = workbook["内容编辑"]
        if (metadata_sheet.max_row or 0) > 100 or (edit_sheet.max_row or 0) > MAX_REVIEW_ROWS + 1:
            raise MarketContentError("review worksheet dimensions exceed supported limits")

        def values(sheet: Any, columns: int) -> Iterator[tuple[object, ...]]:
            for cells in sheet.iter_rows(max_col=columns):
                if any(cell.data_type in {"f", "e"} for cell in cells):
                    raise MarketContentError("formulas and spreadsheet errors are not review text")
                yield tuple(cell.value for cell in cells)

        metadata_rows = values(metadata_sheet, 2)
        if next(metadata_rows, ()) != ("key", "value"):
            raise MarketContentError("包信息 must begin with key,value")
        metadata: dict[str, object] = {}
        for key, value in metadata_rows:
            if key is None and value is None:
                continue
            if not isinstance(key, str) or key in metadata:
                raise MarketContentError("invalid or duplicated workbook metadata key")
            metadata[key] = value
        rows = values(edit_sheet, 6)
        if next(rows, ()) != EDIT_HEADERS:
            raise MarketContentError("内容编辑 must retain the six canonical column headers")
        return metadata, [row for row in rows if any(value is not None for value in row)]
    except MarketContentError:
        raise
    except (KeyError, ValueError, IndexError, ParseError) as exc:
        raise MarketContentError("invalid review worksheet data") from exc
    finally:
        workbook.close()


def _writable_target(path: Path, repo_root: Path) -> Path:
    ensure_no_symlinks(path)
    resolved = path.resolve()
    repository = repo_root.resolve()
    protected = (
        repository / "doc",
        repository / "raw material",
        repository / "data",
        repository / "apps",
        repository / "services",
        repository / "tests",
        repository / "deploy",
        repository / "runtime" / "basic60",
        repository / "runtime" / "market-research",
    )
    if (
        resolved == repository
        or repository.is_relative_to(resolved)
        or any(resolved.is_relative_to(item) or item.is_relative_to(resolved) for item in protected)
    ):
        raise MarketContentError(
            "market outputs cannot overwrite repository source or frozen assets"
        )
    return resolved


def _fsync_directory(path: Path) -> None:
    if os.name == "posix":
        descriptor = os.open(path, os.O_RDONLY | os.O_DIRECTORY)
        try:
            os.fsync(descriptor)
        finally:
            os.close(descriptor)


def _immutable_bytes(path: Path, payload: bytes) -> None:
    ensure_no_symlinks(path)
    path.parent.mkdir(parents=True, exist_ok=True, mode=0o755)
    if path.exists():
        if not path.is_file() or path.read_bytes() != payload:
            raise MarketContentError(
                "an immutable market-content artifact already has different bytes"
            )
        return
    descriptor, name = tempfile.mkstemp(prefix=".market-", dir=path.parent)
    temporary = Path(name)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        temporary.chmod(0o444)
        try:
            os.link(temporary, path)
        except FileExistsError:
            if path.read_bytes() != payload:
                raise MarketContentError("concurrent immutable artifact conflict") from None
        _fsync_directory(path.parent)
    finally:
        temporary.unlink(missing_ok=True)


def _publish_directory(temporary: Path, target: Path) -> None:
    """Atomically create an import without replacing even an empty concurrent target."""
    if sys.platform.startswith("linux"):
        libc = ctypes.CDLL(None, use_errno=True)
        rename = getattr(libc, "renameat2", None)
        if rename is None:
            raise MarketContentError("this host lacks atomic no-replace directory publication")
        rename.argtypes = [
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_int,
            ctypes.c_char_p,
            ctypes.c_uint,
        ]
        rename.restype = ctypes.c_int
        # AT_FDCWD and RENAME_NOREPLACE. WSL/Linux is the supported tooling runtime.
        if rename(-100, os.fsencode(temporary), -100, os.fsencode(target), 1) != 0:
            error = ctypes.get_errno()
            if error in {errno.EEXIST, errno.ENOTEMPTY}:
                raise MarketContentError(
                    "import output already exists, including a concurrent directory"
                )
            raise OSError(error, os.strerror(error))
    elif os.name == "nt":
        # Unlike POSIX rename, Windows rename refuses an existing destination.
        try:
            temporary.rename(target)
        except FileExistsError as exc:
            raise MarketContentError("import output already exists") from exc
    else:
        raise MarketContentError("atomic market review import requires Linux/WSL or Windows")


def import_market_review(
    *,
    candidate_dir: Path,
    workbook_path: Path,
    output_dir: Path,
    profiles_path: Path,
    repo_root: Path,
) -> dict[str, Any]:
    candidate = load_market_candidate(candidate_dir, profiles_path)
    workbook_payload = _workbook_bytes(workbook_path)
    metadata, rows = _read_workbook(workbook_payload)
    imported = apply_market_review_rows(
        candidate, metadata, rows, load_country_scope(profiles_path)
    )
    target = _writable_target(output_dir, repo_root)
    original = candidate_dir.resolve()
    if target.exists() or original.is_relative_to(target) or target.is_relative_to(original):
        raise MarketContentError(
            "import output must be a new directory, not the original candidate"
        )
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = Path(tempfile.mkdtemp(prefix=f".{target.name}-", dir=target.parent))
    receipt = {
        "schema_version": RECEIPT_SCHEMA,
        "package_id": candidate.package_id,
        "source_candidate_sha256": candidate.sha256,
        "candidate_sha256": imported.sha256,
        "workbook_sha256": hashlib.sha256(workbook_payload).hexdigest(),
        "workbook_created_at": metadata["created_at"],
        "imported_at": _now(),
        "country_count": len(imported.countries),
    }
    try:
        _immutable_bytes(
            temporary / "package.json",
            canonical_bytes(
                {
                    "schema_version": CANDIDATE_SCHEMA,
                    "package_id": imported.package_id,
                }
            ),
        )
        for country in imported.countries:
            _immutable_bytes(
                temporary / "countries" / f"{country.country_code}.json",
                canonical_bytes(country.model_dump(mode="json")),
            )
        _immutable_bytes(temporary / "review-import.json", canonical_bytes(receipt))
        _publish_directory(temporary, target)
        _fsync_directory(target.parent)
    finally:
        if temporary.exists():
            # Only this call's same-parent uniquely allocated temporary directory is removed.
            shutil.rmtree(temporary)
    return {"status": "review_imported", "published": False, **receipt, "output_dir": str(target)}


def _confirmation(
    candidate: MarketCandidate, receipt: dict[str, Any], workbook_hash: str, confirmation_path: Path
) -> dict[str, Any]:
    confirmation = read_json(confirmation_path)
    fields = {
        "schema_version",
        "package_id",
        "decision",
        "approved_by",
        "approved_at",
        "workbook_sha256",
        "candidate_sha256",
    }
    if (
        set(confirmation) != fields
        or confirmation.get("schema_version") != CONFIRMATION_SCHEMA
        or confirmation.get("decision") != "approved"
        or confirmation.get("approved_by") != "kevin"
    ):
        raise MarketContentError(
            "publication requires the actual project approver's single confirmation"
        )
    _iso_time(confirmation.get("approved_at"), "approved_at")
    if (
        receipt.get("schema_version") != RECEIPT_SCHEMA
        or receipt.get("package_id") != candidate.package_id
        or receipt.get("candidate_sha256") != candidate.sha256
        or receipt.get("workbook_sha256") != workbook_hash
        or confirmation.get("package_id") != candidate.package_id
        or confirmation.get("candidate_sha256") != candidate.sha256
        or confirmation.get("workbook_sha256") != workbook_hash
        or receipt.get("country_count") != len(candidate.countries)
    ):
        raise MarketContentError(
            "confirmation must bind the exact reviewed workbook and imported candidate"
        )
    checked_sha(receipt.get("source_candidate_sha256"))
    _iso_time(receipt.get("imported_at"), "imported_at")
    _iso_time(receipt.get("workbook_created_at"), "workbook_created_at")
    now = _now()
    approved = confirmation["approved_at"]
    created = receipt["workbook_created_at"]
    imported = receipt["imported_at"]
    if _time_before(approved, created) or _time_before(imported, created):
        raise MarketContentError("confirmation and import cannot precede workbook creation")
    if any(_time_before(now, value) for value in (approved, created, imported)):
        raise MarketContentError("review timestamps cannot be in the future at publication")
    if any(_time_before(approved, item.as_of.isoformat()) for item in candidate.countries):
        raise MarketContentError("content as_of cannot be later than its actual confirmation")
    for item in candidate.countries:
        version_day = item.content_version.split("-")[2]
        if _time_before(approved, f"{version_day[:4]}-{version_day[4:6]}-{version_day[6:]}"):
            raise MarketContentError("content version date cannot be later than its confirmation")
    return confirmation


@contextmanager
def _writer(root: Path, repo_root: Path) -> Iterator[Path]:
    target = _writable_target(root, repo_root)
    target.mkdir(parents=True, exist_ok=True, mode=0o755)
    lock = checked_path(target, ".publish.lock")
    try:
        descriptor = os.open(lock, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
    except FileExistsError as exc:
        raise MarketContentError("another content operation holds the publication lock") from exc
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(f"pid={os.getpid()}\nstarted_at={_now()}\n")
            handle.flush()
            os.fsync(handle.fileno())
        yield target
    finally:
        lock.unlink(missing_ok=True)


def _write_state(
    root: Path, state: dict[str, Any], previous: str | None, event: dict[str, Any]
) -> str:
    updated = {
        "schema_version": "navigator.market-state.v1",
        "generation": state["generation"] + 1,
        "active": state["active"],
        "revoked": sorted(set(state["revoked"])),
        "previous_sha256": previous,
        "event": {**event, "at": _now()},
    }
    digest = content_sha256(updated)
    _immutable_bytes(checked_path(root, "manifests", f"{digest}.json"), canonical_bytes(updated))
    pointer = canonical_bytes(
        {
            "schema_version": "navigator.market-pointer.v1",
            "manifest_sha256": digest,
        }
    )
    descriptor, name = tempfile.mkstemp(prefix=".current-", dir=root)
    temporary = Path(name)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(pointer)
            handle.flush()
            os.fsync(handle.fileno())
        temporary.chmod(0o644)
        os.replace(temporary, checked_path(root, "current.json"))
        _fsync_directory(root)
    finally:
        temporary.unlink(missing_ok=True)
    return digest


def publish_market_content(
    *,
    candidate_dir: Path,
    workbook_path: Path,
    confirmation_path: Path,
    content_root: Path,
    profiles_path: Path,
    repo_root: Path,
) -> dict[str, Any]:
    store_path = _writable_target(content_root, repo_root)
    candidate_path = candidate_dir.resolve()
    if store_path.is_relative_to(candidate_path) or candidate_path.is_relative_to(store_path):
        raise MarketContentError("published content and candidate directories must be separate")
    candidate = load_market_candidate(candidate_dir, profiles_path)
    receipt = read_json(checked_path(candidate_dir, "review-import.json"))
    workbook_payload = _workbook_bytes(workbook_path)
    workbook_hash = hashlib.sha256(workbook_payload).hexdigest()
    confirmation = _confirmation(candidate, receipt, workbook_hash, confirmation_path)
    confirmation_hash = content_sha256(confirmation)
    with _writer(content_root, repo_root) as root:
        state, previous = load_market_state(root)
        releases: list[dict[str, Any]] = []
        for country in candidate.countries:
            version = country.content_version
            if (
                version in state["revoked"]
                or checked_path(root, "revocations", f"{version}.json").exists()
            ):
                raise MarketContentError(
                    "revoked content versions cannot be republished; create a new version"
                )
            active_version = state["active"].get(country.country_code)
            if (
                active_version
                and active_version.startswith("OVERVIEW-")
                and active_version != version
            ):
                incoming_order = (version.split("-")[2], int(version.rsplit("-R", 1)[1]))
                active_order = (
                    active_version.split("-")[2],
                    int(active_version.rsplit("-R", 1)[1]),
                )
                if incoming_order <= active_order:
                    raise MarketContentError(
                        "older content requires explicit rollback, not publication"
                    )
            object_hash = content_sha256(country.model_dump(mode="json"))
            release = {
                "schema_version": "navigator.market-release.v1",
                "country_code": country.country_code,
                "content_version": version,
                "object_sha256": object_hash,
                "confirmation_sha256": confirmation_hash,
                "candidate_sha256": candidate.sha256,
                "workbook_sha256": workbook_hash,
                "package_id": candidate.package_id,
                "published_at": _now(),
            }
            existing = checked_path(root, "releases", f"{version}.json")
            if existing.exists():
                old = read_json(existing)
                if any(
                    old.get(key) != value for key, value in release.items() if key != "published_at"
                ):
                    raise MarketContentError(
                        "a content version cannot be rebound to changed content or review"
                    )
                release = old
            releases.append(release)
        _immutable_bytes(
            checked_path(root, "batches", f"{candidate.sha256}.json"),
            canonical_bytes(candidate.payload()),
        )
        _immutable_bytes(
            checked_path(root, "confirmations", f"{confirmation_hash}.json"),
            canonical_bytes(confirmation),
        )
        _immutable_bytes(checked_path(root, "reviews", f"{workbook_hash}.xlsx"), workbook_payload)
        for country, release in zip(candidate.countries, releases, strict=True):
            _immutable_bytes(
                checked_path(root, "objects", f"{release['object_sha256']}.json"),
                canonical_bytes(country.model_dump(mode="json")),
            )
            _immutable_bytes(
                checked_path(root, "releases", f"{country.content_version}.json"),
                canonical_bytes(release),
            )
            load_published_version(root, country.content_version)
        active = {item.country_code: item.content_version for item in candidate.countries}
        if all(state["active"].get(code) == version for code, version in active.items()):
            digest = previous
        else:
            state["active"].update(active)
            digest = _write_state(
                root,
                state,
                previous,
                {
                    "action": "publish",
                    "package_id": candidate.package_id,
                    "countries": sorted(active),
                    "confirmation_sha256": confirmation_hash,
                },
            )
    return {
        "status": "published",
        "active": active,
        "manifest_sha256": digest,
        "candidate_sha256": candidate.sha256,
    }


def _operation_actor(actor: str, reason: str) -> None:
    if not actor.strip() or not reason.strip() or len(actor) > 100 or len(reason) > 2000:
        raise MarketContentError("content operations require a bounded actor and reason")


def revoke_market_content(
    *,
    content_root: Path,
    country_code: str,
    reason: str,
    actor: str,
    repo_root: Path,
    content_version: str | None = None,
) -> dict[str, Any]:
    code = _country_code(country_code)
    _operation_actor(actor, reason)
    with _writer(content_root, repo_root) as root:
        state, previous = load_market_state(root)
        version = content_version or state["active"].get(code)
        if (
            not isinstance(version, str)
            or not VERSION_PATTERN.fullmatch(version)
            or not version.startswith((f"MARKET-{code}-", f"OVERVIEW-{code}-"))
        ):
            raise MarketContentError("no matching published version to revoke")
        tombstone = checked_path(root, "revocations", f"{version}.json")
        if not tombstone.exists():
            published = load_published_version(root, version)
            if published.country_code != code:
                raise MarketContentError("revocation country mismatch")
            _immutable_bytes(
                tombstone,
                canonical_bytes(
                    {
                        "schema_version": "navigator.market-revocation.v1",
                        "country_code": code,
                        "content_version": version,
                        "revoked_at": _now(),
                        "actor": actor,
                        "reason": reason,
                    }
                ),
            )
        state["revoked"].append(version)
        if state["active"].get(code) == version:
            del state["active"][code]
        digest = _write_state(
            root,
            state,
            previous,
            {
                "action": "revoke",
                "country_code": code,
                "content_version": version,
                "actor": actor,
                "reason": reason,
            },
        )
    return {
        "status": "revoked",
        "country_code": code,
        "content_version": version,
        "manifest_sha256": digest,
    }


def rollback_market_content(
    *,
    content_root: Path,
    country_code: str,
    content_version: str,
    reason: str,
    actor: str,
    repo_root: Path,
) -> dict[str, Any]:
    code = _country_code(country_code)
    _operation_actor(actor, reason)
    with _writer(content_root, repo_root) as root:
        state, previous = load_market_state(root)
        if content_version in state["revoked"]:
            raise MarketContentError("rollback cannot reactivate a revoked version")
        try:
            content = load_published_version(root, content_version)
        except MarketContentUnavailable as exc:
            raise MarketContentError("rollback cannot reactivate a revoked version") from exc
        if content.country_code != code:
            raise MarketContentError("rollback version belongs to a different country")
        if not isinstance(content, CountryMarketOverview):
            raise MarketContentError("retired content types cannot be restored as market overviews")
        historical = state
        visited: set[str] = set()
        while historical["active"].get(code) != content_version:
            predecessor = historical.get("previous_sha256")
            if predecessor is None:
                raise MarketContentError("rollback requires a previously active published version")
            predecessor = checked_sha(predecessor)
            if predecessor in visited or len(visited) >= 10_000:
                raise MarketContentError("invalid or excessive market-content history chain")
            visited.add(predecessor)
            historical = read_json(
                checked_path(root, "manifests", f"{predecessor}.json"),
                expected_sha256=predecessor,
            )
            if historical.get("schema_version") != "navigator.market-state.v1" or not isinstance(
                historical.get("active"), dict
            ):
                raise MarketContentError("invalid historical content manifest")
        if state["active"].get(code) == content_version:
            digest = previous
        else:
            state["active"][code] = content_version
            digest = _write_state(
                root,
                state,
                previous,
                {
                    "action": "rollback",
                    "country_code": code,
                    "content_version": content_version,
                    "actor": actor,
                    "reason": reason,
                },
            )
    return {
        "status": "rolled_back",
        "country_code": code,
        "content_version": content_version,
        "manifest_sha256": digest,
    }


def add_market_content_commands(commands: Any) -> None:
    for name in sorted(MARKET_COMMANDS):
        command = commands.add_parser(
            name, help="Manage independent, human-reviewed market content."
        )
        if name in {"validate-market-content", "import-market-review", "publish-market-content"}:
            command.add_argument("--candidate-dir", type=Path, required=True)
            command.add_argument("--profiles", type=Path, default=None)
        if name in {"import-market-review", "publish-market-content"}:
            command.add_argument("--workbook", type=Path, required=True)
        if name == "import-market-review":
            command.add_argument("--output-dir", type=Path, required=True)
        if name == "publish-market-content":
            command.add_argument("--confirmation", type=Path, required=True)
        if name in {"publish-market-content", "revoke-market-content", "rollback-market-content"}:
            command.add_argument("--content-root", type=Path, required=True)
        if name in {"revoke-market-content", "rollback-market-content"}:
            command.add_argument("--country", required=True)
            command.add_argument("--content-version", required=name == "rollback-market-content")
            command.add_argument("--reason", required=True)
            command.add_argument("--actor", required=True)


def run_market_content_command(args: argparse.Namespace, repo_root: Path) -> int:
    try:
        for field in (
            "candidate_dir",
            "profiles",
            "workbook",
            "output_dir",
            "confirmation",
            "content_root",
        ):
            value = getattr(args, field, None)
            if isinstance(value, Path) and not value.is_absolute():
                setattr(args, field, repo_root / value)
        profiles = (
            getattr(args, "profiles", None)
            or repo_root / "raw material/global_sources/60_country_profiles.csv"
        )
        if args.command == "validate-market-content":
            result = validate_market_content(args.candidate_dir, profiles)
        elif args.command == "import-market-review":
            result = import_market_review(
                candidate_dir=args.candidate_dir,
                workbook_path=args.workbook,
                output_dir=args.output_dir,
                profiles_path=profiles,
                repo_root=repo_root,
            )
        elif args.command == "publish-market-content":
            result = publish_market_content(
                candidate_dir=args.candidate_dir,
                workbook_path=args.workbook,
                confirmation_path=args.confirmation,
                content_root=args.content_root,
                profiles_path=profiles,
                repo_root=repo_root,
            )
        elif args.command == "revoke-market-content":
            result = revoke_market_content(
                content_root=args.content_root,
                country_code=args.country,
                content_version=args.content_version,
                reason=args.reason,
                actor=args.actor,
                repo_root=repo_root,
            )
        elif args.command == "rollback-market-content":
            result = rollback_market_content(
                content_root=args.content_root,
                country_code=args.country,
                content_version=args.content_version,
                reason=args.reason,
                actor=args.actor,
                repo_root=repo_root,
            )
        else:
            raise MarketContentError("unknown market-content command")
    except (MarketContentError, OSError, ValueError, KeyError, zipfile.BadZipFile) as exc:
        print(
            json.dumps({"status": "not_ready", "error": str(exc)}, ensure_ascii=False),
            file=sys.stderr,
        )
        return 1
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0
