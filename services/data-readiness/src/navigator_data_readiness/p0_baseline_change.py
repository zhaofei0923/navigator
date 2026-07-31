from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from copy import copy
from dataclasses import replace
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.worksheet import Worksheet

from .baseline import TECH_SHEETS, extract_contracts, sha256_file
from .models import CheckResult
from .p0_resolution import PROPOSED_CONTRACT_SPECS, validate_p0_resolution_packet
from .p0_traceability import (
    EXPECTED_P0_REQUIREMENT_COUNT,
    build_p0_traceability_report,
    extract_references,
)
from .paths import RepositoryPaths
from .validation import validate_structure
from .workbook import load_read_only, normalized_value

_PERMISSION_SEPARATOR = re.compile(r"[;；,，、\n]+")

_CHANGEABLE_CONTRACTS = {
    "feature_requirements",
    "role_permissions",
    "page_routes",
    "api_catalog",
    "test_cases",
}
_CONTRACT_IDS = {
    "feature_requirements": "需求编号",
    "role_permissions": "权限编号",
    "page_routes": "页面编号",
    "api_catalog": "接口编号",
    "test_cases": "用例编号",
}
_PROPOSED_CONTRACTS = {
    category: (contract_name, identifier_field)
    for category, (
        contract_name,
        identifier_field,
        _required_fields,
        _extension_fields,
    ) in PROPOSED_CONTRACT_SPECS.items()
}
_ROUTE_DIMENSIONS = {
    "requirement": ("关联需求编号", "proposed_requirement_ids", ("FR-",)),
    "api": ("依赖API", "proposed_api_ids", ("API-",)),
    "test": ("测试用例", "proposed_test_case_ids", ("TC-",)),
}
_NEW_ROW_REQUIRED_FIELDS = {
    contract_name: required_fields
    for (
        contract_name,
        _identifier_field,
        required_fields,
        _extension_fields,
    ) in PROPOSED_CONTRACT_SPECS.values()
}
_NOT_APPLICABLE_MARKERS = {"不适用", "无需API", "N/A", "NA"}
_REVIEWED_HEADER_EXTENSIONS = {"role_permissions": {"权限代码"}}
_CONTRACT_SHEETS = {contract_name: sheet_name for sheet_name, contract_name in TECH_SHEETS.items()}


def _payload_sha256(payload: Any) -> str:
    serialized = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    return hashlib.sha256(serialized).hexdigest()


def _text(record: dict[str, Any], field: str) -> str:
    return str(record.get(field, "")).strip()


def _string_list(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _records(contracts: dict[str, Any], name: str) -> list[dict[str, Any]]:
    return cast(list[dict[str, Any]], contracts.get(name, []))


def _index(records: list[dict[str, Any]], identifier_field: str) -> dict[str, dict[str, Any]]:
    return {
        _text(record, identifier_field): record
        for record in records
        if _text(record, identifier_field)
    }


def _permission_codes(value: Any) -> set[str]:
    return {item.strip() for item in _PERMISSION_SEPARATOR.split(str(value or "")) if item.strip()}


def _changed_fields(before: dict[str, Any], after: dict[str, Any]) -> set[str]:
    return {field for field in set(before) | set(after) if before.get(field) != after.get(field)}


def _header_contract(workbook_path: Path) -> dict[str, dict[int, str]]:
    workbook = load_read_only(workbook_path, data_only=False)
    try:
        result: dict[str, dict[int, str]] = {}
        for sheet_name, contract_name in TECH_SHEETS.items():
            row = next(
                workbook[sheet_name].iter_rows(min_row=4, max_row=4),
                (),
            )
            result[contract_name] = {
                int(cell.column): str(cell.value).strip()
                for cell in row
                if cell.value is not None and str(cell.value).strip()
            }
        return result
    finally:
        workbook.close()


def _validate_header_delta(
    checks: list[CheckResult],
    source_workbook: Path,
    candidate_workbook: Path,
) -> None:
    source_headers = _header_contract(source_workbook)
    candidate_headers = _header_contract(candidate_workbook)
    for contract_name in TECH_SHEETS.values():
        before = source_headers[contract_name]
        after = candidate_headers[contract_name]
        changed_existing = {
            column: (header, after.get(column))
            for column, header in before.items()
            if after.get(column) != header
        }
        added = {column: header for column, header in after.items() if column not in before}
        allowed_extensions = _REVIEWED_HEADER_EXTENSIONS.get(contract_name, set())
        expected_extension_columns = {
            max(before, default=0) + offset: header
            for offset, header in enumerate(sorted(allowed_extensions), start=1)
        }
        unexpected_added = {
            column: header
            for column, header in added.items()
            if expected_extension_columns.get(column) != header
        }
        duplicate_extensions = {
            header
            for header in allowed_extensions
            if sum(value == header for value in after.values()) > 1
        }
        if changed_existing or unexpected_added or duplicate_extensions:
            details: list[str] = []
            if changed_existing:
                details.append(
                    "changed="
                    + ", ".join(
                        f"{column}:{old}->{new}"
                        for column, (old, new) in sorted(changed_existing.items())
                    )
                )
            if unexpected_added:
                details.append(
                    "added="
                    + ", ".join(
                        f"{column}:{header}" for column, header in sorted(unexpected_added.items())
                    )
                )
            if duplicate_extensions:
                details.append(f"duplicate={', '.join(sorted(duplicate_extensions))}")
            checks.append(
                CheckResult(
                    code="P0_CHANGE_UNAUTHORIZED_HEADER_CHANGE",
                    message=(
                        f"{contract_name} contains unreviewed header changes: " + "; ".join(details)
                    ),
                    location=contract_name,
                )
            )


def _workbook_row_lookups(workbook_path: Path) -> dict[str, dict[str, int]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        result: dict[str, dict[str, int]] = {}
        for contract_name in _CHANGEABLE_CONTRACTS:
            sheet = workbook[_CONTRACT_SHEETS[contract_name]]
            headers = _editable_headers(sheet)
            identifier_column = headers.get(_CONTRACT_IDS[contract_name])
            if identifier_column is None:
                result[contract_name] = {}
                continue
            result[contract_name] = _editable_rows(sheet, identifier_column)
        return result
    finally:
        workbook.close()


def _cell_values(workbook_path: Path) -> tuple[list[str], dict[str, dict[tuple[int, int], Any]]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        values: dict[str, dict[tuple[int, int], Any]] = {}
        for sheet in workbook.worksheets:
            cells: dict[tuple[int, int], Any] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    value = normalized_value(cell.value)
                    if (
                        value is not None
                        and str(value).strip()
                        and cell.row is not None
                        and cell.column is not None
                    ):
                        cells[(int(cell.row), int(cell.column))] = value
            values[sheet.title] = cells
        return list(workbook.sheetnames), values
    finally:
        workbook.close()


def _validate_cell_delta(
    checks: list[CheckResult],
    source_workbook: Path,
    candidate_workbook: Path,
    resolution: dict[str, Any],
) -> None:
    source_sheets, source_values = _cell_values(source_workbook)
    candidate_sheets, candidate_values = _cell_values(candidate_workbook)
    if source_sheets != candidate_sheets:
        checks.append(
            CheckResult(
                code="P0_CHANGE_SHEET_SET_INVALID",
                message="Candidate workbook sheet names or order differ from the frozen workbook",
                location=str(candidate_workbook),
            )
        )
        return

    headers = _header_contract(candidate_workbook)
    rows = _workbook_row_lookups(candidate_workbook)
    allowed: set[tuple[str, int, int]] = set()

    permission_headers = headers.get("role_permissions", {})
    permission_code_column = next(
        (column for column, header in permission_headers.items() if header == "权限代码"),
        None,
    )
    if permission_code_column is not None:
        allowed.add((_CONTRACT_SHEETS["role_permissions"], 4, permission_code_column))

    proposed_rows = _proposed_rows(resolution)
    for category, (contract_name, _identifier_field) in _PROPOSED_CONTRACTS.items():
        sheet_name = _CONTRACT_SHEETS[contract_name]
        for identifier in proposed_rows[category]:
            row = rows.get(contract_name, {}).get(identifier)
            if row is not None:
                allowed.update((sheet_name, row, column) for column in headers[contract_name])

    requirement_items = resolution.get("requirement_test_resolutions")
    if isinstance(requirement_items, list):
        requirement_column = next(
            (
                column
                for column, header in headers.get("test_cases", {}).items()
                if header == "需求编号"
            ),
            None,
        )
        if requirement_column is not None:
            for item in requirement_items:
                if not isinstance(item, dict):
                    continue
                for test_id in _string_list(item.get("proposed_test_case_ids")):
                    row = rows.get("test_cases", {}).get(test_id)
                    if row is not None:
                        allowed.add((_CONTRACT_SHEETS["test_cases"], row, requirement_column))

    route_items = resolution.get("route_mapping_resolutions")
    if isinstance(route_items, list):
        route_headers = {
            header: column for column, header in headers.get("page_routes", {}).items()
        }
        for item in route_items:
            if not isinstance(item, dict):
                continue
            row = rows.get("page_routes", {}).get(_text(item, "page_id"))
            if row is None:
                continue
            for dimension in _string_list(item.get("unresolved_dimensions")):
                if dimension in _ROUTE_DIMENSIONS:
                    field = _ROUTE_DIMENSIONS[dimension][0]
                    if field in route_headers:
                        allowed.add((_CONTRACT_SHEETS["page_routes"], row, route_headers[field]))

    permission_items = resolution.get("permission_code_resolutions")
    if isinstance(permission_items, list) and permission_code_column is not None:
        for item in permission_items:
            if not isinstance(item, dict):
                continue
            for permission_id in _string_list(item.get("proposed_perm_ids")):
                row = rows.get("role_permissions", {}).get(permission_id)
                if row is not None:
                    allowed.add(
                        (
                            _CONTRACT_SHEETS["role_permissions"],
                            row,
                            permission_code_column,
                        )
                    )

    for sheet_name in source_sheets:
        before = source_values.get(sheet_name, {})
        after = candidate_values.get(sheet_name, {})
        changed = {
            coordinate
            for coordinate in set(before) | set(after)
            if before.get(coordinate) != after.get(coordinate)
            and (sheet_name, coordinate[0], coordinate[1]) not in allowed
        }
        if changed:
            preview = ", ".join(f"R{row}C{column}" for row, column in sorted(changed)[:10])
            checks.append(
                CheckResult(
                    code="P0_CHANGE_UNAUTHORIZED_CELL_CHANGE",
                    message=(
                        f"{sheet_name} contains {len(changed)} unreviewed cell value/formula "
                        f"change(s): {preview}"
                    ),
                    location=sheet_name,
                )
            )


def _proposed_ids(resolution: dict[str, Any]) -> dict[str, set[str]]:
    proposed = resolution.get("proposed_contract_ids")
    if not isinstance(proposed, dict):
        return {category: set() for category in _PROPOSED_CONTRACTS}
    return {category: set(_string_list(proposed.get(category))) for category in _PROPOSED_CONTRACTS}


def _proposed_rows(resolution: dict[str, Any]) -> dict[str, dict[str, dict[str, Any]]]:
    raw_rows = resolution.get("proposed_contract_rows")
    if not isinstance(raw_rows, dict):
        return {category: {} for category in _PROPOSED_CONTRACTS}
    result: dict[str, dict[str, dict[str, Any]]] = {}
    for category, (_contract_name, identifier_field) in _PROPOSED_CONTRACTS.items():
        values = raw_rows.get(category)
        rows: dict[str, dict[str, Any]] = {}
        if isinstance(values, list):
            for value in values:
                if not isinstance(value, dict):
                    continue
                identifier = _text(value, identifier_field)
                if identifier:
                    rows[identifier] = value
        result[category] = rows
    return result


def _allowed_existing_changes(
    resolution: dict[str, Any],
) -> dict[str, dict[str, set[str]]]:
    allowed: dict[str, dict[str, set[str]]] = {contract: {} for contract in _CHANGEABLE_CONTRACTS}

    requirement_resolutions = resolution.get("requirement_test_resolutions")
    if isinstance(requirement_resolutions, list):
        for item in requirement_resolutions:
            if not isinstance(item, dict):
                continue
            for test_id in _string_list(item.get("proposed_test_case_ids")):
                allowed["test_cases"].setdefault(test_id, set()).add("需求编号")

    route_resolutions = resolution.get("route_mapping_resolutions")
    if isinstance(route_resolutions, list):
        for item in route_resolutions:
            if not isinstance(item, dict):
                continue
            page_id = _text(item, "page_id")
            dimensions = set(_string_list(item.get("unresolved_dimensions")))
            allowed_fields = {
                _ROUTE_DIMENSIONS[dimension][0]
                for dimension in dimensions
                if dimension in _ROUTE_DIMENSIONS
            }
            if page_id:
                allowed["page_routes"][page_id] = allowed_fields

    permission_resolutions = resolution.get("permission_code_resolutions")
    if isinstance(permission_resolutions, list):
        for item in permission_resolutions:
            if not isinstance(item, dict):
                continue
            for permission_id in _string_list(item.get("proposed_perm_ids")):
                allowed["role_permissions"].setdefault(permission_id, set()).add("权限代码")
    return allowed


def _validate_contract_delta(
    checks: list[CheckResult],
    original: dict[str, Any],
    candidate: dict[str, Any],
    resolution: dict[str, Any],
) -> None:
    proposed = _proposed_ids(resolution)
    proposed_rows = _proposed_rows(resolution)
    allowed_changes = _allowed_existing_changes(resolution)

    for contract_name in TECH_SHEETS.values():
        original_records = _records(original, contract_name)
        candidate_records = _records(candidate, contract_name)
        if contract_name not in _CHANGEABLE_CONTRACTS:
            if candidate_records != original_records:
                checks.append(
                    CheckResult(
                        code="P0_CHANGE_UNAUTHORIZED_CONTRACT_CHANGE",
                        message=f"{contract_name} changed outside the reviewed P0 proposal",
                        location=contract_name,
                    )
                )
            continue

        identifier_field = _CONTRACT_IDS[contract_name]
        original_index = _index(original_records, identifier_field)
        candidate_index = _index(candidate_records, identifier_field)
        removed = sorted(set(original_index) - set(candidate_index))
        if removed:
            checks.append(
                CheckResult(
                    code="P0_CHANGE_CONTRACT_ROW_REMOVED",
                    message=f"Removed IDs: {', '.join(removed)}",
                    location=contract_name,
                )
            )

        category = next(
            (
                name
                for name, (mapped_contract, _field) in _PROPOSED_CONTRACTS.items()
                if mapped_contract == contract_name
            ),
            None,
        )
        expected_added = proposed[category] if category else set()
        actual_added = set(candidate_index) - set(original_index)
        if actual_added != expected_added:
            checks.append(
                CheckResult(
                    code="P0_CHANGE_ADDED_ID_SET_INVALID",
                    message=(
                        f"Added IDs differ; missing={sorted(expected_added - actual_added)}, "
                        f"unexpected={sorted(actual_added - expected_added)}"
                    ),
                    location=contract_name,
                )
            )

        for identifier in sorted(set(original_index) & set(candidate_index)):
            changed = _changed_fields(original_index[identifier], candidate_index[identifier])
            unauthorized = changed - allowed_changes[contract_name].get(identifier, set())
            if unauthorized:
                checks.append(
                    CheckResult(
                        code="P0_CHANGE_UNAUTHORIZED_FIELD_CHANGE",
                        message=(
                            f"{identifier} changed unreviewed fields: "
                            f"{', '.join(sorted(unauthorized))}"
                        ),
                        location=f"{contract_name}.{identifier}",
                    )
                )

        required_fields = _NEW_ROW_REQUIRED_FIELDS.get(contract_name)
        if required_fields:
            for identifier in sorted(actual_added):
                record = candidate_index[identifier]
                missing = [field for field in required_fields if not _text(record, field)]
                if missing:
                    checks.append(
                        CheckResult(
                            code="P0_CHANGE_NEW_CONTRACT_INCOMPLETE",
                            message=f"{identifier} is missing: {', '.join(missing)}",
                            location=f"{contract_name}.{identifier}",
                        )
                    )
                if category and record != proposed_rows[category].get(identifier):
                    mismatch_fields = sorted(
                        _changed_fields(
                            proposed_rows[category].get(identifier, {}),
                            record,
                        )
                    )
                    checks.append(
                        CheckResult(
                            code="P0_CHANGE_NEW_CONTRACT_ROW_MISMATCH",
                            message=(
                                f"{identifier} differs from its reviewed full-row definition"
                                + (f": {', '.join(mismatch_fields)}" if mismatch_fields else "")
                            ),
                            location=f"{contract_name}.{identifier}",
                        )
                    )

    candidate_requirements = _index(_records(candidate, "feature_requirements"), "需求编号")
    for requirement_id in proposed["requirements"]:
        record = candidate_requirements.get(requirement_id, {})
        if _text(record, "优先级") != "P0":
            checks.append(
                CheckResult(
                    code="P0_CHANGE_NEW_REQUIREMENT_NOT_P0",
                    message=f"{requirement_id} must be a P0 requirement",
                    location=f"feature_requirements.{requirement_id}",
                )
            )

    candidate_apis = _index(_records(candidate, "api_catalog"), "接口编号")
    for api_id in proposed["apis"]:
        if "P1" in _text(candidate_apis.get(api_id, {}), "状态"):
            checks.append(
                CheckResult(
                    code="P0_CHANGE_NEW_API_NOT_P0",
                    message=f"{api_id} is marked as P1",
                    location=f"api_catalog.{api_id}",
                )
            )

    candidate_tests = _index(_records(candidate, "test_cases"), "用例编号")
    for test_id in proposed["tests"]:
        if _text(candidate_tests.get(test_id, {}), "优先级") != "P0":
            checks.append(
                CheckResult(
                    code="P0_CHANGE_NEW_TEST_NOT_P0",
                    message=f"{test_id} must be a P0 test",
                    location=f"test_cases.{test_id}",
                )
            )


def _validate_requirement_test_resolutions(
    checks: list[CheckResult],
    original: dict[str, Any],
    candidate: dict[str, Any],
    resolution: dict[str, Any],
) -> None:
    original_tests = _index(_records(original, "test_cases"), "用例编号")
    reviewed_new_tests = _proposed_rows(resolution)["tests"]
    tests = _index(_records(candidate, "test_cases"), "用例编号")
    items = resolution.get("requirement_test_resolutions")
    if not isinstance(items, list):
        return
    expected_by_test: dict[str, set[str]] = {}
    for item in items:
        if not isinstance(item, dict):
            continue
        requirement_id = _text(item, "requirement_id")
        for test_id in _string_list(item.get("proposed_test_case_ids")):
            expected = expected_by_test.setdefault(
                test_id,
                set(
                    extract_references(
                        reviewed_new_tests.get(test_id, original_tests.get(test_id, {})).get(
                            "需求编号"
                        ),
                        ("FR-",),
                    )
                ),
            )
            expected.add(requirement_id)
    for test_id, expected in sorted(expected_by_test.items()):
        actual = set(
            extract_references(
                tests.get(test_id, {}).get("需求编号"),
                ("FR-",),
            )
        )
        if actual != expected:
            checks.append(
                CheckResult(
                    code="P0_CHANGE_REQUIREMENT_TEST_NOT_APPLIED",
                    message=(
                        f"{test_id} requirement IDs differ; "
                        f"expected={sorted(expected)}, actual={sorted(actual)}"
                    ),
                    location=f"test_cases.{test_id}",
                )
            )


def _validate_route_resolutions(
    checks: list[CheckResult],
    candidate: dict[str, Any],
    resolution: dict[str, Any],
) -> None:
    routes = _index(_records(candidate, "page_routes"), "页面编号")
    items = resolution.get("route_mapping_resolutions")
    if not isinstance(items, list):
        return
    for item in items:
        if not isinstance(item, dict):
            continue
        page_id = _text(item, "page_id")
        route = routes.get(page_id, {})
        dimensions = set(_string_list(item.get("unresolved_dimensions")))
        for dimension in sorted(dimensions):
            if dimension not in _ROUTE_DIMENSIONS:
                continue
            field, proposed_field, prefixes = _ROUTE_DIMENSIONS[dimension]
            current_field = {
                "requirement": "current_requirement_ids",
                "api": "current_api_ids",
                "test": "current_test_case_ids",
            }[dimension]
            expected = set(_string_list(item.get(current_field)))
            expected.update(_string_list(item.get(proposed_field)))
            actual = set(extract_references(route.get(field), prefixes))
            api_not_applicable = dimension == "api" and item.get("api_not_applicable") is True
            if api_not_applicable:
                raw_value = _text(route, field).upper()
                if expected or raw_value not in _NOT_APPLICABLE_MARKERS:
                    checks.append(
                        CheckResult(
                            code="P0_CHANGE_ROUTE_NOT_APPLICABLE_INVALID",
                            message=(
                                f"{page_id} requires an explicit API not-applicable value "
                                "without API references"
                            ),
                            location=f"page_routes.{page_id}.{field}",
                        )
                    )
                continue
            if actual != expected:
                checks.append(
                    CheckResult(
                        code="P0_CHANGE_ROUTE_MAPPING_NOT_APPLIED",
                        message=(
                            f"{page_id} {dimension} IDs differ; "
                            f"expected={sorted(expected)}, actual={sorted(actual)}"
                        ),
                        location=f"page_routes.{page_id}.{field}",
                    )
                )


def _validate_permission_resolutions(
    checks: list[CheckResult],
    candidate: dict[str, Any],
    resolution: dict[str, Any],
) -> None:
    permissions = _index(_records(candidate, "role_permissions"), "权限编号")
    items = resolution.get("permission_code_resolutions")
    if not isinstance(items, list):
        return
    for item in items:
        if not isinstance(item, dict):
            continue
        permission_code = _text(item, "permission_code")
        if item.get("permission_not_applicable") is True:
            checks.append(
                CheckResult(
                    code="P0_CHANGE_PERMISSION_NOT_APPLICABLE_UNRESOLVED",
                    message=(
                        f"{permission_code} is used by a P0 route and requires an explicit "
                        "code-to-PERM mapping"
                    ),
                    location=f"permission_code_resolutions.{permission_code}",
                )
            )
            continue
        expected = set(_string_list(item.get("current_perm_ids")))
        expected.update(_string_list(item.get("proposed_perm_ids")))
        actual = {
            permission_id
            for permission_id, record in permissions.items()
            if permission_code in _permission_codes(record.get("权限代码"))
        }
        if actual != expected:
            checks.append(
                CheckResult(
                    code="P0_CHANGE_PERMISSION_MAPPING_NOT_APPLIED",
                    message=(
                        f"{permission_code} PERM IDs differ; "
                        f"expected={sorted(expected)}, actual={sorted(actual)}"
                    ),
                    location=f"permission_code_resolutions.{permission_code}",
                )
            )


def assess_p0_baseline_change(
    paths: RepositoryPaths,
    resolution: dict[str, Any],
    candidate_workbook: Path,
) -> dict[str, Any]:
    checks = validate_p0_resolution_packet(paths, resolution)
    source_workbook = paths.technical_workbook.resolve()
    candidate_path = candidate_workbook.resolve()
    report: dict[str, Any] = {
        "schema_version": 1,
        "stage": "P0",
        "assessment_scope": "P0 candidate technical workbook baseline-change verification",
        "automated_assessment_only": True,
        "does_not_activate_baseline": True,
        "source_workbook": {
            "path": source_workbook.relative_to(paths.root).as_posix(),
            "sha256": sha256_file(source_workbook),
        },
        "candidate_workbook": {
            "path": str(candidate_path),
            "sha256": None,
        },
        "resolution_sha256": _payload_sha256(resolution),
        "traceability_assessment": None,
        "candidate_ready_for_formal_baseline_review": False,
        "checks": [],
        "warning": (
            "Passing this verifier does not modify or approve the frozen baseline. "
            "The authorized technical workbook must be formally replaced before P0 or D4 "
            "can consume the change."
        ),
    }

    if candidate_path == source_workbook:
        checks.append(
            CheckResult(
                code="P0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN",
                message="The read-only authoritative workbook cannot be used as the change target",
                location=str(candidate_path),
            )
        )
    if not candidate_path.is_file():
        checks.append(
            CheckResult(
                code="P0_CHANGE_WORKBOOK_MISSING",
                message="Candidate technical workbook does not exist",
                location=str(candidate_path),
            )
        )
    if checks:
        report["checks"] = [check.to_dict() for check in checks]
        return report

    candidate_hash = sha256_file(candidate_path)
    cast(dict[str, Any], report["candidate_workbook"])["sha256"] = candidate_hash
    if candidate_hash == report["source_workbook"]["sha256"]:
        checks.append(
            CheckResult(
                code="P0_CHANGE_WORKBOOK_UNCHANGED",
                message="Candidate workbook is byte-identical to the frozen source workbook",
                location=str(candidate_path),
            )
        )

    candidate_paths = replace(paths, technical_workbook=candidate_path)
    try:
        checks.extend(validate_structure(candidate_paths))
        _validate_header_delta(checks, source_workbook, candidate_path)
        original_contracts = extract_contracts(paths, data_only=False)
        candidate_contracts = extract_contracts(candidate_paths, data_only=False)
        _validate_cell_delta(
            checks,
            source_workbook,
            candidate_path,
            resolution,
        )
    except Exception as error:
        checks.append(
            CheckResult(
                code="P0_CHANGE_WORKBOOK_INVALID",
                message=f"Cannot read candidate technical workbook: {error}",
                location=str(candidate_path),
            )
        )
        report["checks"] = [check.to_dict() for check in checks]
        return report

    _validate_contract_delta(checks, original_contracts, candidate_contracts, resolution)
    _validate_requirement_test_resolutions(
        checks,
        original_contracts,
        candidate_contracts,
        resolution,
    )
    _validate_route_resolutions(checks, candidate_contracts, resolution)
    _validate_permission_resolutions(checks, candidate_contracts, resolution)

    traceability = build_p0_traceability_report(
        candidate_paths,
        expected_p0_requirement_count=(
            EXPECTED_P0_REQUIREMENT_COUNT + len(_proposed_ids(resolution)["requirements"])
        ),
    )
    report["traceability_assessment"] = traceability
    if traceability.get("traceability_ready") is not True:
        blocker_count = sum(
            int(item.get("count", 0))
            for item in traceability.get("traceability_blockers", [])
            if isinstance(item, dict)
            and isinstance(item.get("count"), int)
            and not isinstance(item.get("count"), bool)
        )
        checks.append(
            CheckResult(
                code="P0_CHANGE_TRACEABILITY_NOT_READY",
                message=f"Candidate workbook still has {blocker_count} traceability blocker items",
                location=str(candidate_path),
            )
        )

    report["checks"] = [check.to_dict() for check in checks]
    report["candidate_ready_for_formal_baseline_review"] = not checks
    return report


def load_and_assess_p0_baseline_change(
    paths: RepositoryPaths,
    resolution_path: Path,
    candidate_workbook: Path,
) -> dict[str, Any]:
    try:
        payload = json.loads(resolution_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as error:
        return _invalid_resolution_report(
            f"Cannot read resolution packet: {error}", resolution_path
        )
    if not isinstance(payload, dict):
        return _invalid_resolution_report(
            "Resolution packet root must be an object", resolution_path
        )
    return assess_p0_baseline_change(paths, payload, candidate_workbook)


def _invalid_resolution_report(message: str, path: Path) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "stage": "P0",
        "assessment_scope": "P0 candidate technical workbook baseline-change verification",
        "automated_assessment_only": True,
        "does_not_activate_baseline": True,
        "candidate_ready_for_formal_baseline_review": False,
        "checks": [
            CheckResult(
                code="P0_CHANGE_RESOLUTION_INVALID",
                message=message,
                location=str(path),
            ).to_dict()
        ],
    }


def _editable_headers(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): int(cell.column)
        for cell in sheet[4]
        if cell.value is not None and str(cell.value).strip() and cell.column is not None
    }


def _editable_rows(sheet: Worksheet, identifier_column: int) -> dict[str, int]:
    return {
        str(sheet.cell(row, identifier_column).value).strip(): row
        for row in range(5, sheet.max_row + 1)
        if sheet.cell(row, identifier_column).value is not None
        and str(sheet.cell(row, identifier_column).value).strip()
    }


def _copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    target.alignment = copy(source.alignment)
    target.protection = copy(source.protection)


def _ordered_union(*groups: list[str]) -> list[str]:
    return list(dict.fromkeys(value for group in groups for value in group if value))


def _ensure_permission_code_header(sheet: Worksheet) -> dict[str, int]:
    headers = _editable_headers(sheet)
    if "权限代码" in headers:
        return headers

    column = max(headers.values()) + 1
    previous_column = column - 1
    header_cell = sheet.cell(4, column)
    _copy_cell_style(sheet.cell(4, previous_column), header_cell)
    header_cell.value = "权限代码"
    previous_letter = get_column_letter(previous_column)
    column_letter = get_column_letter(column)
    sheet.column_dimensions[column_letter].width = sheet.column_dimensions[previous_letter].width

    identifier_column = headers[_CONTRACT_IDS["role_permissions"]]
    rows = _editable_rows(sheet, identifier_column)
    for row in rows.values():
        _copy_cell_style(sheet.cell(row, previous_column), sheet.cell(row, column))

    for table_name in sheet.tables:
        table = sheet.tables[table_name]
        if not any(table_column.name == "权限代码" for table_column in table.tableColumns):
            next_id = max((table_column.id for table_column in table.tableColumns), default=0) + 1
            table.tableColumns.append(TableColumn(id=next_id, name="权限代码"))
    headers["权限代码"] = column
    return headers


def _resize_sheet_tables(sheet: Worksheet, last_row: int, last_column: int) -> None:
    for table_name in sheet.tables:
        table = sheet.tables[table_name]
        min_column, min_row, _max_column, _max_row = range_boundaries(table.ref)
        if min_column is None or min_row is None:
            raise ValueError(f"Cannot resize invalid table range: {table.ref}")
        table.ref = (
            f"{get_column_letter(min_column)}{min_row}:{get_column_letter(last_column)}{last_row}"
        )


def _append_proposed_rows(
    workbook: Any,
    resolution: dict[str, Any],
) -> dict[str, dict[str, int]]:
    proposed = resolution.get("proposed_contract_rows")
    if not isinstance(proposed, dict):
        proposed = {}

    row_lookups: dict[str, dict[str, int]] = {}
    for category, (contract_name, identifier_field) in _PROPOSED_CONTRACTS.items():
        sheet = workbook[_CONTRACT_SHEETS[contract_name]]
        headers = (
            _ensure_permission_code_header(sheet)
            if contract_name == "role_permissions"
            else _editable_headers(sheet)
        )
        rows = _editable_rows(sheet, headers[identifier_field])
        last_original_row = max(rows.values())
        values = proposed.get(category)
        records = values if isinstance(values, list) else []
        for offset, record in enumerate(records, start=1):
            if not isinstance(record, dict):
                continue
            row = last_original_row + offset
            if sheet.row_dimensions[last_original_row].height is not None:
                sheet.row_dimensions[row].height = sheet.row_dimensions[last_original_row].height
            for field, column in headers.items():
                target = sheet.cell(row, column)
                _copy_cell_style(sheet.cell(last_original_row, column), target)
                target.value = record.get(field)
            identifier = _text(record, identifier_field)
            if identifier:
                rows[identifier] = row
        _resize_sheet_tables(sheet, max(rows.values()), max(headers.values()))
        row_lookups[contract_name] = rows
    for contract_name in _CHANGEABLE_CONTRACTS - set(row_lookups):
        sheet = workbook[_CONTRACT_SHEETS[contract_name]]
        headers = _editable_headers(sheet)
        row_lookups[contract_name] = _editable_rows(
            sheet,
            headers[_CONTRACT_IDS[contract_name]],
        )
    return row_lookups


def _apply_requirement_test_mappings(
    workbook: Any,
    resolution: dict[str, Any],
    rows: dict[str, dict[str, int]],
) -> None:
    sheet = workbook[_CONTRACT_SHEETS["test_cases"]]
    requirement_column = _editable_headers(sheet)["需求编号"]
    items = resolution.get("requirement_test_resolutions")
    if not isinstance(items, list):
        return
    for item in items:
        if not isinstance(item, dict):
            continue
        requirement_id = _text(item, "requirement_id")
        for test_id in _string_list(item.get("proposed_test_case_ids")):
            row = rows["test_cases"][test_id]
            cell = sheet.cell(row, requirement_column)
            current = extract_references(cell.value, ("FR-",))
            if requirement_id not in current:
                raw = str(cell.value or "").strip()
                cell.value = f"{raw};{requirement_id}" if raw else requirement_id


def _apply_route_mappings(
    workbook: Any,
    resolution: dict[str, Any],
    rows: dict[str, dict[str, int]],
) -> None:
    sheet = workbook[_CONTRACT_SHEETS["page_routes"]]
    headers = _editable_headers(sheet)
    current_fields = {
        "requirement": "current_requirement_ids",
        "api": "current_api_ids",
        "test": "current_test_case_ids",
    }
    items = resolution.get("route_mapping_resolutions")
    if not isinstance(items, list):
        return
    for item in items:
        if not isinstance(item, dict):
            continue
        row = rows["page_routes"][_text(item, "page_id")]
        for dimension in _string_list(item.get("unresolved_dimensions")):
            if dimension not in _ROUTE_DIMENSIONS:
                continue
            field, proposed_field, _prefixes = _ROUTE_DIMENSIONS[dimension]
            cell = sheet.cell(row, headers[field])
            if dimension == "api" and item.get("api_not_applicable") is True:
                cell.value = "不适用"
                continue
            identifiers = _ordered_union(
                _string_list(item.get(current_fields[dimension])),
                _string_list(item.get(proposed_field)),
            )
            cell.value = ";".join(identifiers)


def _apply_permission_mappings(
    workbook: Any,
    resolution: dict[str, Any],
    rows: dict[str, dict[str, int]],
) -> None:
    sheet = workbook[_CONTRACT_SHEETS["role_permissions"]]
    permission_code_column = _editable_headers(sheet)["权限代码"]
    items = resolution.get("permission_code_resolutions")
    if not isinstance(items, list):
        return
    for item in items:
        if not isinstance(item, dict) or item.get("permission_not_applicable") is True:
            continue
        permission_code = _text(item, "permission_code")
        for permission_id in _string_list(item.get("proposed_perm_ids")):
            row = rows["role_permissions"][permission_id]
            cell = sheet.cell(row, permission_code_column)
            if permission_code not in _permission_codes(cell.value):
                raw = str(cell.value or "").strip()
                cell.value = f"{raw};{permission_code}" if raw else permission_code


def _apply_resolution_to_workbook(
    candidate_workbook: Path,
    resolution: dict[str, Any],
) -> None:
    workbook = load_workbook(candidate_workbook, read_only=False, data_only=False)
    try:
        rows = _append_proposed_rows(workbook, resolution)
        _apply_requirement_test_mappings(workbook, resolution, rows)
        _apply_route_mappings(workbook, resolution, rows)
        _apply_permission_mappings(workbook, resolution, rows)
        workbook.save(candidate_workbook)
    finally:
        workbook.close()


def _generation_report(
    paths: RepositoryPaths,
    resolution: dict[str, Any],
    output_path: Path,
    checks: list[CheckResult],
) -> dict[str, Any]:
    return {
        "schema_version": 1,
        "stage": "P0",
        "operation": "generate validated P0 candidate technical workbook",
        "automated_assessment_only": True,
        "does_not_activate_baseline": True,
        "source_workbook": {
            "path": paths.technical_workbook.relative_to(paths.root).as_posix(),
            "sha256": sha256_file(paths.technical_workbook),
        },
        "candidate_workbook": {"path": str(output_path), "sha256": None},
        "resolution_sha256": _payload_sha256(resolution),
        "candidate_written": False,
        "candidate_ready_for_formal_baseline_review": False,
        "checks": [check.to_dict() for check in checks],
        "warning": (
            "Generation never overwrites or activates the authoritative technical workbook. "
            "A written candidate still requires formal baseline review and cannot authorize "
            "user-facing development."
        ),
    }


def generate_p0_candidate_workbook(
    paths: RepositoryPaths,
    resolution: dict[str, Any],
    candidate_output: Path,
) -> dict[str, Any]:
    output_path = candidate_output.resolve()
    checks = validate_p0_resolution_packet(paths, resolution)
    source_path = paths.technical_workbook.resolve()
    if output_path == source_path:
        checks.append(
            CheckResult(
                code="P0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN",
                message="The authoritative technical workbook cannot be overwritten",
                location=str(output_path),
            )
        )
    try:
        output_path.relative_to((paths.root / "doc").resolve())
    except ValueError:
        pass
    else:
        checks.append(
            CheckResult(
                code="P0_CHANGE_OUTPUT_LOCATION_FORBIDDEN",
                message="Candidate workbooks must be written outside the authoritative doc tree",
                location=str(output_path),
            )
        )
    if output_path.suffix.lower() != ".xlsx":
        checks.append(
            CheckResult(
                code="P0_CHANGE_OUTPUT_EXTENSION_INVALID",
                message="Candidate output must use the .xlsx extension",
                location=str(output_path),
            )
        )
    if output_path.exists():
        checks.append(
            CheckResult(
                code="P0_CHANGE_OUTPUT_EXISTS",
                message="Candidate output already exists and will not be overwritten",
                location=str(output_path),
            )
        )
    if checks:
        return _generation_report(paths, resolution, output_path, checks)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            dir=output_path.parent,
            prefix=f".{output_path.stem}.",
            suffix=".xlsx",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
        shutil.copy2(source_path, temporary_path)
        _apply_resolution_to_workbook(temporary_path, resolution)
        assessment = assess_p0_baseline_change(paths, resolution, temporary_path)
        assessment["operation"] = "generate validated P0 candidate technical workbook"
        assessment["candidate_written"] = False
        cast(dict[str, Any], assessment["candidate_workbook"])["path"] = str(output_path)
        if assessment.get("candidate_ready_for_formal_baseline_review") is not True:
            return assessment
        os.replace(temporary_path, output_path)
        temporary_path = None
        assessment["candidate_written"] = True
        return assessment
    except Exception as error:
        return _generation_report(
            paths,
            resolution,
            output_path,
            [
                CheckResult(
                    code="P0_CHANGE_GENERATION_FAILED",
                    message=f"Could not generate candidate workbook: {error}",
                    location=str(output_path),
                )
            ],
        )
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)


def load_and_generate_p0_candidate_workbook(
    paths: RepositoryPaths,
    resolution_path: Path,
    candidate_output: Path,
) -> dict[str, Any]:
    try:
        payload = json.loads(resolution_path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as error:
        report = _invalid_resolution_report(
            f"Cannot read resolution packet: {error}", resolution_path
        )
        report["operation"] = "generate validated P0 candidate technical workbook"
        report["candidate_written"] = False
        return report
    if not isinstance(payload, dict):
        report = _invalid_resolution_report(
            "Resolution packet root must be an object", resolution_path
        )
        report["operation"] = "generate validated P0 candidate technical workbook"
        report["candidate_written"] = False
        return report
    return generate_p0_candidate_workbook(paths, payload, candidate_output)
