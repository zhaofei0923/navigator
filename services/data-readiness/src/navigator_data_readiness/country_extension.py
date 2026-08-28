"""Prepare one-country additions without changing or activating the frozen BASIC60 release.

An extension is a replayable, non-importable candidate. It binds the original seed,
new country CSVs, downloaded evidence and optional bilingual overview. No command in
this module signs a review, writes a ready seed, changes a database or publishes data.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import tempfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlsplit
from xml.etree.ElementTree import ParseError
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import openpyxl
from navigator_api.basic60_seed_contract import Basic60SeedArtifact
from navigator_api.market_schemas import CountryMarketOverview
from navigator_api.market_storage import (
    canonical_bytes,
    checked_sha,
    content_sha256,
    ensure_no_symlinks,
    read_json,
)
from openpyxl.utils.exceptions import InvalidFileException

from .basic60_private import ENERGY_METRICS, MACRO_METRICS, PROFILE_METRICS
from .market_content import (
    _immutable_bytes,
    _publish_directory,
    _read_workbook,
    _validated_candidate,
    _workbook_bytes,
    apply_market_review_rows,
)

EXTENSION_SCHEMA = "navigator.country-extension.v1"
EVIDENCE_SCHEMA = "navigator.country-extension-evidence.v1"
REVIEW_SCHEMA = "navigator.country-extension-review.v1"
TARGET_RELEASE_ID = "BASIC61-PRIVATE-R1"
COUNTRY_EXTENSION_COMMANDS = frozenset({"prepare-country-extension", "validate-country-extension"})
REVIEW_HEADERS = (
    "country_code",
    "section",
    "locale",
    "json_pointer",
    "original_value",
    "edited_value",
    "source_ids",
)
SOURCE_HEADERS = ("id", "title", "url", "captured_at", "local_path", "sha256")
MAX_INPUT_BYTES = 128 * 1024 * 1024
PROFILE_FIELDS = frozenset(
    {
        "iso2",
        "iso3",
        "country_name_zh",
        "country_name_en",
        "official_name_en",
        "local_names",
        "capital",
        "admin_level_1_count",
        "admin_level_1_type",
        "population",
        "population_year",
        "official_languages",
        "currency_code",
        "currency_name",
        "time_zones",
        "area_sq_km",
        "area_year",
        "region",
        "collected_at",
    }
)
METRIC_FIELDS = frozenset(
    {
        "iso3",
        "record_type",
        "year",
        "collected_at",
        "electricity_demand_gwh",
        "electricity_demand_year",
        *(field for _code, field, _unit in MACRO_METRICS),
        *(field for _code, field, _unit, _period in ENERGY_METRICS),
        *(period for _code, _field, _unit, period in ENERGY_METRICS),
    }
)


class CountryExtensionError(ValueError):
    """Invalid or unsafe extension input; never a human review outcome."""


def _day(value: object) -> str:
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise CountryExtensionError("dates must be literal YYYY-MM-DD values")
    return date.fromisoformat(value).isoformat()


def _file_reference(repo: Path, path: Path) -> dict[str, Any]:
    ensure_no_symlinks(path)
    resolved = path.resolve()
    if not resolved.is_relative_to(repo) or not resolved.is_file():
        raise CountryExtensionError("extension inputs must be real files within the repository")
    if resolved.stat().st_size > MAX_INPUT_BYTES:
        raise CountryExtensionError("extension input exceeds its size limit")
    with resolved.open("rb") as handle:
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
    return {"path": resolved.relative_to(repo).as_posix(), "sha256": digest}


def _reference_path(repo: Path, reference: object) -> Path:
    if not isinstance(reference, dict) or set(reference) != {"path", "sha256"}:
        raise CountryExtensionError("input reference must contain exactly path and sha256")
    relative = reference["path"]
    if (
        not isinstance(relative, str)
        or "\\" in relative
        or PurePosixPath(relative).is_absolute()
        or any(part in {"", ".", ".."} for part in relative.split("/"))
    ):
        raise CountryExtensionError("input reference must be a canonical repository-relative path")
    checked_sha(reference["sha256"])
    path = repo / relative
    if _file_reference(repo, path) != reference:
        raise CountryExtensionError(f"frozen extension input changed: {relative}")
    return path


def _csv(path: Path, fields: frozenset[str]) -> list[dict[str, str]]:
    reader = csv.DictReader(io.StringIO(path.read_text(encoding="utf-8-sig")))
    headers = reader.fieldnames or []
    if len(headers) != len(set(headers)) or not fields.issubset(headers):
        raise CountryExtensionError("CSV columns are missing or duplicated")
    result = []
    for row in reader:
        if None in row or any(value is None for value in row.values()):
            raise CountryExtensionError("CSV row width differs from its header")
        result.append({key: str(value) for key, value in row.items()})
    return result


def _texts(value: str, field: str) -> list[str]:
    parsed: object = json.loads(value)
    if (
        not isinstance(parsed, list)
        or not parsed
        or any(not isinstance(item, str) or not item.strip() for item in parsed)
        or len({str(item).strip() for item in parsed}) != len(parsed)
    ):
        raise CountryExtensionError(f"{field} must be a nonempty unique JSON string array")
    return [str(item).strip() for item in parsed]


def _source_ids(value: object, allowed: set[str], *, required: bool) -> list[str]:
    if not isinstance(value, list) or any(not isinstance(item, str) for item in value):
        raise CountryExtensionError("source mappings must be arrays of source IDs")
    if (
        len(value) != len(set(value))
        or not set(value).issubset(allowed)
        or (required and not value)
    ):
        raise CountryExtensionError(
            "nonempty known source references are required for supplied values"
        )
    return list(value)


def _evidence(repo: Path, payload: dict[str, Any], code: str, as_of: str) -> dict[str, Any]:
    expected = {
        "schema_version",
        "country_code",
        "sources",
        "identity_source_ids",
        "metric_source_ids",
    }
    if set(payload) != expected or payload["schema_version"] != EVIDENCE_SCHEMA:
        raise CountryExtensionError("invalid country-extension evidence contract")
    if payload["country_code"] != code or not isinstance(payload["sources"], list):
        raise CountryExtensionError("evidence must identify the same new country")
    ids: set[str] = set()
    for source in payload["sources"]:
        if not isinstance(source, dict) or set(source) != {
            "id",
            "title",
            "url",
            "captured_at",
            "local_path",
            "sha256",
        }:
            raise CountryExtensionError("source must bind a title, URL, time and downloaded file")
        for field in ("id", "title", "url", "captured_at", "local_path", "sha256"):
            if not isinstance(source[field], str) or not source[field].strip():
                raise CountryExtensionError("source fields must be nonempty literal text")
        if source["id"] in ids:
            raise CountryExtensionError("source IDs must be unique")
        ids.add(source["id"])
        parsed = urlsplit(source["url"])
        if parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password:
            raise CountryExtensionError("source URLs must be public HTTPS without credentials")
        captured = datetime.fromisoformat(source["captured_at"])
        if (
            captured.tzinfo is None
            or captured.utcoffset() is None
            or captured.date() > date.fromisoformat(as_of)
        ):
            raise CountryExtensionError("source capture requires timezone and cannot follow as_of")
        if not source["local_path"].startswith("raw material/"):
            raise CountryExtensionError("downloaded evidence must remain under raw material")
        _reference_path(repo, {"path": source["local_path"], "sha256": source["sha256"]})
    _source_ids(payload["identity_source_ids"], ids, required=True)
    mappings = payload["metric_source_ids"]
    if not isinstance(mappings, dict):
        raise CountryExtensionError("metric_source_ids must be an object keyed by metric code")
    allowed_metrics = {
        *(item[0] for item in PROFILE_METRICS),
        *(item[0] for item in MACRO_METRICS),
        *(item[0] for item in ENERGY_METRICS),
        "electricity_demand_gwh",
    }
    if not set(mappings).issubset(allowed_metrics):
        raise CountryExtensionError("unknown metric evidence mapping")
    for mapping in mappings.values():
        _source_ids(mapping, ids, required=False)
    return payload


def _observation(
    code: str, raw: str, unit: str, period: str, evidence: dict[str, Any]
) -> dict[str, Any]:
    value = raw.strip()
    if period and (not re.fullmatch(r"\d{4}", period) or not 1900 <= int(period) <= 2200):
        raise CountryExtensionError(
            "statistical periods must be four-digit years or empty for missing data"
        )
    normalized = None
    if value:
        try:
            number = Decimal(value)
        except InvalidOperation as exc:
            raise CountryExtensionError("metric values must be literal decimal numbers") from exc
        if not number.is_finite() or not period:
            raise CountryExtensionError(
                "available metrics require finite values and a statistical year"
            )
        signed = {"gdp_growth_pct", "inflation_cpi_pct", "fdi_net_inflows_usd"}
        if code not in signed and number < 0:
            raise CountryExtensionError("this metric cannot be negative")
        if code.startswith("renewable_share_") and number > 100:
            raise CountryExtensionError("renewable shares cannot exceed 100 percent")
        if code == "population_total" and number != number.to_integral_value():
            raise CountryExtensionError("population must be an integer")
        if code == "official_exchange_rate_lcu_per_usd" and number <= 0:
            raise CountryExtensionError("available official exchange rates must be positive")
        normalized = format(number, "f")
    sources = evidence["metric_source_ids"].get(code, [])
    _source_ids(sources, {item["id"] for item in evidence["sources"]}, required=bool(value))
    return {
        "metric_code": code,
        "period": period or None,
        "original_value": raw if value else None,
        "normalized_value": normalized,
        "unit": unit,
        "value_status": "available" if value else "pending",
        "null_reason": None if value else "source_value_missing",
        "source_ids": sources,
        "quality_status": "machine_validated",
    }


def _country(
    profile: dict[str, str], records: list[dict[str, str]], evidence: dict[str, Any], as_of: str
) -> dict[str, Any]:
    code = profile["iso3"].strip()
    for field in (
        "country_name_zh",
        "country_name_en",
        "official_name_en",
        "capital",
        "admin_level_1_type",
        "currency_name",
        "region",
    ):
        if not profile[field].strip():
            raise CountryExtensionError(f"country identity requires {field}")
    if not re.fullmatch(r"[A-Z]{2}", profile["iso2"]) or not re.fullmatch(
        r"[A-Z]{3}", profile["currency_code"]
    ):
        raise CountryExtensionError("ISO2 and currency codes must use uppercase ASCII codes")
    if not re.fullmatch(r"\d+", profile["admin_level_1_count"]):
        raise CountryExtensionError("administrative unit count must be an integer")
    if _day(profile["collected_at"]) > as_of:
        raise CountryExtensionError("profile collection date cannot follow as_of")
    timezones = _texts(profile["time_zones"], "time_zones")
    try:
        for timezone in timezones:
            ZoneInfo(timezone)
    except (ValueError, ZoneInfoNotFoundError) as exc:
        raise CountryExtensionError("unknown IANA timezone") from exc
    macro = [item for item in records if item["record_type"] == "macro_annual"]
    energy = [item for item in records if item["record_type"] == "energy_latest"]
    if len(records) != 6 or len(macro) != 5 or len(energy) != 1:
        raise CountryExtensionError("extension requires five macro rows and one energy row")
    if {item["year"] for item in macro} != {str(year) for year in range(2020, 2025)}:
        raise CountryExtensionError(
            "macro rows must cover each year 2020 through 2024 exactly once"
        )
    if any(item["iso3"].strip() != code or _day(item["collected_at"]) > as_of for item in records):
        raise CountryExtensionError(
            "metric country or collection date does not match the extension"
        )
    metrics = [
        _observation(
            metric,
            profile[field],
            unit,
            profile["population_year"] if metric == "population_total" else profile["area_year"],
            evidence,
        )
        for metric, field, unit in PROFILE_METRICS
    ]
    metrics.extend(
        _observation(metric, item[field], unit, item["year"], evidence)
        for item in sorted(macro, key=lambda row: row["year"])
        for metric, field, unit in MACRO_METRICS
    )
    metrics.extend(
        _observation(metric, energy[0][field], unit, energy[0][year_field], evidence)
        for metric, field, unit, year_field in ENERGY_METRICS
    )
    metrics.append(
        _observation(
            "electricity_demand_gwh",
            energy[0]["electricity_demand_gwh"],
            "GWH",
            energy[0]["electricity_demand_year"],
            evidence,
        )
    )
    if any(item["period"] and int(item["period"]) > int(as_of[:4]) for item in metrics):
        raise CountryExtensionError("statistical periods cannot follow the information date")
    by_code = {item["metric_code"]: item for item in metrics}
    for suffix, total_code, renewable_code in (
        ("capacity", "electricity_installed_capacity_mw", "renewable_capacity_mw"),
        ("generation", "electricity_generation_gwh", "renewable_generation_gwh"),
    ):
        total, renewable = by_code[total_code], by_code[renewable_code]
        if (
            total["value_status"] != "available"
            or renewable["value_status"] != "available"
            or total["period"] != renewable["period"]
        ):
            continue
        denominator, numerator = (
            Decimal(total["normalized_value"]),
            Decimal(renewable["normalized_value"]),
        )
        if numerator > denominator:
            raise CountryExtensionError("same-year renewable energy cannot exceed the total")
        share = by_code[f"renewable_share_{suffix}_pct"]
        if (
            share["value_status"] == "available"
            and share["period"] == total["period"]
            and (
                denominator == 0
                or abs(Decimal(share["normalized_value"]) - numerator / denominator * 100)
                > Decimal("0.02")
            )
        ):
            raise CountryExtensionError("same-year renewable share disagrees with its components")
    return {
        "iso3": code,
        "iso2": profile["iso2"],
        "identity": {
            "country_name_zh": profile["country_name_zh"],
            "country_name_en": profile["country_name_en"],
            "official_name_en": profile["official_name_en"],
            "local_names": _texts(profile["local_names"], "local_names"),
            "capital": profile["capital"],
            "region_code": profile["region"],
            "admin_level_1_count": int(profile["admin_level_1_count"]),
            "admin_level_1_type": profile["admin_level_1_type"],
            "official_languages": _texts(profile["official_languages"], "official_languages"),
            "currency_code": profile["currency_code"],
            "currency_name": profile["currency_name"],
            "time_zones": timezones,
            "collected_at": profile["collected_at"],
        },
        "identity_source_ids": evidence["identity_source_ids"],
        "metrics": metrics,
        "assessment": {
            "opportunity_level": "pending",
            "policy_friendliness_level": "pending",
            "risk_assessment_status": "unknown",
            "risk_level": None,
        },
    }


def build_country_extension(
    *,
    repo_root: Path,
    parent_seed: Path,
    profile_csv: Path,
    metrics_csv: Path,
    evidence_path: Path,
    as_of: str,
    revision: int = 1,
    overview_path: Path | None = None,
) -> dict[str, Any]:
    repo = repo_root.resolve()
    as_of = _day(as_of)
    if type(revision) is not int or revision < 1:
        raise CountryExtensionError("extension revision must be a positive integer")
    inputs = {
        "parent_seed": parent_seed,
        "profile_csv": profile_csv,
        "metrics_csv": metrics_csv,
        "evidence": evidence_path,
    }
    if overview_path is not None:
        inputs["overview"] = overview_path
    references = {role: _file_reference(repo, path) for role, path in inputs.items()}
    parent_payload = read_json(parent_seed, max_bytes=MAX_INPUT_BYTES)
    parent = Basic60SeedArtifact.model_validate(parent_payload)
    parent_codes = {country.iso3 for country in parent.countries}
    parent_counts = parent.release.counts.model_dump()
    if (
        parent.release.status != "private_trial_ready"
        or len(parent_codes) != 60
        or "CHN" not in parent_codes
        or parent_counts
        != {
            "countries": 60,
            "macro_rows": 300,
            "energy_rows": 60,
            "available_metric_values": 2279,
            "pending_metric_values": 61,
        }
    ):
        raise CountryExtensionError(
            "extension requires the unchanged 60-country BASIC60-PRIVATE-R1 parent"
        )
    profiles = _csv(profile_csv, PROFILE_FIELDS)
    if len(profiles) != 1:
        raise CountryExtensionError("profile CSV must contain exactly the newly added country")
    profile = profiles[0]
    code = profile["iso3"].strip()
    if not re.fullmatch(r"[A-Z]{3}", code) or code in parent_codes or code == "CHN":
        raise CountryExtensionError("extension country must be a new non-China ISO3 code")
    if profile["iso2"] in {country.iso2 for country in parent.countries}:
        raise CountryExtensionError("extension ISO2 code already belongs to a parent country")
    evidence = _evidence(repo, read_json(evidence_path), code, as_of)
    country = _country(profile, _csv(metrics_csv, METRIC_FIELDS), evidence, as_of)
    overview = None
    if overview_path is not None:
        overview = CountryMarketOverview.model_validate(read_json(overview_path)).model_dump(
            mode="json"
        )
        if overview["country_code"] != code or overview["as_of"] != as_of:
            raise CountryExtensionError("overview country and as_of must match the extension")
    available = sum(item["value_status"] == "available" for item in country["metrics"])
    pending = sum(item["value_status"] == "pending" for item in country["metrics"])
    parent_hashes = {item["iso3"]: content_sha256(item) for item in parent_payload["countries"]}
    candidate = {
        "schema_version": EXTENSION_SCHEMA,
        "extension_id": f"COUNTRY-EXT-{code}-{as_of.replace('-', '')}-R{revision}",
        "revision": revision,
        "as_of": as_of,
        "target_release_id": TARGET_RELEASE_ID,
        "status": "awaiting_single_review",
        "published": False,
        "formal_gate_status": "pending",
        "inputs": references,
        "parent": {
            "release_id": parent.release.release_id,
            "seed_sha256": references["parent_seed"]["sha256"],
            "country_sha256": parent_hashes,
            "counts": parent_counts,
        },
        "scope": {
            "raw_country_count": 61,
            "outbound_country_count": 60,
            "country_codes": sorted(parent_codes | {code}),
            "outbound_country_codes": sorted((parent_codes | {code}) - {"CHN"}),
            "added_country_codes": [code],
            "unchanged_parent_country_codes": sorted(parent_codes),
        },
        "projected_counts": {
            "countries": 61,
            "macro_rows": 305,
            "energy_rows": 61,
            "available_metric_values": parent_counts["available_metric_values"] + available,
            "pending_metric_values": parent_counts["pending_metric_values"] + pending,
        },
        "new_country": country,
        "evidence": evidence,
        "overview": overview,
        "release_boundary": {
            "importable_seed": False,
            "requires_actual_single_confirmation": True,
            "requires_ready_seed_and_runtime_authorization": True,
            "review_scope_country": code,
            "existing_countries_require_new_human_review": False,
        },
    }
    for reference in references.values():
        _reference_path(repo, reference)
    return candidate


def _literal(value: object) -> str:
    return (
        value
        if isinstance(value, str)
        else json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    )


def source_review_cell(source: dict[str, Any], field: str) -> str:
    """The review v1 timestamp cell uses one exact, lossless JSON-string encoding."""
    value = source[field]
    if field == "captured_at":
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def extension_review_data(candidate: dict[str, Any]) -> dict[str, Any]:
    country = candidate["new_country"]
    code = country["iso3"]
    rows = [
        [
            code,
            "identity",
            "",
            "/new_country/iso2",
            country["iso2"],
            country["iso2"],
            ",".join(country["identity_source_ids"]),
        ]
    ]
    for key, value in sorted(country["identity"].items()):
        rows.append(
            [
                code,
                "identity",
                "",
                f"/new_country/identity/{key}",
                _literal(value),
                _literal(value),
                ",".join(country["identity_source_ids"]),
            ]
        )
    for index, metric in enumerate(country["metrics"]):
        for field, value in sorted(metric.items()):
            if field == "source_ids":
                continue
            rows.append(
                [
                    code,
                    "metric",
                    "",
                    f"/new_country/metrics/{index}/{field}",
                    _literal(value),
                    _literal(value),
                    ",".join(metric["source_ids"]),
                ]
            )
    return {
        "metadata": {
            "schema_version": REVIEW_SCHEMA,
            "extension_id": candidate["extension_id"],
            "candidate_sha256": content_sha256(candidate),
            "parent_seed_sha256": candidate["parent"]["seed_sha256"],
            "target_release_id": candidate["target_release_id"],
            "review_scope_country": code,
            "as_of": candidate["as_of"],
        },
        "headers": list(REVIEW_HEADERS),
        "rows": rows,
        "source_rows": candidate["evidence"]["sources"],
        "parent_preservation": {
            "country_sha256": candidate["parent"]["country_sha256"],
            "raw_country_count": 60,
            "outbound_country_count": 59,
            "changed_parent_country_codes": [],
        },
    }


def _scope_bytes(candidate: dict[str, Any]) -> bytes:
    handle = io.StringIO(newline="")
    writer = csv.writer(handle, lineterminator="\n")
    writer.writerow(["iso3", "scope_role"])
    for code in candidate["scope"]["country_codes"]:
        writer.writerow([code, "origin_excluded" if code == "CHN" else "outbound"])
    return handle.getvalue().encode("utf-8")


def _output_target(repo: Path, output: Path) -> Path:
    ensure_no_symlinks(output)
    target = output.resolve()
    allowed = repo.resolve() / "runtime" / "country-extensions"
    if not target.is_relative_to(allowed) or target == allowed:
        raise CountryExtensionError(
            "extension candidates must be new children of runtime/country-extensions"
        )
    if target.exists():
        raise CountryExtensionError(
            "extension candidate output already exists; choose a new revision directory"
        )
    return target


def prepare_country_extension(*, output_dir: Path, **inputs: Any) -> dict[str, Any]:
    repo = Path(inputs["repo_root"]).resolve()
    target = _output_target(repo, output_dir)
    candidate = build_country_extension(**inputs)
    target.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=".country-extension-", dir=target.parent) as name:
        staging = Path(name) / "candidate"
        staging.mkdir()
        _immutable_bytes(staging / "candidate.json", canonical_bytes(candidate))
        _immutable_bytes(
            staging / "review-data.json", canonical_bytes(extension_review_data(candidate))
        )
        _immutable_bytes(staging / "scope.csv", _scope_bytes(candidate))
        report = validate_country_extension(candidate_dir=staging, repo_root=repo)
        _immutable_bytes(staging / "validation.json", canonical_bytes(report))
        _publish_directory(staging, target)
    return {**report, "output_dir": target.relative_to(repo).as_posix()}


def _workbook_binding(
    path: Path, review: dict[str, Any], candidate: dict[str, Any]
) -> dict[str, Any]:
    ensure_no_symlinks(path)
    if not path.is_file() or path.stat().st_size > 32 * 1024 * 1024:
        raise CountryExtensionError("review workbook is missing or oversized")
    payload = _workbook_bytes(path)
    try:
        workbook = openpyxl.load_workbook(
            io.BytesIO(payload), read_only=True, data_only=False, keep_links=False
        )
    except (InvalidFileException, ParseError, KeyError, ValueError) as exc:
        raise CountryExtensionError("invalid extension workbook structure") from exc
    try:
        if not {"新增国包信息", "新增国内容", "基础数据依据"}.issubset(workbook.sheetnames):
            raise CountryExtensionError(
                "extension workbook requires 新增国包信息, 新增国内容 and 基础数据依据 worksheets"
            )
        metadata_sheet, content_sheet = workbook["新增国包信息"], workbook["新增国内容"]
        source_sheet = workbook["基础数据依据"]
        if (
            (metadata_sheet.max_row or 0) > 100
            or (content_sheet.max_row or 0) > 1000
            or (source_sheet.max_row or 0) > 1000
        ):
            raise CountryExtensionError(
                "extension review worksheet dimensions exceed supported limits"
            )
        metadata_rows = list(metadata_sheet.iter_rows(max_col=2))
        content_rows = list(content_sheet.iter_rows(max_col=7))
        source_rows = list(source_sheet.iter_rows(max_col=6))
        if any(
            cell.data_type in {"f", "e"}
            for row in [*metadata_rows, *content_rows, *source_rows]
            for cell in row
        ):
            raise CountryExtensionError(
                "review binding requires literal text, not formulas or errors"
            )
        metadata_values = [tuple(cell.value for cell in row) for row in metadata_rows]
        if not metadata_values or metadata_values[0] != ("key", "value"):
            raise CountryExtensionError("review package metadata header was changed")
        metadata: dict[str, object] = {}
        for key, value in metadata_values[1:]:
            if not isinstance(key, str) or key in metadata:
                raise CountryExtensionError("review metadata has invalid or duplicate keys")
            metadata[key] = value
        if metadata != review["metadata"]:
            raise CountryExtensionError("review metadata does not bind this frozen extension")
        expected = [list(REVIEW_HEADERS), *review["rows"]]
        actual = [
            [cell.value if cell.value is not None else "" for cell in row] for row in content_rows
        ]
        if actual != expected:
            raise CountryExtensionError(
                "review content differs; regenerate a new candidate for substantive edits"
            )
        expected_sources = [
            list(SOURCE_HEADERS),
            *[
                [source_review_cell(source, field) for field in SOURCE_HEADERS]
                for source in review["source_rows"]
            ],
        ]
        actual_sources = [
            [cell.value if cell.value is not None else "" for cell in row] for row in source_rows
        ]
        if actual_sources != expected_sources:
            raise CountryExtensionError("review source evidence differs from the frozen extension")
        market_binding = None
        if candidate["overview"] is not None:
            market_metadata, market_rows = _read_workbook(payload)
            market_candidate = _validated_candidate(
                market_metadata.get("package_id"),
                [candidate["overview"]],
                frozenset(candidate["scope"]["outbound_country_codes"]),
            )
            applied = apply_market_review_rows(
                market_candidate,
                market_metadata,
                market_rows,
                frozenset(candidate["scope"]["outbound_country_codes"]),
            )
            if applied.sha256 != market_candidate.sha256:
                raise CountryExtensionError(
                    "overview edits require a rebuilt extension and research package"
                )
            market_binding = {
                "package_id": market_candidate.package_id,
                "candidate_sha256": market_candidate.sha256,
            }
        return {
            "workbook_sha256": hashlib.sha256(payload).hexdigest(),
            "candidate_sha256": metadata["candidate_sha256"],
            "actual_confirmation_required": True,
            "market_overview": market_binding,
        }
    finally:
        workbook.close()


def validate_country_extension(
    *, candidate_dir: Path, repo_root: Path, workbook_path: Path | None = None
) -> dict[str, Any]:
    ensure_no_symlinks(candidate_dir)
    repo = repo_root.resolve()
    candidate = read_json(candidate_dir / "candidate.json")
    if candidate.get("schema_version") != EXTENSION_SCHEMA or not isinstance(
        candidate.get("inputs"), dict
    ):
        raise CountryExtensionError("invalid extension candidate manifest")
    references = candidate["inputs"]
    if set(references) not in (
        {"parent_seed", "profile_csv", "metrics_csv", "evidence"},
        {"parent_seed", "profile_csv", "metrics_csv", "evidence", "overview"},
    ):
        raise CountryExtensionError("extension inputs are missing or unknown")
    paths = {role: _reference_path(repo, reference) for role, reference in references.items()}
    revision = candidate.get("revision")
    if type(revision) is not int:
        raise CountryExtensionError("extension revision must be a positive integer")
    replay = build_country_extension(
        repo_root=repo,
        parent_seed=paths["parent_seed"],
        profile_csv=paths["profile_csv"],
        metrics_csv=paths["metrics_csv"],
        evidence_path=paths["evidence"],
        as_of=_day(candidate.get("as_of")),
        revision=revision,
        overview_path=paths.get("overview"),
    )
    if canonical_bytes(replay) != canonical_bytes(candidate):
        raise CountryExtensionError("candidate does not replay exactly from frozen inputs")
    review = extension_review_data(replay)
    if read_json(candidate_dir / "review-data.json") != review:
        raise CountryExtensionError("review-data differs from the frozen extension candidate")
    scope_path = candidate_dir / "scope.csv"
    ensure_no_symlinks(scope_path)
    if scope_path.read_bytes() != _scope_bytes(candidate):
        raise CountryExtensionError(
            "candidate scope differs from the retained parent and new country"
        )
    return {
        "status": "candidate_valid",
        "extension_id": candidate["extension_id"],
        "candidate_sha256": content_sha256(candidate),
        "parent_seed_sha256": candidate["parent"]["seed_sha256"],
        "country_code": candidate["new_country"]["iso3"],
        "target_release_id": TARGET_RELEASE_ID,
        "raw_country_count": 61,
        "outbound_country_count": 60,
        "changed_parent_country_codes": [],
        "new_metric_count": len(candidate["new_country"]["metrics"]),
        "overview_included": candidate["overview"] is not None,
        "published": False,
        "publication_status": "awaiting_single_review",
        "formal_gate_status": "pending",
        "workbook_binding": None
        if workbook_path is None
        else _workbook_binding(workbook_path, review, candidate),
    }


def add_country_extension_commands(commands: Any) -> None:
    prepare = commands.add_parser(
        "prepare-country-extension",
        help="Prepare a frozen-parent, one-country review candidate; never activate it.",
    )
    for name in ("parent-seed", "profile-csv", "metrics-csv", "evidence", "output-dir"):
        prepare.add_argument(f"--{name}", type=Path, required=True)
    prepare.add_argument("--as-of", required=True)
    prepare.add_argument("--revision", type=int, default=1)
    prepare.add_argument("--overview", type=Path)
    validate = commands.add_parser(
        "validate-country-extension",
        help="Replay an extension and optional workbook binding without publication.",
    )
    validate.add_argument("--candidate-dir", type=Path, required=True)
    validate.add_argument("--workbook", type=Path)


def run_country_extension_command(args: argparse.Namespace, repo_root: Path) -> int:
    def path(value: Path) -> Path:
        return value if value.is_absolute() else repo_root / value

    try:
        if args.command == "prepare-country-extension":
            report = prepare_country_extension(
                repo_root=repo_root,
                parent_seed=path(args.parent_seed),
                profile_csv=path(args.profile_csv),
                metrics_csv=path(args.metrics_csv),
                evidence_path=path(args.evidence),
                as_of=args.as_of,
                revision=args.revision,
                output_dir=path(args.output_dir),
                overview_path=path(args.overview) if args.overview else None,
            )
        else:
            report = validate_country_extension(
                candidate_dir=path(args.candidate_dir),
                repo_root=repo_root,
                workbook_path=path(args.workbook) if args.workbook else None,
            )
    except (OSError, ValueError, KeyError, TypeError) as exc:
        print(
            json.dumps(
                {"status": "not_ready", "published": False, "error": str(exc)}, ensure_ascii=False
            )
        )
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0
