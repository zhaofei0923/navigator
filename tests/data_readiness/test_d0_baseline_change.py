from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Any

from navigator_data_readiness.baseline import sha256_file
from navigator_data_readiness.d0_baseline_change import (
    assess_d0_baseline_change,
    generate_d0_candidate_workbook,
    load_and_assess_d0_baseline_change,
    load_and_generate_d0_candidate_workbook,
)
from navigator_data_readiness.d0_confirmation import (
    build_confirmed_contract_resolution,
    load_confirmation,
)
from navigator_data_readiness.paths import RepositoryPaths, discover_repository
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pytest import MonkeyPatch


def _completed_resolution(paths: RepositoryPaths) -> dict[str, Any]:
    return build_confirmed_contract_resolution(
        paths,
        load_confirmation(paths.root / "data/d0/evidence/kevin_confirmation_20260801.json"),
    )


def _headers(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[4]
        if cell.value is not None and str(cell.value).strip()
    }


def _row_index(sheet: Worksheet, field_id: str) -> int:
    column = _headers(sheet)["字段编号"]
    for row in range(5, sheet.max_row + 1):
        if str(sheet.cell(row, column).value or "").strip() == field_id:
            return row
    raise AssertionError(f"Cannot find {field_id}")


def _next_record_row(sheet: Worksheet) -> int:
    column = _headers(sheet)["字段编号"]
    row = 5
    while str(sheet.cell(row, column).value or "").strip():
        row += 1
    return row


def _candidate_workbook(
    tmp_path: Path,
    packet: dict[str, Any],
) -> Path:
    paths = discover_repository()
    candidate = tmp_path / "d0-candidate.xlsx"
    report = generate_d0_candidate_workbook(paths, packet, candidate)
    assert report["candidate_written"] is True, report["checks"]
    return candidate


def _codes(report: dict[str, Any]) -> set[str]:
    return {
        str(item["code"])
        for item in report.get("checks", [])
        if isinstance(item, dict) and item.get("code")
    }


def test_candidate_workbook_exactly_applying_resolution_is_ready(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    candidate = _candidate_workbook(tmp_path, packet)

    report = assess_d0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is True, report["checks"]
    assert report["does_not_activate_baseline"] is True
    assert report["checks"] == []
    assert report["core_entity_assessment"]["machine_status"] == "pass"
    assert report["core_field_assessment"]["machine_status"] == "pass"
    assert paths.d0_workbook != candidate


def test_candidate_applies_registered_unit_and_composite_key(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    unit = packet["field_unit_resolutions"][0]
    unit.update(
        {
            "unit_applicability": "unit_code",
            "unit_code": "MW",
            "unit_dimension": "power",
            "unit_registry_reference": "UNIT-REGISTRY-SI-POWER",
        }
    )
    composite = next(
        item
        for item in packet["entity_primary_key_resolutions"]
        if len(item["current_field_ids"]) >= 2
    )
    composite["action"] = "model_composite_key"
    composite["primary_key_field_ids"] = composite["current_field_ids"][:2]
    composite["proposed_field_contract"] = None
    candidate = _candidate_workbook(tmp_path, packet)

    report = assess_d0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is True, report["checks"]


def test_unchanged_candidate_cannot_satisfy_resolution(tmp_path: Path) -> None:
    paths = discover_repository()
    candidate = tmp_path / "unchanged.xlsx"
    shutil.copy2(paths.d0_workbook, candidate)

    report = assess_d0_baseline_change(paths, _completed_resolution(paths), candidate)
    codes = _codes(report)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "D0_CHANGE_WORKBOOK_UNCHANGED" in codes
    assert "D0_CHANGE_FIELD_HEADER_INVALID" in codes
    assert "D0_CHANGE_FIELD_EVIDENCE_NOT_READY" in codes


def test_unreviewed_field_change_is_rejected(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["核心字段冻结"]
        row = _row_index(sheet, "DATA-USER-001")
        sheet.cell(row, _headers(sheet)["中文名称"], "未评审名称")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_d0_baseline_change(paths, packet, candidate)

    assert "D0_CHANGE_FIELD_ROW_MISMATCH" in _codes(report)
    assert "D0_CHANGE_UNAUTHORIZED_CELL_CHANGE" in _codes(report)


def test_unit_and_primary_key_must_be_applied_exactly(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    key_item = next(
        item for item in packet["entity_primary_key_resolutions"] if item["current_field_ids"]
    )
    key_item["action"] = "promote_existing_field"
    key_item["primary_key_field_ids"] = [key_item["current_field_ids"][0]]
    key_item["proposed_field_contract"] = None
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["核心字段冻结"]
        unit_row = _row_index(sheet, "DATA-USER-001")
        sheet.cell(unit_row, _headers(sheet)["单位"], "MW")
        key_row = _row_index(sheet, key_item["primary_key_field_ids"][0])
        sheet.cell(key_row, _headers(sheet)["唯一/索引"], "索引")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_d0_baseline_change(paths, packet, candidate)
    codes = _codes(report)

    assert "D0_CHANGE_FIELD_ROW_MISMATCH" in codes
    assert "D0_CHANGE_ENTITY_EVIDENCE_NOT_READY" in codes


def test_candidate_cannot_add_unreviewed_field_or_header(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["核心字段冻结"]
        extra_column = max(_headers(sheet).values()) + 1
        sheet.cell(4, extra_column, "未评审列")
        row = _next_record_row(sheet)
        sheet.cell(row, _headers(sheet)["字段编号"], "DATA-UNREVIEWED-001")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_d0_baseline_change(paths, packet, candidate)
    codes = _codes(report)

    assert "D0_CHANGE_FIELD_HEADER_INVALID" in codes
    assert "D0_CHANGE_FIELD_ID_SET_INVALID" in codes
    assert "D0_CHANGE_FIELD_ID_SEQUENCE_INVALID" in codes


def test_candidate_cannot_change_other_d0_contract_or_formula_cell(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["枚举规则冻结"]
        sheet.cell(5, _headers(sheet)["中文名称"], '="未评审公式"')
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_d0_baseline_change(paths, packet, candidate)
    codes = _codes(report)

    assert "D0_CHANGE_UNAUTHORIZED_CONTRACT_CHANGE" in codes
    assert "D0_CHANGE_UNAUTHORIZED_CELL_CHANGE" in codes


def test_candidate_cannot_change_sheet_set(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        workbook.create_sheet("未评审工作表")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_d0_baseline_change(paths, packet, candidate)

    assert "D0_CHANGE_SHEET_SET_INVALID" in _codes(report)


def test_source_workbook_and_missing_candidate_are_rejected(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)

    source_report = assess_d0_baseline_change(paths, packet, paths.d0_workbook)
    missing_report = assess_d0_baseline_change(paths, packet, tmp_path / "missing.xlsx")

    assert _codes(source_report) == {"D0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN"}
    assert _codes(missing_report) == {"D0_CHANGE_WORKBOOK_MISSING"}


def test_invalid_resolution_and_workbook_are_reported(tmp_path: Path) -> None:
    paths = discover_repository()
    invalid_resolution = tmp_path / "resolution.json"
    invalid_resolution.write_text("[]", encoding="utf-8")
    broken_workbook = tmp_path / "broken.xlsx"
    broken_workbook.write_bytes(b"not an xlsx")

    invalid_report = load_and_assess_d0_baseline_change(
        paths,
        invalid_resolution,
        broken_workbook,
    )
    broken_report = assess_d0_baseline_change(
        paths,
        _completed_resolution(paths),
        broken_workbook,
    )

    assert _codes(invalid_report) == {"D0_CHANGE_RESOLUTION_INVALID"}
    assert "D0_CHANGE_WORKBOOK_INVALID" in _codes(broken_report)


def test_assessment_can_be_serialized(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    candidate = _candidate_workbook(tmp_path, packet)

    payload = json.dumps(
        assess_d0_baseline_change(paths, packet, candidate),
        ensure_ascii=False,
    )

    assert "candidate_ready_for_formal_baseline_review" in payload


def test_generator_writes_only_a_validated_separate_candidate(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    output = tmp_path / "generated-d0-candidate.xlsx"
    source_hash = sha256_file(paths.d0_workbook)

    report = generate_d0_candidate_workbook(paths, packet, output)

    assert report["candidate_written"] is True, report["checks"]
    assert report["candidate_ready_for_formal_baseline_review"] is True
    assert output.is_file()
    assert sha256_file(paths.d0_workbook) == source_hash
    assert assess_d0_baseline_change(paths, packet, output)[
        "candidate_ready_for_formal_baseline_review"
    ]
    workbook = load_workbook(output)
    try:
        sheet = workbook["核心字段冻结"]
        headers = _headers(sheet)
        assert headers["单位"] == max(headers.values())
        assert (
            sheet.cell(4, headers["单位"]).style_id == sheet.cell(4, headers["单位"] - 1).style_id
        )
    finally:
        workbook.close()


def test_generator_refuses_overwrite_authority_doc_tree_and_bad_extension(
    tmp_path: Path,
) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"keep")

    existing_report = generate_d0_candidate_workbook(paths, packet, existing)
    source_report = generate_d0_candidate_workbook(paths, packet, paths.d0_workbook)
    doc_report = generate_d0_candidate_workbook(
        paths,
        packet,
        paths.root / "doc" / "candidate.xlsx",
    )
    extension_report = generate_d0_candidate_workbook(
        paths,
        packet,
        tmp_path / "candidate.xls",
    )

    assert existing.read_bytes() == b"keep"
    assert "D0_CHANGE_OUTPUT_EXISTS" in _codes(existing_report)
    assert "D0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN" in _codes(source_report)
    assert "D0_CHANGE_OUTPUT_LOCATION_FORBIDDEN" in _codes(source_report)
    assert "D0_CHANGE_OUTPUT_LOCATION_FORBIDDEN" in _codes(doc_report)
    assert "D0_CHANGE_OUTPUT_EXTENSION_INVALID" in _codes(extension_report)


def test_generator_does_not_publish_a_failed_candidate(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    output = tmp_path / "failed-candidate.xlsx"
    monkeypatch.setattr(
        "navigator_data_readiness.d0_baseline_change._apply_resolution_to_workbook",
        lambda *_args: None,
    )

    report = generate_d0_candidate_workbook(paths, packet, output)

    assert report["candidate_written"] is False
    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert not output.exists()
    assert list(tmp_path.glob(".failed-candidate.*.xlsx")) == []


def test_generator_cleans_up_after_generation_exception(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    paths = discover_repository()
    packet = _completed_resolution(paths)
    output = tmp_path / "broken-candidate.xlsx"

    def fail_generation(*_args: Any) -> None:
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(
        "navigator_data_readiness.d0_baseline_change._apply_resolution_to_workbook",
        fail_generation,
    )

    report = generate_d0_candidate_workbook(paths, packet, output)

    assert _codes(report) == {"D0_CHANGE_GENERATION_FAILED"}
    assert report["candidate_written"] is False
    assert not output.exists()
    assert list(tmp_path.glob(".broken-candidate.*.xlsx")) == []


def test_generator_loader_rejects_invalid_resolution_without_writing(tmp_path: Path) -> None:
    paths = discover_repository()
    resolution = tmp_path / "resolution.json"
    resolution.write_text("[]", encoding="utf-8")
    output = tmp_path / "candidate.xlsx"

    report = load_and_generate_d0_candidate_workbook(paths, resolution, output)

    assert _codes(report) == {"D0_CHANGE_RESOLUTION_INVALID"}
    assert report["candidate_written"] is False
    assert not output.exists()
