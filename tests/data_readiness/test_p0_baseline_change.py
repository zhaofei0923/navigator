from __future__ import annotations

import copy
import json
import shutil
from pathlib import Path
from typing import Any

from navigator_data_readiness.baseline import sha256_file
from navigator_data_readiness.p0_baseline_change import (
    assess_p0_baseline_change,
    generate_p0_candidate_workbook,
    load_and_assess_p0_baseline_change,
    load_and_generate_p0_candidate_workbook,
)
from navigator_data_readiness.p0_resolution import (
    READY_FOR_CHANGE_REVIEW,
    build_p0_resolution_template,
)
from navigator_data_readiness.paths import discover_repository
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.worksheet import Worksheet
from pytest import MonkeyPatch


def _completed_resolution() -> dict[str, Any]:
    packet = copy.deepcopy(build_p0_resolution_template(discover_repository()))
    packet["template_only"] = False
    proposed_test_ids = [
        f"TC-P0-GAP-{index:03d}"
        for index, _item in enumerate(packet["requirement_test_resolutions"], start=1)
    ]
    packet["proposed_contract_ids"]["tests"] = proposed_test_ids
    packet["proposed_contract_rows"]["tests"] = [
        {
            "优先级": "P0",
            "前置条件": "测试前置条件",
            "场景/目标": f"{item['requirement_id']}追踪覆盖",
            "模块": item["module"],
            "步骤摘要": "执行需求验收路径",
            "测试类型": "契约/E2E",
            "状态": "已确认；待执行",
            "用例编号": test_id,
            "负责人": "测试",
            "需求编号": item["requirement_id"],
            "预期结果": "满足冻结验收条件",
        }
        for item, test_id in zip(
            packet["requirement_test_resolutions"],
            proposed_test_ids,
            strict=True,
        )
    ]
    metadata = {
        "resolution_status": "proposed",
        "change_request_id": "CR-P0-TRACE-20260731",
        "rationale": "Candidate workbook integration test.",
        "reviewer": "kevin",
        "reviewed_at": "2026-07-31",
    }
    for item, test_id in zip(
        packet["requirement_test_resolutions"],
        proposed_test_ids,
        strict=True,
    ):
        item.update(metadata)
        item["proposed_test_case_ids"] = [test_id]
    for item in packet["route_mapping_resolutions"]:
        item.update(metadata)
        dimensions = set(item["unresolved_dimensions"])
        if "requirement" in dimensions:
            item["proposed_requirement_ids"] = ["FR-AUTH-001"]
        if "api" in dimensions:
            item["proposed_api_ids"] = ["API-AUTH-001"]
        if "test" in dimensions:
            item["proposed_test_case_ids"] = ["TC-AUTH-001"]
    for item in packet["permission_code_resolutions"]:
        item.update(metadata)
        item["proposed_perm_ids"] = ["PERM-001"]
    packet["final_review"] = {
        "status": READY_FOR_CHANGE_REVIEW,
        "change_set_id": "P0-TRACE-CHANGESET-20260731",
        "reviewed_by": "kevin",
        "reviewed_at": "2026-07-31",
        "evidence_ids": ["EVD-D0-MAPPING-DECISIONS-20260731"],
    }
    return packet


def _headers(sheet: Worksheet) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in sheet[4]
        if cell.value is not None and str(cell.value).strip()
    }


def _row_index(sheet: Worksheet, identifier_field: str, identifier: str) -> int:
    headers = _headers(sheet)
    column = headers[identifier_field]
    for row in range(5, sheet.max_row + 1):
        if str(sheet.cell(row, column).value or "").strip() == identifier:
            return row
    raise AssertionError(f"Cannot find {identifier}")


def _next_record_row(sheet: Worksheet, identifier_field: str) -> int:
    headers = _headers(sheet)
    column = headers[identifier_field]
    row = 5
    while str(sheet.cell(row, column).value or "").strip():
        row += 1
    return row


def _ordered_union(*values: list[str]) -> list[str]:
    return list(dict.fromkeys(item for items in values for item in items))


def _apply_resolution(candidate: Path, packet: dict[str, Any]) -> None:
    workbook = load_workbook(candidate)
    try:
        test_sheet = workbook["测试用例"]
        test_headers = _headers(test_sheet)
        next_test_row = _next_record_row(test_sheet, "用例编号")
        for values in packet["proposed_contract_rows"]["tests"]:
            for field, value in values.items():
                test_sheet.cell(next_test_row, test_headers[field], value)
            next_test_row += 1

        route_sheet = workbook["页面路由清单"]
        route_headers = _headers(route_sheet)
        dimension_fields = {
            "requirement": (
                "关联需求编号",
                "current_requirement_ids",
                "proposed_requirement_ids",
            ),
            "api": ("依赖API", "current_api_ids", "proposed_api_ids"),
            "test": ("测试用例", "current_test_case_ids", "proposed_test_case_ids"),
        }
        for item in packet["route_mapping_resolutions"]:
            row = _row_index(route_sheet, "页面编号", item["page_id"])
            for dimension in item["unresolved_dimensions"]:
                field, current_field, proposed_field = dimension_fields[dimension]
                identifiers = _ordered_union(item[current_field], item[proposed_field])
                route_sheet.cell(row, route_headers[field], ";".join(identifiers))

        permission_sheet = workbook["角色权限矩阵"]
        permission_headers = _headers(permission_sheet)
        permission_code_column = permission_headers.get("权限代码")
        if permission_code_column is None:
            permission_code_column = max(permission_headers.values()) + 1
            permission_sheet.cell(4, permission_code_column, "权限代码")
        permission_row = _row_index(permission_sheet, "权限编号", "PERM-001")
        permission_codes = [
            item["permission_code"] for item in packet["permission_code_resolutions"]
        ]
        permission_sheet.cell(permission_row, permission_code_column, ";".join(permission_codes))
        workbook.save(candidate)
    finally:
        workbook.close()


def _candidate_workbook(tmp_path: Path, packet: dict[str, Any]) -> Path:
    paths = discover_repository()
    candidate = tmp_path / "technical-candidate.xlsx"
    shutil.copy2(paths.technical_workbook, candidate)
    _apply_resolution(candidate, packet)
    return candidate


def _codes(report: dict[str, Any]) -> set[str]:
    return {
        str(item["code"])
        for item in report.get("checks", [])
        if isinstance(item, dict) and item.get("code")
    }


def test_candidate_workbook_exactly_applying_resolution_is_ready(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)

    report = assess_p0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is True, report["checks"]
    assert report["does_not_activate_baseline"] is True
    assert report["checks"] == []
    assert report["traceability_assessment"]["traceability_ready"] is True
    assert report["traceability_assessment"]["delivery_ready"] is False
    assert paths.technical_workbook != candidate


def test_unchanged_candidate_does_not_satisfy_resolution(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = tmp_path / "unchanged.xlsx"
    shutil.copy2(paths.technical_workbook, candidate)

    report = assess_p0_baseline_change(paths, packet, candidate)
    codes = _codes(report)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "P0_CHANGE_WORKBOOK_UNCHANGED" in codes
    assert "P0_CHANGE_ADDED_ID_SET_INVALID" in codes
    assert "P0_CHANGE_TRACEABILITY_NOT_READY" in codes


def test_unreviewed_candidate_field_change_is_rejected(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["页面路由清单"]
        row = _row_index(sheet, "页面编号", "PC-PUB-001")
        sheet.cell(row, _headers(sheet)["页面名称"], "未评审名称")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_p0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "P0_CHANGE_UNAUTHORIZED_FIELD_CHANGE" in _codes(report)


def test_unreviewed_non_contract_cell_change_is_rejected(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        workbook["页面路由清单"]["A1"] = "未评审标题"
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_p0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "P0_CHANGE_UNAUTHORIZED_CELL_CHANGE" in _codes(report)


def test_new_test_cannot_add_unreviewed_requirement_reference(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["测试用例"]
        test_id = packet["requirement_test_resolutions"][0]["proposed_test_case_ids"][0]
        row = _row_index(sheet, "用例编号", test_id)
        requirement_column = _headers(sheet)["需求编号"]
        current = str(sheet.cell(row, requirement_column).value)
        sheet.cell(row, requirement_column, f"{current};FR-AUTH-001")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_p0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "P0_CHANGE_REQUIREMENT_TEST_NOT_APPLIED" in _codes(report)


def test_new_contract_row_cannot_contain_unreviewed_content(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["测试用例"]
        test_id = packet["requirement_test_resolutions"][0]["proposed_test_case_ids"][0]
        row = _row_index(sheet, "用例编号", test_id)
        sheet.cell(row, _headers(sheet)["步骤摘要"], "未在评审包中批准的步骤")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_p0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "P0_CHANGE_NEW_CONTRACT_ROW_MISMATCH" in _codes(report)


def test_candidate_cannot_add_an_unreviewed_contract_header(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)
    workbook = load_workbook(candidate)
    try:
        sheet = workbook["测试用例"]
        sheet.cell(4, sheet.max_column + 1, "未评审列")
        workbook.save(candidate)
    finally:
        workbook.close()

    report = assess_p0_baseline_change(paths, packet, candidate)

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert "P0_CHANGE_UNAUTHORIZED_HEADER_CHANGE" in _codes(report)


def test_source_workbook_cannot_be_used_as_candidate() -> None:
    paths = discover_repository()

    report = assess_p0_baseline_change(
        paths,
        _completed_resolution(),
        paths.technical_workbook,
    )

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert _codes(report) == {"P0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN"}


def test_invalid_resolution_file_is_reported(tmp_path: Path) -> None:
    paths = discover_repository()
    resolution = tmp_path / "resolution.json"
    resolution.write_text("[]", encoding="utf-8")

    report = load_and_assess_p0_baseline_change(
        paths,
        resolution,
        tmp_path / "missing.xlsx",
    )

    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert _codes(report) == {"P0_CHANGE_RESOLUTION_INVALID"}


def test_assessment_can_be_serialized(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    candidate = _candidate_workbook(tmp_path, packet)

    payload = json.dumps(
        assess_p0_baseline_change(paths, packet, candidate),
        ensure_ascii=False,
    )

    assert "candidate_ready_for_formal_baseline_review" in payload


def test_generator_writes_only_a_validated_separate_candidate(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    output = tmp_path / "generated-p0-candidate.xlsx"
    source_hash = sha256_file(paths.technical_workbook)

    report = generate_p0_candidate_workbook(paths, packet, output)

    assert report["candidate_written"] is True, report["checks"]
    assert report["candidate_ready_for_formal_baseline_review"] is True
    assert output.is_file()
    assert sha256_file(paths.technical_workbook) == source_hash
    assert assess_p0_baseline_change(paths, packet, output)[
        "candidate_ready_for_formal_baseline_review"
    ]
    workbook = load_workbook(output)
    try:
        permission_sheet = workbook["角色权限矩阵"]
        permission_headers = _headers(permission_sheet)
        assert permission_headers["权限代码"] == max(permission_headers.values())
        assert (
            permission_sheet.cell(4, permission_headers["权限代码"]).style_id
            == permission_sheet.cell(4, permission_headers["权限代码"] - 1).style_id
        )
        permission_table = permission_sheet.tables["PermissionMatrixTable"]
        assert range_boundaries(permission_table.ref)[2] == permission_headers["权限代码"]

        test_sheet = workbook["测试用例"]
        last_test_id = packet["proposed_contract_ids"]["tests"][-1]
        last_test_row = _row_index(test_sheet, "用例编号", last_test_id)
        test_table = test_sheet.tables["TestCaseTable"]
        assert range_boundaries(test_table.ref)[3] == last_test_row
    finally:
        workbook.close()


def test_generator_materializes_every_declared_contract_category(tmp_path: Path) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    new_requirement_id = "FR-P0-GAP-001"
    new_api_id = "API-P0-GAP-001"
    new_permission_id = "PERM-P0-GAP-001"
    packet["proposed_contract_ids"]["requirements"] = [new_requirement_id]
    packet["proposed_contract_rows"]["requirements"] = [
        {
            "需求编号": new_requirement_id,
            "模块": "追踪",
            "功能": "新增追踪合同",
            "优先级": "P0",
            "需求状态": "需求已确认",
        }
    ]
    packet["proposed_contract_rows"]["tests"][0]["需求编号"] += f";{new_requirement_id}"
    packet["proposed_contract_ids"]["apis"] = [new_api_id]
    packet["proposed_contract_rows"]["apis"] = [
        {
            "接口编号": new_api_id,
            "模块": "追踪",
            "用途": "新增追踪接口",
            "方法": "GET",
            "路径": "/api/v1/traceability-gap",
            "状态": "已确认；待实现",
        }
    ]
    api_item = next(
        item
        for item in packet["route_mapping_resolutions"]
        if "api" in item["unresolved_dimensions"]
    )
    api_item["proposed_api_ids"] = [new_api_id]
    permission_code = packet["permission_code_resolutions"][0]["permission_code"]
    packet["proposed_contract_ids"]["permissions"] = [new_permission_id]
    packet["proposed_contract_rows"]["permissions"] = [
        {
            "权限编号": new_permission_id,
            "模块/对象": "追踪",
            "操作": "查看",
            "权限代码": permission_code,
        }
    ]
    packet["permission_code_resolutions"][0]["proposed_perm_ids"] = [new_permission_id]
    output = tmp_path / "all-contract-categories.xlsx"

    report = generate_p0_candidate_workbook(paths, packet, output)

    assert report["candidate_written"] is True, report["checks"]
    assert report["traceability_assessment"]["traceability_ready"] is True
    assert report["traceability_assessment"]["counts"]["p0_requirements"] == 91
    workbook = load_workbook(output)
    try:
        assert _row_index(workbook["功能优先级"], "需求编号", new_requirement_id) > 0
        assert _row_index(workbook["API清单"], "接口编号", new_api_id) > 0
        permission_sheet = workbook["角色权限矩阵"]
        permission_row = _row_index(permission_sheet, "权限编号", new_permission_id)
        assert (
            permission_sheet.cell(
                permission_row,
                _headers(permission_sheet)["权限代码"],
            ).value
            == permission_code
        )
    finally:
        workbook.close()


def test_generator_refuses_overwrite_authority_doc_tree_and_bad_extension(
    tmp_path: Path,
) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    existing = tmp_path / "existing.xlsx"
    existing.write_bytes(b"keep")

    existing_report = generate_p0_candidate_workbook(paths, packet, existing)
    source_report = generate_p0_candidate_workbook(paths, packet, paths.technical_workbook)
    doc_report = generate_p0_candidate_workbook(
        paths,
        packet,
        paths.root / "doc" / "candidate.xlsx",
    )
    extension_report = generate_p0_candidate_workbook(
        paths,
        packet,
        tmp_path / "candidate.xls",
    )

    assert existing.read_bytes() == b"keep"
    assert "P0_CHANGE_OUTPUT_EXISTS" in _codes(existing_report)
    assert "P0_CHANGE_SOURCE_WORKBOOK_FORBIDDEN" in _codes(source_report)
    assert "P0_CHANGE_OUTPUT_LOCATION_FORBIDDEN" in _codes(source_report)
    assert "P0_CHANGE_OUTPUT_LOCATION_FORBIDDEN" in _codes(doc_report)
    assert "P0_CHANGE_OUTPUT_EXTENSION_INVALID" in _codes(extension_report)


def test_generator_does_not_publish_a_failed_candidate(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    output = tmp_path / "failed-candidate.xlsx"
    monkeypatch.setattr(
        "navigator_data_readiness.p0_baseline_change._apply_resolution_to_workbook",
        lambda *_args: None,
    )

    report = generate_p0_candidate_workbook(paths, packet, output)

    assert report["candidate_written"] is False
    assert report["candidate_ready_for_formal_baseline_review"] is False
    assert not output.exists()
    assert list(tmp_path.glob(".failed-candidate.*.xlsx")) == []


def test_generator_cleans_up_after_generation_exception(
    tmp_path: Path,
    monkeypatch: MonkeyPatch,
) -> None:
    paths = discover_repository()
    packet = _completed_resolution()
    output = tmp_path / "broken-candidate.xlsx"

    def fail_generation(*_args: Any) -> None:
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(
        "navigator_data_readiness.p0_baseline_change._apply_resolution_to_workbook",
        fail_generation,
    )

    report = generate_p0_candidate_workbook(paths, packet, output)

    assert _codes(report) == {"P0_CHANGE_GENERATION_FAILED"}
    assert report["candidate_written"] is False
    assert not output.exists()
    assert list(tmp_path.glob(".broken-candidate.*.xlsx")) == []


def test_generator_loader_rejects_invalid_resolution_without_writing(tmp_path: Path) -> None:
    paths = discover_repository()
    resolution = tmp_path / "resolution.json"
    resolution.write_text("[]", encoding="utf-8")
    output = tmp_path / "candidate.xlsx"

    report = load_and_generate_p0_candidate_workbook(paths, resolution, output)

    assert _codes(report) == {"P0_CHANGE_RESOLUTION_INVALID"}
    assert report["candidate_written"] is False
    assert not output.exists()
