from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from navigator_data_readiness.baseline import sha256_file
from navigator_data_readiness.basic60_private import (
    ACCEPTANCE_TEMPLATE_NAME,
    EXCEL_REVIEW_APPROVAL_SCHEMA,
    EXCEL_REVIEW_MACHINE_SHEET,
    EXCEL_REVIEW_SCOPE,
    EXCEL_REVIEW_WORKBOOK_SCHEMA,
    RAW_MANIFEST_NAME,
    READY_SEED_NAME,
    SEED_NAME,
    _canonical_sha256,
    build_basic60_excel_review_gate,
    build_basic60_excel_review_payload,
    load_and_assess_basic60_private,
    prepare_basic60_excel_review_bundle,
    prepare_basic60_private,
    write_basic60_ready_seed,
    write_basic60_release_authorization,
)
from navigator_data_readiness.paths import RepositoryPaths
from openpyxl import Workbook
from test_basic60_private import (
    _authorize_pbd_and_d1,
)
from test_basic60_private import (
    basic60_fixture as _base_basic60_fixture,
)


@pytest.fixture
def basic60_fixture(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> tuple[RepositoryPaths, Path, Path, Path]:
    return _base_basic60_fixture.__wrapped__(tmp_path, monkeypatch)


def _write_review_workbook(
    paths: RepositoryPaths,
    raw_dir: Path,
    runtime_dir: Path,
    candidate_dir: Path,
) -> Path:
    seed_path = runtime_dir / SEED_NAME
    source_manifest_path = candidate_dir / "basic60_country_source_uri_manifest.json"
    seed = json.loads(seed_path.read_text(encoding="utf-8"))
    source_manifest = json.loads(source_manifest_path.read_text(encoding="utf-8"))
    payload = build_basic60_excel_review_payload(
        raw_dir,
        seed,
        seed_sha256=sha256_file(seed_path),
        source_uri_manifest=source_manifest,
        source_uri_manifest_sha256=sha256_file(source_manifest_path),
    )
    payload_sha256 = _canonical_sha256(payload)
    workbook = Workbook()
    workbook.active.title = "审核说明"
    machine = workbook.create_sheet(EXCEL_REVIEW_MACHINE_SHEET)
    metadata: dict[str, str | int] = {
        "schema_version": EXCEL_REVIEW_WORKBOOK_SCHEMA,
        "release_id": "BASIC60-PRIVATE-R1",
        "profile_id": "basic60_private",
        "canonical_payload_sha256": payload_sha256,
        "country_count": 60,
        "macro_row_count": 300,
        "energy_row_count": 60,
        "available_observation_count": 2279,
        "pending_observation_count": 61,
        "seed_sha256": sha256_file(seed_path),
        "source_uri_manifest_sha256": sha256_file(source_manifest_path),
    }
    for row_index, item in enumerate(metadata.items(), start=1):
        machine.cell(row_index, 1, item[0])
        machine.cell(row_index, 2, item[1])
    output = paths.root / "outputs/basic60/BASIC60-PRIVATE-R1_60国基础数据集中审核.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    return output


def _single_excel_bundle(
    paths: RepositoryPaths,
    raw_dir: Path,
    runtime_dir: Path,
    candidate_dir: Path,
) -> tuple[dict[str, Any], Path]:
    prepare_basic60_private(paths, raw_dir, runtime_dir, candidate_dir)
    bundle = json.loads((candidate_dir / ACCEPTANCE_TEMPLATE_NAME).read_text(encoding="utf-8"))
    _authorize_pbd_and_d1(paths, bundle)
    workbook_path = _write_review_workbook(paths, raw_dir, runtime_dir, candidate_dir)
    expected_gate = build_basic60_excel_review_gate(
        paths,
        raw_dir,
        runtime_dir / SEED_NAME,
        candidate_dir / "basic60_country_source_uri_manifest.json",
        workbook_path,
    )
    input_bundle_path = paths.root / "data/basic60/review/authorized-input.json"
    input_bundle_path.parent.mkdir(parents=True, exist_ok=True)
    input_bundle_path.write_text(json.dumps(bundle, ensure_ascii=False), encoding="utf-8")
    bundle_path = paths.root / "data/basic60/review/single-excel-review.json"
    _output, bundle = prepare_basic60_excel_review_bundle(
        paths,
        raw_dir,
        runtime_dir / SEED_NAME,
        candidate_dir / "basic60_country_source_uri_manifest.json",
        input_bundle_path,
        workbook_path,
        bundle_path,
    )
    assert bundle["single_excel_review"] == expected_gate
    assert "collection" not in bundle
    assert "acceptance" not in bundle
    return bundle, bundle_path


def test_single_excel_review_has_only_one_human_blocker(
    basic60_fixture: tuple[RepositoryPaths, Path, Path, Path],
) -> None:
    paths, raw_dir, runtime_dir, candidate_dir = basic60_fixture
    _bundle, bundle_path = _single_excel_bundle(
        paths,
        raw_dir,
        runtime_dir,
        candidate_dir,
    )

    assessment = load_and_assess_basic60_private(
        paths,
        raw_dir,
        candidate_dir / RAW_MANIFEST_NAME,
        runtime_dir / SEED_NAME,
        bundle_path,
    )

    assert assessment["status"] == "not_ready"
    assert [item["code"] for item in assessment["checks"]] == ["B60_EXCEL_REVIEW_APPROVAL_PENDING"]


def test_exact_kevin_workbook_approval_reaches_private_trial_ready(
    basic60_fixture: tuple[RepositoryPaths, Path, Path, Path],
) -> None:
    paths, raw_dir, runtime_dir, candidate_dir = basic60_fixture
    bundle, bundle_path = _single_excel_bundle(
        paths,
        raw_dir,
        runtime_dir,
        candidate_dir,
    )
    gate = bundle["single_excel_review"]
    approval_path = paths.root / "data/basic60/evidence/basic60_excel_review_approval.json"
    approval_path.parent.mkdir(parents=True, exist_ok=True)
    approval_path.write_text(
        json.dumps(
            {
                "schema_version": EXCEL_REVIEW_APPROVAL_SCHEMA,
                "record_type": "basic60_single_excel_review_approval",
                "evidence_id": "EVD-B60-EXCEL-APPROVAL-001",
                "release_id": "BASIC60-PRIVATE-R1",
                "profile_id": "basic60_private",
                "decision": "approved",
                "person_name": "kevin",
                "role": "项目批准人",
                "approved_at": "2026-08-26T12:00:00+08:00",
                "review_scope": EXCEL_REVIEW_SCOPE,
                "authorized_outcome": "private_trial_ready",
                "workbook": gate["workbook"],
                "canonical_payload_sha256": gate["canonical_payload_sha256"],
                "formal_gate_status": "pending",
                "does_not_complete_formal_d1_d4": True,
                "does_not_authorize_p0_or_production": True,
                "authorization_basis": "explicit_user_approval_after_workbook_review",
                "template_only": False,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    gate["approval"] = {
        "path": approval_path.relative_to(paths.root).as_posix(),
        "sha256": sha256_file(approval_path),
    }
    bundle_path.write_text(json.dumps(bundle, ensure_ascii=False), encoding="utf-8")

    assessment = load_and_assess_basic60_private(
        paths,
        raw_dir,
        candidate_dir / RAW_MANIFEST_NAME,
        runtime_dir / SEED_NAME,
        bundle_path,
    )

    assert assessment["status"] == "private_trial_ready"
    assert assessment["checks"] == []
    validation_sha256 = "a" * 64
    ready_seed_path = runtime_dir / READY_SEED_NAME
    write_basic60_ready_seed(
        paths,
        candidate_dir / RAW_MANIFEST_NAME,
        runtime_dir / SEED_NAME,
        bundle_path,
        ready_seed_path,
        validation_report_sha256=validation_sha256,
    )
    authorization_path = runtime_dir / "basic60_release_authorization.json"
    write_basic60_release_authorization(
        paths,
        bundle_path,
        ready_seed_path,
        authorization_path,
        validation_report_sha256=validation_sha256,
    )
    authorization = json.loads(authorization_path.read_text(encoding="utf-8"))
    assert authorization["formal_gate_status"] == "pending"
    assert set(authorization["signoffs"]) == {"single_excel_review"}
    assert authorization["signoffs"]["single_excel_review"]["signature"]["person_id"] == ("kevin")


def test_workbook_drift_invalidates_single_excel_approval(
    basic60_fixture: tuple[RepositoryPaths, Path, Path, Path],
) -> None:
    paths, raw_dir, runtime_dir, candidate_dir = basic60_fixture
    bundle, bundle_path = _single_excel_bundle(
        paths,
        raw_dir,
        runtime_dir,
        candidate_dir,
    )
    workbook_path = paths.root / bundle["single_excel_review"]["workbook"]["path"]
    with workbook_path.open("ab") as destination:
        destination.write(b"drift")
    bundle_path.write_text(json.dumps(bundle, ensure_ascii=False), encoding="utf-8")

    assessment = load_and_assess_basic60_private(
        paths,
        raw_dir,
        candidate_dir / RAW_MANIFEST_NAME,
        runtime_dir / SEED_NAME,
        bundle_path,
    )

    assert assessment["status"] == "not_ready"
    assert "B60_EXCEL_REVIEW_WORKBOOK_INVALID" in {item["code"] for item in assessment["checks"]}
